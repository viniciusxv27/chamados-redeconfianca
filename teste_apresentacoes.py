"""Assistente de Apresentações — backend: acesso, documento, IA (dublê), mídias, módulos e templates.

Roda dentro de uma transação desfeita no fim: não grava nada no banco. A OpenAI
é um dublê (nenhuma chamada sai), as threads não rodam (as tarefas são
executadas aqui mesmo) e os arquivos vão para um armazenamento em memória.
"""
import base64
import io
import json
import os
import sys
from datetime import timedelta
from types import SimpleNamespace
from unittest import mock

import django

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'redeconfianca.settings')
django.setup()

from django.conf import settings

if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS.append('testserver')

from django.test.utils import setup_test_environment

setup_test_environment()

import numpy as np
from django.contrib.auth import get_user_model
from django.contrib.staticfiles import finders
from django.core.cache import caches
from django.core.files.storage import InMemoryStorage
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import transaction
from django.test import Client
from django.utils import timezone
from PIL import Image, ImageDraw

from apresentacoes import formato, ia, importacao, limpeza, montagem, modulos, roteiro, tarefas
from apresentacoes.models import (Apresentacao, ConfiguracaoApresentacoes, MensagemIA, Midia, TarefaIA,
                                  TemplateApresentacao, VersaoApresentacao)
from apresentacoes.padrao import NOME as NOME_PADRAO
from apresentacoes.padrao import garantir_template_padrao
from users.models import Sector

User = get_user_model()
ok = fail = 0


def t(nome, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'  OK   {nome}')
    else:
        fail += 1
        print(f'  FALHA {nome} {extra}')


memoria = InMemoryStorage()
Midia._meta.get_field('arquivo').storage = memoria


def png(largura=320, altura=180, cor=(40, 10, 60), texto=None):
    img = Image.new('RGB', (largura, altura), cor)
    if texto:
        ImageDraw.Draw(img).rectangle((20, 20, 120, 50), fill=(255, 255, 255))
    saida = io.BytesIO()
    img.save(saida, 'PNG')
    return saida.getvalue()


def pdf_de_duas_paginas():
    paginas = [Image.new('RGB', (960, 540), (20, 20, 20)), Image.new('RGB', (960, 540), (200, 30, 90))]
    saida = io.BytesIO()
    paginas[0].save(saida, 'PDF', save_all=True, append_images=paginas[1:])
    return saida.getvalue()


# ---------------------------------------------------------------------------
# Dublê da OpenAI
# ---------------------------------------------------------------------------
class ChatFalso:
    def __init__(self):
        self.respostas = []
        self.chamadas = []

    def create(self, **parametros):
        self.chamadas.append(parametros)
        if not self.respostas:
            raise AssertionError('Chamada à IA sem resposta preparada')
        resposta = self.respostas.pop(0)
        if isinstance(resposta, Exception):
            raise resposta
        if isinstance(resposta, SimpleNamespace):
            return resposta
        return SimpleNamespace(choices=[SimpleNamespace(
            message=SimpleNamespace(content=json.dumps(resposta, ensure_ascii=False), refusal=None),
            finish_reason='stop')])


class ClienteFalso:
    def __init__(self):
        self.chat = SimpleNamespace(completions=ChatFalso())
        self.imagens = []
        self.vozes = []
        self.images = SimpleNamespace(generate=self._imagem)
        self.audio = SimpleNamespace(speech=SimpleNamespace(create=self._voz))

    def _imagem(self, **parametros):
        self.imagens.append(parametros)
        return SimpleNamespace(data=[SimpleNamespace(b64_json=base64.b64encode(png(64, 40)).decode(), url=None)])

    def _voz(self, **parametros):
        self.vozes.append(parametros)
        return SimpleNamespace(content=b'ID3-mp3-falso')


def slide(layout, **campos):
    base = {'layout': layout, 'rotulo': None, 'titulo': 'Título', 'titulo_destaque': None, 'subtitulo': None,
            'texto': None, 'itens': None, 'numero': None, 'colunas': None, 'tabela': None, 'imagem': None,
            'marcacoes': None, 'botao': None, 'notas': 'Fala do slide.'}
    base.update(campos)
    return base


ROTEIRO = {
    'tipo': 'apresentacao', 'perguntas': [], 'titulo': 'Campanha Vivo Total',
    'slides': [
        slide('capa', rotulo='campanha', titulo='vivo total', titulo_destaque='setembro', subtitulo='O que muda',
              botao='vamos juntos'),
        slide('topicos', rotulo='oferta', titulo='O que muda', texto='Três mudanças importantes para a loja.',
              itens=[{'titulo': 'Preço novo', 'texto': 'R$ 199 no combo', 'icone': 'fa-solid fa-tag'},
                     {'titulo': 'Mais dados', 'texto': None, 'icone': 'bolt'},
                     {'titulo': 'Bônus', 'texto': 'Para quem migra', 'icone': 'fa-regular fa-invalido<x>'}]),
        slide('passo_a_passo', titulo='Como vender', itens=[{'titulo': f'Passo {i}', 'texto': 'Faça isso', 'icone': None}
                                                            for i in range(1, 5)]),
        slide('imagem_texto', titulo='Material', texto='Veja o print.',
              imagem={'fonte': 'anexo', 'indice': 0, 'prompt': None, 'legenda': 'Print da Vivo'},
              itens=[{'titulo': 'Ponto', 'texto': 'detalhe', 'icone': None}]),
        slide('tela_anotada', titulo='A tela', imagem={'fonte': 'anexo', 'indice': 1, 'prompt': None, 'legenda': None},
              marcacoes=[{'x': 0.1, 'y': 0.2, 'texto': 'Clique aqui'}, {'x': 0.8, 'y': 0.9, 'texto': 'Confira'}]),
        slide('duas_colunas', titulo='Antes e depois', colunas=[{'titulo': 'Antes', 'itens': ['a', 'b']},
                                                                 {'titulo': 'Depois', 'itens': ['c']}]),
        slide('numero_destaque', titulo='Meta', numero={'valor': '120%', 'legenda': 'da meta'}, texto='Vamos bater.'),
        slide('tabela', titulo='Preços', tabela={'cabecalho': ['Plano', 'Preço'], 'linhas': [['Total', 'R$ 199']]}),
        slide('citacao', titulo='Frase', texto='Confiança se constrói todo dia.', subtitulo='Diretoria'),
        slide('imagem_texto', titulo='Ilustração', texto='Com imagem da IA.',
              imagem={'fonte': 'ia', 'indice': None, 'prompt': 'a phone glowing in neon', 'legenda': None}),
        slide('encerramento', rotulo='obrigado', titulo='bora', titulo_destaque='vender', subtitulo='Dúvidas? Fale com o gerente.'),
    ],
}

marcador = transaction.atomic()
marcador.__enter__()
patches = []
try:
    cliente = ClienteFalso()
    disparadas = []
    sino = mock.MagicMock()
    for alvo, novo in (
        ('apresentacoes.ia._cliente', lambda: cliente),
        ('apresentacoes.ia._http', mock.MagicMock(side_effect=AssertionError('HTTP real não pode sair no teste'))),
        ('apresentacoes.tarefas.disparar', lambda tarefa: disparadas.append(tarefa.pk)),
        ('apresentacoes.tarefas.duracao_de_audio', lambda *a, **k: 4.2),
        ('notifications.services.notification_service._send_in_app', sino),
    ):
        p = mock.patch(alvo, novo)
        p.start()
        patches.append(p)

    area = Sector.objects.create(name='ZZ Setor Apresentacoes')

    def novo(u, nome, **kw):
        return User.objects.create_user(username=u, email=f'{u}@exemplo-teste.local', password='S3nha!teste',
                                        sector=area, first_name=nome, last_name='Teste', **kw)

    chefe = novo('zzap.super', 'ZZApSuper', hierarchy='SUPERADMIN')
    ana = novo('zzap.ana', 'ZZApAna')
    bia = novo('zzap.bia', 'ZZApBia')
    sem = novo('zzap.sem', 'ZZApSem')
    cfg = ConfiguracaoApresentacoes.get()
    cfg.liberados.add(ana, bia)
    caches['local'].delete_many([f'apresentacoes:menu:{u.pk}' for u in (chefe, ana, bia, sem)])

    c_chefe, c_ana, c_bia, c_sem = Client(), Client(), Client(), Client()
    c_chefe.force_login(chefe)
    c_ana.force_login(ana)
    c_bia.force_login(bia)
    c_sem.force_login(sem)

    print('== ACESSO ==')
    r = c_sem.get('/apresentacoes/')
    t('sem liberação não entra (volta para a home)', r.status_code == 302, r.status_code)
    r = c_sem.get('/apresentacoes/1/documento/')
    t('e a API responde 403 em JSON', r.status_code == 403 and r.json()['ok'] is False, r.status_code)
    t('liberada entra', c_ana.get('/apresentacoes/').status_code == 200)
    t('SUPERADMIN entra sem estar na lista', c_chefe.get('/apresentacoes/').status_code == 200)
    html = c_ana.get('/apresentacoes/').content.decode()
    t('o menu mostra Apresentações para quem usa', 'Apresentações</span>' in html and '/apresentacoes/' in html)
    html = c_sem.get('/', follow=True).content.decode()
    t('e esconde de quem não usa', 'Apresentações</span>' not in html)
    t('configurações: só SUPERADMIN', c_ana.get('/apresentacoes/configuracoes/').status_code == 302
      and c_chefe.get('/apresentacoes/configuracoes/').status_code == 200)
    r = c_chefe.post('/apresentacoes/configuracoes/', {'secao': 'acesso', 'liberados': [ana.pk, bia.pk, sem.pk]})
    t('SUPERADMIN libera mais alguém', cfg.liberados.filter(pk=sem.pk).exists())
    t('e o menu de quem ganhou acesso atualiza na hora', c_sem.get('/apresentacoes/').status_code == 200)
    c_chefe.post('/apresentacoes/configuracoes/', {'secao': 'acesso', 'liberados': [ana.pk, bia.pk]})
    t('tirar a liberação vale na hora', c_sem.get('/apresentacoes/').status_code == 302)
    c_chefe.post('/apresentacoes/configuracoes/', {'secao': 'canva', 'canva_client_id': 'cli-123',
                                                   'canva_client_secret': 'segredo-super-secreto'})
    cfg.refresh_from_db()
    t('segredo do Canva guardado cifrado', cfg.canva_client_secret_cifrado and 'segredo-super-secreto' not in
      cfg.canva_client_secret_cifrado and cfg.get_canva_secret() == 'segredo-super-secreto')
    html = c_chefe.get('/apresentacoes/configuracoes/').content.decode()
    t('e nunca volta para a tela', 'segredo-super-secreto' not in html and cfg.canva_client_secret_cifrado not in html)
    c_chefe.post('/apresentacoes/configuracoes/', {'secao': 'ia', 'modelo_texto': 'gpt-5; drop', 'voz': 'coral'})
    cfg.refresh_from_db()
    t('modelo de IA com caractere estranho é ignorado', cfg.modelo_texto == 'gpt-4.1' and cfg.voz == 'coral', cfg.modelo_texto)

    print('\n== SANEADOR ==')
    t('script some', formato.sanear_html('<b>a</b><script>x()</script>') == '<b>a</b>')
    t('onerror e img somem', formato.sanear_html('<img src=x onerror=alert(1)>t') == 't')
    t('javascript: no link some', 'javascript' not in formato.sanear_html('<a href="javascript:alert(1)">x</a>'))
    t('url() no estilo some', 'url' not in formato.sanear_html('<span style="background: url(x); color: #fff">a</span>'))
    t('HTML aberto fecha dentro da caixa', formato.sanear_html('<b><i>x') == '<b><i>x</i></b>')
    doc = formato.sanear_documento({'slides': [
        {'id': 'mesmo123', 'fundo': {'imagem': 'https://externo/x.png'}, 'elementos': [
            {'id': 'e-igual1', 'tipo': 'imagem', 'src': 'https://externo/y.png'},
            {'id': 'e-igual1', 'tipo': 'texto', 'html': 'a', 'estilo': {'tamanho': 'NaN', 'cor': 'red;x'}},
            {'tipo': 'desconhecido'}]},
        {'id': 'mesmo123'}]})
    t('mídia externa descartada', doc['slides'][0]['fundo']['imagem'] == '' and doc['slides'][0]['elementos'][0]['src'] == '')
    t('tipo desconhecido descartado', len(doc['slides'][0]['elementos']) == 2)
    t('ids repetidos ganham outro', doc['slides'][0]['id'] != doc['slides'][1]['id']
      and doc['slides'][0]['elementos'][0]['id'] != doc['slides'][0]['elementos'][1]['id'])
    t('número inválido e cor inválida viram padrão', doc['slides'][0]['elementos'][1]['estilo']['tamanho'] == 36
      and doc['slides'][0]['elementos'][1]['estilo']['cor'] == 'tema:texto')
    t('gradiente com url() recusado', formato.gradiente('linear-gradient(90deg, #000, url(x))') == '')
    t('gradiente normal aceito', formato.gradiente('linear-gradient(90deg, #000 0%, rgba(255,0,0,.5) 100%)') != '')

    print('\n== TEMPLATE PADRÃO ==')
    padrao = garantir_template_padrao()
    t('template Rede Confiança existe e é da rede', padrao.nome == NOME_PADRAO and padrao.padrao_da_rede)
    t('garantir não duplica', garantir_template_padrao().pk == padrao.pk)
    t('5 layouts', [s['layout'] for s in padrao.layouts] == ['capa', 'secao', 'conteudo', 'quadro', 'encerramento'])
    fundos = {s['fundo']['imagem'] for s in padrao.layouts}
    t('os fundos limpos existem nos estáticos', all(finders.find(f[len('static:'):]) for f in fundos), fundos)
    t('usuário comum não edita o template da rede', not c_ana.get(f'/apresentacoes/templates/{padrao.pk}/').status_code == 200)

    print('\n== MONTAGEM ==')
    anexo_a = tarefas.salvar_midia(png(1600, 900), 'print-vivo.png', dono=ana, tipo=Midia.Tipo.IMAGEM,
                                   origem=Midia.Origem.PRINT)
    anexo_b = tarefas.salvar_midia(png(1440, 900), 'tela.png', dono=ana, tipo=Midia.Tipo.IMAGEM,
                                   origem=Midia.Origem.CAPTURA)
    documento = montagem.montar_documento(ROTEIRO, padrao, {3: anexo_a, 4: anexo_b})
    slides = documento['slides']
    t('um slide por item do roteiro', len(slides) == len(ROTEIRO['slides']))
    capa = slides[0]
    textos_capa = {e['slot']: e for e in capa['elementos'] if e['tipo'] == 'texto'}
    t('capa com rótulo, título em duas cores, subtítulo e botão',
      textos_capa['rotulo']['html'] == 'campanha' and 'vivo total<br><span style="color: #FEDB63">setembro</span>'
      in textos_capa['titulo']['html'] and textos_capa['subtitulo']['html'] == 'O que muda'
      and textos_capa['botao']['html'] == 'vamos juntos')
    t('capa usa o fundo do template', capa['fundo']['imagem'].endswith('capa.jpg'))
    t('paginação fica por cima de tudo', capa['elementos'][-1]['slot'] == 'paginacao')
    conteudo = slides[1]
    t('sem área de conteúdo tracejada na apresentação', not any(e['slot'] == 'area_conteudo' for s in slides for e in s['elementos']))
    t('slide sem botão não leva a pílula', not any(e['slot'] == 'botao' for e in conteudo['elementos']))
    icones = [e['icone'] for e in conteudo['elementos'] if e['tipo'] == 'icone']
    t('ícones: válido, curto completado e inválido trocado', icones == ['fa-solid fa-tag', 'fa-solid fa-bolt',
                                                                         'fa-solid fa-circle-check'], icones)
    t('tópicos entram em cascata (auto) com o conteúdo junto',
      any(e['animacao']['gatilho'] == 'auto' for e in conteudo['elementos'])
      and any(e['animacao']['gatilho'] == 'junto' for e in conteudo['elementos']))
    passos = slides[2]
    t('passo a passo numerado 01..04', [formato.html_para_texto(e['html']) for e in passos['elementos']
                                         if e['tipo'] == 'texto' and formato.html_para_texto(e['html']).startswith('0')]
      == ['01', '02', '03', '04'])
    com_print = slides[3]
    img = next(e for e in com_print['elementos'] if e['tipo'] == 'imagem')
    t('print anexado vira imagem inteira (contain) com a proporção certa',
      img['midia'] == anexo_a.pk and img['ajuste'] == 'contain' and abs(img['w'] / img['h'] - 16 / 9) < 0.02,
      (img['w'], img['h']))
    anotada = slides[4]
    img_tela = next(e for e in anotada['elementos'] if e['tipo'] == 'imagem')
    bolas = [e for e in anotada['elementos'] if e['tipo'] == 'forma' and e['forma'] == 'circulo'
             and e['preenchimento'] == 'tema:destaque' and e['w'] == 56]
    centro_x = bolas[0]['x'] + 28
    t('marcações ficam sobre a imagem na posição pedida',
      len(bolas) == 2 and abs(centro_x - (img_tela['x'] + 0.1 * img_tela['w'])) < 1.5, centro_x)
    t('tabela nativa com cabeçalho', any(e['tipo'] == 'tabela' and e['linhas'][0] == ['Plano', 'Preço'] for e in slides[7]['elementos']))
    t('número em destaque usa o quadro', slides[6]['fundo']['imagem'].endswith('quadro.jpg'))
    sem_imagem = slides[9]
    t('imagem da IA ausente: o slide vira tópicos sem perder o texto',
      not any(e['tipo'] == 'imagem' for e in sem_imagem['elementos'])
      and any('Com imagem da IA.' in e.get('html', '') for e in sem_imagem['elementos']))
    t('notas vão para o slide', all(s['notas'] == 'Fala do slide.' for s in slides))
    t('texto longo encolhe a fonte', montagem.ajustar_tamanho('palavra ' * 60, 400, 120, 40) < 40)
    t('texto curto mantém', montagem.ajustar_tamanho('Oi', 400, 120, 40) == 40)
    t('documento montado já passa no saneador sem mudar', formato.sanear_documento(documento) == documento)
    capa_longa = montagem.montar_slide(slide('capa', titulo='Vini Renova', titulo_destaque='Nova entrega'),
                                       padrao.documento, padrao.tema)
    titulo_longo = next(e for e in capa_longa['elementos'] if e['slot'] == 'titulo')
    est = titulo_longo['estilo']
    t('título grande: cada parte cabe numa linha (sem "Nova / entrega")', all(
        montagem.estimar_linhas(parte, est['tamanho'], titulo_longo['w'], est['peso'], est['espacamento'] * est['tamanho'] / 126)
        == 1 for parte in ('Vini Renova', 'Nova entrega')) and est['tamanho'] < 126, est['tamanho'])
    ilustracao = tarefas.salvar_midia(png(1536, 1024), 'ia.png', dono=ana, tipo=Midia.Tipo.IMAGEM,
                                      origem=Midia.Origem.IA_IMAGEM)
    secao_img = montagem.montar_slide(slide('secao', rotulo='Atenção!', titulo='Dicas e cuidados'), padrao.documento,
                                      padrao.tema, ilustracao)
    img_secao = next(e for e in secao_img['elementos'] if e['tipo'] == 'imagem')
    textos_secao = [e for e in secao_img['elementos'] if e['tipo'] == 'texto' and e['slot'] in ('titulo', 'rotulo')]
    t('seção com ilustração: os textos param antes da imagem', textos_secao and all(
        e['x'] + e['w'] <= img_secao['x'] for e in textos_secao), [(e['x'], e['w']) for e in textos_secao])

    print('\n== NOVA APRESENTAÇÃO (PEDIDO) ==')
    r = c_ana.post('/apresentacoes/nova/', {
        'pedido': 'Explique a campanha do print para as lojas.', 'template': padrao.pk, 'publico': 'Lojas',
        'tom': 'Profissional e motivador', 'quantidade': '8', 'fonte_titulo': 'Poppins', 'fonte_texto': 'Inter',
        'imagens_ia': 'on', 'narracao': 'on',
        'anexos': [SimpleUploadedFile('print.png', png(800, 450), content_type='image/png'),
                   SimpleUploadedFile('material.pdf', pdf_de_duas_paginas(), content_type='application/pdf'),
                   SimpleUploadedFile('virus.exe', b'MZ', content_type='application/octet-stream'),
                   SimpleUploadedFile('falsa.png', b'nao sou imagem', content_type='image/png')]})
    ap = Apresentacao.objects.filter(dono=ana).order_by('-pk').first()
    t('cria a apresentação e abre o editor', r.status_code == 302 and ap and r['Location'].endswith(f'/apresentacoes/{ap.pk}/'))
    t('status gerando, opções guardadas', ap.status == 'GERANDO' and ap.opcoes['quantidade'] == 8
      and ap.opcoes['fonte_titulo'] == 'Poppins' and ap.opcoes['imagens_ia'] is True and ap.opcoes['narracao'] is True)
    prints = list(Midia.objects.filter(apresentacao=ap, origem=Midia.Origem.PRINT))
    t('imagem + 2 páginas do PDF viram anexos; exe e imagem falsa ficam de fora', len(prints) == 3, len(prints))
    tarefa = TarefaIA.objects.get(apresentacao=ap, tipo=TarefaIA.Tipo.GERAR)
    t('tarefa GERAR criada com os anexos e disparada', sorted(tarefa.parametros['anexos']) == sorted(m.pk for m in prints)
      and tarefa.pk in disparadas)
    t('o pedido fica no diálogo', MensagemIA.objects.filter(apresentacao=ap, papel='USUARIO').exists())

    print('\n== IA PERGUNTA ANTES DE MONTAR ==')
    cliente.chat.completions.respostas = [{'tipo': 'perguntas', 'titulo': '', 'slides': [], 'perguntas': [
        {'id': 'p1', 'pergunta': 'Qual o período da campanha?', 'opcoes': ['Setembro', 'Outubro'], 'multipla': False}]}]
    tarefas.executar(tarefa.pk)
    tarefa.refresh_from_db()
    ap.refresh_from_db()
    t('tarefa para em PERGUNTAS', tarefa.status == 'PERGUNTAS' and tarefa.perguntas[0]['id'] == 'p1')
    t('apresentação também', ap.status == 'PERGUNTAS')
    chamada = cliente.chat.completions.chamadas[-1]
    partes = chamada['messages'][1]['content']
    t('a IA recebeu os 3 anexos como imagem', sum(1 for p in partes if p['type'] == 'image_url') == 3)
    t('com JSON Schema estrito', chamada['response_format']['json_schema']['strict'] is True
      and chamada['model'] == 'gpt-4.1')
    r = c_ana.post(f'/apresentacoes/tarefas/{tarefa.pk}/responder/', data=json.dumps({'respostas': {'p1': 'Setembro'}}),
                   content_type='application/json')
    tarefa.refresh_from_db()
    t('responder devolve a tarefa à fila', r.status_code == 200 and tarefa.status == 'PENDENTE'
      and tarefa.parametros['historico'][0]['respostas'] == {'p1': 'Setembro'})
    roteiro_final = json.loads(json.dumps(ROTEIRO))
    roteiro_final['slides'][3]['imagem'] = {'fonte': 'anexo', 'indice': 0, 'prompt': None, 'legenda': None}
    roteiro_final['slides'][4]['imagem'] = {'fonte': 'anexo', 'indice': 9, 'prompt': None, 'legenda': None}
    cliente.chat.completions.respostas = [roteiro_final]
    tarefas.executar(tarefa.pk)
    tarefa.refresh_from_db()
    ap.refresh_from_db()
    t('geração concluída', tarefa.status == 'CONCLUIDA' and ap.status == 'PRONTA', (tarefa.status, tarefa.erro))
    ultima = cliente.chat.completions.chamadas[-1]['messages']
    t('a resposta do usuário foi para a IA', any('Setembro' in (m['content'] if isinstance(m['content'], str) else '')
                                                 for m in ultima))
    t('documento gravado com todos os slides', len(ap.slides) == len(ROTEIRO['slides']) and ap.revisao >= 1)
    t('título da IA', ap.titulo == 'Campanha Vivo Total')
    t('fontes escolhidas no pedido valem sobre as do template',
      ap.documento['tema']['fonte_titulo'] == 'Poppins' and ap.documento['tema']['fonte_texto'] == 'Inter')
    t('imagem pedida à IA foi gerada e salva', Midia.objects.filter(apresentacao=ap, origem='IA_IMAGEM').count() == 1
      and len(cliente.imagens) == 1 and 'no text' in cliente.imagens[0]['prompt'].lower())
    t('anexo com índice inexistente é ignorado (sem quebrar)', not any(
        e['tipo'] == 'imagem' for e in ap.slides[4]['elementos']))
    t('narração gravada em cada slide', all((s.get('narracao') or {}).get('duracao') == 4.2 for s in ap.slides)
      and len(cliente.vozes) == len(ap.slides))
    t('versão guardada', VersaoApresentacao.objects.filter(apresentacao=ap, motivo='Gerada pela IA').exists())
    t('aviso no sino (só no portal)', sino.called and 'pronta' in sino.call_args[0][1].lower())

    print('\n== PERGUNTAS TÊM LIMITE ==')
    ap2 = Apresentacao.objects.create(titulo='Limite', pedido='Algo', dono=bia, template=padrao, status='GERANDO')
    tarefa2 = TarefaIA.objects.create(tipo='GERAR', usuario=bia, apresentacao=ap2, parametros={
        'anexos': [], 'historico': [{'perguntas': [{'id': 'a', 'pergunta': 'x'}], 'respostas': {'a': '1'}}] * 2})
    cliente.chat.completions.respostas = [{'tipo': 'perguntas', 'titulo': '', 'slides': [], 'perguntas': [
        {'id': 'b', 'pergunta': 'de novo?', 'opcoes': [], 'multipla': False}]}]
    tarefas.executar(tarefa2.pk)
    tarefa2.refresh_from_db()
    ap2.refresh_from_db()
    t('depois de 2 rodadas não pergunta mais: erro claro', tarefa2.status == 'ERRO' and ap2.status == 'ERRO'
      and 'slide' in tarefa2.erro.lower(), tarefa2.erro)
    t('e o pedido mandou a IA gerar sem perguntar', 'sem novas perguntas' in
      cliente.chat.completions.chamadas[-1]['messages'][-1]['content'])

    print('\n== DOCUMENTO (SALVAR) ==')
    url_doc = f'/apresentacoes/{ap.pk}/documento/'
    r = c_ana.get(url_doc)
    dados = r.json()
    t('GET devolve documento e revisão', r.status_code == 200 and dados['revisao'] == ap.revisao)
    novo_doc = dados['documento']
    novo_doc['slides'][0]['elementos'].append({'tipo': 'texto', 'html': '<b>ok</b><script>1</script>', 'x': 1, 'y': 1})
    r = c_ana.post(url_doc, data=json.dumps({'documento': novo_doc, 'revisao': dados['revisao'], 'titulo': 'Novo título'}),
                   content_type='application/json')
    t('POST salva saneado e sobe a revisão', r.status_code == 200 and r.json()['revisao'] == dados['revisao'] + 1
      and r.json()['documento']['slides'][0]['elementos'][-1]['html'] == '<b>ok</b>')
    r = c_ana.post(url_doc, data=json.dumps({'documento': novo_doc, 'revisao': dados['revisao']}),
                   content_type='application/json')
    t('revisão velha dá 409', r.status_code == 409 and r.json()['revisao'] == dados['revisao'] + 1)
    t('outra pessoa não lê', c_bia.get(url_doc).status_code == 403)
    t('SUPERADMIN lê e edita', c_chefe.get(url_doc).status_code == 200)
    midia_da_bia = tarefas.salvar_midia(png(), 'bia.png', dono=bia, tipo='IMAGEM', origem='UPLOAD')
    doc_roubo = r.json() if False else c_ana.get(url_doc).json()
    doc_roubo['documento']['slides'][0]['elementos'].append(
        {'tipo': 'imagem', 'src': midia_da_bia.url, 'x': 0, 'y': 0, 'w': 10, 'h': 10})
    r = c_ana.post(url_doc, data=json.dumps({'documento': doc_roubo['documento'], 'revisao': doc_roubo['revisao']}),
                   content_type='application/json')
    t('não dá para usar mídia de outra pessoa no documento', r.status_code == 403, r.status_code)

    print('\n== MÍDIAS ==')
    r = c_ana.post('/apresentacoes/midias/', {'arquivo': SimpleUploadedFile('foto.png', png(), 'image/png'),
                                              'apresentacao': ap.pk})
    midia = Midia.objects.get(pk=r.json()['midia']['id'])
    t('upload de imagem', r.status_code == 200 and midia.apresentacao_id == ap.pk and midia.largura == 320)
    r = c_ana.post('/apresentacoes/midias/', {'arquivo': SimpleUploadedFile('tela.jpg', png(), 'image/jpeg'),
                                              'apresentacao': ap.pk, 'origem': 'CAPTURA', 'nome': 'Início'})
    captura_sem_ext = Midia.objects.get(pk=r.json()['midia']['id'])
    t('captura com nome sem extensão guarda o arquivo com extensão', captura_sem_ext.nome == 'Início'
      and captura_sem_ext.arquivo.name.endswith('.jpg'), captura_sem_ext.arquivo.name)
    r = c_ana.post('/apresentacoes/midias/', {'arquivo': SimpleUploadedFile('x.png', b'texto', 'image/png')})
    t('imagem falsa recusada', r.status_code == 400)
    r = c_ana.post('/apresentacoes/midias/', {'arquivo': SimpleUploadedFile('x.svg', b'<svg onload=1>', 'image/svg+xml')})
    t('SVG recusado (pode carregar script)', r.status_code == 400)
    r = c_bia.post('/apresentacoes/midias/', {'arquivo': SimpleUploadedFile('f.png', png(), 'image/png'),
                                              'apresentacao': ap.pk})
    t('upload na apresentação de outra pessoa: 403', r.status_code == 403)
    r = c_ana.get(midia.url)
    t('dona baixa a mídia', r.status_code == 200 and b''.join(r.streaming_content)[:4] == b'\x89PNG'
      and r['X-Content-Type-Options'] == 'nosniff')
    t('outra pessoa recebe 404', c_bia.get(midia.url).status_code == 404)
    t('quem não usa o módulo também não', c_sem.get(midia.url).status_code in (302, 404))
    audio = tarefas.salvar_midia(b'0123456789' * 10, 'fala.mp3', dono=ana, tipo='AUDIO', origem='IA_VOZ',
                                 apresentacao=ap, mime='audio/mpeg')
    r = c_ana.get(audio.url, HTTP_RANGE='bytes=10-19')
    t('áudio responde por faixa (206)', r.status_code == 206 and b''.join(r.streaming_content) == b'0123456789'
      and r['Content-Range'] == 'bytes 10-19/100', r.status_code)
    lista = c_ana.get(f'/apresentacoes/{ap.pk}/midias/').json()['midias']
    t('biblioteca da apresentação', any(m['id'] == midia.pk for m in lista))

    print('\n== IA NO EDITOR ==')
    r = c_ana.post(f'/apresentacoes/{ap.pk}/ia/', data=json.dumps({'acao': 'editar', 'instrucao': ''}),
                   content_type='application/json')
    t('editar sem pedido é recusado', r.status_code == 400)
    doc_atual = c_ana.get(url_doc).json()['documento']
    ids = [s['id'] for s in doc_atual['slides']]
    titulo_capa = next(e for e in doc_atual['slides'][0]['elementos'] if e['slot'] == 'titulo')
    r = c_ana.post(f'/apresentacoes/{ap.pk}/ia/', data=json.dumps({
        'acao': 'editar', 'instrucao': 'Encurte a capa, remova o slide de citação e troque a fonte', 'documento': doc_atual}),
        content_type='application/json')
    tarefa_edicao = TarefaIA.objects.get(pk=r.json()['tarefa']['id'])
    t('pedido de edição vira tarefa com o documento da tela', r.status_code == 200 and tarefa_edicao.tipo == 'EDITAR'
      and tarefa_edicao.parametros.get('documento'))
    cliente.chat.completions.respostas = [{'resumo': 'Capa encurtada.', 'perguntas': [], 'tema': {
        'fonte_titulo': 'Oswald', 'fonte_texto': None, 'primaria': '#123456', 'destaque': 'vermelho'}, 'operacoes': [
        {'op': 'editar_texto', 'slide_id': ids[0], 'elemento_id': titulo_capa['id'], 'texto': 'vivo\n<b>total</b>',
         'posicao': None, 'slide': None},
        {'op': 'remover_slide', 'slide_id': ids[8], 'elemento_id': None, 'texto': None, 'posicao': None, 'slide': None},
        {'op': 'mover_slide', 'slide_id': ids[1], 'elemento_id': None, 'texto': None, 'posicao': 3, 'slide': None},
        {'op': 'novo_slide', 'slide_id': None, 'elemento_id': None, 'texto': None, 'posicao': 1,
         'slide': slide('topicos', titulo='Novo', itens=[{'titulo': 'X', 'texto': None, 'icone': None}])},
        {'op': 'editar_texto', 'slide_id': 's-naoexiste', 'elemento_id': 'e-x', 'texto': 'nada', 'posicao': None,
         'slide': None}]}]
    tarefas.executar(tarefa_edicao.pk)
    tarefa_edicao.refresh_from_db()
    resultado = tarefa_edicao.resultado
    novo = resultado['documento']
    t('edição concluída com 5 mudanças válidas (a inválida é ignorada)', tarefa_edicao.status == 'CONCLUIDA'
      and resultado['operacoes'] == 5, (tarefa_edicao.status, resultado.get('operacoes'), tarefa_edicao.erro))
    capa_nova = next(s for s in novo['slides'] if s['id'] == ids[0])
    t('texto editado escapado, com quebra de linha', next(e for e in capa_nova['elementos'] if e['id'] == titulo_capa['id'])
      ['html'] == 'vivo<br>&lt;b&gt;total&lt;/b&gt;')
    t('slide removido', ids[8] not in [s['id'] for s in novo['slides']])
    t('novo slide no início (depois da posição 1)', novo['slides'][1]['nome'] == 'Novo')
    t('fonte e cor do tema trocadas; cor inválida ignorada', novo['tema']['fonte_titulo'] == 'Oswald'
      and novo['tema']['cores']['primaria'] == '#123456' and novo['tema']['cores']['destaque'] == '#FEDB63')
    ap.refresh_from_db()
    t('a edição da IA não grava sozinha (a tela aplica)', len(ap.slides) == len(doc_atual['slides']))

    r = c_ana.post(f'/apresentacoes/{ap.pk}/ia/', data=json.dumps({'acao': 'refazer_slide', 'slide_id': ids[2],
                                                                  'instrucao': 'mais visual', 'documento': doc_atual}),
                   content_type='application/json')
    refazer = TarefaIA.objects.get(pk=r.json()['tarefa']['id'])
    cliente.chat.completions.respostas = [{'tipo': 'apresentacao', 'perguntas': [], 'titulo': '', 'slides': [
        slide('numero_destaque', titulo='Resultado', numero={'valor': '98%', 'legenda': 'satisfação'})]}]
    tarefas.executar(refazer.pk)
    refazer.refresh_from_db()
    t('refazer slide mantém o id do slide', refazer.status == 'CONCLUIDA' and refazer.resultado['slide']['id'] == ids[2]
      and refazer.resultado['slide_id'] == ids[2], refazer.erro)

    r = c_ana.post(f'/apresentacoes/{ap.pk}/ia/', data=json.dumps({'acao': 'novo_slide', 'slide_id': ids[0],
                                                                  'instrucao': 'slide sobre metas'}),
                   content_type='application/json')
    novo_slide = TarefaIA.objects.get(pk=r.json()['tarefa']['id'])
    cliente.chat.completions.respostas = [{'tipo': 'apresentacao', 'perguntas': [], 'titulo': '', 'slides': [
        slide('topicos', titulo='Metas', itens=[{'titulo': 'Meta 1', 'texto': None, 'icone': None}])]}]
    tarefas.executar(novo_slide.pk)
    novo_slide.refresh_from_db()
    t('novo slide vem com a posição', novo_slide.status == 'CONCLUIDA' and novo_slide.resultado['depois_de'] == ids[0]
      and novo_slide.resultado['slides'][0]['nome'] == 'Metas')

    r = c_ana.post(f'/apresentacoes/{ap.pk}/ia/', data=json.dumps({'acao': 'texto', 'html': '<b>texto antigo</b>',
                                                                  'instrucao': 'mais curto'}),
                   content_type='application/json')
    reescrita = TarefaIA.objects.get(pk=r.json()['tarefa']['id'])
    cliente.chat.completions.respostas = [{'texto': 'Curto & <direto>'}]
    tarefas.executar(reescrita.pk)
    reescrita.refresh_from_db()
    t('reescrever devolve HTML seguro', reescrita.resultado.get('html') == 'Curto &amp; &lt;direto&gt;')

    r = c_ana.post(f'/apresentacoes/{ap.pk}/ia/', data=json.dumps({'acao': 'imagem', 'instrucao': 'loja moderna',
                                                                  'opcoes': {'formato': 'quadrado'}}),
                   content_type='application/json')
    imagem_tarefa = TarefaIA.objects.get(pk=r.json()['tarefa']['id'])
    tarefas.executar(imagem_tarefa.pk)
    imagem_tarefa.refresh_from_db()
    t('imagem avulsa gerada no formato pedido', imagem_tarefa.status == 'CONCLUIDA'
      and imagem_tarefa.resultado['midia']['tipo'] == 'IMAGEM' and cliente.imagens[-1]['size'] == '1024x1024')

    with mock.patch('apresentacoes.ia.gerar_video', return_value=b'\x00\x00\x00\x18ftypmp42') as video_falso:
        r = c_ana.post(f'/apresentacoes/{ap.pk}/ia/', data=json.dumps({'acao': 'video', 'instrucao': 'abertura',
                                                                      'opcoes': {'segundos': '99'}}),
                       content_type='application/json')
        video_tarefa = TarefaIA.objects.get(pk=r.json()['tarefa']['id'])
        tarefas.executar(video_tarefa.pk)
        video_tarefa.refresh_from_db()
    t('vídeo com IA (dublê): segundos inválidos viram 8', video_tarefa.status == 'CONCLUIDA'
      and video_falso.call_args.kwargs['segundos'] == '8' and video_tarefa.resultado['midia']['tipo'] == 'VIDEO')

    doc_sem_notas = c_ana.get(url_doc).json()['documento']
    for s in doc_sem_notas['slides']:
        s['notas'] = ''
    alvo = doc_sem_notas['slides'][0]['id']
    r = c_ana.post(f'/apresentacoes/{ap.pk}/ia/', data=json.dumps({'acao': 'narracao', 'documento': doc_sem_notas,
                                                                  'slides': [alvo]}),
                   content_type='application/json')
    narracao = TarefaIA.objects.get(pk=r.json()['tarefa']['id'])
    cliente.chat.completions.respostas = [{'notas': [{'slide_id': alvo, 'texto': 'Bem-vindos à campanha.'}]}]
    vozes_antes = len(cliente.vozes)
    tarefas.executar(narracao.pk)
    narracao.refresh_from_db()
    t('narração escreve as notas que faltam e grava só o slide pedido', narracao.status == 'CONCLUIDA'
      and list(narracao.resultado['narracoes']) == [alvo] and narracao.resultado['notas'][alvo] == 'Bem-vindos à campanha.'
      and len(cliente.vozes) == vozes_antes + 1, narracao.erro)

    print('\n== TAREFAS ==')
    t('outra pessoa não vê a tarefa', c_bia.get(f'/apresentacoes/tarefas/{narracao.pk}/').status_code == 404)
    t('SUPERADMIN vê', c_chefe.get(f'/apresentacoes/tarefas/{narracao.pk}/').status_code == 200)
    parada = TarefaIA.objects.create(tipo='IMAGEM', usuario=bia, status='RODANDO')
    TarefaIA.objects.filter(pk=parada.pk).update(atualizado_em=timezone.now() - timedelta(minutes=40))
    parada.refresh_from_db()
    tarefas.conferir_parada(parada)
    t('tarefa sem sinal de vida vira erro explicado', parada.status == 'ERRO' and 'reiniciou' in parada.erro)
    for _ in range(4):
        TarefaIA.objects.create(tipo='IMAGEM', usuario=bia, status='PENDENTE')
    ap_bia = Apresentacao.objects.create(titulo='Bia', dono=bia, template=padrao)
    r = c_bia.post(f'/apresentacoes/{ap_bia.pk}/ia/', data=json.dumps({'acao': 'imagem', 'instrucao': 'x'}),
                   content_type='application/json')
    t('mais de 4 pedidos abertos por pessoa: 429', r.status_code == 429, r.status_code)
    with mock.patch.object(ia, '_cliente', side_effect=ia.IAIndisponivel('A chave da OpenAI (OPENAI_API_KEY) não está configurada no portal.')):
        falha = TarefaIA.objects.create(tipo='TEXTO', usuario=ana, apresentacao=ap, parametros={'html': 'a'})
        tarefas.executar(falha.pk)
        falha.refresh_from_db()
    t('sem chave: erro em português, sem stack', falha.status == 'ERRO' and 'OPENAI_API_KEY' in falha.erro)

    print('\n== CLIENTE DA IA ==')
    erro_modelo = type('NotFoundError', (Exception,), {})('model_not_found: gpt-4.1')
    cliente.chat.completions.respostas = [erro_modelo, {'texto': 'ok'}]
    resposta = ia.chamar_json([{'role': 'user', 'content': 'x'}], roteiro.SCHEMA_TEXTO, 'x', 'gpt-4.1')
    t('modelo inexistente cai para o reserva', resposta == {'texto': 'ok'}
      and cliente.chat.completions.chamadas[-1]['model'] == ia.MODELO_RESERVA)
    cliente.chat.completions.respostas = [{'texto': 'ok'}]
    ia.chamar_json([{'role': 'user', 'content': 'x'}], roteiro.SCHEMA_TEXTO, 'x', 'gpt-5')
    chamada = cliente.chat.completions.chamadas[-1]
    t('gpt-5 usa max_completion_tokens e sem temperatura', 'max_completion_tokens' in chamada
      and 'temperature' not in chamada and 'max_tokens' not in chamada)
    cliente.chat.completions.respostas = [SimpleNamespace(choices=[SimpleNamespace(
        message=SimpleNamespace(content='{}', refusal='não posso'), finish_reason='stop')])]
    try:
        ia.chamar_json([{'role': 'user', 'content': 'x'}], roteiro.SCHEMA_TEXTO, 'x', 'gpt-4o')
        recusou = False
    except ia.IAIndisponivel:
        recusou = True
    t('recusa da IA vira erro legível', recusou)
    t('mensagens de erro sem segredo', ia.mensagem_de_erro(Exception('sk-abc insufficient_quota')) ==
      'A conta da OpenAI está sem crédito. Avise o responsável pelo portal.')

    def schemas_estritos(schema, caminho='raiz'):
        problemas = []
        if isinstance(schema, dict):
            if schema.get('type') == 'object':
                if schema.get('additionalProperties') is not False:
                    problemas.append(caminho)
                if set(schema.get('required', [])) != set(schema.get('properties', {})):
                    problemas.append(caminho + ' (required)')
            for chave, valor in schema.items():
                problemas += schemas_estritos(valor, f'{caminho}.{chave}')
        elif isinstance(schema, list):
            for i, valor in enumerate(schema):
                problemas += schemas_estritos(valor, f'{caminho}[{i}]')
        return problemas
    todos = [roteiro.SCHEMA_ROTEIRO, roteiro.SCHEMA_EDICAO, roteiro.SCHEMA_TEXTO, roteiro.SCHEMA_NOTAS,
             roteiro.SCHEMA_TEMPLATE]
    problemas = sum((schemas_estritos(s) for s in todos), [])
    t('todos os JSON Schemas seguem o modo estrito da OpenAI', not problemas, problemas[:5])

    print('\n== MÓDULOS DO PORTAL ==')
    requisicao = c_chefe.get('/apresentacoes/modulos/?atualizar=1')
    catalogo = requisicao.context['modulos']
    labels = {m['label'] for m in catalogo}
    t('catálogo sai do menu do SUPERADMIN', requisicao.status_code == 200 and 'renova' in labels and len(labels) >= 10,
      labels)
    t('o próprio módulo de apresentações fica fora', 'apresentacoes' not in labels)
    contexto = modulos.contexto_do_modulo('renova', 'Vini Renova')
    t('contexto do módulo traz modelos, rotas e frases dos testes', 'Modelos e campos' in contexto
      and '/renova/' in contexto and 'Comportamentos garantidos pelos testes' in contexto)
    paginas = modulos.paginas_para_capturar('renova', [{'url': '/renova/', 'nome': 'Vini Renova'}])
    t('páginas para capturar: sem parâmetro e sem ação', paginas and paginas[0]['url'] == '/renova/'
      and all('<' not in p['url'] and 'csv' not in p['url'] for p in paginas) and len(paginas) <= 6, paginas)
    r = c_chefe.post('/apresentacoes/modulos/renova/gerar/', data=json.dumps({
        'publico': 'Vendedores', 'instrucoes': 'foque na aprovação', 'imagens_ia': True, 'narracao': True}),
        content_type='application/json')
    criado = r.json()
    ap_mod = Apresentacao.objects.get(pk=criado['id'])
    t('gerar do módulo cria rascunho e devolve as telas', r.status_code == 200 and ap_mod.origem == 'MODULO'
      and ap_mod.modulo == 'renova' and criado['paginas'] and ap_mod.status == 'RASCUNHO')
    t('pedido explica a entrega', 'Vini Renova' in ap_mod.pedido and 'foque na aprovação' in ap_mod.pedido)
    r = c_ana.post('/apresentacoes/modulos/modulo-que-nao-existe/gerar/', data='{}', content_type='application/json')
    t('módulo fora do menu: 404', r.status_code == 404)
    captura = c_chefe.post('/apresentacoes/midias/', {'arquivo': SimpleUploadedFile('tela.jpg', png(1440, 900), 'image/png'),
                                                      'apresentacao': ap_mod.pk, 'origem': 'CAPTURA', 'nome': 'Início'}).json()
    r = c_chefe.post(criado['iniciar'], data=json.dumps({'midias': [captura['midia']['id'], midia_da_bia.pk]}),
                     content_type='application/json')
    tarefa_mod = TarefaIA.objects.filter(apresentacao=ap_mod, tipo='GERAR').first()
    t('iniciar leva só as capturas da própria apresentação', r.status_code == 200 and tarefa_mod
      and tarefa_mod.parametros['anexos'] == [captura['midia']['id']])
    cliente.chat.completions.respostas = [{'tipo': 'apresentacao', 'perguntas': [], 'titulo': 'Vini Renova — como usar',
                                           'slides': [slide('capa', titulo='vini', titulo_destaque='renova'),
                                                      slide('tela_anotada', titulo='Tela inicial', imagem={
                                                          'fonte': 'anexo', 'indice': 0, 'prompt': None, 'legenda': None},
                                                            marcacoes=[{'x': 0.5, 'y': 0.5, 'texto': 'Nova avaliação'}]),
                                                      slide('encerramento', titulo='pronto')]}]
    tarefas.executar(tarefa_mod.pk)
    tarefa_mod.refresh_from_db()
    pedido_mod = cliente.chat.completions.chamadas[-1]['messages'][1]['content'][0]['text']
    t('a IA recebe o contexto do módulo e a captura', 'Contexto do módulo' in pedido_mod and 'captura de tela do portal'
      in pedido_mod and tarefa_mod.status == 'CONCLUIDA', tarefa_mod.erro)

    print('\n== CAPTURA EM IFRAME (MIDDLEWARE) ==')
    r = c_ana.get('/apresentacoes/?captura_apresentacao=1')
    t('com o marcador, quem usa o módulo pode emoldurar (SAMEORIGIN)', r['X-Frame-Options'] == 'SAMEORIGIN')
    r = c_ana.get('/apresentacoes/')
    t('sem o marcador continua DENY', r['X-Frame-Options'] == 'DENY')
    r = c_sem.get('/?captura_apresentacao=1', follow=True)
    t('quem não usa o módulo continua DENY', r['X-Frame-Options'] == 'DENY')

    print('\n== LIMPEZA DO FUNDO ==')
    fundo = np.zeros((200, 400, 3), dtype=np.uint8)
    fundo[:, :, 0] = np.linspace(20, 60, 400, dtype=np.uint8)
    fundo[:, :, 2] = 40
    imagem_fundo = Image.fromarray(fundo)
    ImageDraw.Draw(imagem_fundo).text((110, 90), 'TEXTO DE EXEMPLO', fill=(255, 255, 255))
    antes = np.asarray(imagem_fundo, dtype=np.float32)[85:110, 105:230].max()
    limpa = np.asarray(limpeza.limpar_areas(imagem_fundo, [{'x': 100, 'y': 80, 'w': 140, 'h': 35}]), dtype=np.float32)
    depois = limpa[85:110, 105:230].max()
    t('texto branco some do fundo escuro', antes > 200 and depois < 110, (antes, depois))
    t('o resto da imagem fica intacto', np.abs(limpa[150:, :] - fundo[150:, :]).max() < 1)

    print('\n== IMPORTAR TEMPLATE ==')
    tela_modelo = Image.new('RGB', (1920, 1080), (15, 5, 25))
    desenho = ImageDraw.Draw(tela_modelo)
    desenho.rectangle((150, 160, 700, 240), fill=(255, 255, 255))
    buffer = io.BytesIO()
    tela_modelo.save(buffer, 'PNG')
    r = c_ana.post('/apresentacoes/templates/novo/', {'nome': 'ZZ Template', 'arquivos': [
        SimpleUploadedFile('capa.png', buffer.getvalue(), 'image/png'),
        SimpleUploadedFile('conteudo.png', png(1280, 720), 'image/png')]})
    importado = TemplateApresentacao.objects.get(nome='ZZ Template')
    t('importar cria template em análise e abre o editor de layouts', r.status_code == 302
      and importado.status == 'ANALISANDO' and not importado.padrao_da_rede)
    imagens_template = list(Midia.objects.filter(template=importado).order_by('pk'))
    t('imagens normalizadas para 1920×1080', len(imagens_template) == 2 and all(
        (m.largura, m.altura) == (1920, 1080) for m in imagens_template))
    analise = TarefaIA.objects.get(template=importado)
    cliente.chat.completions.respostas = [{
        'nome': 'Novo', 'estilo': 'Escuro com destaque branco.', 'fonte_titulo': 'poppins', 'fonte_texto': 'Fonte Inventada',
        'cores': {'fundo': '#0F0519', 'superficie': '#1A1022', 'primaria': '#FF0080', 'secundaria': '#7700FF',
                  'destaque': '#FFCC00', 'texto': '#FFFFFF', 'texto_suave': '#CCCCCC'},
        'layouts': [{'indice': 0, 'papel': 'capa', 'textos': [
            {'slot': 'titulo', 'texto': 'Título', 'x': 0.075, 'y': 0.14, 'w': 0.3, 'h': 0.08, 'tamanho': 90, 'peso': 800,
             'cor': '#FFFFFF', 'alinhamento': 'left', 'maiusculas': False, 'espacamento': 0}],
            'formas': [], 'area_conteudo': {'x': 0.08, 'y': 0.3, 'w': 0.8, 'h': 0.5}},
            {'indice': 7, 'papel': 'conteudo', 'textos': [], 'formas': [], 'area_conteudo': {'x': 0, 'y': 0, 'w': 1, 'h': 1}}]}]
    tarefas.executar(analise.pk)
    analise.refresh_from_db()
    importado.refresh_from_db()
    t('análise concluída', analise.status == 'CONCLUIDA' and importado.status == 'PRONTO', analise.erro)
    t('layout com índice inexistente é ignorado', len(importado.layouts) == 1)
    layout = importado.layouts[0]
    titulo_el = next(e for e in layout['elementos'] if e['slot'] == 'titulo')
    t('caixa do título refinada pelos pixels', abs(titulo_el['x'] - 146) < 10 and titulo_el['estilo']['tamanho'] > 40,
      (titulo_el['x'], titulo_el['estilo']['tamanho']))
    fundo_limpo = Midia.objects.get(pk=formato.id_midia(layout['fundo']['imagem']))
    with fundo_limpo.arquivo.open('rb') as arq:
        pixels = np.asarray(Image.open(io.BytesIO(arq.read())).convert('RGB'), dtype=np.float32)
    t('o texto de exemplo saiu do fundo', pixels[180:220, 200:650].mean() < 60, pixels[180:220, 200:650].mean())
    t('fonte conhecida normalizada e inventada vira Montserrat',
      importado.tema['fonte_titulo'] == 'Poppins' and importado.tema['fonte_texto'] == 'Montserrat')
    t('estilo da IA guardado', importado.estilo == 'Escuro com destaque branco.')
    r = c_bia.post(f'/apresentacoes/templates/{importado.pk}/documento/', data=json.dumps({
        'documento': importado.documento, 'revisao': importado.revisao}), content_type='application/json')
    t('outra pessoa não edita o template de Ana', r.status_code == 403)
    r = c_ana.post(f'/apresentacoes/templates/{importado.pk}/limpar-fundo/', data=json.dumps({
        'imagem': layout['fundo']['imagem'], 'areas': [{'x': 0, 'y': 0, 'w': 50, 'h': 50}]}), content_type='application/json')
    t('limpar fundo pelo editor devolve mídia nova', r.status_code == 200 and r.json()['midia']['id'] != fundo_limpo.pk)
    r = c_ana.post(f'/apresentacoes/templates/{padrao.pk}/excluir/')
    padrao.refresh_from_db()
    t('template da rede não é excluído', padrao.ativo)
    r = c_chefe.post(f'/apresentacoes/templates/{padrao.pk}/limpar-fundo/', data=json.dumps({
        'imagem': padrao.layouts[3]['fundo']['imagem'], 'areas': [{'x': 1600, 'y': 950, 'w': 200, 'h': 80}]}),
        content_type='application/json')
    t('SUPERADMIN limpa fundo do template da rede (estático)', r.status_code == 200, r.status_code)

    from pptx import Presentation
    from pptx.util import Inches, Pt
    arquivo_pptx = Presentation()
    arquivo_pptx.slide_width, arquivo_pptx.slide_height = Inches(13.333), Inches(7.5)
    for titulo_slide in ('Capa do PPT', 'Conteúdo do PPT', 'Fim'):
        s = arquivo_pptx.slides.add_slide(arquivo_pptx.slide_layouts[1])
        s.shapes.title.text = titulo_slide
        s.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
        s.placeholders[1].text = 'Texto de exemplo'
    saida_pptx = io.BytesIO()
    arquivo_pptx.save(saida_pptx)
    r = c_ana.post('/apresentacoes/templates/novo/', {'nome': 'ZZ PPT', 'arquivos': [
        SimpleUploadedFile('modelo.pptx', saida_pptx.getvalue(),
                           'application/vnd.openxmlformats-officedocument.presentationml.presentation')]})
    ppt = TemplateApresentacao.objects.get(nome='ZZ PPT')
    tarefa_ppt = TarefaIA.objects.get(template=ppt)
    chamadas_antes = len(cliente.chat.completions.chamadas)
    tarefas.executar(tarefa_ppt.pk)
    tarefa_ppt.refresh_from_db()
    ppt.refresh_from_db()
    papeis = [l['layout'] for l in ppt.layouts]
    t('PowerPoint vira layouts sem chamar a IA', tarefa_ppt.status == 'CONCLUIDA' and papeis ==
      ['capa', 'conteudo', 'encerramento'] and len(cliente.chat.completions.chamadas) == chamadas_antes,
      (papeis, tarefa_ppt.erro))
    titulo_ppt = next(e for e in ppt.layouts[0]['elementos'] if e['slot'] == 'titulo')
    t('título do PowerPoint com tamanho convertido (40 pt → 80 px no quadro)', titulo_ppt['estilo']['tamanho'] == 80
      and titulo_ppt['html'] == 'Capa do PPT', titulo_ppt['estilo'])
    corpo_ppt = [e for e in ppt.layouts[1]['elementos'] if e['slot'] == 'area_conteudo']
    t('o corpo do slide de conteúdo vira a área de conteúdo', len(corpo_ppt) == 1 and not any(
        e['tipo'] == 'texto' and 'Texto de exemplo' in e['html'] for e in ppt.layouts[1]['elementos']))

    print('\n== CANVA E VÍDEO ==')
    r = c_ana.post(f'/apresentacoes/{ap.pk}/exportar/canva/')
    t('Canva configurado mas pessoa não conectada: manda conectar', r.status_code == 200 and 'conectar' in r.json()
      and f'volta=/apresentacoes/{ap.pk}/' in r.json()['conectar'])
    cfg.set_canva_secret('')
    cfg.save()
    r = c_ana.post(f'/apresentacoes/{ap.pk}/exportar/canva/')
    t('Canva não configurado: explica e oferece o PowerPoint', r.status_code == 400 and r.json()['pptx'].endswith('/exportar/pptx/'))
    with mock.patch('apresentacoes.views._converter_para_mp4', return_value=b'\x00\x00\x00\x18ftypisom'):
        r = c_ana.post(f'/apresentacoes/{ap.pk}/video/', {'arquivo': SimpleUploadedFile('video.webm', b'webm', 'video/webm')})
    t('vídeo narrado sai em mp4 com link de download', r.status_code == 200 and r.json()['midia']['nome'].endswith('.mp4')
      and r.json()['midia']['download'].endswith('?download=1'))

    print('\n== LISTA, DUPLICAR, VERSÕES E EXCLUIR ==')
    html = c_chefe.get('/apresentacoes/').content.decode()
    t('SUPERADMIN vê as apresentações de todos, com filtro por pessoa', 'Todas as apresentações' in html
      and 'name="pessoa"' in html and 'Novo título' in html)
    html = c_bia.get('/apresentacoes/').content.decode()
    t('cada um vê só as suas', 'Novo título' not in html)
    r = c_ana.get(f'/apresentacoes/{ap.pk}/versoes/')
    versao = r.json()['versoes'][0]
    r = c_ana.post(versao['restaurar'])
    t('restaurar versão guarda a atual antes', r.status_code == 200 and VersaoApresentacao.objects.filter(
        apresentacao=ap, motivo='Antes de restaurar uma versão').exists())
    r = c_ana.post(f'/apresentacoes/{ap.pk}/duplicar/')
    copia = Apresentacao.objects.filter(dono=ana, titulo__endswith='(cópia)').first()
    t('duplicar cria cópia editável', copia and copia.slides and r['Location'].endswith(f'/apresentacoes/{copia.pk}/'))
    t('outra pessoa não exclui', c_bia.post(f'/apresentacoes/{ap.pk}/excluir/').status_code == 302
      and Apresentacao.objects.filter(pk=ap.pk).exists())
    nomes_arquivos = [m.arquivo.name for m in Midia.objects.filter(apresentacao=ap)]
    c_ana.post(f'/apresentacoes/{ap.pk}/excluir/')
    t('dona exclui e os arquivos saem do storage', not Apresentacao.objects.filter(pk=ap.pk).exists()
      and nomes_arquivos and not any(memoria.exists(n) for n in nomes_arquivos))

    print('\n== PÁGINAS ==')
    for url in ('/apresentacoes/nova/', '/apresentacoes/modulos/', '/apresentacoes/templates/',
                '/apresentacoes/templates/novo/'):
        t(f'{url} abre', c_ana.get(url).status_code == 200)
    html = c_ana.get('/apresentacoes/nova/').content.decode()
    t('nova: template da rede marcado e seletor de fontes', f'value="{padrao.pk}" data-fonte-titulo' in html
      and 'data-seletor-fonte' in html and 'fontes.json' in html)

finally:
    for p in patches:
        p.stop()
    transaction.set_rollback(True)
    marcador.__exit__(None, None, None)
    print('\nrollback: nada deste teste foi gravado no banco.')

print(f'\n{ok} OK / {fail} falhas')
sys.exit(1 if fail else 0)
