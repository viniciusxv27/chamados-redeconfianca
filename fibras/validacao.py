"""Acompanhamento completo de fibras: Base de Vendas (portal) x Esteira de Vendas Vivo.

Mesmo cruzamento da ferramenta "Validação de Vendas Go" (fibras/cruzamento.py),
com as três bases vindo do portal:

- Base de Vendas: as fibras que o /fibras já sincroniza do MySQL do Vivo GO
  (cada fibra é uma linha; FILIAL = PDV da venda).
- Esteira de Vendas Vivo: a planilha ANALÍTICO DE VENDAS importada no /fibras
  (só a importação mais recente fica guardada).
- Dados Lojas: os setores com ADABAS cadastrado (PDV = nome do setor).

O resultado é gravado em cada fibra (ORDEM vinculada, assertividade, motivo,
conferência de endereço...) e a ORDEM vinculada passa a ditar o status do
kanban. Os ajustes manuais (confirmar, rejeitar, vincular, comentário) ficam na
própria fibra e valem para todos os próximos cruzamentos.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime, timedelta

import numpy as np
import pandas as pd
from django.db import transaction
from django.utils import timezone

from . import cruzamento as cz
from .models import (Cruzamento, EsteiraImportacao, EsteiraVenda, Fibra, FibraStatusHistory, FilialAdabas,
                     PlanilhaOrdemInconsistente)

# Colunas extras da Esteira que o kanban já usava (SLA e motivo das vendas com problema).
EXTRAS_ESTEIRA = ['SLA AGENDA', 'MOTIVO']
OBRIGATORIAS_ESTEIRA = ['ORDEM', 'ADABAS', 'DATA_CRIAÇÃO', 'NOME VENDEDOR']
CORES = {cz.ALTA: 'C6EFCE', cz.MEDIA: 'FFEB9C', cz.MUITO_BAIXA: 'F8CBAD', cz.MANUAL: 'BDD7EE',
         cz.NAO_LOCALIZADA: 'FFC7CE', cz.NAO_LANCADA: 'D9D9D9', cz.B2B: 'E4DFEC',
         cz.END_CONFERE: 'C6EFCE', cz.END_PARCIAL: 'FFEB9C', cz.END_DIVERGE: 'FFC7CE'}


class ErroValidacao(Exception):
    """Erro com mensagem para mostrar na tela."""


# ---------------------------------------------------------------------------
# Esteira de Vendas Vivo (planilha ANALÍTICO)
# ---------------------------------------------------------------------------

def _texto(valor) -> str:
    if valor is None or (isinstance(valor, float) and np.isnan(valor)) or valor is pd.NaT:
        return ''
    if isinstance(valor, (pd.Timestamp, datetime)):
        return valor.strftime('%d/%m/%Y')
    if isinstance(valor, date):
        return valor.strftime('%d/%m/%Y')
    return cz.para_texto_id(valor)


def _ler_planilha(file_obj) -> pd.DataFrame:
    """Lê o .xlsx achando a linha de cabeçalho (a que tem ORDEM) nas 15 primeiras."""
    conteudo = file_obj.read() if hasattr(file_obj, 'read') else file_obj
    topo = pd.read_excel(io.BytesIO(conteudo), header=None, nrows=15)
    cabecalho = 0
    for i, linha in topo.iterrows():
        chaves = {cz.chave_coluna(v) for v in linha if isinstance(v, str)}
        if 'ORDEM' in chaves:
            cabecalho = i
            break
    return pd.read_excel(io.BytesIO(conteudo), header=cabecalho)


def importar_esteira(file_obj, *, nome='', by_user=None) -> EsteiraImportacao:
    """Guarda a Esteira (só as colunas do cruzamento) no lugar da importação anterior."""
    bruto = _ler_planilha(file_obj)
    tratada, mapa, ausentes = cz.tratar_base(bruto, cz.COLUNAS_ANALITICO)
    faltam = [c for c in OBRIGATORIAS_ESTEIRA if c in ausentes]
    if faltam:
        raise ErroValidacao('Esteira de Vendas Vivo: coluna(s) ' + ', '.join(faltam) + ' não encontrada(s). '
                            'Envie o ANALÍTICO DE VENDAS.')
    extras, _ = cz.resolver_colunas(bruto, EXTRAS_ESTEIRA)
    datas = pd.to_datetime(tratada[mapa['DATA_CRIAÇÃO']], errors='coerce', dayfirst=True)

    def col(nome_col, i):
        return _texto(bruto.at[i, mapa[nome_col]]) if nome_col in mapa else ''

    with transaction.atomic():
        imp = EsteiraImportacao.objects.create(
            nome_arquivo=(nome or getattr(file_obj, 'name', '') or 'ANALITICO DE VENDAS.xlsx')[:255],
            linhas=len(tratada), colunas_originais=bruto.shape[1],
            colunas=[[k, str(v)] for k, v in mapa.items()], ausentes=ausentes, importado_por=by_user)
        novas = []
        for pos, i in enumerate(tratada.index, start=1):
            data = datas.at[i]
            dados = {c: col(c, i) for c in cz.COLUNAS_ANALITICO if c in mapa}
            dados['DATA_CRIAÇÃO'] = data.strftime('%Y-%m-%d') if pd.notna(data) else ''
            novas.append(EsteiraVenda(
                importacao=imp, linha=pos,
                ordem=re.sub(r'\D', '', col('ORDEM', i))[:40],
                adabas=col('ADABAS', i)[:40],
                data_criacao=data.date() if pd.notna(data) else None,
                etapa_instalacao=col('ETAPA INSTALACAO', i)[:120],
                status_instalacao=col('STATUS INSTALACAO', i)[:120],
                produto=col('PRODUTO', i)[:200],
                nome_vendedor=col('NOME VENDEDOR', i)[:200],
                sla_agenda=(_texto(bruto.at[i, extras['SLA AGENDA']]) if 'SLA AGENDA' in extras else '')[:120],
                motivo=(_texto(bruto.at[i, extras['MOTIVO']]) if 'MOTIVO' in extras else '')[:255],
                dados=dados))
        EsteiraVenda.objects.bulk_create(novas, batch_size=1000)
        EsteiraVenda.objects.exclude(importacao=imp).delete()
    return imp


def ultima_importacao():
    return EsteiraImportacao.objects.filter(vendas__isnull=False).distinct().order_by('-importado_em').first()


def base_esteira(imp=None) -> pd.DataFrame:
    """A Esteira guardada, no formato que o motor espera (índice = id da EsteiraVenda)."""
    imp = imp or ultima_importacao()
    if not imp:
        return pd.DataFrame(columns=cz.COLUNAS_ANALITICO)
    linhas = list(EsteiraVenda.objects.filter(importacao=imp).order_by('linha').values('id', 'dados', 'data_criacao'))
    df = pd.DataFrame([dict(l['dados'], **{'DATA_CRIAÇÃO': pd.Timestamp(l['data_criacao']) if l['data_criacao'] else pd.NaT})
                       for l in linhas], index=[l['id'] for l in linhas])
    for c in cz.COLUNAS_ANALITICO:
        if c not in df:
            df[c] = ''
    return df[cz.COLUNAS_ANALITICO]


# ---------------------------------------------------------------------------
# Base de Vendas (fibras do portal) e lojas
# ---------------------------------------------------------------------------

def base_vendas(ini=None, fim=None, janela=1) -> pd.DataFrame:
    """As fibras no formato do BASE Go (índice = id da fibra)."""
    qs = Fibra.objects.all()
    if ini and fim:
        qs = qs.filter(data_da_venda__gte=ini - timedelta(days=janela), data_da_venda__lte=fim + timedelta(days=janela))
    linhas = list(qs.order_by('data_da_venda', 'id').values(
        'id', 'pdv', 'servico_tecnico', 'numero_da_venda', 'data_da_venda', 'vendedor', 'endereco', 'plano',
        'numero_protocolo', 'valor', 'pilar', 'cpf', 'cliente'))
    df = pd.DataFrame({
        'FILIAL': [l['pdv'] for l in linhas],
        'SERVIÇO': [l['servico_tecnico'] for l in linhas],
        'Nº DA VENDA': [l['numero_da_venda'] for l in linhas],
        'DATA': [pd.Timestamp(l['data_da_venda']) if l['data_da_venda'] else pd.NaT for l in linhas],
        'VENDEDOR': [l['vendedor'] for l in linhas],
        'ENDEREÇO': [l['endereco'] for l in linhas],
        'PLANO': [l['plano'] for l in linhas],
        'NUMERO PROTOCOLO': [l['numero_protocolo'] for l in linhas],
        'NOVA RECEITA': [float(l['valor'] or 0) for l in linhas],
        'PILAR': [l['pilar'] for l in linhas],
        'CPF/CNPJ': [l['cpf'] for l in linhas],
        'CLIENTE': [l['cliente'] for l in linhas],
    }, index=[l['id'] for l in linhas])
    return df


MAPA_GO = {c: c for c in ['FILIAL', 'SERVIÇO', 'Nº DA VENDA', 'DATA', 'VENDEDOR', 'ENDEREÇO', 'PLANO',
                          'NUMERO PROTOCOLO', 'NOVA RECEITA', 'PILAR', 'CPF/CNPJ']}
MAPA_AN = {c: c for c in cz.COLUNAS_ANALITICO}


def dados_lojas() -> pd.DataFrame:
    """Setores com ADABAS: o "Dados Lojas" da ferramenta (PDV = nome do setor)."""
    from users.models import Sector
    linhas = list(Sector.objects.exclude(adabas='').order_by('name').values_list('adabas', 'name'))
    return pd.DataFrame(linhas, columns=['ADABAS', 'PDV']).drop_duplicates('ADABAS')


def nomes_lojas() -> dict:
    lj = dados_lojas()
    return dict(zip(lj['ADABAS'].astype(str), lj['PDV'].astype(str)))


def mapeamento_lojas(filiais=None) -> pd.DataFrame:
    """FILIAL -> ADABAS: o que foi salvo na tela vale; o resto é sugerido pelo nome do setor."""
    if filiais is None:
        filiais = Fibra.objects.exclude(pdv='').values_list('pdv', flat=True).distinct()
    salvo = dict(FilialAdabas.objects.values_list('filial', 'adabas'))
    return cz.sugerir_mapeamento(list(filiais), dados_lojas(), 'ADABAS', 'PDV', salvo)


def mapa_lojas(filiais=None) -> dict:
    m = mapeamento_lojas(filiais)
    return {r['FILIAL (Base de Vendas)']: r['ADABAS'] or '' for _, r in m.iterrows()}


def salvar_lojas(mapa: dict, by_user=None):
    for filial, adabas in mapa.items():
        filial = str(filial).strip()[:160]
        if filial:
            FilialAdabas.objects.update_or_create(
                filial=filial, defaults={'adabas': str(adabas or '').strip()[:40], 'atualizado_por': by_user})


# ---------------------------------------------------------------------------
# Escopo e cruzamento
# ---------------------------------------------------------------------------

def escopo_padrao(an: pd.DataFrame | None = None) -> dict:
    """Período da Esteira, lojas que têm FILIAL na Base de Vendas, todas as etapas, ±1 dia."""
    an = base_esteira() if an is None else an
    datas = pd.to_datetime(an['DATA_CRIAÇÃO'], errors='coerce').dropna()
    hoje = timezone.localdate()
    ini = datas.min().date() if len(datas) else hoje.replace(day=1)
    fim = datas.max().date() if len(datas) else hoje
    adabas_an = sorted(an['ADABAS'].dropna().astype(str).unique()) if len(an) else []
    mapeadas = set(mapa_lojas().values())
    etapas = sorted(x for x in an['ETAPA INSTALACAO'].dropna().astype(str).unique() if x) if len(an) else []
    return {'ini': ini.isoformat(), 'fim': fim.isoformat(), 'janela': 1,
            'lojas': [a for a in adabas_an if a in mapeadas], 'etapas': etapas}


def _param_limpo(param: dict | None) -> dict:
    p = dict(escopo_padrao(), **(param or {}))
    try:
        p['janela'] = max(0, min(1, int(p.get('janela') or 0)))
    except (TypeError, ValueError):
        p['janela'] = 1
    for k in ('ini', 'fim'):
        if not p.get(k):
            raise ErroValidacao('Informe o período.')
        p[k] = str(p[k])[:10]
    if p['ini'] > p['fim']:
        raise ErroValidacao('A data inicial é depois da final.')
    p['lojas'] = [str(x) for x in (p.get('lojas') or [])]
    p['etapas'] = [str(x) for x in (p.get('etapas') or [])]
    return p


def calcular(param: dict | None = None) -> dict:
    """Roda o motor sem gravar nada. Devolve os DataFrames e o contexto."""
    imp = ultima_importacao()
    if not imp:
        raise ErroValidacao('Importe a Esteira de Vendas Vivo (ANALÍTICO DE VENDAS) primeiro.')
    p = _param_limpo(param)
    ini, fim = date.fromisoformat(p['ini']), date.fromisoformat(p['fim'])
    an = base_esteira(imp)
    go = base_vendas(ini, fim, p['janela'])
    mapa = mapa_lojas(go['FILIAL'].unique()) if len(go) else {}

    datas_an = pd.to_datetime(an['DATA_CRIAÇÃO'], errors='coerce').dt.normalize()
    escopo = (datas_an >= pd.Timestamp(ini)) & (datas_an <= pd.Timestamp(fim))
    escopo &= an['ADABAS'].astype(str).isin(p['lojas'])
    if p['etapas']:
        escopo &= an['ETAPA INSTALACAO'].astype(str).isin(p['etapas'])

    # Ajustes manuais guardados na fibra -> chaves de pedido do motor
    linhas = cz.montar_pedidos(go, MAPA_GO, mapa) if len(go) else pd.DataFrame(columns=['linha_go', 'pid'])
    pid_da_fibra = dict(zip(linhas['linha_go'], linhas['pid']))
    manuais, rejeitados = {}, set()
    for fid, manual, rejeit in Fibra.objects.filter(id__in=list(go.index)).values_list(
            'id', 'ordem_manual', 'ordens_rejeitadas'):
        pid = pid_da_fibra.get(fid)
        if not pid:
            continue
        if manual:
            manuais[pid] = manual
        for o in rejeit or []:
            rejeitados.add((pid, str(o)))

    if len(go):
        rg, nl, rs, ped = cz.cruzar(go, MAPA_GO, an, MAPA_AN, mapa, escopo, p['janela'], manuais, rejeitados,
                                    nomes_lojas())
    else:
        rg, nl, rs, ped = (pd.DataFrame(), an.loc[escopo].copy(), pd.DataFrame(), pd.DataFrame())
        nl['STATUS CRUZAMENTO'] = cz.NAO_LANCADA
        nl['OBSERVAÇÃO'] = cz.NAO_LANCADA
    nomes = nomes_lojas()
    if len(rs):
        rs.insert(1, 'PDV', rs['ADABAS'].map(lambda a: nomes.get(a, '')))
    return dict(param=p, imp=imp, go=go, an=an, mapa=mapa, escopo=escopo, linhas=linhas, rg=rg, nl=nl, resumo=rs,
                ped=ped)


def _kpis(rg: pd.DataFrame, nl: pd.DataFrame) -> dict:
    if not len(rg):
        return dict(linhas=0, pedidos=0, alta=0, media=0, muito_baixa=0, manual=0, nao_localizada=0, b2b=0,
                    nao_lancada=len(nl), exato=0, end_diverge=0, end_confere=0)
    ped = rg.drop_duplicates('ID PEDIDO')
    cont = ped['ASSERTIVIDADE'].value_counts()
    st = ped['STATUS CRUZAMENTO'].value_counts()
    return dict(linhas=len(rg), pedidos=len(ped), alta=int(cont.get(cz.ALTA, 0)), media=int(cont.get(cz.MEDIA, 0)),
                muito_baixa=int(cont.get(cz.MUITO_BAIXA, 0)), manual=int(cont.get(cz.MANUAL, 0)),
                nao_localizada=int(st.get(cz.NAO_LOCALIZADA, 0)), b2b=int(st.get(cz.B2B, 0)), nao_lancada=len(nl),
                exato=int((ped['MÉTODO'] == '1. Protocolo exato').sum()),
                end_diverge=int(ped['CONFERÊNCIA ENDEREÇO'].isin([cz.END_DIVERGE, cz.END_PARCIAL]).sum()),
                end_confere=int((ped['CONFERÊNCIA ENDEREÇO'] == cz.END_CONFERE).sum()))


def valor_json(v):
    if v is None or v is pd.NaT or (isinstance(v, float) and np.isnan(v)):
        return ''
    if isinstance(v, (pd.Timestamp, datetime, date)):
        return '' if pd.isna(v) else v.strftime('%d/%m/%Y')
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.floating):
        return float(v)
    if isinstance(v, np.bool_):
        return bool(v)
    return v


def tabela(df: pd.DataFrame, limite: int | None = None) -> dict:
    if limite:
        df = df.head(limite)
    return {'colunas': [str(c) for c in df.columns],
            'linhas': [[valor_json(v) for v in linha] for linha in df.itertuples(index=False, name=None)]}


def executar(param: dict | None = None, *, by_user=None, origem='manual', notificar=True) -> Cruzamento:
    """Roda o cruzamento e grava o resultado nas fibras e na Esteira."""
    from .services import _map_planilha_status, _notify_vendor

    c = calcular(param)
    rg, nl, linhas, an = c['rg'], c['nl'], c['linhas'], c['an']
    p = c['param']
    agora = timezone.now()
    with transaction.atomic():
        cruz = Cruzamento.objects.create(
            executado_por=by_user, origem=origem, importacao=c['imp'], ini=p['ini'], fim=p['fim'], janela=p['janela'],
            lojas=p['lojas'], etapas=p['etapas'], kpis=_kpis(rg, nl), resumo=tabela(c['resumo']) if len(c['resumo']) else {})

        esteira = {e.ordem: e for e in EsteiraVenda.objects.filter(importacao=c['imp']) if e.ordem}
        fibras = {f.id: f for f in Fibra.objects.filter(id__in=list(c['go'].index))}
        campos = ['cruzamento', 'adabas', 'protocolo_tratado', 'formato_protocolo', 'status_cruzamento',
                  'assertividade', 'ordem_vinculada', 'metodo_cruzamento', 'motivo_cruzamento',
                  'observacao_cruzamento', 'conferencia_endereco', 'detalhe_endereco', 'etapa_instalacao',
                  'last_planilha_at', 'ordem_planilha', 'status_planilha_raw', 'sla_agenda', 'motivo_planilha',
                  'status']
        mudancas = []
        vinculos = {}
        for i, r in rg.iterrows():
            f = fibras.get(linhas.at[i, 'linha_go'])
            if not f:
                continue
            ordem = str(r['ORDEM VINCULADA'] or '')
            f.cruzamento = cruz
            f.adabas = str(r['LOJA ADABAS'] or '')[:40]
            f.protocolo_tratado = str(r['PROTOCOLO TRATADO'] or '')[:60]
            f.formato_protocolo = str(r['FORMATO PROTOCOLO'] or '')[:60]
            f.status_cruzamento = str(r['STATUS CRUZAMENTO'] or '')[:30]
            f.assertividade = '' if r['ASSERTIVIDADE'] == '-' else str(r['ASSERTIVIDADE'])[:20]
            f.ordem_vinculada = ordem[:40]
            f.metodo_cruzamento = '' if r['MÉTODO'] == '-' else str(r['MÉTODO'])[:80]
            f.motivo_cruzamento = str(r['MOTIVO'] or '')
            f.observacao_cruzamento = str(r['OBSERVAÇÃO'] or '')
            f.conferencia_endereco = str(r['CONFERÊNCIA ENDEREÇO'] or '')[:40]
            f.detalhe_endereco = str(r['DETALHE ENDEREÇO'] or '')
            e = esteira.get(ordem) if ordem else None
            f.etapa_instalacao = (e.etapa_instalacao if e else '')[:120]
            if e:
                vinculos[e.id] = f.id
                # A ORDEM vinculada dita o status do kanban, como a planilha já fazia.
                f.last_planilha_at = agora
                f.ordem_planilha = e.linha
                f.status_planilha_raw = e.status_instalacao[:120]
                f.sla_agenda = e.sla_agenda[:120]
                f.motivo_planilha = e.motivo[:255]
                novo = _map_planilha_status(e.status_instalacao)
                if novo and novo != f.status:
                    mudancas.append((f, f.status, novo, e.status_instalacao))
                    f.status = novo
        Fibra.objects.bulk_update(list(fibras.values()), campos, batch_size=500)
        for f, antigo, novo, bruto in mudancas:
            FibraStatusHistory.objects.create(fibra=f, status_anterior=antigo, status_novo=novo,
                                              retorno=f'Esteira Vivo: {bruto[:80]}', alterado_por=by_user)

        # Esteira: escopo e vínculo
        EsteiraVenda.objects.filter(importacao=c['imp']).update(no_escopo=False, fibra=None)
        dentro = [int(i) for i in an.index[c['escopo'].to_numpy()]]
        EsteiraVenda.objects.filter(id__in=dentro).update(no_escopo=True)
        for eid, fid in vinculos.items():
            EsteiraVenda.objects.filter(id=eid).update(fibra_id=fid)

        # "Ordens inconsistentes" do kanban = vendas não lançadas do escopo
        nao_lancadas = {o: st for o, st in EsteiraVenda.objects.filter(id__in=[int(i) for i in nl.index])
                        .values_list('ordem', 'status_instalacao') if o}
        PlanilhaOrdemInconsistente.objects.exclude(ordem__in=list(nao_lancadas)).delete()
        for ordem, st in nao_lancadas.items():
            obj, criado = PlanilhaOrdemInconsistente.objects.get_or_create(ordem=ordem[:60],
                                                                           defaults={'status_raw': st[:120]})
            if not criado and obj.status_raw != st[:120]:
                obj.status_raw = st[:120]
                obj.save(update_fields=['status_raw', 'last_seen_at'])

    if notificar:
        for f, antigo, novo, _ in mudancas:
            _notify_vendor(f, antigo, novo)
    return cruz


def ultimo_cruzamento():
    return Cruzamento.objects.order_by('-executado_em').first()


def parametros_atuais() -> dict:
    """Os parâmetros do último cruzamento (o escopo padrão se não houver)."""
    ult = ultimo_cruzamento()
    if ult:
        return {'ini': ult.ini.isoformat(), 'fim': ult.fim.isoformat(), 'janela': ult.janela,
                'lojas': ult.lojas, 'etapas': ult.etapas}
    return escopo_padrao()


# ---------------------------------------------------------------------------
# Ajustes manuais
# ---------------------------------------------------------------------------

def aplicar_ajuste(fibra: Fibra, acao: str, ordem: str = '', comentario=None):
    ordem = re.sub(r'\D', '', str(ordem or ''))
    rejeit = [str(o) for o in (fibra.ordens_rejeitadas or [])]
    if acao == 'confirmar' and ordem:
        fibra.ordem_manual = ordem
    elif acao == 'vincular':
        if not ordem:
            raise ErroValidacao('Informe a ORDEM.')
        fibra.ordem_manual = ordem
        rejeit = [o for o in rejeit if o != ordem]
    elif acao == 'rejeitar' and ordem:
        fibra.ordem_manual = ''
        if ordem not in rejeit:
            rejeit.append(ordem)
    elif acao == 'desfazer':
        fibra.ordem_manual = ''
        rejeit = []
    fibra.ordens_rejeitadas = rejeit
    if comentario is not None:
        fibra.comentario_cruzamento = str(comentario).strip()
    fibra.save(update_fields=['ordem_manual', 'ordens_rejeitadas', 'comentario_cruzamento'])


def limpar_ajustes():
    Fibra.objects.exclude(ordem_manual='', ordens_rejeitadas=[], comentario_cruzamento='').update(
        ordem_manual='', ordens_rejeitadas=[], comentario_cruzamento='')


def contagem_ajustes() -> dict:
    return {'manuais': Fibra.objects.exclude(ordem_manual='').count(),
            'rejeitados': sum(len(x or []) for x in Fibra.objects.exclude(ordens_rejeitadas=[])
                              .values_list('ordens_rejeitadas', flat=True)),
            'comentarios': Fibra.objects.exclude(comentario_cruzamento='').count()}


# ---------------------------------------------------------------------------
# JSON para a tela
# ---------------------------------------------------------------------------

def resultado_json(cruz: Cruzamento | None = None) -> dict | None:
    cruz = cruz or ultimo_cruzamento()
    if not cruz:
        return None
    fibras = Fibra.objects.filter(cruzamento=cruz).order_by('data_da_venda', 'id')
    go = [{
        'pid': f.id, 'filial': f.pdv, 'data': f.data_da_venda.strftime('%d/%m/%Y') if f.data_da_venda else '',
        'servico': f.servico_tecnico, 'venda': f.numero_da_venda, 'vendedor': f.vendedor, 'plano': f.plano,
        'cliente': f.cliente, 'protocolo': f.numero_protocolo, 'tratado': f.protocolo_tratado,
        'formato': f.formato_protocolo, 'status': f.status_cruzamento, 'assert': f.assertividade or '-',
        'ordem': f.ordem_vinculada, 'etapa': f.etapa_instalacao if f.ordem_vinculada else '',
        'status_inst': f.status_planilha_raw if f.ordem_vinculada else '',
        'metodo': f.metodo_cruzamento, 'motivo': f.motivo_cruzamento, 'obs': f.observacao_cruzamento,
        'comentario': f.comentario_cruzamento, 'manual': bool(f.ordem_manual),
        'endereco': f.conferencia_endereco, 'end_det': f.detalhe_endereco, 'kanban': f.get_status_display(),
    } for f in fibras]
    nl = [{
        'data': e.data_criacao.strftime('%d/%m/%Y') if e.data_criacao else '', 'adabas': e.adabas, 'ordem': e.ordem,
        'produto': e.produto, 'vendedor': e.nome_vendedor, 'etapa': e.etapa_instalacao,
        'status_inst': e.status_instalacao,
    } for e in EsteiraVenda.objects.filter(importacao=cruz.importacao, no_escopo=True, fibra__isnull=True)
        .order_by('data_criacao', 'linha')] if cruz.importacao_id else []
    quem = cruz.executado_por.get_full_name() if cruz.executado_por else ''
    return {'quando': timezone.localtime(cruz.executado_em).strftime('%d/%m/%Y %H:%M:%S'), 'por': quem,
            'origem': cruz.get_origem_display(), 'kpis': cruz.kpis, 'go': go, 'nl': nl,
            'resumo': cruz.resumo or {'colunas': [], 'linhas': []}, 'ajustes': contagem_ajustes()}


def estado_json() -> dict:
    imp = ultima_importacao()
    an = base_esteira(imp)
    n_fibras = Fibra.objects.count()
    lj = dados_lojas()
    arquivos = {
        'go': {'rotulo': 'BASE DE VENDAS', 'ok': n_fibras > 0, 'linhas': n_fibras,
               'nome': 'Vendas de Fixa do Vivo GO (portal)', 'origem': 'portal',
               'detalhe': 'Sincronizadas do MySQL do Vivo GO'},
        'an': {'rotulo': 'ESTEIRA DE VENDAS VIVO', 'ok': bool(imp), 'linhas': len(an),
               'nome': imp.nome_arquivo if imp else '', 'origem': 'importada',
               'detalhe': (f'Importada em {timezone.localtime(imp.importado_em):%d/%m/%Y %H:%M}'
                           + (f' por {imp.importado_por.get_full_name()}' if imp.importado_por else '')) if imp else '',
               'colunas': imp.colunas if imp else [], 'ausentes': imp.ausentes if imp else [],
               'colunas_orig': imp.colunas_originais if imp else 0},
        'lj': {'rotulo': 'DADOS LOJAS', 'ok': len(lj) > 0, 'linhas': len(lj), 'nome': 'Setores com ADABAS',
               'origem': 'portal', 'detalhe': 'Cadastro de setores do portal'},
    }
    r = {'arquivos': arquivos, 'ajustes': contagem_ajustes()}
    nomes = nomes_lojas()
    m = mapeamento_lojas()
    r['lojas'] = [{'filial': x['FILIAL (Base de Vendas)'], 'adabas': x['ADABAS'] or '', 'pdv': x['PDV'],
                   'origem': x['Origem']} for _, x in m.iterrows()]
    adabas_an = sorted(an['ADABAS'].dropna().astype(str).unique()) if len(an) else []
    opcoes = dict(nomes)
    for a in adabas_an:
        opcoes.setdefault(a, '')
    r['opcoes_lojas'] = [{'adabas': a, 'pdv': p} for a, p in sorted(opcoes.items()) if a]
    mapeadas = {x['adabas'] for x in r['lojas'] if x['adabas']}
    r['fora_go'] = [{'adabas': a, 'pdv': nomes.get(a, '')} for a in adabas_an if a not in mapeadas]
    if imp:
        esc = parametros_atuais()
        esc['lojas_todas'] = [{'adabas': a, 'pdv': nomes.get(a, '')} for a in adabas_an]
        esc['etapas_todas'] = sorted(x for x in an['ETAPA INSTALACAO'].dropna().astype(str).unique() if x)
        r['escopo'] = esc
    else:
        r['erro'] = 'Importe a Esteira de Vendas Vivo (ANALÍTICO DE VENDAS) para cruzar.'
    res = resultado_json()
    if res:
        r['resultado'] = res
    return r


def detalhe_json(fibra: Fibra) -> dict:
    an = base_esteira()
    nomes = nomes_lojas()
    pedido = pd.Series({'adabas': fibra.adabas, 'prot': fibra.protocolo_tratado or cz.limpar_protocolo(fibra.numero_protocolo),
                        'data': pd.Timestamp(fibra.data_da_venda) if fibra.data_da_venda else pd.NaT,
                        'vendedor': cz.normalizar_pessoa(fibra.vendedor)})
    usadas = set(Fibra.objects.filter(cruzamento=fibra.cruzamento_id).exclude(ordem_vinculada='')
                 .values_list('ordem_vinculada', flat=True)) if fibra.cruzamento_id else set()
    janela = fibra.cruzamento.janela if fibra.cruzamento_id else 1
    sug = cz.candidatos(pedido, an, MAPA_AN, usadas, janela) if len(an) else pd.DataFrame()
    if len(sug):
        sug.insert(2, 'PDV', sug['ADABAS'].map(lambda a: nomes.get(a, '')))
    analitico = {}
    if fibra.ordem_vinculada:
        e = EsteiraVenda.objects.filter(ordem=fibra.ordem_vinculada).order_by('-importacao__importado_em').first()
        if e:
            analitico = dict(e.dados)
            analitico['DATA_CRIAÇÃO'] = e.data_criacao.strftime('%d/%m/%Y') if e.data_criacao else ''
            analitico['SLA AGENDA'] = e.sla_agenda
            analitico['MOTIVO'] = e.motivo
    combo = Fibra.objects.filter(cruzamento=fibra.cruzamento_id, adabas=fibra.adabas, data_da_venda=fibra.data_da_venda,
                                 protocolo_tratado=fibra.protocolo_tratado).exclude(protocolo_tratado='') \
        if fibra.cruzamento_id else Fibra.objects.none()
    combo = list(combo) or [fibra]
    return {
        'pid': fibra.id, 'ordem': fibra.ordem_vinculada, 'assertividade': fibra.assertividade or '-',
        'status': fibra.status_cruzamento, 'motivo': fibra.motivo_cruzamento, 'metodo': fibra.metodo_cruzamento or '-',
        'obs': fibra.observacao_cruzamento, 'endereco': fibra.conferencia_endereco, 'end_det': fibra.detalhe_endereco,
        'comentario': fibra.comentario_cruzamento, 'manual': bool(fibra.ordem_manual),
        'rejeitadas': list(fibra.ordens_rejeitadas or []),
        'venda': {'filial': fibra.pdv, 'data': fibra.data_da_venda.strftime('%d/%m/%Y') if fibra.data_da_venda else '',
                  'vendedor': fibra.vendedor, 'protocolo': fibra.numero_protocolo, 'formato': fibra.formato_protocolo,
                  'adabas': fibra.adabas, 'loja': nomes.get(fibra.adabas, ''), 'endereco': fibra.endereco,
                  'cliente': fibra.cliente, 'kanban': fibra.get_status_display(), 'url': f'/fibras/{fibra.id}/'},
        'linhas': [{'servico': f.servico_tecnico, 'venda': f.numero_da_venda, 'plano': f.plano, 'cpf': f.cpf}
                   for f in combo],
        'analitico': analitico, 'sugestoes': tabela(sug),
    }


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def excel_bytes(abas: dict) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as w:
        for nome, df in abas.items():
            aba = nome[:31]
            df.to_excel(w, sheet_name=aba, index=False)
            ws = w.sheets[aba]
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for cel in ws[1]:
                cel.font = Font(bold=True, color='FFFFFF')
                cel.fill = PatternFill('solid', fgColor='FF6B35')
                cel.alignment = Alignment(wrap_text=True, vertical='center')
            for i, col in enumerate(df.columns, 1):
                letra = get_column_letter(i)
                serie = df[col]
                if pd.api.types.is_datetime64_any_dtype(serie):
                    for c in ws[letra][1:]:
                        c.number_format = 'dd/mm/yyyy'
                largura = max([len(str(col))] + [len(str(v)) for v in serie.head(300)])
                ws.column_dimensions[letra].width = min(max(10, largura + 2), 45)
            for alvo in ('ASSERTIVIDADE', 'STATUS CRUZAMENTO', 'CONFERÊNCIA ENDEREÇO'):
                if alvo in df.columns:
                    j = list(df.columns).index(alvo) + 1
                    for c in ws[get_column_letter(j)][1:]:
                        cor = CORES.get(c.value)
                        if cor:
                            c.fill = PatternFill('solid', fgColor=cor)
    return buf.getvalue()


def arquivo(o: str) -> tuple[str, bytes]:
    """Excel para baixar: as bases tratadas (go/an/lj) ou o resultado do último cruzamento."""
    if o == 'go':
        p = parametros_atuais()
        go = base_vendas(date.fromisoformat(p['ini']), date.fromisoformat(p['fim']), int(p.get('janela') or 0))
        return 'BASE DE VENDAS - TRATADO.xlsx', excel_bytes({'BASE DE VENDAS': go})
    if o == 'an':
        return 'ESTEIRA DE VENDAS VIVO - TRATADO.xlsx', excel_bytes({'ESTEIRA DE VENDAS VIVO': base_esteira()})
    if o == 'lj':
        return 'DADOS LOJAS - TRATADO.xlsx', excel_bytes({'DADOS LOJAS': dados_lojas()})
    c = calcular(parametros_atuais())
    rg, nl, p = c['rg'], c['nl'], c['param']
    if len(rg):
        rg = rg.copy()
        comentarios = dict(Fibra.objects.exclude(comentario_cruzamento='')
                           .values_list('numero_da_venda', 'comentario_cruzamento'))
        rg['COMENTÁRIO MANUAL'] = rg['Nº DA VENDA'].map(comentarios).fillna('')
        rg = rg.drop(columns=['ID PEDIDO'])
    consolidado = pd.concat([rg, nl.add_prefix('ESTEIRA VIVO: ').rename(columns={
        'ESTEIRA VIVO: STATUS CRUZAMENTO': 'STATUS CRUZAMENTO', 'ESTEIRA VIVO: OBSERVAÇÃO': 'OBSERVAÇÃO'})],
        ignore_index=True)
    aj = contagem_ajustes()
    parametros = pd.DataFrame({
        'Parâmetro': ['Gerado em', 'Período Esteira Vivo', 'Lojas', 'Etapas', 'Tolerância de data (± dias)',
                      'Ajustes manuais', 'Rejeitados'],
        'Valor': [f'{timezone.localtime():%d/%m/%Y %H:%M}',
                  f'{date.fromisoformat(p["ini"]):%d/%m/%Y} a {date.fromisoformat(p["fim"]):%d/%m/%Y}',
                  ', '.join(p['lojas']), ', '.join(p['etapas']), p['janela'], aj['manuais'], aj['rejeitados']]})
    abas = {'Resultado Base de Vendas': rg, 'Vendas não lançadas': nl, 'Consolidado': consolidado,
            'Resumo loja-dia': c['resumo'], 'Parâmetros': parametros}
    return f'VALIDACAO BASE DE VENDAS x ESTEIRA VIVO {timezone.localtime():%Y-%m-%d %H%M}.xlsx', excel_bytes(abas)


def previa(tipo: str) -> dict:
    if tipo == 'go':
        p = parametros_atuais()
        return tabela(base_vendas(date.fromisoformat(p['ini']), date.fromisoformat(p['fim']),
                                  int(p.get('janela') or 0)), 200)
    if tipo == 'an':
        return tabela(base_esteira(), 200)
    if tipo == 'lj':
        return tabela(dados_lojas(), 200)
    return {'colunas': [], 'linhas': []}
