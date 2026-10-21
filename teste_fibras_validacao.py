"""/fibras: acompanhamento completo (Base de Vendas x Esteira de Vendas Vivo).

Mesmo cruzamento da ferramenta "Validação de Vendas Go", com a Base de Vendas
vindo das fibras do portal e a Esteira importada no /fibras. Tudo roda numa
transação desfeita no fim; as datas são de 2031 para não cruzar com vendas reais
e o MySQL do Vivo GO não é chamado.
"""
import io
import json
import os
import sys
from datetime import date
from unittest import mock

import django

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'redeconfianca.settings')
django.setup()

from django.conf import settings

if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS.append('testserver')

import pandas as pd
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection, transaction
from django.db.migrations.executor import MigrationExecutor
from django.test import Client

User = get_user_model()
ok = fail = 0


def t(nome, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'  OK   {nome}')
    else:
        fail += 1
        print(f'  FALHA {nome} {extra}')


def planilha(linhas):
    """ANALÍTICO DE VENDAS com uma linha de título antes do cabeçalho."""
    cols = ['DATA_CRIAÇÃO', 'ETAPA INSTALACAO', 'STATUS INSTALACAO', 'ORDEM', 'PRODUTO', 'ADABAS', 'LOGRADOURO',
            'CEP', 'NUMERO', 'LOGIN', 'NOME VENDEDOR', 'VENDA COMISSIONADA', 'SLA AGENDA', 'MOTIVO', 'OUTRA COLUNA']
    df = pd.DataFrame(linhas, columns=cols)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as w:
        pd.DataFrame([['RELATÓRIO ANALÍTICO DE VENDAS']]).to_excel(w, index=False, header=False, startrow=0)
        df.to_excel(w, index=False, startrow=1)
    return buf.getvalue()


marcador = transaction.atomic()
marcador.__enter__()
try:
    executor = MigrationExecutor(connection)
    ultima = executor.loader.graph.leaf_nodes('fibras')
    if executor.migration_plan(ultima):
        executor.migrate(ultima)

    from fibras import services, validacao as val
    from fibras.models import EsteiraVenda, Fibra, FilialAdabas, PlanilhaOrdemInconsistente
    from users.models import Sector

    Sector.objects.create(name='ZZ LOJA CENTRO', adabas='ZZA001')
    Sector.objects.create(name='ZZ LOJA PRAIA', adabas='ZZB002')
    ilha = User.objects.create_user(username='zzfib.ilha', email='zzfib.ilha@exemplo-teste.local', password='S3nha!teste',
                                    first_name='ZZ', last_name='Ilha', hierarchy='SUPERVISOR')
    comum = User.objects.create_user(username='zzfib.cons', email='zzfib.cons@exemplo-teste.local', password='S3nha!teste',
                                     first_name='ZZ', last_name='Consultor', hierarchy='PADRAO')
    D = date(2031, 3, 10)

    def fibra(venda, prot, pdv='ZZ LOJA CENTRO', dia=D, cpf='12345678909', endereco='', vendedor='ZZ VENDEDOR UM'):
        return Fibra.objects.create(numero_da_venda=venda, numero_protocolo=prot, pdv=pdv, data_da_venda=dia,
                                    cpf=cpf, cliente=f'Cliente {venda}', endereco=endereco, vendedor=vendedor,
                                    pilar='FIXA', servico_tecnico='Alta Banda Larga', valor=100,
                                    status=Fibra.STATUS_PENDENTE)

    f1 = fibra('ZZV1', "'7123456789A", endereco='Rua das Flores, 100, Bairro Centro , CEP 29160-054')
    f2 = fibra('ZZV2', '7123456700')
    f3 = fibra('ZZV3', '7999999999')
    f4 = fibra('ZZV4', 'ABC123X', cpf='11.222.333/0001-81')
    f5 = fibra('ZZV5', '7555555555')
    f6 = fibra('ZZV6', '7444444444', pdv='ZZ OUTRA LOJA')

    def linha(dia, ordem, adabas, status='3.1 PENDENTE ENRIQUECIMENTO', etapa='PENDENTE', **kw):
        return [dia, etapa, status, ordem, 'BANDA LARGA', adabas, kw.get('log', ''), kw.get('cep', ''),
                kw.get('num', ''), 'zz', 'ZZ VENDEDOR UM', 'SIM', kw.get('sla', ''), kw.get('motivo', ''), 'x']

    arquivo = planilha([
        linha('10/03/2031', 7123456789, 'ZZA001', '1.1 INSTALADA', 'INSTALADA', log='R DAS FLORES', cep='29160054', num=100),
        linha('10/03/2031', 7123456709, 'ZZA001', '3.5 AGENDADA', 'AGENDADA', sla='D+2'),
        linha('10/03/2031', 7999999999, 'ZZB002'),
        linha('12/03/2031', 7888888888, 'ZZA001'),
        linha('10/03/2031', 7444444444, 'ZZB002', '3.4 PENDENTE TECNICA', motivo='SEM VIABILIDADE'),
    ])

    ci, cc = Client(), Client()
    ci.force_login(ilha)
    cc.force_login(comum)

    print('== ACESSO ==')
    t('a Ilha abre a validação', ci.get('/fibras/validacao/').status_code == 200)
    r = cc.get('/fibras/validacao/')
    t('consultor volta para o kanban', r.status_code == 302 and r['Location'].endswith('/fibras/'))
    t('e não usa a API', cc.get('/fibras/validacao/api/estado/').status_code == 403)
    t('o kanban da Ilha tem o botão da validação', 'Validação Base × Esteira' in ci.get('/fibras/').content.decode())

    print('\n== IMPORTAR A ESTEIRA (cruza na hora) ==')
    with mock.patch.object(services, 'sync_fibras', return_value={'created': 0, 'updated': 0, 'total_in_source': 0}) as sinc:
        r = ci.post('/fibras/validacao/api/importar/', {'arquivo': SimpleUploadedFile('ANALITICO DE VENDAS.xlsx', arquivo)})
    j = r.json()
    t('importa e devolve o estado', r.status_code == 200 and 'resultado' in j, r.content[:300])
    t('sincroniza a Base de Vendas antes de cruzar', sinc.called)
    t('acha o cabeçalho depois da linha de título', EsteiraVenda.objects.count() == 5, EsteiraVenda.objects.count())
    t('guarda só as colunas do cruzamento (+SLA e motivo)',
      'OUTRA COLUNA' not in EsteiraVenda.objects.first().dados and EsteiraVenda.objects.get(ordem='7123456709').sla_agenda == 'D+2')
    esc = j['escopo']
    t('escopo padrão = período da Esteira e lojas com FILIAL', esc['ini'] == '2031-03-10' and esc['fim'] == '2031-03-12'
      and esc['lojas'] == ['ZZA001'], esc)
    for f in (f1, f2, f3, f4, f5, f6):
        f.refresh_from_db()
    t('1. protocolo exato (apóstrofo e A do final tratados) → Alta', f1.ordem_vinculada == '7123456789'
      and f1.assertividade == 'Alta' and f1.metodo_cruzamento == '1. Protocolo exato', (f1.ordem_vinculada, f1.assertividade))
    t('endereço confere', f1.conferencia_endereco == 'Confere', f1.detalhe_endereco)
    t('a ORDEM vinculada move o kanban (instalada)', f1.status == Fibra.STATUS_INSTALADO and f1.etapa_instalacao == 'INSTALADA')
    t('2. mesma loja/dia, protocolo com 1 caractere diferente → Alta', f2.ordem_vinculada == '7123456709'
      and f2.assertividade == 'Alta' and f2.metodo_cruzamento.startswith('2.'), (f2.ordem_vinculada, f2.metodo_cruzamento))
    t('e traz SLA e status (agendada)', f2.status == Fibra.STATUS_AGENDADO and f2.sla_agenda == 'D+2')
    t('4. protocolo exato em outra loja → Muito baixa com as duas lojas', f3.assertividade == 'Muito baixa'
      and 'lojas divergentes' in f3.motivo_cruzamento, f3.motivo_cruzamento)
    t('B2B fica fora do pareamento', f4.status_cruzamento == 'Venda B2B' and not f4.ordem_vinculada)
    t('sem par → Não localizada com observação', f5.status_cruzamento == 'Não localizada'
      and 'Não localizada' in f5.observacao_cruzamento, f5.observacao_cruzamento)
    t('filial sem ADABAS ainda casa pelo protocolo exato (Muito baixa)', f6.assertividade == 'Muito baixa'
      and f6.ordem_vinculada == '7444444444', (f6.assertividade, f6.motivo_cruzamento))
    res = j['resultado']
    t('KPIs do cruzamento', res['kpis']['pedidos'] == 6 and res['kpis']['alta'] == 2 and res['kpis']['muito_baixa'] == 2
      and res['kpis']['b2b'] == 1 and res['kpis']['nao_localizada'] == 1 and res['kpis']['nao_lancada'] == 1, res['kpis'])
    t('venda não lançada na lista e no kanban (ordens inconsistentes)', [x['ordem'] for x in res['nl']] == ['7888888888']
      and PlanilhaOrdemInconsistente.objects.filter(ordem='7888888888').exists())
    t('resumo loja/dia com PDV', res['resumo']['colunas'][:3] == ['ADABAS', 'PDV', 'DATA'] and res['resumo']['linhas'])
    t('histórico de status registra a Esteira', f1.status_history.filter(retorno__startswith='Esteira Vivo:').exists())

    print('\n== DETALHE DO PEDIDO ==')
    d = ci.get(f'/fibras/validacao/api/pedido/?pid={f5.pk}').json()
    t('mostra a venda e sugestões de ORDEM', d['venda']['filial'] == 'ZZ LOJA CENTRO' and d['sugestoes']['linhas']
      and 'PDV' in d['sugestoes']['colunas'], d.get('sugestoes'))
    d = ci.get(f'/fibras/validacao/api/pedido/?pid={f2.pk}').json()
    t('a ORDEM vinculada vem com os dados da Esteira', d['analitico'].get('STATUS INSTALACAO') == '3.5 AGENDADA'
      and d['analitico'].get('SLA AGENDA') == 'D+2')

    print('\n== AJUSTES MANUAIS ==')
    def ajuste(**corpo):
        return ci.post('/fibras/validacao/api/ajuste/', json.dumps(corpo), content_type='application/json')
    r = ajuste(pid=f2.pk, acao='rejeitar', ordem='7123456709', comentario='não é dessa venda')
    f2.refresh_from_db()
    t('rejeitar recruza: a ORDEM não volta', r.status_code == 200 and f2.ordem_vinculada == ''
      and f2.status_cruzamento == 'Não localizada' and f2.ordens_rejeitadas == ['7123456709'], (f2.ordem_vinculada, r.content[:200]))
    t('e guarda o comentário', f2.comentario_cruzamento == 'não é dessa venda')
    r = ajuste(pid=f2.pk, acao='vincular', ordem='7123456709')
    f2.refresh_from_db()
    t('vincular à mão → Manual (e tira da rejeição)', f2.ordem_vinculada == '7123456709' and f2.assertividade == 'Manual'
      and f2.ordens_rejeitadas == [])
    r = ajuste(pid=f3.pk, acao='confirmar', ordem='7999999999')
    f3.refresh_from_db()
    t('confirmar vira Manual', f3.assertividade == 'Manual' and f3.ordem_manual == '7999999999')
    t('contagem de ajustes', r.json()['ajustes'] == {'manuais': 2, 'rejeitados': 0, 'comentarios': 1}, r.json()['ajustes'])
    r = ajuste(pid=f5.pk, acao='comentar', comentario='cliente desistiu')
    t('comentário sozinho não recruza', r.status_code == 200 and Fibra.objects.get(pk=f5.pk).comentario_cruzamento == 'cliente desistiu')
    t('ação inválida é recusada', ajuste(pid=f5.pk, acao='apagar').status_code == 400)
    r = ajuste(pid=f2.pk, acao='desfazer')
    f2.refresh_from_db()
    t('desfazer volta ao automático', f2.ordem_manual == '' and f2.assertividade == 'Alta')

    print('\n== LOJAS ==')
    lojas = {x['filial']: x for x in ci.get('/fibras/validacao/api/estado/').json()['lojas']}
    t('FILIAL com nome igual ao setor é sugerida', lojas['ZZ LOJA CENTRO']['adabas'] == 'ZZA001'
      and lojas['ZZ LOJA CENTRO']['origem'] == 'Nome idêntico')
    t('FILIAL sem setor parecido fica sem correspondência', lojas['ZZ OUTRA LOJA']['adabas'] == '')
    r = ci.post('/fibras/validacao/api/lojas/', json.dumps({'mapa': {'ZZ OUTRA LOJA': 'ZZB002'}}), content_type='application/json')
    f6.refresh_from_db()
    t('salvar a loja recruza: protocolo exato na mesma loja vira Alta', r.status_code == 200
      and FilialAdabas.objects.filter(filial='ZZ OUTRA LOJA', adabas='ZZB002').exists()
      and f6.assertividade == 'Alta' and f6.status == Fibra.STATUS_PROBLEMA, (f6.assertividade, f6.status))
    t('a loja nova entra no escopo', 'ZZB002' in r.json()['escopo']['lojas'])

    print('\n== CRUZAR COM OUTRO ESCOPO ==')
    r = ci.post('/fibras/validacao/api/cruzar/', json.dumps({'ini': '2031-03-10', 'fim': '2031-03-10', 'janela': 0,
                                                             'lojas': ['ZZA001', 'ZZB002'], 'etapas': []}),
                content_type='application/json')
    t('cruza com o período escolhido', r.status_code == 200 and r.json()['kpis']['nao_lancada'] == 0, r.content[:200])
    r = ci.post('/fibras/validacao/api/cruzar/', json.dumps({'ini': '2031-03-12', 'fim': '2031-03-10'}),
                content_type='application/json')
    t('período invertido é recusado', r.status_code == 400 and 'depois' in r.json()['erro'])

    print('\n== BASES E EXCEL ==')
    for tipo in ('go', 'an', 'lj'):
        p = ci.get(f'/fibras/validacao/api/previa/?tipo={tipo}').json()
        t(f'prévia da base {tipo}', p['colunas'] and p['linhas'], p)
    r = ci.get('/fibras/validacao/api/baixar/?o=resultado')
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    t('Excel do resultado com as 5 abas da ferramenta', wb.sheetnames == ['Resultado Base de Vendas', 'Vendas não lançadas',
      'Consolidado', 'Resumo loja-dia', 'Parâmetros'], wb.sheetnames)
    t('e o comentário manual', 'COMENTÁRIO MANUAL' in [c.value for c in wb['Resultado Base de Vendas'][1]])
    r = ci.get('/fibras/validacao/api/baixar/?o=an')
    t('baixar a Esteira tratada', r.status_code == 200 and r['Content-Disposition'].startswith('attachment'))

    print('\n== TELAS ==')
    html = ci.get('/fibras/?ordem=7123456789&periodo=todos').content.decode()
    t('kanban filtra pela ORDEM vinculada e mostra a assertividade', 'ZZV1' in html and 'ORDEM <b>7123456789</b>' in html)
    html = ci.get(f'/fibras/{f1.pk}/').content.decode()
    t('a fibra mostra o cruzamento', 'Cruzamento com a Esteira de Vendas Vivo' in html and '7123456789' in html)

    print('\n== ERROS ==')
    ruim = planilha([])
    df = pd.read_excel(io.BytesIO(ruim), header=1).drop(columns=['ADABAS'])
    buf = io.BytesIO()
    df.to_excel(buf, index=False)
    with mock.patch.object(services, 'sync_fibras', return_value={}):
        r = ci.post('/fibras/validacao/api/importar/', {'arquivo': SimpleUploadedFile('errado.xlsx', buf.getvalue())})
    t('planilha sem ADABAS é recusada com explicação', r.status_code == 400 and 'ADABAS' in r.json()['erro'], r.content[:200])
    t('e a Esteira anterior continua', EsteiraVenda.objects.count() == 5)
    r = ci.post('/fibras/validacao/api/ajustes/limpar/', '{}', content_type='application/json')
    t('limpar ajustes', r.status_code == 200 and r.json()['ajustes'] == {'manuais': 0, 'rejeitados': 0, 'comentarios': 0})
finally:
    transaction.set_rollback(True)
    marcador.__exit__(None, None, None)
    print('\nrollback: nada deste teste foi gravado; o MySQL do Vivo GO não foi chamado.')

print(f'\n{ok} OK / {fail} falhas')
sys.exit(1 if fail else 0)
