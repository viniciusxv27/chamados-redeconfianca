"""Planilha de PIS/Cofins da competência e o pacote .zip com os arquivos.

As alíquotas ficam em células do Resumo (1,65% e 7,6%, as do regime não
cumulativo): o PIS e a Cofins de cada documento são fórmulas que apontam para
elas, então a contabilidade troca a alíquota ali e a planilha inteira refaz a
conta. O crédito é estimativa — quem decide o que dá crédito é a contabilidade.
"""
import os
import re
import zipfile
from decimal import Decimal
from io import BytesIO

from django.utils.timezone import localtime

DINHEIRO = 'R$ #,##0.00'
DATA_BR = 'DD/MM/YYYY'
VERDE = 'FF065F46'
ALIQUOTA_PIS = Decimal('0.0165')
ALIQUOTA_COFINS = Decimal('0.076')
MESES = ['janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho', 'julho', 'agosto',
         'setembro', 'outubro', 'novembro', 'dezembro']


def _mes_por_extenso(competencia):
    return f'{MESES[competencia.month - 1]}/{competencia.year}'


def nome_no_pacote(documento):
    """Nome do arquivo dentro do .zip: tipo, CNPJ, número e o id (que garante ser único)."""
    ext = os.path.splitext(documento.arquivo.name or '')[1].lower() or '.pdf'
    numero = re.sub(r'[^0-9A-Za-z-]+', '', documento.numero or '')[:30] or 'sn'
    return f'{documento.tipo.lower()}_{documento.fornecedor.cnpj}_{numero}_{documento.id}{ext}'


def _atencao(documento):
    f = documento.fornecedor
    avisos = []
    if f.situacao and f.situacao != 'ATIVA':
        avisos.append(f'CNPJ {f.situacao.lower()} na Receita')
    if f.consultado_em is None:
        avisos.append('cadastro não conferido na Receita')
    if f.simples:
        avisos.append('optante pelo Simples')
    if not documento.numero and documento.tipo != 'ALUGUEL':
        avisos.append('sem número')
    return '; '.join(avisos)


def planilha_piscofins(competencia, documentos):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from .models import DocumentoPisCofins

    negrito = Font(bold=True)
    cabecalho = Font(bold=True, color='FFFFFFFF')
    fundo = PatternFill('solid', fgColor=VERDE)

    livro = Workbook()
    resumo = livro.active
    resumo.title = 'Resumo'
    docs = livro.create_sheet('Documentos')
    forn = livro.create_sheet('Por fornecedor')

    # ── Documentos ───────────────────────────────────────────────────────
    colunas = [('Competência', 12), ('Tipo', 18), ('Número', 14), ('Emissão', 12), ('CNPJ', 20),
               ('Razão social', 40), ('Situação na Receita', 18), ('Simples', 9), ('Loja / unidade', 24),
               ('Descrição', 36), ('Valor (base)', 15), ('PIS', 13), ('Cofins', 13), ('Atenção', 34),
               ('Arquivo no pacote', 44), ('Lançado por', 24), ('Lançado em', 16)]
    for i, (titulo, largura) in enumerate(colunas, start=1):
        c = docs.cell(row=1, column=i, value=titulo)
        c.font, c.fill = cabecalho, fundo
        docs.column_dimensions[get_column_letter(i)].width = largura
    docs.freeze_panes = 'A2'
    for linha, d in enumerate(documentos, start=2):
        f = d.fornecedor
        valores = [
            competencia, d.get_tipo_display(), d.numero, d.data_documento, f.cnpj_formatado, f.razao_social,
            f.situacao or 'não consultado', {True: 'Sim', False: 'Não'}.get(f.simples, ''),
            d.loja.name if d.loja else '', d.descricao, float(d.valor),
            f'=ROUND(K{linha}*Resumo!$B$5,2)', f'=ROUND(K{linha}*Resumo!$B$6,2)',
            _atencao(d), nome_no_pacote(d),
            getattr(d.registrado_por, 'full_name', '') or '',
            localtime(d.criado_em).replace(tzinfo=None) if d.criado_em else None,
        ]
        for col, v in enumerate(valores, start=1):
            c = docs.cell(row=linha, column=col, value=v)
            if col in (1, 4):
                c.number_format = DATA_BR if col == 4 else 'MM/YYYY'
            elif col in (11, 12, 13):
                c.number_format = DINHEIRO
            elif col == 17:
                c.number_format = 'DD/MM/YYYY HH:MM'
    ultima = len(documentos) + 1
    if documentos:
        docs.auto_filter.ref = f'A1:{get_column_letter(len(colunas))}{ultima}'
    total_linha = ultima + 2
    docs.cell(row=total_linha, column=10, value='TOTAL (filtro)').font = negrito
    for col, letra in ((11, 'K'), (12, 'L'), (13, 'M')):
        c = docs.cell(row=total_linha, column=col, value=f'=SUBTOTAL(109,{letra}2:{letra}{max(ultima, 2)})')
        c.number_format, c.font = DINHEIRO, negrito

    # ── Resumo ───────────────────────────────────────────────────────────
    resumo.column_dimensions['A'].width = 34
    for letra in 'BCDE':
        resumo.column_dimensions[letra].width = 17
    resumo['A1'] = f'PIS/Cofins — competência {_mes_por_extenso(competencia)}'
    resumo['A1'].font = Font(bold=True, size=13)
    resumo['A2'] = ('Valores lançados no portal (Contas a pagar · PIS/Cofins). O crédito é estimativa '
                    'pelas alíquotas abaixo — a contabilidade confirma o que dá crédito.')
    resumo['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    resumo.merge_cells('A2:E3')
    resumo.row_dimensions[2].height = 30
    resumo['A5'], resumo['B5'] = 'Alíquota PIS', float(ALIQUOTA_PIS)
    resumo['A6'], resumo['B6'] = 'Alíquota Cofins', float(ALIQUOTA_COFINS)
    for ref in ('B5', 'B6'):
        resumo[ref].number_format = '0.00%'
        resumo[ref].fill = PatternFill('solid', fgColor='FFFEF3C7')
    resumo['C5'] = '← altere aqui se o regime for outro'
    resumo['C5'].font = Font(italic=True, color='FF6B7280')

    linha = 8
    for i, titulo in enumerate(('Tipo de documento', 'Quantidade', 'Valor (base)', 'PIS', 'Cofins'), start=1):
        c = resumo.cell(row=linha, column=i, value=titulo)
        c.font, c.fill = cabecalho, fundo
    intervalo = f'Documentos!$B$2:$B${max(ultima, 2)}'
    primeira = linha + 1
    for _chave, rotulo in DocumentoPisCofins.TIPOS:
        linha += 1
        resumo.cell(row=linha, column=1, value=rotulo)
        resumo.cell(row=linha, column=2, value=f'=COUNTIF({intervalo},A{linha})')
        for col, letra in ((3, 'K'), (4, 'L'), (5, 'M')):
            c = resumo.cell(row=linha, column=col,
                            value=f'=SUMIF({intervalo},A{linha},Documentos!${letra}$2:${letra}${max(ultima, 2)})')
            c.number_format = DINHEIRO
    linha += 1
    resumo.cell(row=linha, column=1, value='Total').font = negrito
    for col, letra in ((2, 'B'), (3, 'C'), (4, 'D'), (5, 'E')):
        c = resumo.cell(row=linha, column=col, value=f'=SUM({letra}{primeira}:{letra}{linha - 1})')
        c.font = negrito
        if col > 2:
            c.number_format = DINHEIRO

    alertas = [(d, _atencao(d)) for d in documentos if _atencao(d)]
    if alertas:
        linha += 2
        resumo.cell(row=linha, column=1, value=f'Pontos de atenção ({len(alertas)})').font = negrito
        for d, aviso in alertas[:50]:
            linha += 1
            resumo.cell(row=linha, column=1, value=f'{d.fornecedor.razao_social} — {d.get_tipo_display()} {d.numero}')
            resumo.cell(row=linha, column=2, value=aviso)

    # ── Por fornecedor ───────────────────────────────────────────────────
    for i, (titulo, largura) in enumerate((('CNPJ', 20), ('Razão social', 42), ('Situação', 14),
                                           ('Documentos', 12), ('Valor (base)', 15), ('PIS', 13),
                                           ('Cofins', 13)), start=1):
        c = forn.cell(row=1, column=i, value=titulo)
        c.font, c.fill = cabecalho, fundo
        forn.column_dimensions[get_column_letter(i)].width = largura
    fornecedores = {}
    for d in documentos:
        fornecedores.setdefault(d.fornecedor.cnpj, d.fornecedor)
    cnpjs = f'Documentos!$E$2:$E${max(ultima, 2)}'
    for linha, f in enumerate(sorted(fornecedores.values(), key=lambda f: f.razao_social), start=2):
        forn.cell(row=linha, column=1, value=f.cnpj_formatado)
        forn.cell(row=linha, column=2, value=f.razao_social)
        forn.cell(row=linha, column=3, value=f.situacao or 'não consultado')
        forn.cell(row=linha, column=4, value=f'=COUNTIF({cnpjs},A{linha})')
        for col, letra in ((5, 'K'), (6, 'L'), (7, 'M')):
            c = forn.cell(row=linha, column=col,
                          value=f'=SUMIF({cnpjs},A{linha},Documentos!${letra}$2:${letra}${max(ultima, 2)})')
            c.number_format = DINHEIRO
    forn.freeze_panes = 'A2'

    buffer = BytesIO()
    livro.save(buffer)
    return buffer.getvalue()


def pacote_piscofins(competencia, documentos):
    """O .zip da competência: a planilha na raiz e os arquivos em ``documentos/``."""
    buffer = BytesIO()
    faltando = []
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(f'pis_cofins_{competencia:%Y-%m}.xlsx', planilha_piscofins(competencia, documentos))
        for d in documentos:
            try:
                with d.arquivo.open('rb') as arquivo:
                    zf.writestr(f'documentos/{nome_no_pacote(d)}', arquivo.read())
            except Exception as exc:  # noqa: BLE001 — o pacote sai, e diz o que faltou
                faltando.append(f'{nome_no_pacote(d)}: {exc}')
        if faltando:
            zf.writestr('LEIA-ME_arquivos_faltando.txt',
                        'Estes arquivos não puderam ser lidos do armazenamento:\n\n' + '\n'.join(faltando))
    return buffer.getvalue()
