"""Relatório de sangrias em Excel: resumo, lista e os achados (recorrentes, fora do padrão)."""
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse

DINHEIRO = 'R$ #,##0.00'
DATA_BR = 'DD/MM/YYYY'
ROXO = 'FF4C1D95'


def _cabecalho(aba, linha, colunas):
    from openpyxl.styles import Font, PatternFill
    for i, (titulo, largura) in enumerate(colunas, start=1):
        celula = aba.cell(row=linha, column=i, value=titulo)
        celula.font = Font(bold=True, color='FFFFFFFF')
        celula.fill = PatternFill('solid', fgColor=ROXO)
        aba.column_dimensions[celula.column_letter].width = max(aba.column_dimensions[celula.column_letter].width or 0, largura)


def _escrever(aba, linha, valores, formatos=None):
    for i, valor in enumerate(valores, start=1):
        if isinstance(valor, Decimal):
            valor = float(valor)
        celula = aba.cell(row=linha, column=i, value=valor)
        formato = (formatos or {}).get(i)
        if formato:
            celula.number_format = formato


def planilha_sangrias(filtro, sangrias, numeros, categorias, lojas, recorrentes, fora):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    livro = Workbook()
    resumo = livro.active
    resumo.title = 'Resumo'
    resumo.column_dimensions['A'].width = 30
    resumo.column_dimensions['B'].width = 18
    resumo.cell(row=1, column=1, value='Sangrias de caixa').font = Font(bold=True, size=13)
    linha = 3
    for rotulo, valor in (('Período', f"{filtro['inicio']:%d/%m/%Y} a {filtro['fim']:%d/%m/%Y}"),
                          ('Total retirado', numeros['total']), ('Sangrias', numeros['quantidade']),
                          ('Valor médio', numeros['media']), ('A conferir', numeros['a_conferir']),
                          ('Valor a conferir', numeros['valor_a_conferir']),
                          ('Sem comprovante', numeros['sem_comprovante'])):
        resumo.cell(row=linha, column=1, value=rotulo).font = Font(bold=True)
        _escrever_celula = resumo.cell(row=linha, column=2, value=float(valor) if isinstance(valor, Decimal) else valor)
        if isinstance(valor, Decimal):
            _escrever_celula.number_format = DINHEIRO
        linha += 1

    for titulo, linhas in (('Por categoria', categorias), ('Por loja', lojas)):
        if not linhas:
            continue
        linha += 1
        _cabecalho(resumo, linha, [(titulo, 30), ('Total', 18), ('Qtd.', 8), ('% do total', 12)])
        for item in linhas:
            linha += 1
            _escrever(resumo, linha, [item['nome'], item['total'], item['n'], item['percentual'] / 100],
                      {2: DINHEIRO, 4: '0.0%'})
        linha += 1

    aba = livro.create_sheet('Sangrias')
    colunas = [('Data', 12), ('Loja', 26), ('Categoria', 22), ('Descrição', 44), ('Pago a', 24),
               ('Valor', 14), ('Comprovante', 13), ('Registrada por', 24), ('Conferida', 11),
               ('Conferida por', 24), ('Conferida em', 16)]
    _cabecalho(aba, 1, colunas)
    for i, s in enumerate(sangrias, start=2):
        _escrever(aba, i, [
            s.data, s.loja.name, s.categoria.nome, s.descricao, s.favorecido, s.valor,
            'Sim' if s.comprovante else 'Não',
            getattr(s.registrada_por, 'full_name', '') or '', 'Sim' if s.conferida else 'Não',
            getattr(s.conferida_por, 'full_name', '') or '',
            s.conferida_em.replace(tzinfo=None) if s.conferida_em else None,
        ], {1: DATA_BR, 6: DINHEIRO, 11: 'DD/MM/YYYY HH:MM'})
    ultima = len(sangrias) + 1
    aba.freeze_panes = 'A2'
    if sangrias:
        aba.auto_filter.ref = f'A1:{get_column_letter(len(colunas))}{ultima}'
        aba.cell(row=ultima + 2, column=5, value='TOTAL (filtro)').font = Font(bold=True)
        total = aba.cell(row=ultima + 2, column=6, value=f'=SUBTOTAL(109,F2:F{ultima})')
        total.number_format, total.font = DINHEIRO, Font(bold=True)

    if recorrentes:
        aba = livro.create_sheet('Recorrentes')
        _cabecalho(aba, 1, [('Gasto', 36), ('Categoria', 22), ('Vezes', 8), ('Total', 14),
                            ('Lojas', 40), ('Última', 12)])
        for i, g in enumerate(recorrentes, start=2):
            _escrever(aba, i, [g['rotulo'], g['categoria'], g['n'], g['total'], ', '.join(g['lojas']), g['ultima']],
                      {4: DINHEIRO, 6: DATA_BR})
    if fora:
        aba = livro.create_sheet('Fora do padrão')
        _cabecalho(aba, 1, [('Data', 12), ('Loja', 26), ('Categoria', 22), ('Descrição', 40), ('Valor', 14),
                            ('Mediana da categoria', 20), ('Vezes a mediana', 16)])
        for i, s in enumerate(fora, start=2):
            _escrever(aba, i, [s.data, s.loja.name, s.categoria.nome, s.descricao, s.valor,
                               s.mediana_categoria, s.vezes_mediana], {1: DATA_BR, 5: DINHEIRO, 6: DINHEIRO})

    buffer = BytesIO()
    livro.save(buffer)
    resposta = HttpResponse(buffer.getvalue(),
                            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resposta['Content-Disposition'] = (
        f"attachment; filename=\"sangrias_{filtro['inicio']:%Y%m%d}-{filtro['fim']:%Y%m%d}.xlsx\"")
    return resposta
