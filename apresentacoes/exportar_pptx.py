"""Exporta uma apresentação do portal para PowerPoint (.pptx nativo e editável).

O documento é o JSON de `Apresentacao.documento` (quadro lógico de 1920×1080 px).
Cada elemento vira um objeto nativo — caixa de texto, forma, figura, tabela,
vídeo — para quem baixa continuar editando no PowerPoint, no Keynote, no Google
Slides ou no Canva. Transições e animações de entrada saem no mesmo XML que o
PowerPoint grava: é o formato que esses programas abrem sem pedir "reparo".

Escalas: 1 px do quadro = 6350 EMU (1920 px ↔ 12.192.000 EMU, slide 16:9) e,
para fonte, 1 px = 0,5 pt (1920 px ↔ 960 pt, a largura do slide em pontos).

Uso:
    from apresentacoes.exportar_pptx import gerar_pptx, nome_do_arquivo
    conteudo = gerar_pptx(apresentacao.documento, apresentacao.titulo)
"""
import colorsys
import io
import json
import logging
import math
import mimetypes
import re
import struct
import unicodedata
from datetime import datetime, timezone as dt_timezone
from functools import lru_cache
from html.parser import HTMLParser
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageFilter, ImageFont, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu

logger = logging.getLogger(__name__)

AUTOR = 'Portal Rede Confiança'

QUADRO_LARGURA = 1920
QUADRO_ALTURA = 1080
SLIDE_LARGURA_EMU = 12192000
SLIDE_ALTURA_EMU = 6858000
EMU_POR_PX = 6350
PT_POR_PX = 0.5

RECURSOS = Path(__file__).resolve().parent / 'recursos'
MAPA_ICONES = RECURSOS / 'fa-6.0.0-icones.json'
FONTES_ICONE = {'s': 'fa-solid-900.ttf', 'r': 'fa-regular-400.ttf', 'b': 'fa-brands-400.ttf'}

NS_MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
NS_P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
NS_P14 = 'http://schemas.microsoft.com/office/powerpoint/2010/main'

# Tabela sem estilo e sem grade: as cores de cada célula vêm do documento, não do tema do PowerPoint.
ESTILO_TABELA_NENHUM = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'

TEMA_PADRAO = {
    'fonte_titulo': 'Montserrat',
    'fonte_texto': 'Montserrat',
    'cores': {'fundo': '#0B0612', 'superficie': '#1A0F24', 'primaria': '#E23FCF', 'secundaria': '#8B3DFF',
              'destaque': '#FFD15C', 'texto': '#FFFFFF', 'texto_suave': '#D9C9E8'},
}

# Sombras do editor em px do quadro (desfoque, deslocamento para baixo, opacidade do preto).
SOMBRAS = {
    'suave': {'desfoque': 24, 'distancia': 8, 'alfa': 0.28},
    'forte': {'desfoque': 48, 'distancia': 18, 'alfa': 0.5},
}
BRILHO = {'raio': 18, 'alfa': 0.65}          # "brilho" = glow na cor primária do tema

ICONE_ESCALA_SEM_FUNDO = 0.8                 # glifo ocupa 80% do menor lado da caixa
ICONE_ESCALA_COM_FUNDO = 0.56                # com fundo, sobra respiro até a borda

# Entrada → (presetID, presetSubtype) exatamente como o PowerPoint grava cada efeito.
PRESETS_ENTRADA = {
    'aparecer': (1, 0),            # Appear
    'fade': (10, 0),               # Fade
    'subir': (42, 0),              # Float In (Ascend): fade + sobe 10% do slide
    'descer': (47, 0),             # Float Down (Descend)
    'entrar-esquerda': (2, 8),     # Fly In, From Left (8 = esquerda)
    'entrar-direita': (2, 2),      # Fly In, From Right (2 = direita)
    'zoom': (53, 16),              # Zoom (Fade Zoom), Object Center
    'revelar': (22, 8),            # Wipe, From Left → filter="wipe(right)"
}

# Transição → (elemento clássico de p:transition, atributos). Direção = sentido do movimento.
TRANSICOES = {
    'fade': ('fade', {}),
    'empurrar-esquerda': ('push', {'dir': 'l'}),
    'empurrar-direita': ('push', {'dir': 'r'}),
    'empurrar-cima': ('push', {'dir': 'u'}),
    'empurrar-baixo': ('push', {'dir': 'd'}),
    'revelar': ('wipe', {'dir': 'r'}),
    'zoom': ('zoom', {'dir': 'in'}),
    'cobrir': ('cover', {'dir': 'l'}),
}

FORMAS = {
    'retangulo': MSO_SHAPE.RECTANGLE,
    'retangulo-arredondado': MSO_SHAPE.ROUNDED_RECTANGLE,
    'pilula': MSO_SHAPE.ROUNDED_RECTANGLE,
    'circulo': MSO_SHAPE.OVAL,
    'seta': MSO_SHAPE.RIGHT_ARROW,
    'triangulo': MSO_SHAPE.ISOSCELES_TRIANGLE,
    'estrela': MSO_SHAPE.STAR_5_POINT,
    'hexagono': MSO_SHAPE.HEXAGON,
}

EXTENSOES_VIDEO = {
    'video/mp4': 'mp4', 'video/quicktime': 'mov', 'video/webm': 'webm', 'video/x-m4v': 'm4v',
    'video/x-msvideo': 'avi', 'video/x-ms-wmv': 'wmv', 'video/mpeg': 'mpg', 'video/ogg': 'ogv',
}

_RE_MIDIA = re.compile(r'^/apresentacoes/midia/(\d+)/?$')
_RE_CONTROLE = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f￾￿\ud800-\udfff]')
_RE_LINK_SLIDE = re.compile(r'^#slide-(\d+)$')
_RE_LINK_WEB = re.compile(r'^https?://[^\s<>"]+$', re.I)


# ----------------------------------------------------------------------------------------------
# utilidades
# ----------------------------------------------------------------------------------------------

def _num(valor, padrao=0.0):
    try:
        numero = float(valor)
    except (TypeError, ValueError):
        return padrao
    if math.isnan(numero) or math.isinf(numero):
        return padrao
    return numero


def _dict(valor):
    return valor if isinstance(valor, dict) else {}


def _texto_xml(texto):
    """Tira caracteres que o XML não aceita (controle, substitutos soltos)."""
    return _RE_CONTROLE.sub('', str(texto or ''))


def _limitar(valor, minimo, maximo):
    return max(minimo, min(maximo, valor))


def _sub(pai, tag, **attrs):
    elemento = etree.SubElement(pai, qn(tag))
    for chave, valor in attrs.items():
        if valor is not None:
            elemento.set(chave, str(valor))
    return elemento


def _dividir_topo(texto, separador=','):
    """Divide no separador fora de parênteses: 'rgba(0,0,0,.5) 10%, red' → 2 partes."""
    partes, nivel, atual = [], 0, []
    for caractere in texto:
        if caractere == '(':
            nivel += 1
        elif caractere == ')':
            nivel = max(0, nivel - 1)
        if nivel == 0 and (caractere == separador or (separador == ' ' and caractere.isspace())):
            partes.append(''.join(atual))
            atual = []
        else:
            atual.append(caractere)
    partes.append(''.join(atual))
    return [parte.strip() for parte in partes if parte.strip()]


# ----------------------------------------------------------------------------------------------
# cores e gradientes
# ----------------------------------------------------------------------------------------------

_CORES_NOMEADAS = {
    'white': 'FFFFFF', 'black': '000000', 'red': 'FF0000', 'green': '008000', 'blue': '0000FF',
    'yellow': 'FFFF00', 'gray': '808080', 'grey': '808080', 'orange': 'FFA500', 'purple': '800080',
    'pink': 'FFC0CB', 'silver': 'C0C0C0',
}


def _canal(texto, maximo=255):
    texto = texto.strip()
    if texto.endswith('%'):
        return _limitar(_num(texto[:-1]) / 100 * maximo, 0, maximo)
    return _limitar(_num(texto), 0, maximo)


def _alfa(texto):
    texto = texto.strip()
    if texto.endswith('%'):
        return _limitar(_num(texto[:-1], 100) / 100, 0, 1)
    return _limitar(_num(texto, 1), 0, 1)


def _cor_bruta(valor, cores_tema=None, profundidade=0):
    """Cor CSS/tema → ('RRGGBB', alfa) — 'transparent' vira alfa 0; vazio ou inválido → None."""
    if valor is None:
        return None
    texto = str(valor).strip()
    if not texto or texto.lower() in ('none', 'inherit', 'initial', 'currentcolor'):
        return None
    if texto.lower() == 'transparent':
        return ('000000', 0.0)
    if texto.lower().startswith('tema:'):
        if profundidade > 3:
            return None
        return _cor_bruta(_dict(cores_tema).get(texto[5:].strip()), cores_tema, profundidade + 1)
    if texto.startswith('#'):
        hexa = texto[1:]
        if not re.fullmatch(r'[0-9a-fA-F]{3,8}', hexa) or len(hexa) not in (3, 4, 6, 8):
            return None
        if len(hexa) in (3, 4):
            hexa = ''.join(c * 2 for c in hexa)
        alfa = int(hexa[6:8], 16) / 255 if len(hexa) == 8 else 1.0
        return (hexa[:6].upper(), alfa)
    funcao = re.fullmatch(r'(rgba?|hsla?)\((.*)\)', texto, re.I | re.S)
    if funcao:
        partes = [p for p in re.split(r'[\s,/]+', funcao.group(2).strip()) if p]
        if len(partes) < 3:
            return None
        alfa = _alfa(partes[3]) if len(partes) > 3 else 1.0
        if funcao.group(1).lower().startswith('rgb'):
            r, g, b = (int(round(_canal(p))) for p in partes[:3])
        else:
            matiz = _num(partes[0].replace('deg', '')) % 360 / 360
            saturacao = _canal(partes[1], 100) / 100
            luz = _canal(partes[2], 100) / 100
            r, g, b = (int(round(c * 255)) for c in colorsys.hls_to_rgb(matiz, luz, saturacao))
        return (f'{r:02X}{g:02X}{b:02X}', alfa)
    if texto.lower() in _CORES_NOMEADAS:
        return (_CORES_NOMEADAS[texto.lower()], 1.0)
    return None


def resolver_cor(valor, cores_tema=None):
    """Cor do documento → ('RRGGBB', alfa 0..1) ou None quando não pinta nada."""
    cor = _cor_bruta(valor, cores_tema)
    if cor is None or cor[1] <= 0.001:
        return None
    return cor


_RE_GRADIENTE = re.compile(r'^\s*(?:repeating-)?(linear|radial)-gradient\s*\((.*)\)\s*$', re.I | re.S)


def _angulo_para_lados(lados, largura, altura):
    lados = set(lados.lower().split())
    horizontal = 'left' if 'left' in lados else 'right' if 'right' in lados else None
    vertical = 'top' if 'top' in lados else 'bottom' if 'bottom' in lados else None
    if horizontal and vertical:
        # No CSS o "to top right" é perpendicular à diagonal: depende da proporção da caixa.
        base = math.degrees(math.atan2(altura, largura))
        return {('top', 'right'): base, ('bottom', 'right'): 180 - base,
                ('bottom', 'left'): 180 + base, ('top', 'left'): 360 - base}[(vertical, horizontal)]
    return {'top': 0.0, 'right': 90.0, 'bottom': 180.0, 'left': 270.0}.get(vertical or horizontal, 180.0)


def _posicao_css(texto):
    palavras = {'left': 0.0, 'top': 0.0, 'center': 0.5, 'right': 1.0, 'bottom': 1.0}
    tokens = texto.lower().split()
    x, y = 0.5, 0.5
    valores = []
    for token in tokens:
        if token in palavras:
            if token in ('top', 'bottom'):
                y = palavras[token]
            elif token in ('left', 'right'):
                x = palavras[token]
            else:
                valores.append(0.5)
        elif token.endswith('%'):
            valores.append(_limitar(_num(token[:-1], 50) / 100, 0, 1))
    if valores and not any(t in ('left', 'right', 'top', 'bottom') for t in tokens):
        x = valores[0]
        y = valores[1] if len(valores) > 1 else 0.5
    return x, y


def resolver_gradiente(css, cores_tema=None, largura=QUADRO_LARGURA, altura=QUADRO_ALTURA):
    """'linear-gradient(135deg, #a 0%, #b 100%)' → {'tipo', 'angulo' (CSS), 'centro', 'paradas'} ou None.

    `paradas` = [(posição 0..1, ('RRGGBB', alfa))]. Radial é aproximado por um gradiente de caminho.
    """
    if not isinstance(css, str) or 'gradient' not in css.lower():
        return None
    for camada in _dividir_topo(css):
        combinou = _RE_GRADIENTE.match(camada)
        if combinou:
            break
    else:
        return None
    tipo = combinou.group(1).lower()
    argumentos = _dividir_topo(combinou.group(2))
    angulo, centro = 180.0, (0.5, 0.5)
    largura, altura = max(_num(largura, 1), 1), max(_num(altura, 1), 1)
    if argumentos:
        primeiro = argumentos[0].strip().lower()
        if tipo == 'linear':
            medida = re.fullmatch(r'(-?[\d.]+)\s*(deg|rad|grad|turn)?', primeiro)
            if medida and (medida.group(2) or _num(medida.group(1)) == 0):
                valor = _num(medida.group(1))
                unidade = medida.group(2) or 'deg'
                angulo = {'deg': valor, 'rad': math.degrees(valor), 'grad': valor * 0.9, 'turn': valor * 360}[unidade]
                argumentos = argumentos[1:]
            elif primeiro.startswith('to '):
                angulo = _angulo_para_lados(primeiro[3:], largura, altura)
                argumentos = argumentos[1:]
        else:
            tokens = _dividir_topo(primeiro, ' ')
            if not tokens or _cor_bruta(tokens[0], cores_tema) is None:
                posicao = re.search(r'\bat\s+(.+)$', primeiro)
                if posicao:
                    centro = _posicao_css(posicao.group(1))
                argumentos = argumentos[1:]
    if tipo == 'linear':
        radianos = math.radians(angulo)
        comprimento = abs(largura * math.sin(radianos)) + abs(altura * math.cos(radianos))
    else:
        cx, cy = centro[0] * largura, centro[1] * altura
        comprimento = math.hypot(max(cx, largura - cx), max(cy, altura - cy))
    comprimento = max(comprimento, 1)

    paradas = []
    for argumento in argumentos:
        tokens = _dividir_topo(argumento, ' ')
        if not tokens:
            continue
        cor = _cor_bruta(tokens[0], cores_tema)
        if cor is None:
            continue                               # dica de interpolação ou lixo: ignora
        posicoes = []
        for token in tokens[1:3]:
            token = token.lower()
            if token.endswith('%'):
                posicoes.append(_num(token[:-1]) / 100)
            elif token.endswith('px'):
                posicoes.append(_num(token[:-2]) / comprimento)
            elif re.fullmatch(r'-?[\d.]+', token):
                posicoes.append(_num(token) / comprimento)
        if posicoes:
            paradas.extend([cor, p] for p in posicoes)
        else:
            paradas.append([cor, None])
    if not paradas:
        return None
    if len(paradas) == 1:
        paradas = [[paradas[0][0], 0.0], [paradas[0][0], 1.0]]
    if paradas[0][1] is None:
        paradas[0][1] = 0.0
    if paradas[-1][1] is None:
        paradas[-1][1] = 1.0
    maior = paradas[0][1]
    for parada in paradas:                         # regra do CSS: posição nunca volta para trás
        if parada[1] is not None:
            parada[1] = max(parada[1], maior)
            maior = parada[1]
    indice = 1
    while indice < len(paradas):                   # paradas sem posição: distribui por igual
        if paradas[indice][1] is None:
            fim = indice
            while paradas[fim][1] is None:
                fim += 1
            inicio_pos, fim_pos = paradas[indice - 1][1], paradas[fim][1]
            passos = fim - indice + 1
            for k in range(indice, fim):
                paradas[k][1] = inicio_pos + (fim_pos - inicio_pos) * (k - indice + 1) / passos
            indice = fim
        indice += 1
    return {
        'tipo': tipo,
        'angulo': angulo % 360,
        'centro': centro,
        'paradas': [(_limitar(p, 0.0, 1.0), cor) for cor, p in paradas],
    }


# ----------------------------------------------------------------------------------------------
# HTML restrito → parágrafos e trechos
# ----------------------------------------------------------------------------------------------

_BLOCOS = {'p', 'div', 'li', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'blockquote', 'section', 'article',
           'header', 'footer'}
_FAMILIAS_GENERICAS = {'serif', 'sans-serif', 'monospace', 'cursive', 'fantasy', 'system-ui', 'inherit',
                       'initial', 'ui-sans-serif', 'ui-serif', 'ui-monospace'}


def _tamanho_css(valor, atual, base):
    texto = valor.strip().lower()
    if texto.endswith('px'):
        return _num(texto[:-2], atual) or atual
    if texto.endswith('rem'):
        return _num(texto[:-3], 1) * base
    if texto.endswith('em'):
        return _num(texto[:-2], 1) * atual
    if texto.endswith('%'):
        return _num(texto[:-1], 100) / 100 * atual
    if texto.endswith('pt'):
        return _num(texto[:-2], atual * 0.75) * 4 / 3
    numero = _num(texto, 0)
    return numero if numero > 0 else atual


def _aplicar_css(estilo, css, base):
    for declaracao in str(css or '').split(';'):
        if ':' not in declaracao:
            continue
        propriedade, valor = declaracao.split(':', 1)
        propriedade = propriedade.strip().lower()
        valor = re.sub(r'\s*!important\s*$', '', valor.strip(), flags=re.I)
        if not valor:
            continue
        if propriedade == 'color':
            estilo['cor'] = valor
        elif propriedade in ('background-color', 'background'):
            estilo['realce'] = valor
        elif propriedade == 'font-size':
            estilo['tamanho'] = _limitar(_tamanho_css(valor, estilo['tamanho'], base), 1, 4000)
        elif propriedade == 'font-weight':
            nome = valor.lower()
            estilo['peso'] = {'bold': 700, 'bolder': 800, 'normal': 400, 'lighter': 300}.get(nome, _num(nome, 400))
        elif propriedade == 'font-style':
            estilo['italico'] = valor.lower() in ('italic', 'oblique')
        elif propriedade in ('text-decoration', 'text-decoration-line'):
            nome = valor.lower()
            estilo['sublinhado'] = 'underline' in nome
            estilo['tachado'] = 'line-through' in nome
        elif propriedade == 'font-family':
            familia = valor.split(',')[0].strip().strip('\'"').strip()
            if familia and familia.lower() not in _FAMILIAS_GENERICAS:
                estilo['fonte'] = familia


class _LeitorHtml(HTMLParser):
    """Lê o HTML permitido do editor e separa em parágrafos com trechos estilizados."""

    def __init__(self, estilo_base):
        super().__init__(convert_charrefs=True)
        self.base = estilo_base['tamanho']
        self.pilha = [(None, dict(estilo_base))]
        self.listas = []
        self.paragrafos = []
        self.atual = None

    @property
    def estilo(self):
        return self.pilha[-1][1]

    def _abrir(self, lista=None, nivel=0):
        if self.atual is not None and not self.atual['trechos']:
            # Bloco dentro de <li> ainda vazio: continua o mesmo parágrafo (mantém o marcador).
            if lista is not None:
                self.atual.update(lista=lista, nivel=nivel)
            return
        self._fechar()
        self.atual = {'trechos': [], 'lista': lista, 'nivel': nivel}

    def _fechar(self):
        if self.atual is not None and self.atual['trechos']:
            self.paragrafos.append(self.atual)
        self.atual = None

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        atributos = {nome.lower(): (valor or '') for nome, valor in attrs}
        if tag == 'br':
            if self.atual is None:
                self._abrir()
            self.atual['trechos'].append({'quebra': True, 'estilo': dict(self.estilo)})
            return
        estilo = dict(self.estilo)
        if tag in ('b', 'strong'):
            estilo['peso'] = max(_num(estilo.get('peso'), 400), 700)
        elif tag in ('i', 'em'):
            estilo['italico'] = True
        elif tag == 'u':
            estilo['sublinhado'] = True
        elif tag in ('s', 'strike', 'del'):
            estilo['tachado'] = True
        elif tag == 'a':
            href = atributos.get('href', '').strip()
            if _RE_LINK_WEB.match(href) or _RE_LINK_SLIDE.match(href):
                estilo['link'] = href
        if atributos.get('style'):
            _aplicar_css(estilo, atributos['style'], self.base)
        if tag in ('ul', 'ol'):
            self._fechar()
            self.listas.append(tag)
        elif tag == 'li':
            self._abrir(lista=self.listas[-1] if self.listas else 'ul', nivel=max(len(self.listas) - 1, 0))
        elif tag in _BLOCOS:
            if self.atual is not None and not self.atual['trechos'] and self.atual.get('lista'):
                pass                                   # <li><p>…: o parágrafo do item continua
            else:
                self._abrir()
        self.pilha.append((tag, estilo))

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag.lower() != 'br':
            self.handle_endtag(tag)

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag == 'br':
            return
        if any(nome == tag for nome, _ in self.pilha[1:]):
            while len(self.pilha) > 1:
                nome, _ = self.pilha.pop()
                if nome == tag:
                    break
        if tag in ('ul', 'ol'):
            if self.listas:
                self.listas.pop()
            self._fechar()
        elif tag == 'li' or tag in _BLOCOS:
            self._fechar()

    def handle_data(self, dados):
        if not dados:
            return
        if self.atual is None:
            if not dados.strip(' \t\r\n\f'):
                return                                 # espaço entre blocos não vira parágrafo
            self._abrir()
        self.atual['trechos'].append({'texto': dados, 'estilo': dict(self.estilo)})

    def resultado(self):
        self.close()
        self._fechar()
        return [_normalizar_paragrafo(p) for p in self.paragrafos]


def _normalizar_paragrafo(paragrafo):
    """Espaços como no CSS (white-space: normal) e <br> final sem linha fantasma."""
    trechos = []
    anterior_espaco = True                              # começo de linha engole espaço
    for trecho in paragrafo['trechos']:
        if trecho.get('quebra'):
            if trechos and not trechos[-1].get('quebra'):
                trechos[-1]['texto'] = trechos[-1]['texto'].rstrip(' ')
            trechos.append(trecho)
            anterior_espaco = True
            continue
        texto = re.sub(r'[ \t\r\n\f]+', ' ', trecho['texto'])
        if anterior_espaco:
            texto = texto.lstrip(' ')
        if not texto:
            continue
        anterior_espaco = texto.endswith(' ')
        trechos.append({'texto': texto, 'estilo': trecho['estilo']})
    if trechos and not trechos[-1].get('quebra'):
        trechos[-1]['texto'] = trechos[-1]['texto'].rstrip(' ')
    trechos = [t for t in trechos if t.get('quebra') or t['texto']]
    estilo_final = paragrafo['trechos'][-1]['estilo'] if paragrafo['trechos'] else {}
    if trechos and trechos[-1].get('quebra') and len(trechos) > 1:
        trechos.pop()                                   # "a<br>" mostra uma linha só, como no navegador
    if len(trechos) == 1 and trechos[0].get('quebra'):
        estilo_final = trechos[0]['estilo']
        trechos = []                                    # "<div><br></div>" = linha em branco
    return {'trechos': trechos, 'lista': paragrafo['lista'], 'nivel': paragrafo['nivel'],
            'estilo_final': estilo_final}


def ler_html(html, estilo_base):
    """HTML restrito do editor → [{'trechos': [{'texto'|'quebra', 'estilo'}], 'lista', 'nivel', 'estilo_final'}]."""
    leitor = _LeitorHtml(estilo_base)
    try:
        leitor.feed(str(html or ''))
        return leitor.resultado()
    except Exception:                                   # noqa: BLE001 — HTML quebrado vira texto puro
        logger.warning('Exportação PPTX: HTML de texto ilegível, usando texto puro', exc_info=True)
        texto = re.sub(r'<[^>]+>', ' ', str(html or ''))
        return [{'trechos': [{'texto': texto, 'estilo': dict(estilo_base)}], 'lista': None, 'nivel': 0,
                 'estilo_final': dict(estilo_base)}]


# ----------------------------------------------------------------------------------------------
# mídias
# ----------------------------------------------------------------------------------------------

def ler_midia_padrao(src, permitir=None):
    """Resolve a fonte de mídia do documento em (bytes, mime) — None se não der.

    `/apresentacoes/midia/<id>/` lê o arquivo da `Midia`; `static:apresentacoes/...` lê o estático do
    portal. Qualquer outra URL é ignorada (o saneador do servidor já não deixa passar). `permitir`
    (opcional) recebe a `Midia` e diz se quem exporta pode usá-la.
    """
    src = str(src or '').strip()
    if not src:
        return None
    combinou = _RE_MIDIA.match(src.split('?')[0].split('#')[0])
    if combinou:
        from .models import Midia
        midia = Midia.objects.filter(pk=int(combinou.group(1))).first()
        if midia is None or not midia.arquivo:
            return None
        if permitir is not None and not permitir(midia):
            return None
        try:
            with midia.arquivo.open('rb') as arquivo:
                dados = arquivo.read()
        except Exception:                               # noqa: BLE001 — storage fora do ar ou arquivo sumido
            logger.warning('Exportação PPTX: não foi possível ler a mídia %s', midia.pk, exc_info=True)
            return None
        mime = midia.mime or mimetypes.guess_type(midia.arquivo.name or '')[0] or ''
        return dados, mime
    if src.startswith('static:'):
        caminho = src[len('static:'):].lstrip('/')
        partes = caminho.split('/')
        if not caminho.startswith('apresentacoes/') or '..' in partes or '\\' in caminho or '\x00' in caminho:
            return None
        from django.contrib.staticfiles import finders
        achado = finders.find(caminho)
        try:
            if achado:
                dados = Path(achado).read_bytes()
            else:
                from django.contrib.staticfiles.storage import staticfiles_storage
                with staticfiles_storage.open(caminho, 'rb') as arquivo:
                    dados = arquivo.read()
        except Exception:                               # noqa: BLE001
            logger.warning('Exportação PPTX: estático indisponível: %s', caminho)
            return None
        return dados, mimetypes.guess_type(caminho)[0] or ''
    return None


def _abrir_imagem(dados):
    try:
        imagem = Image.open(io.BytesIO(dados))
        imagem.size                                     # noqa: B018 — força ler o cabeçalho
        return imagem
    except Exception:                                   # noqa: BLE001 — bytes que não são imagem
        return None


def _preparar_imagem(dados, desfocar_raio=0.0, max_lado=4096):
    """Deixa a imagem num formato que o PowerPoint lê (PNG/JPEG/GIF), já girada pelo EXIF.

    Retorna (bytes, largura, altura) ou None.
    """
    imagem = _abrir_imagem(dados)
    if imagem is None:
        return None
    formato = (imagem.format or '').upper()
    recodificar = formato not in ('PNG', 'JPEG', 'GIF')
    try:
        orientacao = imagem.getexif().get(0x0112, 1)
    except Exception:                                   # noqa: BLE001
        orientacao = 1
    try:
        if orientacao not in (None, 1):
            imagem = ImageOps.exif_transpose(imagem)     # o navegador gira pelo EXIF, o PowerPoint não
            recodificar = True
        if desfocar_raio > 0.3:
            imagem = imagem.convert('RGBA').filter(ImageFilter.GaussianBlur(desfocar_raio))
            recodificar = True
        if max(imagem.size) > max_lado:
            imagem = imagem.copy()
            imagem.thumbnail((max_lado, max_lado))
            recodificar = True
        if not recodificar:
            return dados, imagem.size[0], imagem.size[1]
        saida = io.BytesIO()
        tem_alfa = imagem.mode in ('RGBA', 'LA', 'PA') or (imagem.mode == 'P' and 'transparency' in imagem.info)
        if formato == 'JPEG' and not tem_alfa:
            imagem.convert('RGB').save(saida, 'JPEG', quality=90)
        else:
            if imagem.mode not in ('RGB', 'RGBA', 'L', 'LA'):
                imagem = imagem.convert('RGBA')
            imagem.save(saida, 'PNG')
        return saida.getvalue(), imagem.size[0], imagem.size[1]
    except Exception:                                   # noqa: BLE001 — imagem corrompida no meio
        logger.warning('Exportação PPTX: imagem não pôde ser convertida', exc_info=True)
        return None


def _caixas_mp4(dados, inicio, fim):
    posicao = inicio
    while posicao + 8 <= fim:
        tamanho, tipo = struct.unpack('>I4s', dados[posicao:posicao + 8])
        cabecalho = 8
        if tamanho == 1:
            if posicao + 16 > fim:
                return
            tamanho = struct.unpack('>Q', dados[posicao + 8:posicao + 16])[0]
            cabecalho = 16
        elif tamanho == 0:
            tamanho = fim - posicao
        if tamanho < cabecalho or posicao + tamanho > fim:
            return
        yield tipo, posicao + cabecalho, posicao + tamanho
        posicao += tamanho


def info_video(dados):
    """MP4/MOV → {'duracao_ms', 'largura', 'altura'} lendo moov/mvhd/tkhd (sem ffmpeg). None se não der."""
    try:
        info = {}
        for tipo, inicio, fim in _caixas_mp4(dados, 0, len(dados)):
            if tipo != b'moov':
                continue
            for tipo2, inicio2, fim2 in _caixas_mp4(dados, inicio, fim):
                if tipo2 == b'mvhd':
                    if dados[inicio2] == 1:
                        escala, duracao = struct.unpack('>IQ', dados[inicio2 + 20:inicio2 + 32])
                    else:
                        escala, duracao = struct.unpack('>II', dados[inicio2 + 12:inicio2 + 20])
                    if escala:
                        info['duracao_ms'] = int(duracao * 1000 / escala)
                elif tipo2 == b'trak' and 'largura' not in info:
                    for tipo3, inicio3, _fim3 in _caixas_mp4(dados, inicio2, fim2):
                        if tipo3 != b'tkhd':
                            continue
                        deslocamento = 88 if dados[inicio3] == 1 else 76
                        matriz = inicio3 + (52 if dados[inicio3] == 1 else 40)
                        largura, altura = struct.unpack('>II', dados[inicio3 + deslocamento:inicio3 + deslocamento + 8])
                        a, b = struct.unpack('>ii', dados[matriz:matriz + 8])
                        largura, altura = largura / 65536, altura / 65536
                        if largura > 0 and altura > 0:
                            if b != 0 and a == 0:           # vídeo de celular gravado "em pé"
                                largura, altura = altura, largura
                            info['largura'], info['altura'] = largura, altura
            break
        return info or None
    except Exception:                                   # noqa: BLE001 — arquivo que não é MP4
        return None


def _mime_video(dados, mime):
    mime = str(mime or '').split(';')[0].strip().lower()
    if mime.startswith('video/'):
        return mime
    if dados[4:8] == b'ftyp':
        return 'video/quicktime' if dados[8:12] == b'qt  ' else 'video/mp4'
    if dados[:4] == b'\x1a\x45\xdf\xa3':
        return 'video/webm'
    return 'video/mp4'


def _encaixe(largura_img, altura_img, largura, altura, ajuste, foco_x=0.5, foco_y=0.5):
    """Cortes (l, t, r, b em fração) e sub-caixa (dx, dy, w, h) para cover/contain como no CSS."""
    if largura_img <= 0 or altura_img <= 0 or largura <= 0 or altura <= 0 or ajuste == 'fill':
        return (0.0, 0.0, 0.0, 0.0), (0.0, 0.0, largura, altura)
    proporcao_img, proporcao_caixa = largura_img / altura_img, largura / altura
    if ajuste == 'contain':
        if proporcao_img > proporcao_caixa:
            nova_l, nova_a = largura, largura / proporcao_img
        else:
            nova_l, nova_a = altura * proporcao_img, altura
        return (0.0, 0.0, 0.0, 0.0), ((largura - nova_l) * foco_x, (altura - nova_a) * foco_y, nova_l, nova_a)
    if proporcao_img > proporcao_caixa:                 # cover: sobra na horizontal
        visivel = proporcao_caixa / proporcao_img
        return ((1 - visivel) * foco_x, 0.0, (1 - visivel) * (1 - foco_x), 0.0), (0.0, 0.0, largura, altura)
    visivel = proporcao_img / proporcao_caixa
    return (0.0, (1 - visivel) * foco_y, 0.0, (1 - visivel) * (1 - foco_y)), (0.0, 0.0, largura, altura)


# ----------------------------------------------------------------------------------------------
# ícones (Font Awesome 6.0.0 free → PNG)
# ----------------------------------------------------------------------------------------------

_ESTILOS_CLASSE = {'fa-solid': 's', 'fas': 's', 'fa-regular': 'r', 'far': 'r', 'fa-brands': 'b', 'fab': 'b',
                   'fa-light': 's', 'fal': 's', 'fa-thin': 's', 'fat': 's', 'fa-duotone': 's', 'fad': 's'}


@lru_cache(maxsize=1)
def _mapa_icones():
    try:
        return json.loads(MAPA_ICONES.read_text(encoding='utf-8')).get('icones') or {}
    except Exception:                                   # noqa: BLE001
        logger.warning('Exportação PPTX: mapa de ícones indisponível em %s', MAPA_ICONES, exc_info=True)
        return {}


@lru_cache(maxsize=64)
def _fonte_icone(estilo, tamanho):
    return ImageFont.truetype(str(RECURSOS / FONTES_ICONE[estilo]), tamanho)


def glifo_do_icone(classe):
    """'fa-solid fa-circle-check' → (estilo 's'|'r'|'b', caractere) ou None se o ícone não existe."""
    mapa = _mapa_icones()
    estilo, nome = None, None
    for token in str(classe or '').lower().split():
        if token in _ESTILOS_CLASSE:
            estilo = _ESTILOS_CLASSE[token]
        elif token == 'fa':
            estilo = estilo or 's'
        elif token.startswith('fa-') and nome is None and token[3:] in mapa:
            nome = token[3:]
    if not nome:
        return None
    codigo, estilos = mapa[nome]
    if estilo not in estilos:                            # ex.: fa-regular de ícone que só existe sólido
        estilo = next(e for e in 'srb' if e in estilos)
    return estilo, chr(int(codigo, 16))


def _mascara_arredondada(largura, altura, raio):
    fator = 4 if largura * altura <= 300_000 else 2 if largura * altura <= 1_500_000 else 1
    grande = Image.new('L', (largura * fator, altura * fator), 0)
    raio = _limitar(raio * fator, 0, min(largura, altura) * fator / 2)
    ImageDraw.Draw(grande).rounded_rectangle(
        [0, 0, largura * fator - 1, altura * fator - 1], radius=int(round(raio)), fill=255)
    return grande.resize((largura, altura), Image.LANCZOS) if fator > 1 else grande


def _pintar_camada(base, mascara, cor, alfa):
    r, g, b = (int(cor[i:i + 2], 16) for i in (0, 2, 4))
    camada = Image.new('RGBA', base.size, (r, g, b, 0))
    fator = _limitar(alfa, 0, 1)
    camada.putalpha(mascara.point(lambda v: int(round(v * fator))))
    return Image.alpha_composite(base, camada)


def png_do_icone(classe, largura_px, altura_px, cor, fundo=None, raio_px=0.0, opacidade=1.0, escala=2.0):
    """Desenha o ícone num PNG transparente (2× o tamanho da caixa). Retorna bytes ou None."""
    glifo = glifo_do_icone(classe)
    if glifo is None:
        return None
    estilo, caractere = glifo
    largura = int(_limitar(round(largura_px * escala), 8, 2048))
    altura = int(_limitar(round(altura_px * escala), 8, 2048))
    fator = largura / max(largura_px, 1)
    imagem = Image.new('RGBA', (largura, altura), (0, 0, 0, 0))
    if fundo:
        imagem = _pintar_camada(imagem, _mascara_arredondada(largura, altura, raio_px * fator), fundo[0],
                                fundo[1] * opacidade)
    lado = min(largura, altura)
    tamanho = max(4, int(round(lado * (ICONE_ESCALA_COM_FUNDO if fundo else ICONE_ESCALA_SEM_FUNDO))))
    fonte = _fonte_icone(estilo, tamanho)
    esquerda, topo, direita, base = fonte.getbbox(caractere)
    tinta = max(direita - esquerda, base - topo, 1)
    limite = lado * (0.8 if fundo else 0.98)
    if tinta > limite:                                   # glifo largo (ex.: fa-arrows-left-right)
        tamanho = max(4, int(tamanho * limite / tinta))
        fonte = _fonte_icone(estilo, tamanho)
        esquerda, topo, direita, base = fonte.getbbox(caractere)
    mascara = Image.new('L', (largura, altura), 0)
    ImageDraw.Draw(mascara).text(
        (largura / 2 - (esquerda + direita) / 2, altura / 2 - (topo + base) / 2), caractere, font=fonte, fill=255)
    cor = cor or ('FFFFFF', 1.0)
    imagem = _pintar_camada(imagem, mascara, cor[0], cor[1] * opacidade)
    saida = io.BytesIO()
    imagem.save(saida, 'PNG')
    return saida.getvalue()


def _poster_padrao(largura, altura):
    """Quadro escuro com o botão de play, na proporção do vídeo (ou da caixa)."""
    base = 1280
    proporcao = (largura / altura) if largura > 0 and altura > 0 else 16 / 9
    if proporcao >= 1:
        tamanho = (base, int(_limitar(round(base / proporcao), 16, 4096)))
    else:
        tamanho = (int(_limitar(round(base * proporcao), 16, 4096)), base)
    imagem = Image.new('RGBA', tamanho, (17, 17, 24, 255))
    lado = int(min(tamanho) * 0.3)
    icone = png_do_icone('fa-solid fa-circle-play', lado, lado, ('FFFFFF', 0.9), escala=1.0)
    if icone:
        figura = Image.open(io.BytesIO(icone))
        imagem.alpha_composite(figura, ((tamanho[0] - figura.size[0]) // 2, (tamanho[1] - figura.size[1]) // 2))
    else:
        cx, cy, r = tamanho[0] / 2, tamanho[1] / 2, lado / 2
        ImageDraw.Draw(imagem).polygon([(cx - r * 0.4, cy - r * 0.6), (cx - r * 0.4, cy + r * 0.6),
                                        (cx + r * 0.6, cy)], fill=(255, 255, 255, 230))
    saida = io.BytesIO()
    imagem.convert('RGB').save(saida, 'PNG')
    return saida.getvalue()


# ----------------------------------------------------------------------------------------------
# DrawingML: preenchimentos, linhas, efeitos
# ----------------------------------------------------------------------------------------------

def _srgb(pai, cor, opacidade=1.0):
    elemento = _sub(pai, 'a:srgbClr', val=cor[0])
    alfa = _limitar(cor[1] * opacidade, 0, 1)
    if alfa < 0.9995:
        _sub(elemento, 'a:alpha', val=int(round(alfa * 100000)))
    return elemento


def _solido(pai, cor, opacidade=1.0):
    preenchimento = _sub(pai, 'a:solidFill')
    _srgb(preenchimento, cor, opacidade)
    return preenchimento


def _gradiente_xml(pai, gradiente, opacidade=1.0):
    preenchimento = _sub(pai, 'a:gradFill', rotWithShape='1')
    lista = _sub(preenchimento, 'a:gsLst')
    for posicao, cor in gradiente['paradas']:
        parada = _sub(lista, 'a:gs', pos=int(round(posicao * 100000)))
        _srgb(parada, cor, opacidade)
    if gradiente['tipo'] == 'linear':
        # CSS: 0deg aponta para cima e gira no horário; DrawingML: 0 = da esquerda para a direita.
        angulo = (gradiente['angulo'] - 90) % 360
        _sub(preenchimento, 'a:lin', ang=int(round(angulo * 60000)) % 21600000, scaled='0')
    else:
        caminho = _sub(preenchimento, 'a:path', path='circle')
        cx, cy = gradiente['centro']
        _sub(caminho, 'a:fillToRect', l=int(round(cx * 100000)), t=int(round(cy * 100000)),
             r=int(round((1 - cx) * 100000)), b=int(round((1 - cy) * 100000)))
    return preenchimento


def _efeitos_xml(pai, sombra, cores_tema, opacidade=1.0, escala=EMU_POR_PX):
    """"suave" | "forte" → outerShdw; "brilho" → glow na cor primária. Retorna o a:effectLst ou None."""
    sombra = str(sombra or '').strip().lower()
    if sombra in SOMBRAS:
        conf = SOMBRAS[sombra]
        efeitos = _sub(pai, 'a:effectLst')
        externa = _sub(efeitos, 'a:outerShdw', blurRad=int(conf['desfoque'] * escala),
                       dist=int(conf['distancia'] * escala), dir=5400000, algn='ctr', rotWithShape='0')
        _srgb(externa, ('000000', conf['alfa']), opacidade)
        return efeitos
    if sombra == 'brilho':
        cor = resolver_cor('tema:primaria', cores_tema) or ('E23FCF', 1.0)
        efeitos = _sub(pai, 'a:effectLst')
        brilho = _sub(efeitos, 'a:glow', rad=int(BRILHO['raio'] * escala))
        _srgb(brilho, (cor[0], BRILHO['alfa']), opacidade)
        return efeitos
    return None


def _linha_xml(pai, cor, largura_emu, estilo='solid', opacidade=1.0, tag='a:ln'):
    linha = _sub(pai, tag, w=int(_limitar(largura_emu, 0, 20116800)))
    if cor is None or largura_emu <= 0:
        _sub(linha, 'a:noFill')
        return linha
    _solido(linha, cor, opacidade)
    traco = {'dashed': 'dash', 'dotted': 'sysDot', 'tracejado': 'dash', 'pontilhado': 'sysDot'}.get(
        str(estilo or '').lower())
    if traco:
        _sub(linha, 'a:prstDash', val=traco)
    return linha


def _limpar_spPr(spPr):
    """Mantém só xfrm e geometria: o resto (preenchimento, linha, efeitos) é escrito na ordem do schema."""
    for filho in list(spPr):
        if filho.tag not in (qn('a:xfrm'), qn('a:prstGeom'), qn('a:custGeom')):
            spPr.remove(filho)


def _sem_estilo_do_tema(elemento):
    # p:style faria a forma herdar preenchimento, linha e sombra do tema do PowerPoint.
    for estilo in elemento.findall(qn('p:style')):
        elemento.remove(estilo)


def _geometria_arredondada(spPr, raio_px, largura_px, altura_px, prst='roundRect'):
    geometria = spPr.find(qn('a:prstGeom'))
    if geometria is None:
        return
    geometria.set('prst', prst)
    lista = geometria.find(qn('a:avLst'))
    if lista is None:
        lista = _sub(geometria, 'a:avLst')
    for filho in list(lista):
        lista.remove(filho)
    menor = max(min(largura_px, altura_px), 1)
    ajuste = int(round(_limitar(raio_px / menor, 0, 0.5) * 100000))
    _sub(lista, 'a:gd', name='adj', fmla=f'val {ajuste}')


# ----------------------------------------------------------------------------------------------
# transições e animações
# ----------------------------------------------------------------------------------------------

def _velocidade(duracao):
    return 'fast' if duracao < 625 else 'med' if duracao < 875 else 'slow'


def xml_transicao(transicao):
    """p:transition com p14:dur dentro de mc:AlternateContent (como o PowerPoint 2010+ grava)."""
    transicao = _dict(transicao)
    tipo = str(transicao.get('tipo') or 'nenhuma').strip().lower()
    if tipo not in TRANSICOES:
        return None
    duracao = int(_limitar(_num(transicao.get('duracao'), 700), 50, 60000))
    tag, atributos = TRANSICOES[tipo]
    efeito = f'<p:{tag}' + ''.join(f' {k}="{v}"' for k, v in atributos.items()) + '/>'
    velocidade = _velocidade(duracao)
    xml = (
        f'<mc:AlternateContent xmlns:mc="{NS_MC}" xmlns:p="{NS_P}">'
        f'<mc:Choice xmlns:p14="{NS_P14}" Requires="p14">'
        f'<p:transition spd="{velocidade}" p14:dur="{duracao}">{efeito}</p:transition>'
        f'</mc:Choice>'
        f'<mc:Fallback><p:transition spd="{velocidade}">{efeito}</p:transition></mc:Fallback>'
        f'</mc:AlternateContent>'
    )
    return etree.fromstring(xml)


class _Ids:
    """ids de p:cTn: únicos no slide e crescentes na ordem do documento."""

    def __init__(self):
        self.ultimo = 0

    def __call__(self):
        self.ultimo += 1
        return self.ultimo


def _alvo(pai, spid):
    alvo = _sub(pai, 'p:tgtEl')
    _sub(alvo, 'p:spTgt', spid=spid)


def _visivel(filhos, spid, ids):
    conjunto = _sub(filhos, 'p:set')
    comportamento = _sub(conjunto, 'p:cBhvr')
    no = _sub(comportamento, 'p:cTn', id=ids(), dur=1, fill='hold')
    condicoes = _sub(no, 'p:stCondLst')
    _sub(condicoes, 'p:cond', delay=0)
    _alvo(comportamento, spid)
    nomes = _sub(comportamento, 'p:attrNameLst')
    _sub(nomes, 'p:attrName').text = 'style.visibility'
    destino = _sub(conjunto, 'p:to')
    _sub(destino, 'p:strVal', val='visible')


def _efeito_filtro(filhos, spid, duracao, filtro, ids):
    efeito = _sub(filhos, 'p:animEffect', transition='in', filter=filtro)
    comportamento = _sub(efeito, 'p:cBhvr')
    _sub(comportamento, 'p:cTn', id=ids(), dur=duracao)
    _alvo(comportamento, spid)


def _animar(filhos, spid, duracao, atributo, de, ate, ids, aditivo=None):
    animacao = _sub(filhos, 'p:anim', calcmode='lin', valueType='num')
    comportamento = _sub(animacao, 'p:cBhvr', additive=aditivo)
    _sub(comportamento, 'p:cTn', id=ids(), dur=duracao, fill='hold')
    _alvo(comportamento, spid)
    nomes = _sub(comportamento, 'p:attrNameLst')
    _sub(nomes, 'p:attrName').text = atributo
    valores = _sub(animacao, 'p:tavLst')
    for tempo, valor in ((0, de), (100000, ate)):
        quadro = _sub(valores, 'p:tav', tm=tempo)
        variante = _sub(quadro, 'p:val')
        if isinstance(valor, (int, float)):
            _sub(variante, 'p:fltVal', val=valor)
        else:
            _sub(variante, 'p:strVal', val=valor)


def _escrever_efeito(par, efeito, tipo_no, ids):
    spid = efeito['spid']
    if efeito['tipo'] == 'midia':
        no = _sub(par, 'p:cTn', id=ids(), presetID=1, presetClass='mediacall', presetSubtype=0, fill='hold',
                  nodeType=tipo_no)
        condicoes = _sub(no, 'p:stCondLst')
        _sub(condicoes, 'p:cond', delay=efeito['atraso'])
        filhos = _sub(no, 'p:childTnLst')
        comando = _sub(filhos, 'p:cmd', type='call', cmd='playFrom(0.0)')
        comportamento = _sub(comando, 'p:cBhvr')
        _sub(comportamento, 'p:cTn', id=ids(), dur=efeito['duracao'], fill='hold')
        _alvo(comportamento, spid)
        return
    animacao = efeito['animacao']
    duracao = efeito['duracao']
    preset, subtipo = PRESETS_ENTRADA[animacao]
    atributos = {'presetID': preset, 'presetClass': 'entr', 'presetSubtype': subtipo, 'fill': 'hold'}
    if efeito.get('grupo'):
        atributos['grpId'] = 0
    atributos['nodeType'] = tipo_no
    no = _sub(par, 'p:cTn', id=ids(), **atributos)
    condicoes = _sub(no, 'p:stCondLst')
    _sub(condicoes, 'p:cond', delay=efeito['atraso'])
    filhos = _sub(no, 'p:childTnLst')
    _visivel(filhos, spid, ids)
    if animacao == 'fade':
        _efeito_filtro(filhos, spid, duracao, 'fade', ids)
    elif animacao in ('subir', 'descer'):
        _efeito_filtro(filhos, spid, duracao, 'fade', ids)
        _animar(filhos, spid, duracao, 'ppt_x', '#ppt_x', '#ppt_x', ids)
        _animar(filhos, spid, duracao, 'ppt_y', '#ppt_y+.1' if animacao == 'subir' else '#ppt_y-.1', '#ppt_y', ids)
    elif animacao in ('entrar-esquerda', 'entrar-direita'):
        inicio = '0-#ppt_w/2' if animacao == 'entrar-esquerda' else '1+#ppt_w/2'
        _animar(filhos, spid, duracao, 'ppt_x', inicio, '#ppt_x', ids, aditivo='base')
        _animar(filhos, spid, duracao, 'ppt_y', '#ppt_y', '#ppt_y', ids, aditivo='base')
    elif animacao == 'zoom':
        _animar(filhos, spid, duracao, 'ppt_w', 0, '#ppt_w', ids)
        _animar(filhos, spid, duracao, 'ppt_h', 0, '#ppt_h', ids)
        _efeito_filtro(filhos, spid, duracao, 'fade', ids)
    elif animacao == 'revelar':
        _efeito_filtro(filhos, spid, duracao, 'wipe(right)', ids)


class SequenciaAnimacoes:
    """Organiza as entradas do slide em grupos de clique, como a sequência principal do PowerPoint.

    - `clique`: abre um grupo novo (clickEffect) que espera o avanço;
    - `junto`: entra no mesmo passo do anterior (withEffect);
    - `auto`: começa quando o passo anterior termina (afterEffect). Se for o primeiro do slide,
      o grupo inicia sozinho depois da transição (cond onBegin da mainSeq).
    """

    def __init__(self):
        self.grupos = []
        self.midias_iniciais = []
        self.construcoes = {}                            # spid → 'sp' | 'grafico' (p:bldLst)
        self.videos = []

    def _grupo(self, inicia_sozinho):
        grupo = {'inicia_sozinho': inicia_sozinho, 'etapas': [{'inicio': 0, 'fim': 0, 'efeitos': []}]}
        self.grupos.append(grupo)
        return grupo

    def adicionar(self, efeito, gatilho):
        gatilho = gatilho if gatilho in ('clique', 'junto', 'auto') else 'auto'
        if gatilho == 'clique':
            etapa = self._grupo(False)['etapas'][-1]
            tipo_no = 'clickEffect'
        elif not self.grupos:
            etapa = self._grupo(True)['etapas'][-1]
            tipo_no = 'withEffect' if gatilho == 'junto' else 'afterEffect'
        elif gatilho == 'junto':
            etapa = self.grupos[-1]['etapas'][-1]
            tipo_no = 'withEffect'
        else:
            ultima = self.grupos[-1]['etapas'][-1]
            if ultima['efeitos']:
                etapa = {'inicio': ultima['inicio'] + ultima['fim'], 'fim': 0, 'efeitos': []}
                self.grupos[-1]['etapas'].append(etapa)
            else:
                etapa = ultima
            tipo_no = 'afterEffect'
        etapa['efeitos'].append((tipo_no, efeito))
        if efeito['tipo'] != 'midia':                    # tocar o vídeo não segura a fila de animações
            etapa['fim'] = max(etapa['fim'], efeito['atraso'] + efeito['duracao_fila'])

    def adicionar_midia_inicial(self, efeito):
        self.midias_iniciais.append(efeito)

    def _encaixar_midias_iniciais(self):
        if not self.midias_iniciais:
            return
        if self.grupos and self.grupos[0]['inicia_sozinho']:
            etapa = self.grupos[0]['etapas'][0]
            for efeito in self.midias_iniciais:
                etapa['efeitos'].append(('withEffect', efeito))
        else:
            grupo = {'inicia_sozinho': True, 'etapas': [{'inicio': 0, 'fim': 0, 'efeitos': []}]}
            for indice, efeito in enumerate(self.midias_iniciais):
                grupo['etapas'][0]['efeitos'].append(('afterEffect' if indice == 0 else 'withEffect', efeito))
            self.grupos.insert(0, grupo)
        self.midias_iniciais = []

    def xml(self):
        """p:timing completo (ou None quando o slide não tem animação nem vídeo)."""
        self._encaixar_midias_iniciais()
        if not self.grupos and not self.videos:
            return None
        ids = _Ids()
        timing = etree.Element(qn('p:timing'), nsmap={'p': NS_P})
        lista = _sub(timing, 'p:tnLst')
        par_raiz = _sub(lista, 'p:par')
        raiz = _sub(par_raiz, 'p:cTn', id=ids(), dur='indefinite', restart='never', nodeType='tmRoot')
        filhos_raiz = _sub(raiz, 'p:childTnLst')
        if self.grupos:
            sequencia = _sub(filhos_raiz, 'p:seq', concurrent=1, nextAc='seek')
            principal_id = ids()
            principal = _sub(sequencia, 'p:cTn', id=principal_id, dur='indefinite', nodeType='mainSeq')
            filhos_principal = _sub(principal, 'p:childTnLst')
            for grupo in self.grupos:
                par_grupo = _sub(filhos_principal, 'p:par')
                no_grupo = _sub(par_grupo, 'p:cTn', id=ids(), fill='hold')
                condicoes = _sub(no_grupo, 'p:stCondLst')
                _sub(condicoes, 'p:cond', delay='indefinite')
                if grupo['inicia_sozinho']:
                    condicao = _sub(condicoes, 'p:cond', evt='onBegin', delay=0)
                    _sub(condicao, 'p:tn', val=principal_id)
                filhos_grupo = _sub(no_grupo, 'p:childTnLst')
                for etapa in grupo['etapas']:
                    par_etapa = _sub(filhos_grupo, 'p:par')
                    no_etapa = _sub(par_etapa, 'p:cTn', id=ids(), fill='hold')
                    condicoes_etapa = _sub(no_etapa, 'p:stCondLst')
                    _sub(condicoes_etapa, 'p:cond', delay=int(etapa['inicio']))
                    filhos_etapa = _sub(no_etapa, 'p:childTnLst')
                    for tipo_no, efeito in etapa['efeitos']:
                        _escrever_efeito(_sub(filhos_etapa, 'p:par'), efeito, tipo_no, ids)
            anteriores = _sub(sequencia, 'p:prevCondLst')
            condicao = _sub(anteriores, 'p:cond', evt='onPrev', delay=0)
            _sub(_sub(condicao, 'p:tgtEl'), 'p:sldTgt')
            proximos = _sub(sequencia, 'p:nextCondLst')
            condicao = _sub(proximos, 'p:cond', evt='onNext', delay=0)
            _sub(_sub(condicao, 'p:tgtEl'), 'p:sldTgt')
        for video in self.videos:
            no_video = _sub(filhos_raiz, 'p:video')
            midia = _sub(no_video, 'p:cMediaNode', vol=80000, mute='1' if video.get('mudo') else None)
            no = _sub(midia, 'p:cTn', id=ids(), repeatCount='indefinite' if video.get('loop') else None,
                      fill='hold', display=0)
            condicoes = _sub(no, 'p:stCondLst')
            _sub(condicoes, 'p:cond', delay='indefinite')
            _alvo(midia, video['spid'])
        if self.construcoes:
            construcoes = _sub(timing, 'p:bldLst')
            for spid, tipo in self.construcoes.items():
                if tipo == 'sp':
                    _sub(construcoes, 'p:bldP', spid=spid, grpId=0, animBg=1)
                else:
                    grafico = _sub(construcoes, 'p:bldGraphic', spid=spid, grpId=0)
                    _sub(grafico, 'p:bldAsOne')
        return timing


# ----------------------------------------------------------------------------------------------
# exportador
# ----------------------------------------------------------------------------------------------

class _Exportador:
    def __init__(self, documento, titulo, ler_midia):
        self.documento = _dict(documento)
        self.titulo = _texto_xml(titulo).strip()
        self.ler_midia = ler_midia or ler_midia_padrao
        tema = _dict(self.documento.get('tema'))
        cores = dict(TEMA_PADRAO['cores'])
        cores.update({k: v for k, v in _dict(tema.get('cores')).items() if isinstance(v, str) and v.strip()})
        self.cores = cores
        self.fonte_titulo = str(tema.get('fonte_titulo') or TEMA_PADRAO['fonte_titulo'])
        self.fonte_texto = str(tema.get('fonte_texto') or TEMA_PADRAO['fonte_texto'])
        largura = _num(self.documento.get('largura'), QUADRO_LARGURA) or QUADRO_LARGURA
        altura = _num(self.documento.get('altura'), QUADRO_ALTURA) or QUADRO_ALTURA
        self.quadro = (largura, altura)
        self.ex = SLIDE_LARGURA_EMU / largura
        self.ey = SLIDE_ALTURA_EMU / altura
        self.pt_por_px = PT_POR_PX * self.ex / EMU_POR_PX
        self._midias = {}
        self.slides = []

    # -- conversões -------------------------------------------------------------------------
    def emu_x(self, px):
        return int(round(px * self.ex))

    def emu_y(self, px):
        return int(round(px * self.ey))

    def cor(self, valor):
        return resolver_cor(valor, self.cores)

    def midia(self, src):
        src = str(src or '').strip()
        if not src:
            return None
        if src not in self._midias:
            try:
                resultado = self.ler_midia(src)
            except Exception:                           # noqa: BLE001 — leitor externo não derruba a exportação
                logger.warning('Exportação PPTX: falha ao ler a mídia %s', src, exc_info=True)
                resultado = None
            if resultado and isinstance(resultado, (tuple, list)) and resultado[0]:
                self._midias[src] = (bytes(resultado[0]), str(resultado[1] or '') if len(resultado) > 1 else '')
            else:
                self._midias[src] = None
        return self._midias[src]

    # -- documento --------------------------------------------------------------------------
    def gerar(self):
        prs = Presentation()
        self._preparar_modelo(prs)
        layout = prs.slide_layouts[6]                    # "Blank": sem placeholders no slide
        slides_doc = [s for s in (self.documento.get('slides') or []) if isinstance(s, dict)] \
            if isinstance(self.documento.get('slides'), list) else []
        # Todos os slides existem antes do conteúdo: link "#slide-N" pode apontar para frente.
        self.slides = [prs.slides.add_slide(layout) for _ in slides_doc]
        total_visiveis = sum(1 for s in slides_doc if not s.get('oculto'))
        numero = 0
        for indice, (slide_doc, slide) in enumerate(zip(slides_doc, self.slides)):
            if not slide_doc.get('oculto'):
                numero += 1
            try:
                self._slide(slide, slide_doc, indice, max(numero, 1) if slide_doc.get('oculto') else numero,
                            total_visiveis)
            except Exception:                           # noqa: BLE001
                logger.exception('Exportação PPTX: slide %s não pôde ser montado', indice + 1)
        self._metadados(prs, len(slides_doc), len(slides_doc) - total_visiveis)
        saida = io.BytesIO()
        prs.save(saida)
        return saida.getvalue()

    def _preparar_modelo(self, prs):
        fator = SLIDE_LARGURA_EMU / prs.slide_width      # modelo do python-pptx é 4:3
        prs.slide_width = Emu(SLIDE_LARGURA_EMU)
        prs.slide_height = Emu(SLIDE_ALTURA_EMU)
        tamanho = prs.part._element.find(qn('p:sldSz'))
        if tamanho is not None:
            tamanho.attrib.pop('type', None)              # "screen4x3" deixaria o PowerPoint confuso
        # Placeholders do mestre e dos layouts acompanham a largura nova (slide novo no PowerPoint).
        for parte in [prs.slide_master, *prs.slide_layouts]:
            for xfrm in parte._element.iter(qn('a:xfrm')):
                deslocamento, extensao = xfrm.find(qn('a:off')), xfrm.find(qn('a:ext'))
                if deslocamento is not None:
                    deslocamento.set('x', str(int(round(int(deslocamento.get('x', '0')) * fator))))
                if extensao is not None:
                    extensao.set('cx', str(int(round(int(extensao.get('cx', '0')) * fator))))
        try:
            self._aplicar_tema(prs)
        except Exception:                               # noqa: BLE001 — tema é conforto, não requisito
            logger.warning('Exportação PPTX: não foi possível aplicar o tema ao mestre', exc_info=True)

    def _aplicar_tema(self, prs):
        parte = prs.slide_master.part.part_related_by(RT.THEME)
        raiz = etree.fromstring(parte.blob)
        ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
        for caminho, familia in (('a:themeElements/a:fontScheme/a:majorFont/a:latin', self.fonte_titulo),
                                 ('a:themeElements/a:fontScheme/a:minorFont/a:latin', self.fonte_texto)):
            elemento = raiz.find(caminho, ns)
            if elemento is not None and familia:
                elemento.set('typeface', _texto_xml(familia)[:120])
        esquema = raiz.find('a:themeElements/a:clrScheme', ns)
        if esquema is not None:
            for chave, nome in (('accent1', 'primaria'), ('accent2', 'secundaria'), ('accent3', 'destaque'),
                                ('hlink', 'destaque'), ('folHlink', 'secundaria')):
                cor = self.cor(self.cores.get(nome))
                alvo = esquema.find(f'a:{chave}', ns)
                if cor is None or alvo is None:
                    continue
                for filho in list(alvo):
                    alvo.remove(filho)
                etree.SubElement(alvo, qn('a:srgbClr')).set('val', cor[0])
        parte.blob = etree.tostring(raiz, xml_declaration=True, encoding='UTF-8', standalone=True)

    def _metadados(self, prs, total, ocultos):
        propriedades = prs.core_properties
        agora = datetime.now(dt_timezone.utc).replace(tzinfo=None, microsecond=0)
        propriedades.title = (self.titulo or 'Apresentação')[:255]
        propriedades.author = AUTOR
        propriedades.last_modified_by = AUTOR
        propriedades.comments = 'Exportado pelo Assistente de Apresentações'
        propriedades.subject = ''
        propriedades.keywords = ''
        propriedades.category = ''
        propriedades.revision = 1
        propriedades.created = agora
        propriedades.modified = agora
        try:
            app = prs.part.package.part_related_by(RT.EXTENDED_PROPERTIES)
            texto = app.blob.decode('utf-8')
            texto = re.sub(r'<PresentationFormat>[^<]*</PresentationFormat>',
                           '<PresentationFormat>Widescreen</PresentationFormat>', texto)
            texto = re.sub(r'<Slides>\d+</Slides>', f'<Slides>{total}</Slides>', texto)
            texto = re.sub(r'<HiddenSlides>\d+</HiddenSlides>', f'<HiddenSlides>{ocultos}</HiddenSlides>', texto)
            app.blob = texto.encode('utf-8')
        except Exception:                               # noqa: BLE001
            logger.debug('Exportação PPTX: docProps/app.xml não ajustado', exc_info=True)

    # -- slide ------------------------------------------------------------------------------
    def _slide(self, slide, slide_doc, indice, numero, total):
        sld = slide._element
        if slide_doc.get('oculto'):
            sld.set('show', '0')
        try:
            self._fundo(slide, _dict(slide_doc.get('fundo')))
        except Exception:                               # noqa: BLE001
            logger.warning('Exportação PPTX: fundo inválido no slide %s', indice + 1, exc_info=True)
        sequencia = SequenciaAnimacoes()
        arvore = slide.shapes._spTree
        elementos = slide_doc.get('elementos') if isinstance(slide_doc.get('elementos'), list) else []
        contexto = {'numero': numero, 'total': total, 'indice': indice}
        for elemento in elementos:
            if not isinstance(elemento, dict):
                continue
            if elemento.get('tipo') == 'apagar' or elemento.get('slot') == 'area_conteudo':
                continue                                  # só existem no modo template
            quantos = len(arvore)
            try:
                info = self._elemento(slide, elemento, contexto)
            except Exception:                           # noqa: BLE001
                logger.warning('Exportação PPTX: elemento %s (%s) ignorado no slide %s',
                               elemento.get('id'), elemento.get('tipo'), indice + 1, exc_info=True)
                while len(arvore) > quantos:
                    arvore.remove(arvore[-1])
                continue
            if info:
                self._registrar_animacao(sequencia, elemento, info)
        notas = _texto_xml(slide_doc.get('notas')).strip()
        if notas:
            try:
                slide.notes_slide.notes_text_frame.text = notas
            except Exception:                           # noqa: BLE001
                logger.warning('Exportação PPTX: notas do slide %s não gravadas', indice + 1, exc_info=True)
        for antigo in sld.findall(qn('p:timing')):       # o python-pptx cria um p:timing ao inserir vídeo
            sld.remove(antigo)
        ancora = sld.find(qn('p:clrMapOvr'))
        if ancora is None:
            ancora = sld.find(qn('p:cSld'))
        try:
            transicao = xml_transicao(slide_doc.get('transicao'))
        except Exception:                               # noqa: BLE001
            logger.warning('Exportação PPTX: transição inválida no slide %s', indice + 1, exc_info=True)
            transicao = None
        if transicao is not None:
            ancora.addnext(transicao)
            ancora = transicao
        try:
            timing = sequencia.xml()
        except Exception:                               # noqa: BLE001
            logger.warning('Exportação PPTX: animações do slide %s descartadas', indice + 1, exc_info=True)
            timing = None
        if timing is not None:
            ancora.addnext(timing)

    def _registrar_animacao(self, sequencia, elemento, info):
        animacao = _dict(elemento.get('animacao'))
        tipo = str(animacao.get('tipo') or 'nenhuma').strip().lower()
        gatilho = str(animacao.get('gatilho') or 'auto').strip().lower()
        atraso = int(_limitar(_num(animacao.get('atraso'), 0), 0, 600000))
        video = info.get('video')
        if video:
            sequencia.videos.append(video)
        efeito_entrada = None
        if tipo in PRESETS_ENTRADA:
            duracao = int(_limitar(_num(animacao.get('duracao'), 600), 1, 60000))
            construcao = info.get('construcao')
            efeito_entrada = {
                'tipo': 'entrada', 'animacao': tipo, 'spid': info['spid'], 'atraso': atraso,
                'duracao': duracao, 'duracao_fila': 0 if tipo == 'aparecer' else duracao,
                'grupo': bool(construcao),
            }
            sequencia.adicionar(efeito_entrada, gatilho)
            if construcao:
                sequencia.construcoes[info['spid']] = construcao
        if video and video.get('autoplay'):
            tocar = {'tipo': 'midia', 'spid': info['spid'], 'duracao': video.get('duracao_ms') or 1000,
                     'duracao_fila': 0, 'atraso': atraso if efeito_entrada else 0}
            if efeito_entrada:
                sequencia.adicionar(tocar, 'junto')       # começa a tocar quando aparece
            else:
                sequencia.adicionar_midia_inicial(tocar)  # sem entrada: toca ao abrir o slide

    # -- fundo ------------------------------------------------------------------------------
    def _fundo(self, slide, fundo):
        largura, altura = self.quadro
        cor = self.cor(fundo.get('cor'))
        cor = (cor[0], 1.0) if cor else (self.cor('tema:fundo') or ('FFFFFF', 1.0))
        gradiente = resolver_gradiente(fundo.get('gradiente'), self.cores, largura, altura)
        gradiente_opaco = gradiente and all(c[1] >= 0.999 for _, c in gradiente['paradas'])
        csld = slide._element.find(qn('p:cSld'))
        for antigo in csld.findall(qn('p:bg')):
            csld.remove(antigo)
        bg = etree.Element(qn('p:bg'))
        propriedades = _sub(bg, 'p:bgPr')
        if gradiente_opaco:
            _gradiente_xml(propriedades, gradiente)
        else:
            _solido(propriedades, cor)
        _sub(propriedades, 'a:effectLst')
        csld.insert(0, bg)
        if gradiente and not gradiente_opaco:
            # Gradiente com transparência pinta por cima da cor, como as camadas do CSS.
            forma = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_LARGURA_EMU, SLIDE_ALTURA_EMU)
            forma.name = 'Fundo (gradiente)'
            _sem_estilo_do_tema(forma._element)
            sp_pr = forma._element.spPr
            _limpar_spPr(sp_pr)
            _gradiente_xml(sp_pr, gradiente)
            _linha_xml(sp_pr, None, 0)
        imagem = fundo.get('imagem')
        if imagem:
            ajuste = str(fundo.get('ajuste') or 'cover').lower()
            elemento = {'tipo': 'imagem', 'id': '', 'src': imagem, 'x': 0, 'y': 0, 'w': largura, 'h': altura,
                        'ajuste': ajuste if ajuste in ('cover', 'contain', 'fill') else 'cover'}
            info = self._imagem(slide, elemento, nome='Fundo (imagem)', ausente_visivel=False)
            if info is None:
                logger.warning('Exportação PPTX: imagem de fundo indisponível: %s', imagem)

    # -- elementos --------------------------------------------------------------------------
    def _caixa(self, elemento):
        x = _num(elemento.get('x'))
        y = _num(elemento.get('y'))
        w = max(_num(elemento.get('w'), 1), 1)
        h = max(_num(elemento.get('h'), 1), 1)
        return x, y, w, h

    def _elemento(self, slide, elemento, contexto):
        tipo = elemento.get('tipo')
        if tipo == 'texto':
            return self._texto(slide, elemento, contexto)
        if tipo == 'imagem':
            return self._imagem(slide, elemento)
        if tipo == 'forma':
            return self._forma(slide, elemento)
        if tipo == 'icone':
            return self._icone(slide, elemento)
        if tipo == 'video':
            return self._video(slide, elemento)
        if tipo == 'tabela':
            return self._tabela(slide, elemento)
        logger.info('Exportação PPTX: tipo de elemento desconhecido ignorado: %r', tipo)
        return None

    @staticmethod
    def _opacidade(elemento):
        return _limitar(_num(elemento.get('opacidade'), 1), 0, 1)

    @staticmethod
    def _rotacionar(forma, elemento):
        rotacao = _num(elemento.get('rotacao')) % 360
        if rotacao:
            forma.rotation = rotacao

    def _nomear(self, forma, rotulo, elemento):
        identificador = _texto_xml(elemento.get('id') or '')[:40]
        forma.name = f'{rotulo} {identificador}'.strip()

    def _link(self, forma, link):
        link = str(link or '').strip()
        if not link:
            return
        salto = _RE_LINK_SLIDE.match(link)
        if salto:
            indice = int(salto.group(1)) - 1
            if 0 <= indice < len(self.slides):
                forma.click_action.target_slide = self.slides[indice]
        elif _RE_LINK_WEB.match(link) and len(link) <= 2000:
            forma.click_action.hyperlink.address = link

    # -- texto ------------------------------------------------------------------------------
    def _estilo_base(self, estilo):
        papel = str(estilo.get('papel') or 'texto').lower()
        fonte = str(estilo.get('fonte') or '').strip() or (self.fonte_titulo if papel == 'titulo' else self.fonte_texto)
        return {
            'fonte': fonte,
            'tamanho': _limitar(_num(estilo.get('tamanho'), 32) or 32, 1, 4000),
            'peso': _num(estilo.get('peso'), 400),
            'italico': bool(estilo.get('italico')),
            'sublinhado': bool(estilo.get('sublinhado')),
            'tachado': bool(estilo.get('tachado')),
            'cor': estilo.get('cor') or 'tema:texto',
            'realce': '',
            'link': '',
        }

    def _texto(self, slide, elemento, contexto):
        estilo = _dict(elemento.get('estilo'))
        opacidade = self._opacidade(elemento)
        x, y, w, h = self._caixa(elemento)
        fundo = self.cor(estilo.get('fundo'))
        borda = _dict(estilo.get('borda'))
        cor_borda = self.cor(borda.get('cor'))
        largura_borda = max(_num(borda.get('largura')), 0)
        tem_borda = cor_borda is not None and largura_borda > 0
        raio = max(_num(estilo.get('raio')), 0)
        caixa_desenhada = fundo is not None or tem_borda or raio > 0
        esquerda, topo, largura, altura = self.emu_x(x), self.emu_y(y), self.emu_x(w), self.emu_y(h)
        if caixa_desenhada:
            forma = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if raio > 0 else MSO_SHAPE.RECTANGLE,
                                           esquerda, topo, largura, altura)
            _sem_estilo_do_tema(forma._element)
        else:
            forma = slide.shapes.add_textbox(esquerda, topo, largura, altura)
        self._nomear(forma, 'Texto', elemento)
        sp_pr = forma._element.spPr
        _limpar_spPr(sp_pr)
        if raio > 0:
            _geometria_arredondada(sp_pr, raio, w, h)
        if fundo is not None:
            _solido(sp_pr, fundo, opacidade)
        else:
            _sub(sp_pr, 'a:noFill')
        if tem_borda:
            _linha_xml(sp_pr, cor_borda, self.emu_x(largura_borda), borda.get('estilo'), opacidade)
        elif caixa_desenhada:
            _linha_xml(sp_pr, None, 0)
        sombra = estilo.get('sombra')
        sombra_no_texto = not caixa_desenhada        # sem caixa, a sombra do CSS cai sobre as letras
        if not sombra_no_texto:
            _efeitos_xml(sp_pr, sombra, self.cores, opacidade, self.ex)
        self._rotacionar(forma, elemento)

        html = str(elemento.get('html') or '')
        if elemento.get('slot') == 'paginacao':
            html = html.replace('{n}', f"{contexto['numero']:02d}").replace('{total}', f"{contexto['total']:02d}")
        base = self._estilo_base(estilo)
        paragrafos = ler_html(html, base)
        corpo = forma._element.txBody
        for filho in list(corpo):
            corpo.remove(filho)
        preenchimento = max(_num(estilo.get('preenchimento')), 0) + (largura_borda if tem_borda else 0)
        margem_x, margem_y = self.emu_x(preenchimento), self.emu_y(preenchimento)
        ancora = {'middle': 'ctr', 'center': 'ctr', 'bottom': 'b'}.get(str(estilo.get('vertical') or '').lower(), 't')
        _sub(corpo, 'a:bodyPr', wrap='square', lIns=margem_x, tIns=margem_y, rIns=margem_x, bIns=margem_y,
             anchor=ancora, rtlCol=0)
        _sub(corpo.find(qn('a:bodyPr')), 'a:noAutofit')   # o editor já gravou o tamanho que coube
        _sub(corpo, 'a:lstStyle')
        alinhamento = {'center': 'ctr', 'right': 'r', 'justify': 'just'}.get(
            str(estilo.get('alinhamento') or '').lower(), 'l')
        entrelinha = _limitar(_num(estilo.get('entrelinha'), 1.2) or 1.2, 0.5, 5)
        espacamento = _num(estilo.get('espacamento'))
        extras = {
            'maiusculas': bool(estilo.get('maiusculas')),
            'espacamento': espacamento,
            'sombra': sombra if sombra_no_texto else '',
            'opacidade': opacidade,
            'slide': slide,
        }
        if not paragrafos:
            paragrafos = [{'trechos': [], 'lista': None, 'nivel': 0, 'estilo_final': base}]
        for paragrafo in paragrafos:
            self._paragrafo(corpo, paragrafo, alinhamento, entrelinha, extras)
        self._link(forma, elemento.get('link'))
        return {'spid': forma.shape_id, 'construcao': 'sp'}

    def _paragrafo(self, corpo, paragrafo, alinhamento, entrelinha, extras):
        p = _sub(corpo, 'a:p')
        trechos = paragrafo['trechos']
        estilos = [t['estilo'] for t in trechos] or [paragrafo['estilo_final']]
        maior = max(_num(e.get('tamanho'), 32) for e in estilos)
        atributos = {}
        if paragrafo['lista']:
            recuo = self.emu_x(max(maior * 1.2, 24))
            atributos.update(marL=recuo * (paragrafo['nivel'] + 1), indent=-recuo)
        else:
            atributos.update(marL=0, indent=0)
        pPr = _sub(p, 'a:pPr', algn=alinhamento, **atributos)
        linha = _sub(pPr, 'a:lnSpc')
        _sub(linha, 'a:spcPts', val=int(_limitar(round(maior * entrelinha * self.pt_por_px * 100), 0, 158400)))
        _sub(_sub(pPr, 'a:spcBef'), 'a:spcPts', val=0)
        _sub(_sub(pPr, 'a:spcAft'), 'a:spcPts', val=0)
        if paragrafo['lista'] == 'ol':
            _sub(pPr, 'a:buFont', typeface='+mj-lt')
            _sub(pPr, 'a:buAutoNum', type='arabicPeriod')
        elif paragrafo['lista']:
            _sub(pPr, 'a:buFont', typeface='Arial')
            _sub(pPr, 'a:buChar', char='•')
        else:
            _sub(pPr, 'a:buNone')
        for trecho in trechos:
            if trecho.get('quebra'):
                quebra = _sub(p, 'a:br')
                self._propriedades_trecho(quebra, 'a:rPr', trecho['estilo'], extras, com_link=False)
            else:
                run = _sub(p, 'a:r')
                self._propriedades_trecho(run, 'a:rPr', trecho['estilo'], extras)
                _sub(run, 'a:t').text = _texto_xml(trecho['texto'])
        self._propriedades_trecho(p, 'a:endParaRPr', paragrafo['estilo_final'] or estilos[-1], extras,
                                  com_link=False)

    def _propriedades_trecho(self, pai, tag, estilo, extras, com_link=True):
        tamanho = _num(estilo.get('tamanho'), 32)
        opacidade = extras['opacidade']
        atributos = {
            'lang': 'pt-BR',
            'sz': int(_limitar(round(tamanho * self.pt_por_px * 100), 100, 400000)),
            'b': 1 if _num(estilo.get('peso'), 400) >= 600 else 0,
            'i': 1 if estilo.get('italico') else 0,
        }
        if estilo.get('sublinhado'):
            atributos['u'] = 'sng'
        if estilo.get('tachado'):
            atributos['strike'] = 'sngStrike'
        if extras['maiusculas']:
            atributos['cap'] = 'all'
        if extras['espacamento']:
            atributos['spc'] = int(_limitar(round(extras['espacamento'] * self.pt_por_px * 100), -400000, 400000))
        atributos['dirty'] = 0
        rPr = _sub(pai, tag, **atributos)
        cor = self.cor(estilo.get('cor')) or self.cor('tema:texto') or ('000000', 1.0)
        _solido(rPr, cor, opacidade)
        if extras['sombra']:
            _efeitos_xml(rPr, extras['sombra'], self.cores, opacidade, self.ex)
        realce = self.cor(estilo.get('realce'))
        if realce is not None:
            destaque = _sub(rPr, 'a:highlight')
            _srgb(destaque, realce, opacidade)
        fonte = _texto_xml(estilo.get('fonte') or self.fonte_texto)[:120]
        _sub(rPr, 'a:latin', typeface=fonte)
        _sub(rPr, 'a:cs', typeface=fonte)
        link = estilo.get('link') if com_link else ''
        if link:
            self._link_trecho(rPr, link, extras['slide'])
        return rPr

    def _link_trecho(self, rPr, link, slide):
        salto = _RE_LINK_SLIDE.match(link)
        if salto:
            indice = int(salto.group(1)) - 1
            if not 0 <= indice < len(self.slides):
                return
            rid = slide.part.relate_to(self.slides[indice].part, RT.SLIDE)
            hiperlink = _sub(rPr, 'a:hlinkClick', action='ppaction://hlinksldjump')
        elif _RE_LINK_WEB.match(link) and len(link) <= 2000:
            rid = slide.part.relate_to(link, RT.HYPERLINK, is_external=True)
            hiperlink = _sub(rPr, 'a:hlinkClick')
        else:
            return
        hiperlink.set(qn('r:id'), rid)

    # -- imagem -----------------------------------------------------------------------------
    def _imagem(self, slide, elemento, nome='Imagem', ausente_visivel=True):
        opacidade = self._opacidade(elemento)
        x, y, w, h = self._caixa(elemento)
        ajuste = str(elemento.get('ajuste') or 'cover').lower()
        foco = _dict(elemento.get('foco'))
        foco_x = _limitar(_num(foco.get('x'), 0.5), 0, 1)
        foco_y = _limitar(_num(foco.get('y'), 0.5), 0, 1)
        filtro = str(elemento.get('filtro') or '').lower()
        dados = self.midia(elemento.get('src'))
        imagem = None
        if dados:
            desfoque = 0.0
            if filtro == 'desfocar':
                previa = _abrir_imagem(dados[0])
                if previa is not None:
                    escala_css = max(w / max(previa.size[0], 1), h / max(previa.size[1], 1))
                    desfoque = 6 / max(escala_css, 1e-3)     # blur de ~6 px no quadro
            imagem = _preparar_imagem(dados[0], desfocar_raio=desfoque)
        if imagem is None:
            if not ausente_visivel:
                return None
            return self._quadro_ausente(slide, elemento, 'Imagem indisponível')
        conteudo, largura_img, altura_img = imagem
        cortes, (dx, dy, nova_l, nova_a) = _encaixe(largura_img, altura_img, w, h, ajuste, foco_x, foco_y)
        rotacao = _num(elemento.get('rotacao')) % 360
        if ajuste == 'contain' and rotacao:
            # A sub-caixa gira em torno do centro da caixa original, como no CSS.
            ox, oy = dx + nova_l / 2 - w / 2, dy + nova_a / 2 - h / 2
            radianos = math.radians(rotacao)
            cx = x + w / 2 + ox * math.cos(radianos) - oy * math.sin(radianos)
            cy = y + h / 2 + ox * math.sin(radianos) + oy * math.cos(radianos)
            px, py = cx - nova_l / 2, cy - nova_a / 2
        else:
            px, py = x + dx, y + dy
        figura = slide.shapes.add_picture(io.BytesIO(conteudo), self.emu_x(px), self.emu_y(py),
                                          self.emu_x(nova_l), self.emu_y(nova_a))
        self._nomear(figura, nome, elemento)
        alt = _texto_xml(elemento.get('alt')).strip()
        if alt:
            figura._element.nvPicPr.cNvPr.set('descr', alt[:1000])
        if any(cortes):
            figura.crop_left, figura.crop_top, figura.crop_right, figura.crop_bottom = cortes
        blip = figura._element.blipFill.find(qn('a:blip'))
        if opacidade < 0.9995:
            _sub(blip, 'a:alphaModFix', amt=int(round(opacidade * 100000)))
        if filtro == 'cinza':
            _sub(blip, 'a:grayscl')
        elif filtro == 'escurecer':
            _sub(blip, 'a:lum', bright=-35000)
        sp_pr = figura._element.spPr
        _limpar_spPr(sp_pr)
        raio = max(_num(elemento.get('raio')), 0)
        if raio > 0:
            _geometria_arredondada(sp_pr, raio, nova_l, nova_a)
        borda = _dict(elemento.get('borda'))
        cor_borda = self.cor(borda.get('cor'))
        largura_borda = max(_num(borda.get('largura')), 0)
        if cor_borda is not None and largura_borda > 0:
            _linha_xml(sp_pr, cor_borda, self.emu_x(largura_borda), borda.get('estilo'), opacidade)
        _efeitos_xml(sp_pr, elemento.get('sombra'), self.cores, opacidade, self.ex)
        if rotacao:
            figura.rotation = rotacao
        self._link(figura, elemento.get('link'))
        return {'spid': figura.shape_id, 'construcao': None}

    def _quadro_ausente(self, slide, elemento, legenda):
        """Mídia que não carregou: quadro cinza com legenda, no mesmo lugar (animação continua valendo)."""
        x, y, w, h = self._caixa(elemento)
        opacidade = self._opacidade(elemento)
        forma = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, self.emu_x(x), self.emu_y(y), self.emu_x(w), self.emu_y(h))
        _sem_estilo_do_tema(forma._element)
        self._nomear(forma, 'Mídia indisponível', elemento)
        sp_pr = forma._element.spPr
        _limpar_spPr(sp_pr)
        _solido(sp_pr, ('D9D9D9', 1.0), opacidade)
        _linha_xml(sp_pr, ('9E9E9E', 1.0), self.emu_x(2), 'dashed', opacidade)
        self._rotacionar(forma, elemento)
        corpo = forma._element.txBody
        for filho in list(corpo):
            corpo.remove(filho)
        margem = self.emu_x(8)
        _sub(corpo, 'a:bodyPr', wrap='square', lIns=margem, tIns=margem, rIns=margem, bIns=margem, anchor='ctr')
        _sub(corpo.find(qn('a:bodyPr')), 'a:noAutofit')
        _sub(corpo, 'a:lstStyle')
        tamanho = _limitar(min(w, h) / 8, 14, 48)
        estilo = dict(self._estilo_base({}), tamanho=tamanho, cor='#4A4A4A')
        paragrafo = {'trechos': [{'texto': legenda, 'estilo': estilo}], 'lista': None, 'nivel': 0,
                     'estilo_final': estilo}
        self._paragrafo(corpo, paragrafo, 'ctr', 1.2, {'maiusculas': False, 'espacamento': 0, 'sombra': '',
                                                       'opacidade': opacidade, 'slide': slide})
        return {'spid': forma.shape_id, 'construcao': 'sp'}

    # -- forma ------------------------------------------------------------------------------
    def _forma(self, slide, elemento):
        opacidade = self._opacidade(elemento)
        x, y, w, h = self._caixa(elemento)
        tipo = str(elemento.get('forma') or 'retangulo').lower()
        borda = _dict(elemento.get('borda'))
        cor_borda = self.cor(borda.get('cor'))
        largura_borda = max(_num(borda.get('largura')), 0)
        preenchimento = self.cor(elemento.get('preenchimento'))
        if tipo == 'linha':
            espessura = largura_borda or 4
            cor = cor_borda or preenchimento or self.cor('tema:texto') or ('000000', 1.0)
            meio = y + h / 2
            conector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, self.emu_x(x), self.emu_y(meio),
                                                  self.emu_x(x + w), self.emu_y(meio))
            _sem_estilo_do_tema(conector._element)
            self._nomear(conector, 'Linha', elemento)
            sp_pr = conector._element.spPr
            _limpar_spPr(sp_pr)
            _linha_xml(sp_pr, cor, self.emu_x(espessura), borda.get('estilo'), opacidade)
            _efeitos_xml(sp_pr, elemento.get('sombra'), self.cores, opacidade, self.ex)
            self._rotacionar(conector, elemento)
            return {'spid': conector.shape_id, 'construcao': None}
        raio = max(_num(elemento.get('raio')), 0)
        if tipo == 'retangulo' and raio > 0:
            tipo = 'retangulo-arredondado'
        forma = slide.shapes.add_shape(FORMAS.get(tipo, MSO_SHAPE.RECTANGLE), self.emu_x(x), self.emu_y(y),
                                       self.emu_x(w), self.emu_y(h))
        _sem_estilo_do_tema(forma._element)
        self._nomear(forma, 'Forma', elemento)
        sp_pr = forma._element.spPr
        _limpar_spPr(sp_pr)
        if tipo == 'pilula':
            _geometria_arredondada(sp_pr, min(w, h) / 2, w, h)
        elif tipo == 'retangulo-arredondado':
            _geometria_arredondada(sp_pr, raio, w, h)
        gradiente = resolver_gradiente(elemento.get('gradiente'), self.cores, w, h)
        if gradiente:
            _gradiente_xml(sp_pr, gradiente, opacidade)
        elif preenchimento is not None:
            _solido(sp_pr, preenchimento, opacidade)
        else:
            _sub(sp_pr, 'a:noFill')
        if cor_borda is not None and largura_borda > 0:
            _linha_xml(sp_pr, cor_borda, self.emu_x(largura_borda), borda.get('estilo'), opacidade)
        else:
            _linha_xml(sp_pr, None, 0)
        _efeitos_xml(sp_pr, elemento.get('sombra'), self.cores, opacidade, self.ex)
        self._rotacionar(forma, elemento)
        self._link(forma, elemento.get('link'))
        return {'spid': forma.shape_id, 'construcao': 'sp'}

    # -- ícone ------------------------------------------------------------------------------
    def _icone(self, slide, elemento):
        opacidade = self._opacidade(elemento)
        x, y, w, h = self._caixa(elemento)
        cor = self.cor(elemento.get('cor')) or self.cor('tema:texto') or ('FFFFFF', 1.0)
        fundo = self.cor(elemento.get('fundo'))
        png = png_do_icone(elemento.get('icone'), w, h, cor, fundo, max(_num(elemento.get('raio')), 0), opacidade)
        if png is None:
            logger.info('Exportação PPTX: ícone desconhecido %r', elemento.get('icone'))
            return self._quadro_ausente(slide, elemento, 'Ícone indisponível')
        figura = slide.shapes.add_picture(io.BytesIO(png), self.emu_x(x), self.emu_y(y), self.emu_x(w), self.emu_y(h))
        self._nomear(figura, 'Ícone', elemento)
        figura._element.nvPicPr.cNvPr.set('descr', _texto_xml(elemento.get('icone'))[:200])
        self._rotacionar(figura, elemento)
        self._link(figura, elemento.get('link'))
        return {'spid': figura.shape_id, 'construcao': None}

    # -- vídeo ------------------------------------------------------------------------------
    def _video(self, slide, elemento):
        x, y, w, h = self._caixa(elemento)
        dados = self.midia(elemento.get('src'))
        if not dados:
            return self._quadro_ausente(slide, elemento, 'Vídeo indisponível')
        conteudo, mime = dados
        mime = _mime_video(conteudo, mime)
        info = info_video(conteudo) or {}
        largura_video, altura_video = info.get('largura') or 0, info.get('altura') or 0
        poster = None
        if elemento.get('poster'):
            dados_poster = self.midia(elemento.get('poster'))
            poster = _preparar_imagem(dados_poster[0]) if dados_poster else None
        if poster is not None and not (largura_video and altura_video):
            largura_video, altura_video = poster[1], poster[2]
        if poster is None:
            poster = (_poster_padrao(largura_video or w, altura_video or h),)
        ajuste = str(elemento.get('ajuste') or 'cover').lower()
        cortes, (dx, dy, nova_l, nova_a) = _encaixe(largura_video, altura_video, w, h, ajuste)
        filme = slide.shapes.add_movie(io.BytesIO(conteudo), self.emu_x(x + dx), self.emu_y(y + dy),
                                       self.emu_x(nova_l), self.emu_y(nova_a),
                                       poster_frame_image=io.BytesIO(poster[0]), mime_type=mime)
        self._nomear(filme, 'Vídeo', elemento)
        self._ajustar_extensao_video(slide, filme, mime)
        if any(cortes):
            preenchimento_blip = filme._element.blipFill
            recorte = preenchimento_blip.find(qn('a:srcRect'))
            if recorte is None:
                recorte = etree.Element(qn('a:srcRect'))
                preenchimento_blip.find(qn('a:blip')).addnext(recorte)
            for chave, valor in zip(('l', 't', 'r', 'b'), cortes):
                if valor:
                    recorte.set(chave, str(int(round(valor * 100000))))
        raio = max(_num(elemento.get('raio')), 0)
        if raio > 0:
            _geometria_arredondada(filme._element.spPr, raio, nova_l, nova_a)
        self._rotacionar(filme, elemento)
        return {
            'spid': filme.shape_id, 'construcao': None,
            'video': {'spid': filme.shape_id, 'autoplay': bool(elemento.get('autoplay')),
                      'loop': bool(elemento.get('loop')), 'mudo': bool(elemento.get('mudo')),
                      'duracao_ms': info.get('duracao_ms')},
        }

    @staticmethod
    def _ajustar_extensao_video(slide, filme, mime):
        """O python-pptx grava .vid para tipos que não conhece (webm): acerta a extensão da parte."""
        extensao = EXTENSOES_VIDEO.get(mime)
        if not extensao:
            return
        rid = filme._element.xpath('./p:nvPicPr/p:nvPr/a:videoFile/@r:link')
        if not rid:
            return
        parte = slide.part.related_part(rid[0])
        if parte.partname.ext.lower() != extensao:
            parte.partname = slide.part.package.next_media_partname(extensao)

    # -- tabela -----------------------------------------------------------------------------
    def _tabela(self, slide, elemento):
        linhas_doc = elemento.get('linhas')
        if not isinstance(linhas_doc, list):
            raise ValueError('tabela sem linhas')
        linhas = [[_texto_xml(c if isinstance(c, str) else '' if c is None else str(c)) for c in linha]
                  if isinstance(linha, list) else [_texto_xml(linha)] for linha in linhas_doc[:200]]
        colunas = min(max((len(linha) for linha in linhas), default=0), 50)
        if not linhas or not colunas:
            raise ValueError('tabela vazia')
        linhas = [(linha + [''] * colunas)[:colunas] for linha in linhas]
        estilo = _dict(elemento.get('estilo'))
        opacidade = self._opacidade(elemento)
        x, y, w, h = self._caixa(elemento)
        quadro = slide.shapes.add_table(len(linhas), colunas, self.emu_x(x), self.emu_y(y),
                                        self.emu_x(w), self.emu_y(h))
        self._nomear(quadro, 'Tabela', elemento)
        tabela = quadro.table
        largura_total, altura_total = self.emu_x(w), self.emu_y(h)
        for indice, coluna in enumerate(tabela.columns):
            coluna.width = Emu(largura_total // colunas + (largura_total % colunas if indice == colunas - 1 else 0))
        for indice, linha in enumerate(tabela.rows):
            linha.height = Emu(altura_total // len(linhas) + (altura_total % len(linhas) if indice == len(linhas) - 1 else 0))
        cabecalho = bool(elemento.get('cabecalho', True))
        propriedades = quadro._element.graphic.graphicData.tbl.tblPr
        propriedades.set('firstRow', '1' if cabecalho else '0')
        propriedades.set('bandRow', '0')
        identificador = propriedades.find(qn('a:tableStyleId'))
        if identificador is None:
            identificador = _sub(propriedades, 'a:tableStyleId')
        identificador.text = ESTILO_TABELA_NENHUM
        tamanho = _limitar(_num(estilo.get('tamanho'), 28) or 28, 4, 400)
        fonte = str(estilo.get('fonte') or '').strip() or self.fonte_texto
        cor_texto = estilo.get('cor') or 'tema:texto'
        cor_borda = self.cor(estilo.get('borda'))
        alinhamento = {'center': 'ctr', 'right': 'r', 'justify': 'just'}.get(
            str(estilo.get('alinhamento') or '').lower(), 'l')
        margem_x, margem_y = self.emu_x(tamanho * 0.6), self.emu_y(tamanho * 0.35)
        extras = {'maiusculas': False, 'espacamento': 0, 'sombra': '', 'opacidade': opacidade, 'slide': slide}
        for i, valores in enumerate(linhas):
            eh_cabecalho = cabecalho and i == 0
            posicao = i - (1 if cabecalho else 0)
            if eh_cabecalho:
                fundo = self.cor(estilo.get('fundo_cabecalho'))
            elif posicao % 2 == 1 and estilo.get('fundo_alternado'):
                fundo = self.cor(estilo.get('fundo_alternado'))
            else:
                fundo = self.cor(estilo.get('fundo_linhas'))
            base = {
                'fonte': fonte, 'tamanho': tamanho, 'peso': 700 if eh_cabecalho else 400, 'italico': False,
                'sublinhado': False, 'tachado': False, 'realce': '', 'link': '',
                'cor': (estilo.get('cor_cabecalho') or cor_texto) if eh_cabecalho else cor_texto,
            }
            for j, valor in enumerate(valores):
                celula = tabela.cell(i, j)
                tc = celula._tc
                corpo = tc.find(qn('a:txBody'))
                for filho in list(corpo):
                    corpo.remove(filho)
                _sub(corpo, 'a:bodyPr')
                _sub(corpo, 'a:lstStyle')
                for texto_linha in (valor.split('\n') if valor else ['']):
                    trechos = [{'texto': texto_linha, 'estilo': base}] if texto_linha else []
                    self._paragrafo(corpo, {'trechos': trechos, 'lista': None, 'nivel': 0, 'estilo_final': base},
                                    alinhamento, 1.2, extras)
                for antigo in tc.findall(qn('a:tcPr')):
                    tc.remove(antigo)
                tcPr = etree.Element(qn('a:tcPr'))
                corpo.addnext(tcPr)
                for chave, valor_margem in (('marL', margem_x), ('marR', margem_x), ('marT', margem_y),
                                            ('marB', margem_y)):
                    tcPr.set(chave, str(valor_margem))
                tcPr.set('anchor', 'ctr')
                for lado in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
                    _linha_xml(tcPr, cor_borda, self.emu_x(1) if cor_borda else 0, 'solid', opacidade, tag=lado)
                if fundo is not None:
                    _solido(tcPr, fundo, opacidade)
                else:
                    _sub(tcPr, 'a:noFill')
        self._link(quadro, elemento.get('link'))
        return {'spid': quadro.shape_id, 'construcao': 'grafico'}


# ----------------------------------------------------------------------------------------------
# API pública
# ----------------------------------------------------------------------------------------------

def gerar_pptx(documento: dict, titulo: str = '', ler_midia=None) -> bytes:
    """Gera o .pptx (bytes) do documento da apresentação.

    `ler_midia(src) -> (bytes, mime) | None` resolve imagens/vídeos; o padrão é `ler_midia_padrao`.
    Elemento com dado inválido é pulado (fica no log) e mídia ausente vira um quadro cinza.
    """
    return _Exportador(documento, titulo, ler_midia).gerar()


_RESERVADOS_WINDOWS = {'CON', 'PRN', 'AUX', 'NUL', *(f'COM{i}' for i in range(1, 10)), *(f'LPT{i}' for i in range(1, 10))}


def nome_do_arquivo(titulo: str) -> str:
    """Título → nome de arquivo seguro e legível ('Resultados do mês' → 'Resultados do mes.pptx')."""
    base = unicodedata.normalize('NFKD', str(titulo or '')).encode('ascii', 'ignore').decode('ascii')
    base = re.sub(r'[^A-Za-z0-9 ._-]+', ' ', base)
    base = re.sub(r'\s+', ' ', base).strip(' ._-')
    base = base[:80].rstrip(' ._-') or 'apresentacao'
    if base.upper() in _RESERVADOS_WINDOWS:
        base = f'apresentacao-{base}'
    return f'{base}.pptx'
