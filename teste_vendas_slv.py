"""Vendas (/vendas/): SLV — Documento de Requisitos v1.0.

O que este teste cobre, pelos requisitos do documento:

- [NF001/RF002] CPF em qualquer formato, dígitos verificadores;
- [RF001] PDV e vendedor do cadastro (um POST não muda); usuário sem PDV;
- [RF002/RF003] cliente novo, atualização de telefone/endereço com histórico
  (data, vendedor, PDV), "sem alteração" não grava; [RF004] pré-análise;
- [Produtos.RF002] preço por plano (Pré × com plano, aparelho fora do plano,
  eletrônico pelo SKU); [RF003] Renova Vini (código obrigatório e existente) e
  Alied; [RF004] Cartão/Pix e Vivo+ com o percentual da época; [RF005] valor
  final editável, rastreado (sugerido, calculado, quem editou), zero/negativo;
- [Serviços] linha e número fictício; Alta, Reativação (só Pós/Controle),
  Troca de plano e Migração com delta, Troca de SimCard (serial), Seguro/SVA
  (ativação confirmada), Troca de titularidade, Troca de número (cadastro
  atualizado com histórico);
- [Parametrização] planos (inválido não grava, histórico, inativo some),
  Vivo+ (0–100%, histórico), seguros/SVA; só administrador;
- importação da tabela: relatório e tudo-ou-nada; telas e filtros.

Caches em memória, transação desfeita no fim.
"""
import json
import os
import sys
from decimal import Decimal

import django

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'redeconfianca.settings')
os.environ.setdefault('RC_VARREDURA_ROTINA', '0')

from django.conf import settings

settings.CACHES = {
    'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-slv'},
    'local': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-slv-2'},
}
django.setup()

from django.test.utils import setup_test_environment

setup_test_environment()
if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS.append('testserver')

from django.contrib.auth import get_user_model
from django.db import transaction
from django.test import Client

from users.models import Sector
from vendas import slv
from vendas.models import (Cliente, ConfiguracaoVendas, ImportacaoPrecos, ItemPreco, Plano, RegistroAlteracao,
                           ServicoAdicional, Venda, VendaProduto, VendaServico)

User = get_user_model()
D = Decimal
ok = fail = 0


def t(nome, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'  OK   {nome}')
    else:
        fail += 1
        print(f'  FALHA {nome} {extra}')


def gerar_cpf(base9):
    d = [int(c) for c in base9]
    for tamanho in (9, 10):
        soma = sum(d[i] * (tamanho + 1 - i) for i in range(tamanho))
        d.append((soma * 10) % 11 % 10)
    return ''.join(map(str, d))


CPF, CPF2 = gerar_cpf('529982247'), gerar_cpf('111444777')
CPF_FORMATADO = f'{CPF[:3]}.{CPF[3:6]}.{CPF[6:9]}-{CPF[9:]}'

marcador = transaction.atomic()
marcador.__enter__()
try:
    print('== REGRAS ==')
    t('CPF com e sem pontuação', slv.cpf_valido(CPF) and slv.cpf_valido(CPF_FORMATADO) and slv.cpf_valido(f' {CPF_FORMATADO} '))
    t('CPF inválido e repetido recusados', not slv.cpf_valido(CPF[:-1] + str((int(CPF[-1]) + 1) % 10)) and not slv.cpf_valido('11111111111'))
    t('números fictícios', all(slv.numero_fake(n) for n in ('999999999', '9000000000', '27999999999', '2790000000', '(27) 90000-0000'))
      and not slv.numero_fake('27998765432'))
    t('telefone DDD + número', slv.telefone_valido('(27) 99876-5432') and slv.telefone_valido('2733334444') and not slv.telefone_valido('99876'))
    t('delta', slv.delta(D('120'), D('80')) == D('40.00') and slv.delta(D('99.90'), None) == D('99.90'))
    conta = slv.calcular_item(D('1000'), D('200'), D('50'), D('10'))
    t('item: sugerido − Renova − Vivo+ (sobre o que sobra)', conta == {'desconto_vivo_mais': D('75.00'), 'valor_calculado': D('675.00')}, conta)

    aparelho = ItemPreco.objects.create(categoria='SMARTPHONES', nome='ZZ Phone X', valor=D('1999.00'),
                                        extra={'PRÉ': '2499', 'PÓS INDIVIDUAL': '10x', 'CONTROLE ENTRADA': '-', 'ID DPGC': 'ZZPX1'})
    capa = ItemPreco.objects.create(categoria='PRODUTOS', nome='ZZ Capa', valor=D('59.90'), cod_sap='ZZCAPA')
    fone = ItemPreco.objects.create(categoria='ELETRÔNICOS_LP_Conectados', nome='ZZ Fone', valor=D('299.00'))
    t('categorias', (slv.categoria_slv(aparelho), slv.categoria_slv(capa), slv.categoria_slv(fone)) == ('aparelho', 'essencial', 'eletronico'))
    t('aparelho: cliente Pré paga o preço Pré', slv.preco_sugerido(aparelho, 'PRE')[0] == D('2499.00'))
    t('aparelho: com plano paga o preço com plano', slv.preco_sugerido(aparelho, 'POS', 'PÓS INDIVIDUAL')[0] == D('1999.00'))
    t('aparelho fora do plano: sem valor', slv.preco_sugerido(aparelho, 'CONTROLE', 'CONTROLE ENTRADA') == (None, 'não disponível no plano CONTROLE ENTRADA'))
    t('essencial pelo SKU', slv.preco_sugerido(capa, 'PRE')[0] == D('59.90'))
    b2b_fone = ItemPreco.objects.create(categoria='PRODUTOS B2B', nome='ZZ B2B Phone', valor=D('3000'), extra={'Categoria': 'Smartphone', '12': '250'})
    b2b_capa = ItemPreco.objects.create(categoria='PRODUTOS B2B', nome='ZZ B2B Capa', valor=D('99'), extra={'Categoria': 'Capa'})
    b2b_tablet = ItemPreco.objects.create(categoria='PRODUTOS B2B', nome='ZZ B2B Tab', valor=D('1999'), extra={'Categoria': 'Tablet'})
    t('B2B pela coluna Categoria: smartphone, capa, tablet', (slv.categoria_slv(b2b_fone), slv.categoria_slv(b2b_capa),
      slv.categoria_slv(b2b_tablet)) == ('aparelho', 'essencial', 'eletronico'))
    t('aparelho sem coluna Pré (B2B): preço da tabela mesmo para Pré', slv.preco_sugerido(b2b_fone, 'PRE')[0] == D('3000'))

    print('== PARAMETRIZAÇÃO ==')
    loja = Sector.objects.create(name='Loja ZZ SLV')
    admin = User.objects.create_user(username='zz.slv.adm', email='zz.slv.adm@exemplo-teste.local', password='S3nha!teste',
                                     first_name='Ana', last_name='Gestora', hierarchy='SUPERADMIN', is_superuser=True, sector=loja)
    vend = User.objects.create_user(username='zz.slv.vend', email='zz.slv.vend@exemplo-teste.local', password='S3nha!teste',
                                    first_name='Beto', last_name='Vendedor', hierarchy='PADRAO', sector=loja)
    sem_pdv = User.objects.create_user(username='zz.slv.sem', email='zz.slv.sem@exemplo-teste.local', password='S3nha!teste',
                                       first_name='Caio', last_name='SemLoja', hierarchy='PADRAO')
    ca, cv, cs = Client(), Client(), Client()
    ca.force_login(admin)
    cv.force_login(vend)
    cs.force_login(sem_pdv)

    r = cv.get('/vendas/parametros/')
    t('vendedor não entra na parametrização', r.status_code == 302)
    cv.post('/vendas/parametros/plano/', {'nome': 'ZZ hack', 'segmentacao': 'POS', 'valor': '1'})
    t('nem grava plano', not Plano.objects.filter(nome='ZZ hack').exists())
    ca.post('/vendas/parametros/plano/', {'nome': 'ZZ Pós 50GB', 'segmentacao': 'POS', 'valor': '149,90'})
    ca.post('/vendas/parametros/plano/', {'nome': 'ZZ Controle 25GB', 'segmentacao': 'CONTROLE', 'valor': '79,90'})
    ca.post('/vendas/parametros/plano/', {'nome': 'ZZ Pré Turbo', 'segmentacao': 'PRE', 'valor': '29,99'})
    ca.post('/vendas/parametros/plano/', {'nome': 'ZZ inválido', 'segmentacao': 'POS', 'valor': '0'})
    pos, ctrl, pre = (Plano.objects.get(nome=n) for n in ('ZZ Pós 50GB', 'ZZ Controle 25GB', 'ZZ Pré Turbo'))
    t('planos cadastrados; valor zero recusado', pos.valor == D('149.90') and not Plano.objects.filter(nome='ZZ inválido').exists())
    ca.post('/vendas/parametros/plano/', {'id': ctrl.id, 'nome': 'ZZ Controle 25GB', 'segmentacao': 'CONTROLE', 'valor': '89,90', 'ativo': 'on'})
    h = RegistroAlteracao.objects.filter(tipo='PLANO', objeto_id=ctrl.id, campo='valor').first()
    t('alteração de valor no histórico (antes, depois, usuário)', h and h.antes == '79.90' and h.depois == '89.90' and h.usuario == admin, h and (h.antes, h.depois))
    ca.post('/vendas/parametros/vivo-mais/', {'percentual': '150'})
    t('Vivo+ fora de 0–100% recusado', ConfiguracaoVendas.atual().vivo_mais_percentual != D('150'))
    ca.post('/vendas/parametros/vivo-mais/', {'percentual': '10'})
    t('Vivo+ 10% com histórico', ConfiguracaoVendas.atual().vivo_mais_percentual == D('10.00')
      and RegistroAlteracao.objects.filter(tipo='VIVO_MAIS', depois='10.00%').exists())
    ca.post('/vendas/parametros/adicional/', {'tipo': 'SEGURO', 'nome': 'ZZ Seguro Celular', 'valor': '19,90'})
    seguro = ServicoAdicional.objects.get(nome='ZZ Seguro Celular')
    r = ca.get('/vendas/parametros/')
    t('tela da parametrização', r.status_code == 200 and 'ZZ Pós 50GB' in r.content.decode() and 'Desconto Vivo+' in r.content.decode())

    print('== CLIENTE ==')
    r = cv.get(f'/vendas/api/cliente/?cpf={CPF_FORMATADO}')
    t('CPF novo: não encontrado', r.json() == {'ok': True, 'encontrado': False, 'cpf': CPF, 'cpf_formatado': CPF_FORMATADO}, r.json())
    r = cv.get('/vendas/api/cliente/?cpf=123.456.789-00')
    t('CPF inválido avisa', r.json()['ok'] is False)
    r = cv.post('/vendas/api/cliente/salvar/', json.dumps({'cpf': CPF, 'nome': 'Zz Cliente Teste', 'telefone': '(27) 99876-5432',
                                                           'cidade': 'Vitória', 'uf': 'es'}), content_type='application/json')
    t('cadastro do cliente novo', r.json()['ok'] and r.json()['criado'] and Cliente.objects.get(cpf=CPF).uf == 'ES')
    r = cv.post('/vendas/api/cliente/salvar/', json.dumps({'cpf': CPF, 'nome': 'Zz Cliente Teste', 'telefone': '(27) 99876-5432',
                                                           'cidade': 'Vitória', 'uf': 'ES'}), content_type='application/json')
    t('sem alteração: nada gravado', r.json()['mudou'] == 0 and not RegistroAlteracao.objects.filter(tipo='CLIENTE').exists())
    r = cv.post('/vendas/api/cliente/salvar/', json.dumps({'cpf': CPF, 'nome': 'Zz Cliente Teste', 'telefone': '(27) 3333-4444',
                                                           'cidade': 'Serra', 'uf': 'ES'}), content_type='application/json')
    hist = RegistroAlteracao.objects.filter(tipo='CLIENTE', campo='telefone').first()
    t('telefone alterado: anterior no histórico com vendedor e PDV', r.json()['mudou'] == 2 and hist and hist.antes == '27998765432'
      and hist.depois == '2733334444' and hist.usuario == vend and hist.pdv == loja.name, hist and (hist.antes, hist.depois, hist.pdv))
    r = cv.get(f'/vendas/api/cliente/?cpf={CPF}')
    j = r.json()
    t('busca devolve cliente e pré-análise vazia', j['encontrado'] and j['cliente']['nome'] == 'Zz Cliente Teste' and j['pre_analise']['vazio'])

    print('== PDV E VENDEDOR ==')
    r = cs.get('/vendas/nova/')
    t('sem PDV: bloqueia a venda', 'não tem PDV vinculado' in r.content.decode())
    base = {'cliente': {'cpf': CPF, 'nome': 'Zz Cliente Teste'}, 'linhas': ['27998765432'],
            'servicos': [{'tipo': 'ALTA', 'plano_id': pos.id}]}
    r = cs.post('/vendas/nova/', json.dumps(base), content_type='application/json')
    t('POST sem PDV recusado', not r.json()['ok'] and any('PDV' in e for e in r.json()['erros']))
    outra = Sector.objects.create(name='Loja ZZ Outra')
    r = cv.post('/vendas/nova/', json.dumps(dict(base, loja_id=outra.id, vendedor_id=admin.id)), content_type='application/json')
    v = Venda.objects.get(pk=r.json()['id'])
    t('vendedor e PDV do cadastro — o POST não muda', v.loja == loja and v.vendedor == vend and v.created_by == vend)
    t('alta: plano novo, delta = valor (sem plano anterior)', v.servicos.get().delta == D('149.90') and v.servicos.get().plano == pos)
    v.cliente.refresh_from_db()
    t('cliente passa a ter o plano da alta', v.cliente.plano == pos and v.cliente.valor_pago == D('149.90'))

    print('== PRODUTOS ==')
    def vender(dados, cliente=cv):
        r = cliente.post('/vendas/nova/', json.dumps(dados), content_type='application/json')
        return r.json()
    # O Renova de verdade exige o checklist inteiro; aqui basta um código que "existe".
    from types import SimpleNamespace
    renova = SimpleNamespace(pk=4321, codigo='RN-004321', aparelho='ZZ Usado')
    existe_original = slv.renova_existe
    slv.renova_existe = lambda codigo: renova if slv.so_digitos(codigo) and int(slv.so_digitos(codigo)) == 4321 else None
    payload = {'cliente': {'cpf': CPF, 'nome': 'Zz Cliente Teste'}, 'forma_pagamento': 'PIX', 'vivo_mais': True,
               'plano_anterior': {'id': pos.id},
               'produtos': [
                   {'preco_id': aparelho.id, 'qtde': 1, 'grupamento': 'PÓS INDIVIDUAL', 'prateleira_infinita': True,
                    'renova_vini': True, 'renova_vini_valor': '300,00', 'renova_codigo': renova.codigo if renova else 'RN-999999',
                    'renova_alied': True, 'renova_alied_valor': '100'},
                   {'preco_id': capa.id, 'qtde': 2, 'valor_final': '40,00'},
               ]}
    j = vender(payload)
    t('venda de produtos gravada', j.get('ok'), j)
    if j.get('ok'):
        v = Venda.objects.get(pk=j['id'])
        p1, p2 = v.produtos.get(preco=aparelho), v.produtos.get(preco=capa)
        # 1999 − 300 − 100 = 1599; Vivo+ 10% = 159,90 → 1439,10
        t('aparelho no preço com plano, Renova e Vivo+', p1.valor_sugerido == D('1999.00') and p1.desconto_vivo_mais == D('159.90')
          and p1.valor_venda == D('1439.10') and not p1.valor_editado, (p1.valor_sugerido, p1.desconto_vivo_mais, p1.valor_venda))
        t('prateleira infinita e código do Renova gravados', p1.prateleira_infinita and p1.renova_codigo == renova.codigo)
        t('valor final editado: sugerido, calculado e quem editou', p2.valor_editado and p2.valor_venda == D('40.00')
          and p2.valor_calculado == D('53.91') and p2.editado_por == vend, (p2.valor_calculado, p2.valor_venda))
        t('Vivo+ com o percentual da época', v.vivo_mais and v.vivo_mais_percentual == D('10.00') and v.forma_pagamento == 'PIX')
        t('total = itens finais', v.total == D('1439.10') + D('80.00'), v.total)
    ca.post('/vendas/parametros/vivo-mais/', {'percentual': '15'})
    if j.get('ok'):
        t('mudar o Vivo+ não mexe na venda já lançada', Venda.objects.get(pk=j['id']).vivo_mais_percentual == D('10.00'))

    j = vender({'cliente': {'cpf': CPF, 'nome': 'Zz Cliente Teste'}, 'plano_anterior': {'id': ctrl.id},
                'produtos': [{'preco_id': aparelho.id, 'grupamento': 'CONTROLE ENTRADA', 'renova_vini': True, 'renova_vini_valor': '10',
                              'renova_codigo': 'RN-999999999'},
                             {'preco_id': capa.id, 'valor_final': '0'}]})
    erros = ' | '.join(j.get('erros', []))
    t('erros juntos: forma de pagamento, aparelho fora do plano, Renova sem código válido, valor zerado',
      not j['ok'] and 'Cartão ou Pix' in erros and 'sem valor na tabela para este plano' in erros
      and 'Código do Renova' in erros and 'zerado ou negativo' in erros, erros)
    j = vender({'cliente': {'cpf': CPF, 'nome': 'Zz Cliente Teste'}, 'plano_anterior': {'id': ctrl.id}, 'forma_pagamento': 'CARTAO',
                'produtos': [{'preco_id': aparelho.id, 'grupamento': 'CONTROLE ENTRADA', 'valor_sugerido': '1800'}]})
    t('sem valor para o plano: aceita o valor informado', j.get('ok') and
      VendaProduto.objects.get(venda_id=j['id']).regra_preco.endswith('(valor informado)'), j)
    j = vender({'cliente': {'cpf': CPF, 'nome': 'Zz Cliente Teste'}, 'plano_anterior': {'id': pre.id}, 'forma_pagamento': 'CARTAO',
                'produtos': [{'preco_id': aparelho.id}]})
    t('cliente Pré: aparelho no preço Pré', j.get('ok') and VendaProduto.objects.get(venda_id=j['id']).valor_venda == D('2499.00'))

    print('== SERVIÇOS ==')
    j = vender({'cliente': {'cpf': CPF, 'nome': 'Zz Cliente Teste'}, 'linhas': ['27998765432'], 'plano_anterior': {'id': ctrl.id},
                'servicos': [{'tipo': 'TROCA_PLANO', 'plano_id': pos.id}, {'tipo': 'MIGRACAO', 'plano_id': pos.id, 'plano_anterior_id': pre.id}]})
    t('troca de plano e migração gravadas', j.get('ok'), j)
    if j.get('ok'):
        troca = VendaServico.objects.get(venda_id=j['id'], tipo_servico='TROCA_PLANO')
        migra = VendaServico.objects.get(venda_id=j['id'], tipo_servico='MIGRACAO')
        t('troca: anterior, novo e delta', troca.plano_anterior == ctrl and troca.valor_anterior == D('89.90') and troca.delta == D('60.00'),
          (troca.valor_anterior, troca.delta))
        t('migração: mesmo fluxo, com o plano anterior escolhido', migra.plano_anterior == pre and migra.delta == D('119.91'), migra.delta)
    j = vender({'cliente': {'cpf': CPF, 'nome': 'Zz Cliente Teste'}, 'linhas': ['27998765432'],
                'servicos': [{'tipo': 'REATIVACAO', 'plano_id': pre.id}, {'tipo': 'TROCA_SIMCARD', 'plano_id': pos.id, 'serial': '123'},
                             {'tipo': 'SEGURO', 'servico_adicional_id': seguro.id}, {'tipo': 'TROCA_TITULARIDADE', 'plano_id': pos.id,
                              'novo_titular_cpf': CPF, 'novo_titular_nome': ''}, {'tipo': 'TROCA_NUMERO', 'numero_novo': '123'}, {'tipo': ''}]})
    erros = ' | '.join(j.get('erros', []))
    t('erros de serviço: reativação só Pós/Controle, serial, ativação, titular, número, tipo',
      not j['ok'] and 'Reativação: só Pós ou Controle' in erros and 'ICCID' in erros and 'confirme que o serviço foi ativado' in erros
      and 'é o próprio cliente' in erros and 'nome do novo titular' in erros and 'DDD + número' in erros and 'escolha o tipo' in erros, erros)
    j = vender({'cliente': {'cpf': CPF, 'nome': 'Zz Cliente Teste'}, 'linhas': ['999999999'],
                'servicos': [{'tipo': 'TROCA_SIMCARD', 'plano_id': pos.id, 'serial': '8955101234567890123'},
                             {'tipo': 'SEGURO', 'servico_adicional_id': seguro.id, 'ativacao_confirmada': True},
                             {'tipo': 'TROCA_TITULARIDADE', 'plano_id': ctrl.id, 'novo_titular_cpf': CPF2, 'novo_titular_nome': 'Zz Novo Titular'}]})
    t('simcard, seguro e titularidade gravados', j.get('ok'), j)
    if j.get('ok'):
        v = Venda.objects.get(pk=j['id'])
        t('número fictício sinaliza a venda', v.numero_fake and all(s.numero_fake for s in v.servicos.all()))
        t('seguro com ativação confirmada e valor do catálogo', v.servicos.get(tipo_servico='SEGURO').ativacao_confirmada
          and v.servicos.get(tipo_servico='SEGURO').valor_plano == D('19.90'))
        t('titularidade grava o novo titular', v.servicos.get(tipo_servico='TROCA_TITULARIDADE').novo_titular_cpf == CPF2)
    j = vender({'cliente': {'cpf': CPF, 'nome': 'Zz Cliente Teste'}, 'linhas': ['2733334444'],
                'servicos': [{'tipo': 'TROCA_NUMERO', 'numero_novo': '(27) 98888-7777'}]})
    t('troca de número gravada', j.get('ok'), j)
    if j.get('ok'):
        s = VendaServico.objects.get(venda_id=j['id'])
        cli = Cliente.objects.get(cpf=CPF)
        t('número anterior e novo; cadastro atualizado com histórico', s.numero_anterior == '2733334444' and s.numero_novo == '27988887777'
          and cli.telefone == '27988887777' and RegistroAlteracao.objects.filter(tipo='CLIENTE', campo='telefone', depois='27988887777').exists())
    ca.post('/vendas/parametros/plano/', {'id': ctrl.id, 'nome': 'ZZ Controle 25GB', 'segmentacao': 'CONTROLE', 'valor': '89,90'})
    ctrl.refresh_from_db()
    j = vender({'cliente': {'cpf': CPF, 'nome': 'Zz Cliente Teste'}, 'linhas': ['27998765432'], 'servicos': [{'tipo': 'ALTA', 'plano_id': ctrl.id}]})
    t('plano inativo não entra em lançamento novo', not ctrl.ativo and not j['ok'] and 'inativo' in ' '.join(j['erros']))

    print('== PRÉ-ANÁLISE E TELAS ==')
    j = cv.get(f'/vendas/api/cliente/?cpf={CPF}').json()
    t('pré-análise: plano atual e últimas compras', j['pre_analise']['plano']['nome'] == 'ZZ Pós 50GB'
      and len(j['pre_analise']['compras']) >= 2 and not j['pre_analise']['vazio'], j['pre_analise'])
    r = cv.get('/vendas/precos/buscar/?q=ZZ%20Phone&segmentacao=PRE')
    res = r.json()['results']
    t('busca de produto com preço pelo plano', res and res[0]['valor_sugerido'] == '2499.00' and res[0]['categoria_slv'] == 'aparelho', res[:1])
    r = cv.get(f'/vendas/precos/buscar/?ids={aparelho.id}&segmentacao=POS&grupamento=PÓS INDIVIDUAL')
    t('reprecificar pelos ids', r.json()['results'][0]['valor_sugerido'] == '1999.00')
    r = cv.get('/vendas/precos/buscar/?q=ZZPX1')
    t('busca pelo SKU (ID DPGC)', any(x['id'] == aparelho.id for x in r.json()['results']))
    r = cv.get('/vendas/nova/')
    html = r.content.decode()
    t('tela de nova venda', r.status_code == 200 and 'slvDados' in html and loja.name in html and 'Beto Vendedor' in html)
    r = cv.get('/vendas/')
    html = r.content.decode()
    t('painel', r.status_code == 200 and 'Delta dos planos' in html and 'número fictício' in html)
    r = cv.get('/vendas/?fake=1')
    t('filtro: só número fictício', r.status_code == 200 and r.context['paginator'].count == 1, r.context['paginator'].count)
    r = cv.get('/vendas/?servico=TROCA_PLANO')
    t('filtro por tipo de serviço', r.context['paginator'].count == 1)
    venda_editada = VendaProduto.objects.filter(valor_editado=True, venda__vendedor=vend).first()
    if venda_editada:
        r = cv.get(f'/vendas/{venda_editada.venda_id}/')
        t('detalhe mostra o valor editado e quem editou', 'Editado à mão por' in r.content.decode() and 'Beto Vendedor' in r.content.decode())
    if True:
        r = cv.get(f'/vendas/api/renova/?codigo={renova.codigo}')
        t('código do Renova confere e avisa se já usado', r.json()['ok'] and r.json()['ja_usado'])
    r = cv.get('/vendas/api/renova/?codigo=RN-999999999')
    t('código inexistente', not r.json()['ok'])

    print('== IMPORTAÇÃO ==')
    import openpyxl
    from io import BytesIO
    from django.core.files.uploadedfile import SimpleUploadedFile

    def bloco(base):
        """Um grupo da Tabela Regular: Preço, PIX (−10%) e 2x…21x (juros a partir de 13x)."""
        if base == '-':
            return ['-'] * 22
        return [base, round(base * 0.9, 2)] + [base] * 11 + [round(base * (1 + 0.01 * n), 2) for n in range(1, 10)]

    def cabecalho_regular(grupos):
        cab = []
        for g in grupos:
            cab += [g, 'PIX\ne\nVivo Pay'] + [f'{n}x' for n in range(2, 22)]
        return cab

    def aba_regular(wb, nome, grupos, linhas):
        ws = wb.create_sheet(nome)
        ws.append([None] * 4 + ['TITULO'])
        largura = len(cabecalho_regular(grupos))
        ws.append([None] * 4 + ['Vigência', 'TABELA REGULAR'] + [None] * (largura - 1) + ['TABELA RENOVA'])
        ws.append(['PORTFÓLIO', 'CATEGORIA', 'MARCA', 'OFERTA DE COMUNICAÇÃO', 'Nome Comercial']
                  + cabecalho_regular(grupos) + ['DESCONTO', 'PRÉ-RENOVA'])
        for portfolio, categoria, marca, nome, bases in linhas:
            ws.append([portfolio, categoria, marca, 'Oferta', nome] + sum((bloco(b) for b in bases), []) + [0, 999999])

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.create_sheet('PLANOS')
    aba_regular(wb, 'SMARTPHONES', ['PRÉ', 'PÓS INDIVIDUAL'], [
        ('In', 'Smartphone', 'Zz', 'ZZ Phone Novo', [3000, 2500]),
        ('EOL', 'Smartphone', 'Zz', 'ZZ Phone X', [2600, 2200]),
        ('Out', 'Bem Estar', 'Zz', 'ZZ Watch', [1500, '-']),
    ])
    aba_regular(wb, 'ELETRÔNICOS_LP_Conectados', ['PRÉ', 'MULTIVIVO'], [('In', 'Box e Modem e FWT', 'Zz', 'ZZ Modem', [400, 300])])
    ws = wb.create_sheet('ELETRÔNICOS_LP_Não Conectados')
    ws.append([None, None, None, 'TITULO'])
    ws.append([None, None, None, 'Vigência'])
    ws.append(['PORTFÓLIO', 'CATEGORIA', 'MARCA', 'Nome Comercial', 'PVP Base', 'PIX\ne\nVivo Pay', '2x'])
    ws.append(['In', 'Essenciais', 'Zz', 'ZZ Capa Nova', 39.9, 35.91, 39.9])
    ws.append(['EOL', 'Áudio', 'Zz', 'ZZ Fone Novo', 299, 269.1, 299])
    ws.append(['In', 'Essenciais', 'Zz', 'ZZ Película', '#N/A', None, None])
    arquivo = BytesIO()
    wb.save(arquivo)
    usado = VendaProduto.objects.filter(preco=capa).exists()
    ca.post('/vendas/precos/importar/', {'arquivo': SimpleUploadedFile('tabela.xlsx', arquivo.getvalue())})
    imp = ImportacaoPrecos.objects.order_by('-id').first()
    t('importação substitui a tabela: só as 3 abas', imp and not imp.erro
      and set(ItemPreco.objects.filter(ativo=True).values_list('categoria', flat=True))
      == {'SMARTPHONES', 'ELETRÔNICOS_LP_Conectados', 'ELETRÔNICOS_LP_Não Conectados'}, imp and (imp.erro, imp.incluidos))
    t('relatório: 6 produtos, 1 rejeitado (PVP #N/A), o resto fora', imp.incluidos + imp.alterados == 6
      and len(imp.rejeitados) == 1 and 'PVP Base' in imp.rejeitados[0]['motivo'] and imp.removidos >= 3,
      (imp.incluidos, imp.alterados, imp.rejeitados, imp.removidos))
    t('produto já vendido fica inativo (a venda continua apontando)', not usado or ItemPreco.objects.get(pk=capa.pk).ativo is False)
    phone = ItemPreco.objects.get(categoria='SMARTPHONES', nome='ZZ Phone Novo')
    t('tabela regular: grupos e condições', phone.extra['grupos'] == ['PRÉ', 'PÓS INDIVIDUAL']
      and phone.extra['tabela']['PÓS INDIVIDUAL']['PIX e Vivo Pay'] == '2250.00' and phone.extra['tabela']['PRÉ']['21x'] == '3270.00'
      and phone.valor == D('3000.00') and phone.extra['PORTFÓLIO'] == 'In', phone.extra['tabela']['PRÉ'])
    watch = ItemPreco.objects.get(categoria='SMARTPHONES', nome='ZZ Watch')
    t('"-" na planilha: o grupo some do produto', watch.extra['grupos'] == ['PRÉ'])
    fone = ItemPreco.objects.get(nome='ZZ Fone Novo')
    t('Não Conectados: PVP Base é o valor', fone.valor == D('299.00') and 'tabela' not in fone.extra)
    t('categorias da venda', (slv.categoria_slv(phone), slv.categoria_slv(fone),
                              slv.categoria_slv(ItemPreco.objects.get(nome='ZZ Capa Nova'))) == ('aparelho', 'eletronico', 'essencial'))
    t('preço: grupo + condição', slv.preco_sugerido(phone, 'POS', 'PÓS INDIVIDUAL', 'PIX e Vivo Pay') == (D('2250.00'), 'PÓS INDIVIDUAL · PIX e Vivo Pay')
      and slv.preco_sugerido(phone, 'POS', 'PÓS INDIVIDUAL', '18x')[0] == D('2650.00')
      and slv.preco_sugerido(phone, 'PRE')[0] == D('3000.00')
      and slv.preco_sugerido(phone, 'POS') == (None, 'escolha o grupo do plano')
      and slv.preco_sugerido(watch, 'POS', 'PÓS INDIVIDUAL') == (None, 'não disponível no grupo PÓS INDIVIDUAL'))
    j = vender({'cliente': {'cpf': CPF, 'nome': 'Zz Cliente Teste'}, 'forma_pagamento': 'CARTAO', 'plano_anterior': {'id': pos.id},
                'produtos': [{'preco_id': phone.id, 'grupamento': 'PÓS INDIVIDUAL', 'condicao': '12x'}]})
    t('venda com grupo + condição grava o valor da tabela', j.get('ok') and VendaProduto.objects.get(venda_id=j['id']).valor_sugerido == D('2500.00')
      and VendaProduto.objects.get(venda_id=j['id']).regra_preco == 'PÓS INDIVIDUAL · 12x', j)
    j = vender({'cliente': {'cpf': CPF, 'nome': 'Zz Cliente Teste'}, 'forma_pagamento': 'CARTAO', 'plano_anterior': {'id': pos.id},
                'produtos': [{'preco_id': phone.id}]})
    t('sem grupo escolhido (cliente Pós): pede o grupo', not j['ok'] and 'informe o valor' in ' '.join(j['erros']), j)
    r = cv.get('/vendas/precos/buscar/?q=ZZ%20Phone%20Novo&segmentacao=PRE')
    res = r.json()['results'][0]
    t('busca devolve grupos, condições e o grupo PRÉ já escolhido para Pré', res['grupos'] == ['PRÉ', 'PÓS INDIVIDUAL']
      and res['grupo_padrao'] == 'PRÉ' and res['condicoes']['PRÉ'][:3] == ['Preço', 'PIX e Vivo Pay', '2x'] and res['valor_sugerido'] == '3000.00', res)
    r = cv.get('/vendas/precos/?search=ZZ')
    html = r.content.decode()
    t('lista de produtos mostra a grade da tabela regular', 'PÓS INDIVIDUAL' in html and 'ZZ Phone Novo' in html and 'PVP Base' in html)
    wb2 = openpyxl.Workbook()
    wb2.active.title = 'SMARTPHONES'
    arquivo2 = BytesIO()
    wb2.save(arquivo2)
    antes = ItemPreco.objects.count()
    ca.post('/vendas/precos/importar/', {'arquivo': SimpleUploadedFile('incompleta.xlsx', arquivo2.getvalue())})
    imp = ImportacaoPrecos.objects.order_by('-id').first()
    t('planilha sem as abas: falha registrada e tabela intacta', 'não tem a(s) aba(s)' in imp.erro and ItemPreco.objects.count() == antes, imp.erro)
    antes = ItemPreco.objects.count()
    ca.post('/vendas/precos/importar/', {'arquivo': SimpleUploadedFile('ruim.xlsx', b'nao e planilha')})
    imp = ImportacaoPrecos.objects.order_by('-id').first()
    t('arquivo ilegível: falha registrada e tabela intacta', imp.erro and ItemPreco.objects.count() == antes)
    r = ca.get('/vendas/precos/importar/')
    t('tela de importação com relatório', 'Relatório das importações' in r.content.decode())
    slv.renova_existe = existe_original
finally:
    marcador.__exit__(Exception, Exception('rollback'), None)

print(f'\n{ok} OK, {fail} falha(s)')
sys.exit(1 if fail else 0)
