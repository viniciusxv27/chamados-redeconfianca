"""Quiz gamificado — do cadastro ao pódio, com relatórios e Excel.

Pedido: "Novo módulo — Quiz gamificado (estilo Kahoot)": criação só pelo
SUPERADMIN e por quem ele liberar; quiz com perguntas, alternativas (uma
correta) e tempo; editar, duplicar e excluir antes de publicar; salas com
código, data e participantes por pessoa, loja, cargo, setor ou grupo; entrada
com o login e o código; o responsável inicia; uma resposta por pergunta, sem
troca; tempo esgotado = não respondida; pontos por acerto e velocidade (errou,
zero); ranking ao vivo e pódio; resultado do participante e do gestor;
histórico; banco de perguntas; aviso no portal; relatórios com Excel; e o
módulo no Assistente de Apresentações para virar tutorial.

A partida roda num relógio de mentira (``quiz.jogo.agora``) para os tempos e
pontos serem exatos. Transação desfeita no fim; todos os caches em memória
(nada no Redis compartilhado); push, e-mail e WhatsApp trocados por dublês que
o teste confere que nunca foram chamados.
"""
import json
import os
import sys
from datetime import timedelta
from io import BytesIO
from unittest import mock

import django

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'redeconfianca.settings')
os.environ.setdefault('RC_VARREDURA_ROTINA', '0')

from django.conf import settings

# Nada deste teste vai para o Redis compartilhado: todos os caches em memória.
settings.CACHES = {
    'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-quiz'},
    'local': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-quiz-local'},
}
django.setup()

from django.test.utils import setup_test_environment

setup_test_environment()            # o Client guarda o contexto das telas
if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS.append('testserver')

from django.contrib.auth import get_user_model
from django.contrib.messages import get_messages
from django.core.cache import caches
from django.db import connection, transaction
from django.test import Client, RequestFactory
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from openpyxl import load_workbook

from apresentacoes import modulos
from core.models import Notification
from quiz import jogo
from quiz.context_processors import quiz_menu
from quiz.models import Alternativa, Participante, Pergunta, Quiz, Resposta, Sala
from quiz.permissoes import pode_gerenciar
from users.models import Sector, UserModuleAccess
from users.module_access import GATES, MODULOS, tem_acesso

User = get_user_model()
ok = fail = 0


def t(nome, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'  OK   {nome}')
    else:
        fail += 1
        print(f'  FALHA {nome} {extra}')


class Relogio:
    """O "agora" da partida, andando só quando o teste manda."""

    def __init__(self, inicio):
        self.agora = inicio

    def __call__(self):
        return self.agora

    def em(self, base, segundos):
        self.agora = base + timedelta(seconds=segundos)


def cliente(u):
    c = Client()
    c.force_login(u)
    return c


def avisos(r):
    return ' | '.join(str(m) for m in get_messages(r.wsgi_request))


def json_de(r):
    return r.json() if r['Content-Type'].startswith('application/json') else None


XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

# Dublês de tudo que manda algo para fora do portal: o quiz só avisa no sino.
envios = {nome: mock.MagicMock(return_value={}) for nome in (
    'notifications.push_utils.send_push_notification_to_user',
    'notifications.push_utils.send_push_notification_to_users',
    'notifications.services.send_notification',
    'notifications.services.send_notification_all_channels',
    'notifications.onesignal_service.OneSignalService.send_notification',
    'notifications.truepush_service.TruePushService.send_notification',
)}
remendos = []
for caminho, duble in envios.items():
    try:
        remendo = mock.patch(caminho, duble)
        remendo.start()
        remendos.append(remendo)
    except (AttributeError, ImportError):
        pass                                             # esse caminho não existe nesta versão

relogio = Relogio(timezone.now().replace(microsecond=0))
remendo_relogio = mock.patch.object(jogo, 'agora', relogio)
remendo_relogio.start()

marcador = transaction.atomic()
marcador.__enter__()
try:
    assert not User.objects.filter(username__startswith='zzquiz.').exists(), 'usuários do teste já existem'
    caches['local'].clear()

    loja = Sector.objects.create(name='ZZ Loja Quiz Centro')
    escritorio_setor = Sector.objects.create(name='ZZ Escritório Quiz')

    def novo(apelido, hierarquia='PADRAO', setor=loja, cargo='CONSULTOR DE VENDAS'):
        return User.objects.create_user(
            username=f'zzquiz.{apelido}', email=f'zzquiz.{apelido}@exemplo-teste.local', password='S3nha!teste',
            first_name=apelido.title(), last_name='Zz Quiz', hierarchy=hierarquia, sector=setor, job_title=cargo)

    chefe = novo('chefe', 'SUPERADMIN', escritorio_setor, 'DIRETOR')
    gestor = novo('gestor', setor=escritorio_setor, cargo='TREINAMENTO')
    outro_gestor = novo('outrogestor', setor=escritorio_setor, cargo='TREINAMENTO')
    joao, maria = novo('joao'), novo('maria')
    ana = novo('ana', setor=escritorio_setor, cargo='ANALISTA')
    ausente, saira, novato = novo('ausente', setor=None), novo('saira', setor=None), novo('novato', setor=None)
    fora = novo('fora')
    escritorio = novo('escritorio', setor=escritorio_setor, cargo='FINANCEIRO')
    escritorio.sectors.add(loja)                         # escritório com a loja só no M2M

    # =====================================================================
    print('== PERMISSÃO: SUPERADMIN E QUEM ELE LIBERAR ==')
    modulo = next((m for m in MODULOS if m['chave'] == 'quiz'), None)
    t('o Quiz está na tela de edição do usuário, com a permissão "quiz.gestao"',
      modulo is not None and any(p[0] == 'quiz.gestao' for p in modulo['permissoes']), modulo)
    t('a chave tem porteiro', GATES.get('quiz.gestao') == 'quiz.permissoes:pode_gerenciar')
    t('antes da liberação o gestor não administra', not pode_gerenciar(gestor) and not tem_acesso(gestor, 'quiz.gestao'))
    for u in (gestor, outro_gestor):
        UserModuleAccess.objects.create(user=u, module_key='quiz.gestao', granted_by=chefe)
    gestor = User.objects.get(pk=gestor.pk)
    outro_gestor = User.objects.get(pk=outro_gestor.pk)
    t('liberado pelo SUPERADMIN, passa a administrar', pode_gerenciar(gestor) and tem_acesso(gestor, 'quiz.gestao'))
    t('o SUPERADMIN administra sem liberação', pode_gerenciar(chefe))
    t('colaborador comum não administra', not pode_gerenciar(joao))

    cg, cc, co = cliente(gestor), cliente(chefe), cliente(outro_gestor)
    cj, cm, ca, cz, cf, cn = (cliente(u) for u in (joao, maria, ana, ausente, fora, novato))
    r = cj.get('/quiz/quizzes/')
    t('colaborador comum não entra na gestão (volta para o Quiz com aviso)',
      r.status_code == 302 and r.url == '/quiz/' and 'SUPERADMIN' in avisos(r), (r.status_code, getattr(r, 'url', '')))
    r = cj.get('/quiz/')
    t('todo mundo abre o Quiz (a tela de quem joga)', r.status_code == 200 and not r.context['pode_gerenciar'])
    t('e o Quiz está no menu de todos', 'data-rota="/quiz/"' in r.content.decode())
    r = cg.get('/quiz/quizzes/')
    t('o gestor liberado abre a gestão', r.status_code == 200 and r.context['pode_gerenciar'])

    # =====================================================================
    print('\n== QUIZ E PERGUNTAS (RASCUNHO) ==')
    r = cg.post('/quiz/quizzes/novo/', {'titulo': '  ', 'descricao': 'x'})
    t('quiz sem nome não é criado', r.status_code == 200 and not Quiz.objects.filter(descricao='x').exists())
    r = cg.post('/quiz/quizzes/novo/', {'titulo': 'ZZ Quiz Rascunho', 'descricao': 'Para testar a edição',
                                         'categoria': 'ZZ Produtos', 'tempo_padrao': '999'})
    rascunho = Quiz.objects.get(titulo='ZZ Quiz Rascunho')
    t('quiz criado como rascunho, com tema e o tempo padrão no limite (240 s)',
      r.status_code == 302 and rascunho.status == 'RASCUNHO' and rascunho.categoria == 'ZZ Produtos'
      and rascunho.tempo_padrao == 240 and rascunho.criado_por_id == gestor.pk,
      (rascunho.status, rascunho.tempo_padrao))
    r = cg.get(f'/quiz/quizzes/{rascunho.pk}/')
    t('sem perguntas, a tela aponta o que falta para publicar',
      r.status_code == 200 and 'Cadastre pelo menos uma pergunta.' in r.context['problemas'])
    r = cg.post(f'/quiz/quizzes/{rascunho.pk}/publicar/')
    rascunho.refresh_from_db()
    t('e não publica', rascunho.status == 'RASCUNHO' and 'Ainda não dá para publicar' in avisos(r))

    url_nova = f'/quiz/quizzes/{rascunho.pk}/perguntas/nova/'

    def pergunta_por_tela(c, quiz, enunciado, alternativas, correta, tempo=20, **extra):
        dados = {'enunciado': enunciado, 'alternativas': alternativas, 'correta': str(correta),
                 'tempo_limite': str(tempo), 'explicacao': extra.pop('explicacao', '')}
        dados.update(extra)
        return c.post(f'/quiz/quizzes/{quiz.pk}/perguntas/nova/', dados)

    r = pergunta_por_tela(cg, rascunho, '', ['a', 'b'], 0)
    t('pergunta sem enunciado: erro', r.status_code == 200 and r.context['erro'] == 'Escreva a pergunta.')
    r = pergunta_por_tela(cg, rascunho, 'ZZ Só uma?', ['a', '', '', ''], 0)
    t('uma alternativa só: erro (mínimo 2)', r.status_code == 200 and 'pelo menos 2' in r.context['erro'])
    r = pergunta_por_tela(cg, rascunho, 'ZZ Qual?', ['a', 'b', '', ''], 2)
    t('correta marcada numa alternativa vazia: erro', r.status_code == 200 and 'correta' in r.context['erro'])
    t('na volta com erro, o que foi digitado continua na tela',
      r.context['valores']['enunciado'] == 'ZZ Qual?' and r.context['alternativas'][1] == ('b', False))
    t('nada foi gravado com erro', rascunho.perguntas.count() == 0)

    r = pergunta_por_tela(cg, rascunho, 'ZZ Pergunta A', ['A1', 'A2', 'A3', ''], 1, tempo=15,
                          explicacao='Porque sim.', mais='1')
    t('"Salvar e cadastrar outra" volta para uma pergunta nova', r.status_code == 302 and r.url == url_nova)
    pa = rascunho.perguntas.get(enunciado='ZZ Pergunta A')
    alts = list(pa.alternativas.values_list('texto', 'correta'))
    t('a pergunta tem as 3 alternativas preenchidas e só a certa marcada',
      alts == [('A1', False), ('A2', True), ('A3', False)] and pa.tempo_limite == 15 and pa.explicacao == 'Porque sim.',
      alts)
    pergunta_por_tela(cg, rascunho, 'ZZ Pergunta B', ['B1', 'B2'], 0)
    pergunta_por_tela(cg, rascunho, 'ZZ Pergunta C', ['C1', 'C2', 'C3', 'C4'], 3, tempo=3)
    pc = rascunho.perguntas.get(enunciado='ZZ Pergunta C')
    t('o tempo de cada pergunta respeita o mínimo (5 s)', pc.tempo_limite == 5, pc.tempo_limite)

    def ordem():
        return list(rascunho.perguntas.order_by('ordem', 'id').values_list('enunciado', flat=True))

    cg.post(f'/quiz/perguntas/{pa.pk}/duplicar/')
    t('duplicar põe a cópia logo abaixo da original',
      ordem() == ['ZZ Pergunta A', 'ZZ Pergunta A', 'ZZ Pergunta B', 'ZZ Pergunta C'], ordem())
    copia = rascunho.perguntas.filter(enunciado='ZZ Pergunta A').exclude(pk=pa.pk).get()
    t('a cópia leva as alternativas e a certa',
      list(copia.alternativas.values_list('texto', 'correta')) == alts)
    r = cg.post(f'/quiz/perguntas/{pc.pk}/mover/', {'direcao': 'cima'})
    t('mover para cima troca de lugar com a de cima',
      ordem() == ['ZZ Pergunta A', 'ZZ Pergunta A', 'ZZ Pergunta C', 'ZZ Pergunta B'] and r.url.endswith(f'#pergunta-{pc.pk}'),
      ordem())
    cg.post(f'/quiz/perguntas/{pc.pk}/mover/', {'direcao': 'baixo'})
    t('e para baixo, de volta', ordem() == ['ZZ Pergunta A', 'ZZ Pergunta A', 'ZZ Pergunta B', 'ZZ Pergunta C'])
    cg.post(f'/quiz/perguntas/{copia.pk}/excluir/')
    t('excluir tira a pergunta (e as alternativas)',
      ordem() == ['ZZ Pergunta A', 'ZZ Pergunta B', 'ZZ Pergunta C'] and not Alternativa.objects.filter(pergunta_id=copia.pk).exists())
    pb = rascunho.perguntas.get(enunciado='ZZ Pergunta B')
    r = cg.post(f'/quiz/perguntas/{pb.pk}/editar/', {'enunciado': 'ZZ Pergunta B editada', 'alternativas': ['B1', 'B2', 'B3'],
                                                     'correta': '2', 'tempo_limite': '30'})
    pb.refresh_from_db()
    t('editar troca o texto, as alternativas e a certa',
      pb.enunciado == 'ZZ Pergunta B editada' and pb.tempo_limite == 30
      and list(pb.alternativas.values_list('texto', 'correta')) == [('B1', False), ('B2', False), ('B3', True)])
    r = cg.get(f'/quiz/perguntas/{pb.pk}/editar/')
    t('a tela de edição vem preenchida (4 caixas, a certa marcada)',
      r.status_code == 200 and r.context['alternativas'] == [('B1', False), ('B2', False), ('B3', True), ('', False)])

    r = cg.post(f'/quiz/quizzes/{rascunho.pk}/publicar/')
    rascunho.refresh_from_db()
    t('com tudo certo, publica', rascunho.status == 'PUBLICADO' and rascunho.publicado_em is not None, avisos(r))
    antes = ordem()
    r = cg.post(url_nova, {'enunciado': 'ZZ Nova', 'alternativas': ['x', 'y'], 'correta': '0'})
    r2 = cg.post(f'/quiz/perguntas/{pa.pk}/excluir/')
    r3 = cg.get(f'/quiz/perguntas/{pa.pk}/editar/')
    t('publicado, as perguntas travam (nova, excluir e editar voltam com aviso)',
      ordem() == antes and r.status_code == r2.status_code == r3.status_code == 302 and 'duplique' in avisos(r3))
    r = cg.post(f'/quiz/quizzes/{rascunho.pk}/rascunho/')
    rascunho.refresh_from_db()
    t('nunca jogado e sem sala marcada: volta a rascunho', rascunho.status == 'RASCUNHO')
    cg.post(f'/quiz/quizzes/{rascunho.pk}/publicar/')
    r = cg.post(f'/quiz/quizzes/{rascunho.pk}/duplicar/')
    dup = Quiz.objects.get(duplicado_de=rascunho)
    t('duplicar o quiz cria uma cópia em rascunho com as mesmas perguntas',
      dup.status == 'RASCUNHO' and dup.titulo == 'ZZ Quiz Rascunho (cópia)'
      and list(dup.perguntas.values_list('enunciado', flat=True)) == ordem()
      and sum(p.alternativas.filter(correta=True).count() for p in dup.perguntas.all()) == 3)
    r = cg.post(f'/quiz/quizzes/{dup.pk}/excluir/')
    t('quiz sem salas pode ser excluído', not Quiz.objects.filter(pk=dup.pk).exists() and r.url == '/quiz/quizzes/')

    print('\n== BANCO DE PERGUNTAS ==')
    cg.post('/quiz/quizzes/novo/', {'titulo': 'ZZ Quiz Novo do Banco', 'tempo_padrao': '20'})
    alvo = Quiz.objects.get(titulo='ZZ Quiz Novo do Banco')
    r = cg.get(f'/quiz/quizzes/{alvo.pk}/banco/?q=ZZ+Pergunta')
    no_banco = {p.enunciado: p for p in r.context['perguntas']}
    t('o banco lista as perguntas dos outros quizzes', {'ZZ Pergunta A', 'ZZ Pergunta B editada', 'ZZ Pergunta C'} <= set(no_banco),
      sorted(no_banco))
    r = cg.post(f'/quiz/quizzes/{alvo.pk}/banco/', {'perguntas': [no_banco['ZZ Pergunta A'].pk, no_banco['ZZ Pergunta C'].pk]})
    copiadas = list(alvo.perguntas.order_by('ordem'))
    t('reaproveitar copia as escolhidas (com alternativas) e guarda de onde vieram',
      [p.enunciado for p in copiadas] == ['ZZ Pergunta A', 'ZZ Pergunta C']
      and copiadas[0].origem_id == no_banco['ZZ Pergunta A'].pk
      and list(copiadas[0].alternativas.values_list('texto', 'correta')) == alts, avisos(r))
    copiadas[0].enunciado = 'ZZ Pergunta A mudada no novo'
    copiadas[0].save()
    t('a cópia é independente: mudar aqui não muda a original',
      Pergunta.objects.get(pk=no_banco['ZZ Pergunta A'].pk).enunciado == 'ZZ Pergunta A')
    r = cg.get(f'/quiz/quizzes/{alvo.pk}/banco/?q=ZZ+Pergunta+C')
    t('no banco, a pergunta que já está no quiz vem marcada',
      'zz pergunta c' in r.context['ja_no_quiz'] and 'já está neste quiz' in r.content.decode())
    t('e as do próprio quiz não aparecem no banco dele',
      all(p.quiz_id != alvo.pk for p in r.context['perguntas']))

    # =====================================================================
    print('\n== O QUIZ DA PARTIDA ==')
    cg.post('/quiz/quizzes/novo/', {'titulo': 'ZZ Quiz Vivo', 'categoria': 'ZZ Produtos', 'tempo_padrao': '20'})
    vivo = Quiz.objects.get(titulo='ZZ Quiz Vivo')
    pergunta_por_tela(cg, vivo, 'ZZ Qual plano tem mais internet?', ['Controle 15GB', 'Controle 25GB', 'Pós 50GB', 'Pós 100GB'], 3,
                      explicacao='O Pós 100GB é o maior.')
    pergunta_por_tela(cg, vivo, 'ZZ Quanto é 2 + 2?', ['3', '4', '5'], 1)
    pergunta_por_tela(cg, vivo, 'ZZ Qual a cor do logo?', ['Roxo', 'Laranja'], 0)
    cg.post(f'/quiz/quizzes/{vivo.pk}/publicar/')
    vivo.refresh_from_db()
    perguntas = list(vivo.perguntas.order_by('ordem').prefetch_related('alternativas'))
    alt = [{a.texto: a.pk for a in p.alternativas.all()} for p in perguntas]
    t('quiz de 3 perguntas publicado', vivo.status == 'PUBLICADO' and len(perguntas) == 3)

    # =====================================================================
    print('\n== SALA: QUEM PARTICIPA ==')
    r = cg.get(f'/quiz/salas/nova/?quiz={vivo.pk}')
    abas = {chave: itens for chave, _rotulo, itens in r.context['abas']}
    da_loja = next((i for i in abas['lojas'] if i['id'] == loja.pk), None)
    t('a tela da sala abre com o quiz escolhido e as abas Lojas/Setores/Cargos/Grupos/Coordenação',
      r.status_code == 200 and str(r.context['valores']['quiz']) == str(vivo.pk)
      and list(abas) == ['lojas', 'setores', 'cargos', 'grupos', 'coordenacoes'])
    t('a loja traz quem tem a loja como setor principal',
      da_loja is not None and {joao.pk, maria.pk, fora.pk} <= set(da_loja['membros']), da_loja)
    t('e não o pessoal de escritório que só tem a loja no M2M', da_loja and escritorio.pk not in da_loja['membros'])
    t('quem cria a sala não entra em lista nenhuma (é o responsável)',
      all(gestor.pk not in i['membros'] for itens in abas.values() for i in itens))
    t('o escritório aparece em Setores, não em Lojas',
      any(i['id'] == escritorio_setor.pk for i in abas['setores']) and all(i['id'] != escritorio_setor.pk for i in abas['lojas']))

    quando = (timezone.localtime() + timedelta(days=1)).replace(hour=10, minute=0, second=0, microsecond=0)
    base_post = {'quiz': vivo.pk, 'nome': 'ZZ Sala Teste', 'agendada_para': quando.strftime('%Y-%m-%dT%H:%M')}
    r = cg.post('/quiz/salas/nova/', {**base_post, 'quiz': alvo.pk, 'lojas': [loja.pk]})
    t('quiz em rascunho não vira sala', r.status_code == 200 and r.context['erro'] == 'Escolha um quiz publicado.'
      and alvo.status == 'RASCUNHO')
    r = cg.post('/quiz/salas/nova/', {**base_post, 'quiz': '', 'lojas': [loja.pk]})
    t('sem quiz: erro', r.status_code == 200 and r.context['erro'] == 'Escolha um quiz publicado.')
    r = cg.post('/quiz/salas/nova/', {**base_post, 'agendada_para': 'amanhã', 'lojas': [loja.pk]})
    t('sem data e hora válidas: erro', r.status_code == 200 and 'data' in r.context['erro'])
    t('e o que estava marcado continua marcado',
      next(i for i in {c: i for c, _r, i in r.context['abas']}['lojas'] if i['id'] == loja.pk).get('marcado') is True)
    r = cg.post('/quiz/salas/nova/', base_post)
    t('sem ninguém selecionado: erro', r.status_code == 200 and 'Selecione quem vai participar' in r.context['erro'])
    t('nenhuma sala criada nas tentativas com erro', not Sala.objects.filter(quiz__in=[vivo, alvo]).exists())

    r = cg.post('/quiz/salas/nova/', {**base_post, 'lojas': [loja.pk],
                                       'usuarios': [ana.pk, ausente.pk, saira.pk, gestor.pk]})
    sala = Sala.objects.get(quiz=vivo)
    cod = sala.codigo
    convidados = set(sala.participantes.values_list('user_id', flat=True))
    t('sala criada com código de 6 letras e redireciona para o painel',
      r.status_code == 302 and r.url == f'/quiz/salas/{cod}/painel/' and len(cod) == 6 and sala.fase == 'AGENDADA'
      and sala.responsavel_id == gestor.pk)
    t('participantes: a loja (setor principal) + os escolhidos um a um',
      convidados == {joao.pk, maria.pk, fora.pk, ana.pk, ausente.pk, saira.pk},
      sorted(User.objects.filter(pk__in=convidados).values_list('username', flat=True)))
    t('o responsável não vira participante, nem o escritório de M2M',
      gestor.pk not in convidados and escritorio.pk not in convidados)
    pj = sala.participantes.get(user=joao)
    pa_ana = sala.participantes.get(user=ana)
    t('cada um sabe por onde foi chamado', pj.origem == 'LOJA' and pj.rotulo_origem == loja.name
      and pa_ana.origem == 'MANUAL')
    avisos_sala = Notification.objects.filter(user__username__startswith='zzquiz.', title='Você foi chamado para um quiz')
    exemplo = avisos_sala.filter(user=joao).first()
    t('cada participante recebe o aviso no sino, com nome, data, hora, código e o botão da sala',
      avisos_sala.count() == 6 and exemplo and 'ZZ Sala Teste' in exemplo.message
      and quando.strftime('%d/%m às %H:%M') in exemplo.message and cod in exemplo.message
      and exemplo.related_url == f'/quiz/sala/{cod}/',
      (avisos_sala.count(), exemplo.message if exemplo else None))
    t('só aviso no portal: nenhum push/e-mail/WhatsApp', not any(d.called for d in envios.values()))

    r = cg.post(f'/quiz/quizzes/{vivo.pk}/rascunho/')
    t('com sala marcada, o quiz não volta a rascunho', Quiz.objects.get(pk=vivo.pk).status == 'PUBLICADO'
      and 'cancele a sala' in avisos(r))
    r = cg.post(f'/quiz/quizzes/{vivo.pk}/excluir/')
    t('nem é excluído', Quiz.objects.filter(pk=vivo.pk).exists())

    r = co.get(f'/quiz/salas/{cod}/editar/')
    t('outro gestor não mexe na sala dos outros', r.status_code == 302 and 'Só o responsável' in avisos(r))
    r = cg.post(f'/quiz/salas/{cod}/editar/', {**base_post, 'lojas': [loja.pk],
                                              'usuarios': [ana.pk, ausente.pk, novato.pk]})
    convidados = set(Sala.objects.get(pk=sala.pk).participantes.values_list('user_id', flat=True))
    t('editar: sai quem saiu da seleção (e não tinha entrado), entra quem chegou',
      convidados == {joao.pk, maria.pk, fora.pk, ana.pk, ausente.pk, novato.pk}, convidados)
    t('só quem chegou recebe aviso novo',
      Notification.objects.filter(user=novato, title='Você foi chamado para um quiz').count() == 1
      and Notification.objects.filter(user=joao, title='Você foi chamado para um quiz').count() == 1)
    # "fora" veio pela loja; tira ele por pessoa não dá, então a loja sai e os da loja entram um a um.
    cg.post(f'/quiz/salas/{cod}/editar/', {**base_post, 'usuarios': [joao.pk, maria.pk, ana.pk, ausente.pk, novato.pk]})
    convidados = set(Sala.objects.get(pk=sala.pk).participantes.values_list('user_id', flat=True))
    t('sem a loja, fica só quem foi escolhido', convidados == {joao.pk, maria.pk, ana.pk, ausente.pk, novato.pk}, convidados)

    print('\n== ENTRAR NA SALA ==')
    r = cj.get('/quiz/')
    t('a sala aparece em "Suas salas" de quem foi chamado', cod in r.content.decode()
      and any(p.sala_id == sala.pk for p in r.context['abertas']))
    t('e não aparece para quem não foi', cod not in cf.get('/quiz/').content.decode())
    r = cj.post('/quiz/entrar/', {'codigo': f'  {cod.lower()[:3]} {cod.lower()[3:]} '})
    t('entrar com o código (minúsculo e com espaço) leva à sala', r.status_code == 302 and r.url == f'/quiz/sala/{cod}/')
    r = cf.post('/quiz/entrar/', {'codigo': cod})
    t('quem não foi chamado não entra', r.url == '/quiz/' and 'não foi chamado' in avisos(r))
    r = cf.post('/quiz/entrar/', {'codigo': 'QQQQQQ'})
    t('código que não existe: aviso', r.url == '/quiz/' and 'Não existe sala' in avisos(r))
    r = cg.post('/quiz/entrar/', {'codigo': cod})
    t('o responsável que digita o código vai para o painel', r.url == f'/quiz/salas/{cod}/painel/')
    r = cf.get(f'/quiz/sala/{cod}/')
    r2 = cf.get(f'/quiz/sala/{cod}/estado/')
    t('pela URL também não: a tela manda embora e o estado responde 403',
      r.status_code == 302 and r2.status_code == 403 and json_de(r2)['ok'] is False)
    r = cj.get(f'/quiz/sala/{cod}/')
    t('quem foi chamado abre a tela do jogo', r.status_code == 200 and r.context['participante'].user_id == joao.pk)

    req = RequestFactory().get('/')
    req.user = joao
    caches['local'].clear()
    menu = quiz_menu(req)
    t('o menu mostra 1 sala aberta para ele (ainda não ao vivo)', menu.get('quiz_salas_abertas') == 1
      and not menu.get('quiz_ao_vivo'), menu)

    # =====================================================================
    print('\n== A PARTIDA ==')
    T0 = relogio.agora
    ej = json_de(cj.get(f'/quiz/sala/{cod}/estado/'))
    em = json_de(cm.get(f'/quiz/sala/{cod}/estado/'))
    t('na espera: fase AGENDADA e quem abriu a sala conta como presente',
      ej['fase'] == 'AGENDADA' and em['presentes'] == 2 and ej['total_perguntas'] == 3)
    painel = json_de(cg.get(f'/quiz/salas/{cod}/painel/estado/'))
    t('o responsável vê quem entrou e quem está com a sala aberta',
      painel['convidados'] == 5 and painel['presentes'] == 2
      and {p['nome'] for p in painel['pessoas'] if p['entrou']} == {joao.get_full_name(), maria.get_full_name()}
      and all(p['online'] for p in painel['pessoas'] if p['entrou']), painel)

    def acao(c, qual):
        r = c.post(f'/quiz/salas/{cod}/painel/acao/', {'acao': qual})
        return r.status_code, json_de(r)

    st, d = acao(co, 'iniciar')
    t('outro gestor não conduz (403)', st == 403)
    st, d = acao(cg, 'proxima')
    st2, d2 = acao(cg, 'encerrar')
    t('antes de começar não dá para avançar nem encerrar', st == 409 and st2 == 409, (d, d2))
    r = cj.post(f'/quiz/sala/{cod}/responder/', json.dumps({'indice': 0, 'alternativa': alt[0]['Pós 100GB']}),
                content_type='application/json')
    t('responder antes de começar não vale', r.status_code == 409)

    st, d = acao(cg, 'iniciar')
    t('o responsável inicia: primeira pergunta no ar', st == 200 and d['fase'] == 'PERGUNTA' and d['indice'] == 0
      and d['restante_ms'] == 20000, d)
    t('quem ainda não entrou recebe o aviso "O quiz começou!" (no portal)',
      set(Notification.objects.filter(title='O quiz começou!', user__username__startswith='zzquiz.')
          .values_list('user__username', flat=True)) == {'zzquiz.ana', 'zzquiz.ausente', 'zzquiz.novato'})
    st, _ = acao(cg, 'iniciar')
    t('iniciar de novo não vale', st == 409)
    t('o painel ao vivo não mostra qual é a certa', all('correta' not in a for a in d['alternativas']))

    # Pergunta 1 (20 s; certa: Pós 100GB)
    ej = json_de(cj.get(f'/quiz/sala/{cod}/estado/'))
    t('o participante recebe a pergunta, as alternativas e o tempo — sem a resposta certa',
      ej['fase'] == 'PERGUNTA' and ej['enunciado'] == perguntas[0].enunciado and ej['restante_ms'] == 20000
      and len(ej['alternativas']) == 4 and all('correta' not in a for a in ej['alternativas']) and not ej['respondida'], ej)
    relogio.em(T0, 1)
    em = json_de(cm.get(f'/quiz/sala/{cod}/estado/'))
    t('quem recebe a pergunta 1 s depois tem os 20 s dele', em['restante_ms'] == 20000, em.get('restante_ms'))
    with CaptureQueriesContext(connection) as consultas:
        cm.get(f'/quiz/sala/{cod}/estado/')
    print(f'       (uma consulta da tela do participante = {len(consultas)} idas ao banco, contando os middlewares)')
    t('a consulta da tela do participante é enxuta', len(consultas) <= 25, len(consultas))

    def responder(c, indice, alternativa):
        r = c.post(f'/quiz/sala/{cod}/responder/', json.dumps({'indice': indice, 'alternativa': alternativa}),
                   content_type='application/json')
        return r.status_code, json_de(r)

    relogio.em(T0, 2)
    st, d = responder(cj, 0, alt[0]['Pós 100GB'])
    t('João acerta em 2 s', st == 200 and d['ok'])
    st, d = responder(cj, 0, alt[0]['Pós 50GB'])
    t('e não pode trocar a resposta', st == 409 and 'já respondeu' in d['erro'])
    st, d = responder(ca, 0, alt[1]['4'])
    t('alternativa de outra pergunta não vale', st == 409 and 'Escolha uma das alternativas' in d['erro'], d)
    st, d = responder(ca, 1, alt[0]['Pós 100GB'])
    t('responder a pergunta errada (índice) não vale', st == 409)
    r = ca.post(f'/quiz/sala/{cod}/responder/', 'isto não é json', content_type='application/json')
    t('corpo inválido: 400', r.status_code == 400)
    relogio.em(T0, 5)
    ea = json_de(ca.get(f'/quiz/sala/{cod}/estado/'))
    t('Ana entra atrasada (5 s): o tempo dela vai até o fim da pergunta + a folga', ea.get('restante_ms') == 18000,
      ea)
    relogio.em(T0, 6)
    st, d = responder(cm, 0, alt[0]['Controle 25GB'])
    t('Maria erra', st == 200)
    painel = json_de(cg.get(f'/quiz/salas/{cod}/painel/estado/'))
    t('o painel conta quem já respondeu e a pergunta segue aberta (falta a Ana)',
      painel['fase'] == 'PERGUNTA' and painel['respondidas'] == 2 and painel['presentes'] == 3, painel)
    relogio.em(T0, 15)
    st, d = responder(ca, 0, alt[0]['Pós 100GB'])
    sala.refresh_from_db()
    t('Ana acerta 10 s depois de ver; todos os que entraram responderam (quem nunca entrou não segura a pergunta): '
      'a sala vai sozinha para o resultado',
      st == 200 and sala.fase == 'RESULTADO', (st, d, sala.fase))
    rj = Resposta.objects.get(participante__user=joao, pergunta=perguntas[0])
    rm = Resposta.objects.get(participante__user=maria, pergunta=perguntas[0])
    ra = Resposta.objects.get(participante__user=ana, pergunta=perguntas[0])
    t('pontos: acerto em 2 de 20 s = 950; em 10 s (contados de quando viu) = 750; erro = 0',
      (rj.pontos, rj.tempo_ms, ra.pontos, ra.tempo_ms, rm.pontos, rm.correta) == (950, 2000, 750, 10000, 0, False),
      (rj.pontos, rj.tempo_ms, ra.pontos, ra.tempo_ms, rm.pontos))
    ej = json_de(cj.get(f'/quiz/sala/{cod}/estado/'))
    certas = [a for a in ej['alternativas'] if a['correta']]
    t('no resultado, o participante vê se acertou, os pontos e a certa',
      ej['fase'] == 'RESULTADO' and ej['minha'] == {'respondida': True, 'correta': True, 'pontos': 950,
                                                     'alternativa': alt[0]['Pós 100GB']}
      and [a['texto'] for a in certas] == ['Pós 100GB'], ej)
    t('quantos marcaram cada alternativa', {a['texto']: a['total'] for a in ej['alternativas']}
      == {'Controle 15GB': 0, 'Controle 25GB': 1, 'Pós 50GB': 0, 'Pós 100GB': 2})
    t('e o ranking ao vivo (ele em 1º)', ej['posicao'] == 1 and ej['top'][0]['nome'] == joao.get_full_name()
      and [p['pontos'] for p in ej['top']] == [950, 750, 0], ej['top'])
    em = json_de(cm.get(f'/quiz/sala/{cod}/estado/'))
    t('quem errou vê que errou e zero ponto', em['minha']['correta'] is False and em['minha']['pontos'] == 0
      and em['posicao'] == 3)

    # Pergunta 2 (ninguém responde)
    relogio.em(T0, 30)
    st, d = acao(cg, 'proxima')
    t('próxima pergunta', st == 200 and d['fase'] == 'PERGUNTA' and d['indice'] == 1)
    T2 = relogio.agora
    json_de(cj.get(f'/quiz/sala/{cod}/estado/'))
    relogio.em(T2, 20 + 3.5)
    em = json_de(cm.get(f'/quiz/sala/{cod}/estado/'))
    t('acabou o tempo (+ folga): a própria consulta de um participante fecha a pergunta', em['fase'] == 'RESULTADO', em['fase'])
    t('quem não respondeu vê "não respondeu"', em['minha']['respondida'] is False and em['minha']['pontos'] == 0)
    st, d = responder(cj, 1, alt[1]['4'])
    t('resposta depois do tempo não vale', st == 409 and 'acabou' in d['erro'])

    # Pergunta 3 (última)
    relogio.em(T0, 60)
    acao(cg, 'proxima')
    T3 = relogio.agora
    for c in (cj, cm, ca):
        json_de(c.get(f'/quiz/sala/{cod}/estado/'))
    relogio.em(T3, 1)
    responder(cm, 2, alt[2]['Roxo'])
    relogio.em(T3, 20)
    st, d = responder(cj, 2, alt[2]['Roxo'])
    t('acerto no último instante ainda vale (500 pontos)', st == 200
      and Resposta.objects.get(participante__user=joao, pergunta=perguntas[2]).pontos == 500)
    t('e o de 1 s vale 975', Resposta.objects.get(participante__user=maria, pergunta=perguntas[2]).pontos == 975)
    relogio.em(T3, 24)
    st, d = responder(ca, 2, alt[2]['Roxo'])
    t('depois do tempo e da folga, não vale mais (conta como não respondida)', st == 409 and 'acabou' in d['erro'])
    st, d = acao(cg, 'revelar')
    t('o responsável mostra o resultado; é a última', st == 200 and d['fase'] == 'RESULTADO' and d['ultima'] is True)
    st, d = acao(cg, 'proxima')
    t('depois da última: pódio (sala encerrada)', st == 200 and d['fase'] == 'ENCERRADA'
      and [p['nome'] for p in d['podio']] == [joao.get_full_name(), maria.get_full_name(), ana.get_full_name()], d)
    sala.refresh_from_db()
    t('encerrada com a hora do fim', sala.fase == 'ENCERRADA' and sala.encerrada_em is not None)

    placar = {p.user.username: p for p in jogo.ranking(sala)}
    t('totais: João 1450 (2 acertos), Maria 975 (1 acerto, 1 erro), Ana 750 (1 acerto)',
      [(u, p.pontos, p.acertos, p.erros, p.posicao) for u, p in placar.items()]
      == [('zzquiz.joao', 1450, 2, 0, 1), ('zzquiz.maria', 975, 1, 1, 2), ('zzquiz.ana', 750, 1, 0, 3)],
      [(u, p.pontos, p.acertos, p.erros, p.posicao) for u, p in placar.items()])
    ej = json_de(cj.get(f'/quiz/sala/{cod}/estado/'))
    t('na tela do participante: pódio, posição, pontos, acertos, erros e o link do resultado',
      ej['fase'] == 'ENCERRADA' and ej['posicao'] == 1 and ej['pontos_total'] == 1450 and ej['acertos'] == 2
      and ej['erros'] == 0 and ej['url_resultado'] == f'/quiz/sala/{cod}/meu-resultado/', ej)

    r = cz.get(f'/quiz/sala/{cod}/')
    ez = json_de(cz.get(f'/quiz/sala/{cod}/estado/'))
    t('quem não entrou e abre o link depois do fim não vira "participante" com zero',
      r.status_code == 302 and r.url == '/quiz/' and 'já terminou' in avisos(r)
      and Participante.objects.get(sala=sala, user=ausente).entrou_em is None
      and ez['posicao'] is None and 'url_resultado' not in ez, ez)
    r = cz.get(f'/quiz/sala/{cod}/meu-resultado/')
    t('nem tem "meu resultado"', r.status_code == 302 and r.url == '/quiz/')
    r = cj.get(f'/quiz/sala/{cod}/')
    t('quem jogou e reabre a sala vai para o resultado dele', r.status_code == 302 and r.url == f'/quiz/sala/{cod}/meu-resultado/')

    print('\n== RESULTADO DO PARTICIPANTE E HISTÓRICO ==')
    r = cj.get(f'/quiz/sala/{cod}/meu-resultado/')
    res = r.context['r'] if r.status_code == 200 else {}
    t('acertos, erros, sem resposta, pontos e posição', r.status_code == 200
      and (res['acertos'], res['erros'], res['sem_resposta'], res['pontos'], res['posicao'], res['presentes'])
      == (2, 0, 1, 1450, 1, 3), res and (res['acertos'], res['erros'], res['sem_resposta'], res['pontos'], res['posicao']))
    linhas = res.get('linhas', [])
    t('pergunta por pergunta: a dele, a certa e se acertou',
      [(l['correta'], l['respondida'], l['pontos']) for l in linhas] == [(True, True, 950), (False, False, 0), (True, True, 500)]
      and linhas[0]['certa'].texto == 'Pós 100GB' and linhas[0]['escolhida'].texto == 'Pós 100GB')
    html = r.content.decode()
    t('com a explicação da pergunta e o pódio (marcando onde ele está)', 'O Pós 100GB é o maior.' in html
      and 'Pódio' in html and '<strong class="text-orange-600">você</strong>' in html)
    r = cm.get('/quiz/')
    t('no histórico dela: a sala, os acertos e os pontos', any(p.sala_id == sala.pk for p in r.context['historico'])
      and '975 pontos' in r.content.decode())
    caches['local'].clear()
    t('e o menu não conta mais a sala', quiz_menu(req).get('quiz_salas_abertas', 0) == 0)

    print('\n== RESULTADO DO GESTOR ==')
    r = cj.get(f'/quiz/salas/{cod}/resultados/')
    t('colaborador não vê o resultado da sala', r.status_code == 302)
    r = cg.get(f'/quiz/salas/{cod}/resultados/')
    rr = r.context['r']
    t('quem participou e quem não', [p.user.username for p in rr['placar']] == ['zzquiz.joao', 'zzquiz.maria', 'zzquiz.ana']
      and {p.user.username for p in rr['ausentes']} == {'zzquiz.ausente', 'zzquiz.novato'})
    t('participação 3 de 5 e média de acertos 4/3', rr['convidados'] == 5 and abs(rr['participacao'] - 0.6) < 1e-9
      and abs(rr['media_acertos'] - 4 / 3) < 1e-9, (rr['participacao'], rr['media_acertos']))
    t('as perguntas com mais erros (a 2, que ninguém respondeu, em primeiro)',
      [l['n'] for l in rr['mais_erradas']] == [2, 1, 3] and rr['mais_erradas'][0]['taxa_erro'] == 1.0,
      [(l['n'], l['taxa_erro']) for l in rr['mais_erradas']])
    t('sem resposta e aproveitamento por pessoa', [(p.sem_resposta, round(p.aproveitamento, 3)) for p in rr['placar']]
      == [(1, 0.667), (1, 0.333), (2, 0.333)])
    r = cg.get(f'/quiz/salas/{cod}/resultados.xlsx')
    livro = load_workbook(BytesIO(r.content)) if r['Content-Type'] == XLSX else None
    t('Excel da sala: Ranking, Perguntas e Respostas', livro is not None
      and livro.sheetnames == ['Ranking', 'Perguntas', 'Respostas'] and livro['Ranking'].max_row == 1 + 5
      and livro['Ranking']['B2'].value == joao.get_full_name())

    print('\n== RELATÓRIOS ==')
    r = cg.get(f'/quiz/relatorios/?quiz={vivo.pk}')
    rel = r.context['r']
    t('participação, média de acertos e aproveitamento das salas do período',
      rel['resumo']['salas'] == 1 and rel['resumo']['convidados'] == 5 and rel['resumo']['presentes'] == 3
      and abs(rel['resumo']['media_acertos'] - 4 / 3) < 1e-9, rel['resumo'])
    colab = {c['user'].username: c for c in rel['colaboradores']}
    t('desempenho por colaborador (quem faltou aparece com 0 de 1)',
      colab['zzquiz.joao']['acertos'] == 2 and colab['zzquiz.ausente']['salas'] == 0 and colab['zzquiz.ausente']['convites'] == 1)
    t('perguntas com maior índice de erro', rel['mais_erradas'][0]['pergunta'].pk == perguntas[1].pk
      and rel['mais_erradas'][0]['taxa_erro'] == 1.0)
    r = cg.get(f'/quiz/relatorios.xlsx?quiz={vivo.pk}')
    livro = load_workbook(BytesIO(r.content)) if r['Content-Type'] == XLSX else None
    t('Excel dos relatórios: Salas, Colaboradores e Perguntas com mais erro', livro is not None
      and livro.sheetnames == ['Salas', 'Colaboradores', 'Perguntas com mais erro'] and livro['Salas'].max_row == 2)
    t('colaborador não abre os relatórios', cj.get('/quiz/relatorios/').status_code == 302)
    r = cg.post(f'/quiz/quizzes/{vivo.pk}/rascunho/')
    t('quiz já jogado não volta a rascunho (duplique)', Quiz.objects.get(pk=vivo.pk).status == 'PUBLICADO'
      and 'duplique' in avisos(r))

    print('\n== DESEMPATE, PONTOS E CANCELAMENTO ==')
    t('pontos: 1000 na hora, 750 na metade, 500 no fim (e além), 0 no erro',
      [jogo.pontos_por(True, 0, 20), jogo.pontos_por(True, 10000, 20), jogo.pontos_por(True, 20000, 20),
       jogo.pontos_por(True, 99999, 20), jogo.pontos_por(False, 0, 20)] == [1000, 750, 500, 500, 0])
    outra = Sala.objects.create(quiz=vivo, agendada_para=quando, responsavel=gestor)
    Participante.objects.create(sala=outra, user=joao, entrou_em=T0, pontos=900, tempo_acertos_ms=5000)
    Participante.objects.create(sala=outra, user=maria, entrou_em=T0, pontos=900, tempo_acertos_ms=3000)
    t('empate nos pontos: fica na frente quem acertou mais rápido',
      [p.user_id for p in jogo.ranking(outra)] == [maria.pk, joao.pk])
    r = cg.post(f'/quiz/salas/{outra.codigo}/painel/acao/', {'acao': 'cancelar'})
    outra.refresh_from_db()
    t('sala marcada pode ser cancelada', json_de(r)['fase'] == 'CANCELADA' and outra.fase == 'CANCELADA')
    r = cj.get(f'/quiz/sala/{outra.codigo}/')
    t('e quem foi chamado é avisado que foi cancelada', r.status_code == 302 and 'cancelada' in avisos(r))
    t('outro gestor acompanha o painel mas sem os botões', co.get(f'/quiz/salas/{cod}/painel/').context['pode_conduzir'] is False)
    t('o SUPERADMIN conduz qualquer sala', cc.get(f'/quiz/salas/{cod}/painel/').context['pode_conduzir'] is True)

    print('\n== NO ASSISTENTE DE APRESENTAÇÕES ==')
    r = cc.get('/apresentacoes/modulos/?atualizar=1')
    catalogo = {m['label']: m for m in r.context['modulos']} if r.context else {}
    item = catalogo.get('quiz')
    t('o Quiz está na lista de módulos para gerar tutorial', r.status_code == 200 and item is not None, sorted(catalogo))
    t('com o nome e o ícone do menu', item and item['nome'] == 'Quiz' and item['icone'] == 'fa-circle-question', item)
    paginas = modulos.paginas_para_capturar('quiz', item['telas'] if item else [])
    t('as telas a capturar começam pela entrada do Quiz', paginas and paginas[0]['url'] == '/quiz/', paginas)
    contexto = modulos.contexto_do_modulo('quiz', 'Quiz')
    t('o contexto do tutorial explica as regras (quem cria, pontos, fases)',
      '[permissoes.py]' in contexto and 'quiz.gestao' in contexto and '1000' in contexto, contexto[:300])

    t('em todo o teste, nada saiu por push, e-mail ou WhatsApp', not any(d.called for d in envios.values()),
      [n for n, d in envios.items() if d.called])
finally:
    transaction.set_rollback(True)
    marcador.__exit__(None, None, None)
    remendo_relogio.stop()
    for remendo in remendos:
        remendo.stop()
    print('\nrollback: nada deste teste foi gravado no banco.')

print(f'\n{ok} OK / {fail} falhas')
sys.exit(1 if fail else 0)
