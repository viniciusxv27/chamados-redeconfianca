"""Acess (/acess/): o "Confiança Acess" dentro do portal (20/10/2026).

Pedido: novo módulo; só o SUPERADMIN entra e libera cada função; ao aprovar um
pré-cadastro a pessoa já entra no fluxo.

Cobre: acesso e liberação por função; pré-cadastro aprovado → Etapa 1; as 4
etapas (GEFAC, login → cadastro do portal, Academia, checklist com perfil
regional SP/ES, finalizar com chamado de liberação); voltar etapa; conflito de
versão; conclusão em lote; chamados (código único, prazo, reagendar mantém o
SLA, atribuir, concluir); desligamento (datas, confirmar arquiva); excluir com
senha e sem chamado ativo; mensagens (variáveis, link do WhatsApp — nada é
enviado); duplicados; Excel em texto; histórico; todas as telas abrem.

Transação desfeita no fim; nada sai para a rede.
"""
import io
import os
import sys
from datetime import timedelta
from unittest import mock

import django

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'redeconfianca.settings')
from django.conf import settings

settings.CACHES = {'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-ac'},
                   'local': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-ac2'}}
django.setup()
from django.test.utils import setup_test_environment

setup_test_environment()
if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS.append('testserver')

import requests
from django.contrib.auth import get_user_model
from django.db import transaction
from django.test import Client
from django.utils import timezone
from openpyxl import load_workbook

from acess import servicos as sv
from acess.models import AcessoAcess, Chamado, Colaborador, LogAcess
from users.models import Sector

User = get_user_model()
ok = fail = 0


def t(nome, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'  OK   {nome}')
    else:
        fail += 1
        print(f'  FALHA {nome} {extra}')


def sem_rede(*a, **k):
    raise AssertionError('o teste tentou sair para a rede')


CPF1, CPF2, CPF3 = '52998224725', '11144477735', '39053344705'
HOJE = timezone.localdate()

marcador = transaction.atomic()
marcador.__enter__()
try:
    with mock.patch.object(requests, 'post', sem_rede), mock.patch.object(requests, 'get', sem_rede):
        sp = Sector.objects.create(name='Loja ZZ Paulista', adabas='SPZZZ01-001')
        es = Sector.objects.create(name='Loja ZZ Capixaba', adabas='ESZZZ01-001')
        sede = Sector.objects.filter(name__icontains='sede administrativa').first() or Sector.objects.create(name='Sede Administrativa ZZ')
        chefe = User.objects.create_user(username='zzz.acchefe', email='zzz.acchefe@exemplo-teste.local', password='S3nha!teste',
                                         first_name='ZZ', last_name='Chefe', hierarchy='SUPERADMIN')
        op = User.objects.create_user(username='zzz.acop', email='zzz.acop@exemplo-teste.local', password='S3nha!teste',
                                      first_name='ZZ', last_name='Operadora', hierarchy='ADMIN')
        cc, co = Client(REMOTE_ADDR='10.2.3.4'), Client()
        cc.force_login(chefe); co.force_login(op)

        print('== QUEM ENTRA ==')
        t('o SUPERADMIN entra', cc.get('/acess/').status_code == 200)
        t('e vê o item no menu', 'href="/acess/"' in cc.get('/').content.decode())
        t('ADMIN sem liberação não entra', co.get('/acess/').status_code == 302)
        t('nem vê o menu', 'href="/acess/"' not in co.get('/').content.decode())
        cc.post('/acess/acessos/', {'user': op.id, 'permissoes': ['gefac.concluir', 'chamados.criar'], 'ativo': 'on'})
        a = AcessoAcess.objects.get(user=op)
        t('liberar uma função traz as consultas que ela exige',
          set(a.permissoes) == {'gefac.concluir', 'chamados.criar', 'colaboradores.ver', 'chamados.ver'}, a.permissoes)
        co = Client(); co.force_login(op)
        t('liberada, entra', co.get('/acess/').status_code == 200)
        t('vê a esteira', co.get('/acess/pendentes/').status_code == 200)
        t('mas não o histórico', co.get('/acess/historico/').status_code == 302)
        t('nem a tela de acessos', co.get('/acess/acessos/').status_code == 302)

        print('\n== PRÉ-CADASTRO APROVADO ENTRA NO FLUXO ==')
        novato = User.objects.create_user(username='zzz.acnovo', email='zzz.acnovo@exemplo-teste.local', password='x',
                                          first_name='Zz', last_name='Novato', cpf='529.982.247-25', phone='27 99999-0000',
                                          job_title='Consultor de Vendas', sector=es,
                                          pre_registration_status=User.PRE_REG_COMPLETED)
        r = cc.post(f'/users/manage/users/{novato.id}/pre-register/approve/')
        c1 = Colaborador.objects.filter(user=novato).first()
        t('aprovar o pré-cadastro cria a ficha', c1 is not None and c1.origem == 'PRE_CADASTRO', r.status_code)
        t('na Etapa 1, com os 13 sistemas pendentes', c1 and c1.etapa == 1 and c1.etapa_calculada == 1
          and len(c1.pendencias) == 13)
        t('com os dados do cadastro', c1 and c1.cpf == CPF1 and c1.nome == 'ZZ NOVATO' and c1.local_id == es.id)
        cc.post(f'/users/manage/users/{novato.id}/pre-register/approve/')
        t('aprovar de novo não duplica', Colaborador.objects.filter(user=novato).count() == 1)

        print('\n== AS QUATRO ETAPAS ==')
        url = f'/acess/colaborador/{c1.id}/esteira/'
        co.post(url, {'acao': 'gefac', 'versao': c1.versao})
        c1.refresh_from_db()
        t('Etapa 1 → 2: concluir GEFAC (função liberada)', c1.status_de('GEFAC') == 'OK' and c1.etapa == 2, c1.etapa)
        r = co.post(url, {'acao': 'login', 'login': '81000123', 'versao': c1.versao})
        c1.refresh_from_db()
        t('sem "preencher login", a operadora não salva o login', c1.login == '')
        cc.post(url, {'acao': 'login', 'login': 'PENDENTE', 'versao': c1.versao})
        c1.refresh_from_db()
        t('login com PENDENTE é recusado', c1.login == '')
        versao_velha = c1.versao
        cc.post(url, {'acao': 'login', 'login': '81000123.0', 'versao': c1.versao})
        c1.refresh_from_db(); novato.refresh_from_db()
        t('Etapa 2 → 3: login salvo (sem o ".0")', c1.login == '81000123' and c1.etapa == 3, (c1.login, c1.etapa))
        t('e o login vai para o cadastro do portal', novato.login_code == '81000123')
        cc.post(url, {'acao': 'academia', 'confirmo': '1', 'versao': versao_velha})
        c1.refresh_from_db()
        t('versão velha: conflito, nada muda', c1.etapa == 3 and c1.status_de('TREINAMENTO') == 'PENDENTE')
        cc.post(url, {'acao': 'academia', 'confirmo': '1', 'versao': c1.versao})
        c1.refresh_from_db()
        t('Etapa 3 → 4: Academia concluída abre a liberação', c1.etapa == 4 and c1.entrega_login == 'PENDENTE'
          and c1.status_de('TREINAMENTO') == 'OK')
        t('perfil regional do ES = Renova Allied', sv.perfil_regional_requerido(c1) == 'RENOVA_ALLIED')
        cc.post(url, {'acao': 'finalizar', 'confirmo': '1', 'versao': c1.versao})
        c1.refresh_from_db()
        t('sem checklist não finaliza', c1.entrega_login == 'PENDENTE')
        for item in ('pool', 'perfil', 'login', 'sap'):
            c1.refresh_from_db()
            cc.post(url, {'acao': 'checklist', 'item': item, 'marcado': '1', 'versao': c1.versao})
        c1.refresh_from_db()
        t('checklist completo', sv.checklist_completo(c1) and c1.check_perfil == 'RENOVA_ALLIED')
        cc.post(f'/acess/colaborador/{c1.id}/acao/', {'acao': 'transferir', 'local': sp.id, 'loja_vinculo': '', 'versao': c1.versao})
        c1.refresh_from_db()
        t('mudou para SP: o perfil passa a ser Assurant e o checklist cai', sv.perfil_regional_requerido(c1) == 'ASSURANT'
          and not sv.checklist_completo(c1))
        cc.post(url, {'acao': 'checklist', 'item': 'perfil', 'marcado': '1', 'versao': c1.versao})
        c1.refresh_from_db()
        t('remarcado com o perfil novo', c1.check_perfil == 'ASSURANT' and sv.checklist_completo(c1))
        cc.post('/acess/chamados/novo/', {'colaborador': c1.id, 'codigo': 'ZZ-LIB-1', 'motivo': 'LIBERACAO_INICIAL',
                                         'minutos': '1440'})
        ch = Chamado.objects.get(codigo='ZZ-LIB-1')
        cc.post(url, {'acao': 'finalizar', 'confirmo': '1', 'versao': c1.versao, 'chamado_versao': ch.versao})
        c1.refresh_from_db(); ch.refresh_from_db()
        t('finalizar: entrega ENVIADO, sai da esteira', c1.entrega_login == 'ENVIADO' and c1.etapa_calculada == 0 and c1.liberado_em)
        t('os sistemas pendentes viram OK', not c1.pendencias)
        t('e o chamado de liberação é concluído junto', ch.status == Chamado.CONCLUIDO)
        t('o evento LOGIN_ENVIO_CONFIRMADO fica no histórico', LogAcess.objects.filter(acao='LOGIN_ENVIO_CONFIRMADO', colaborador=c1).exists())
        t('aparece em Logins liberados', 'ZZ NOVATO' in cc.get('/acess/liberados/').content.decode())

        print('\n== VOLTAR ETAPA ==')
        c2 = sv.criar_colaborador({'nome': 'zz segundo', 'cpf': CPF2, 'telefone': '(27) 98888-1111'}, es, usuario=chefe)
        sv.concluir_gefac(c2, c2.versao, chefe)
        c2.refresh_from_db(); sv.salvar_login(c2, '81000999', c2.versao, chefe)
        c2.refresh_from_db()
        cc.post(f'/acess/colaborador/{c2.id}/esteira/', {'acao': 'voltar', 'versao': c2.versao})
        c2.refresh_from_db()
        t('3 → 2: o login sai da ficha, GEFAC fica OK', c2.etapa == 2 and c2.login == '' and c2.status_de('GEFAC') == 'OK')
        cc.post(f'/acess/colaborador/{c2.id}/esteira/', {'acao': 'voltar', 'versao': c2.versao})
        c2.refresh_from_db()
        t('2 → 1: GEFAC volta a PENDENTE', c2.etapa == 1 and c2.status_de('GEFAC') == 'PENDENTE')

        print('\n== CONCLUSÃO EM LOTE ==')
        c3 = sv.criar_colaborador({'nome': 'zz terceiro', 'cpf': CPF3, 'login': '81000777'}, sp, usuario=chefe)
        c3.sistemas = dict(c3.sistemas, TREINAMENTO='OK', SAP='BLACKLIST'); c3.save()
        cc.post('/acess/lote/', {'ids': [c3.id, c2.id]})
        c3.refresh_from_db(); c2.refresh_from_db()
        t('com login: pendências viram OK', c3.status_de('CITRIX') == 'OK')
        t('BLACKLIST não muda', c3.status_de('SAP') == 'BLACKLIST' and c3.situacao == 'BLACKLIST')
        t('sem login fica de fora', c2.status_de('CITRIX') == 'PENDENTE')

        print('\n== CHAMADOS ==')
        cc.post('/acess/chamados/novo/', {'colaborador': c2.id, 'codigo': 'ZZ-77', 'motivo': 'TROCA_SENHA', 'minutos': '60',
                                         'sistemas': ['SAP', 'AKAMAI']})
        ch2 = Chamado.objects.get(codigo='ZZ-77')
        t('abre com SLA e próxima verificação iguais', ch2.sla_original == ch2.proxima_verificacao and ch2.sistemas == ['SAP', 'AKAMAI'])
        cc.post('/acess/chamados/novo/', {'colaborador': c3.id, 'codigo': 'zz-77', 'motivo': 'FALHA', 'minutos': '60'})
        t('número repetido em chamado ativo é recusado', Chamado.objects.filter(codigo__iexact='zz-77').count() == 1)
        cc.post('/acess/chamados/novo/', {'colaborador': c3.id, 'codigo': 'ZZ-78', 'motivo': 'FALHA', 'minutos': '50000'})
        t('prazo acima de 30 dias é recusado', not Chamado.objects.filter(codigo='ZZ-78').exists())
        sla = ch2.sla_original
        cc.post(f'/acess/chamados/{ch2.id}/acao/', {'acao': 'reagendar', 'minutos': '1440', 'versao': ch2.versao})
        ch2.refresh_from_db()
        t('reagendar muda a verificação e mantém o SLA original', ch2.sla_original == sla and ch2.proxima_verificacao > sla)
        cc.post(f'/acess/chamados/{ch2.id}/acao/', {'acao': 'atribuir', 'responsavel': 'eu', 'versao': ch2.versao})
        ch2.refresh_from_db()
        t('assumir o chamado', ch2.responsavel_id == chefe.id)
        cc.post(f'/acess/chamados/{ch2.id}/acao/', {'acao': 'atribuir', 'responsavel': op.id, 'versao': ch2.versao})
        ch2.refresh_from_db()
        t('não atribui a quem não trata chamados', ch2.responsavel_id == chefe.id)
        cc.post(f'/acess/chamados/{ch2.id}/acao/', {'acao': 'concluir', 'versao': ch2.versao})
        ch2.refresh_from_db()
        t('concluir guarda quem e quando, sem trocar o responsável', ch2.status == 'CONCLUIDO' and ch2.concluido_por_id == chefe.id
          and ch2.responsavel_id == chefe.id)
        t('o chamado guarda o próprio histórico', ch2.eventos.count() == 4, ch2.eventos.count())

        print('\n== EXCLUIR E DESLIGAR ==')
        cc.post('/acess/chamados/novo/', {'colaborador': c3.id, 'codigo': 'ZZ-79', 'motivo': 'FALHA', 'minutos': '60'})
        cc.post(f'/acess/colaborador/{c3.id}/acao/', {'acao': 'excluir', 'senha': 'S3nha!teste', 'versao': c3.versao})
        c3.refresh_from_db()
        t('com chamado ativo não exclui', c3.status == 'ATIVO')
        cc.post(f'/acess/colaborador/{c2.id}/acao/', {'acao': 'excluir', 'senha': 'errada', 'versao': c2.versao})
        c2.refresh_from_db()
        t('senha errada não exclui', c2.status == 'ATIVO')
        cc.post(f'/acess/desligamentos/{c2.id}/', {'acao': 'agendar', 'saida': str(HOJE), 'execucao': str(HOJE - timedelta(days=1)),
                                                  'versao': c2.versao})
        c2.refresh_from_db()
        t('execução no passado é recusada', c2.desligamento_em is None)
        cc.post(f'/acess/desligamentos/{c2.id}/', {'acao': 'agendar', 'saida': str(HOJE), 'execucao': str(HOJE + timedelta(days=2)),
                                                  'motivo': 'pediu demissão', 'versao': c2.versao})
        c2.refresh_from_db()
        t('agendado', c2.desligamento_situacao == 'Agendado')
        cc.post(f'/acess/desligamentos/{c2.id}/', {'acao': 'confirmar', 'confirmo': '1', 'versao': c2.versao})
        c2.refresh_from_db()
        t('confirmar arquiva (antecipando)', c2.status == 'DESLIGADO' and c2.desligado_em)
        html = cc.get('/acess/colaboradores/').content.decode()
        t('e sai da lista de ativos', f'/acess/colaborador/{c2.id}/' not in html, [l for l in html.splitlines() if 'ZZ SEGUNDO' in l][:3])

        print('\n== MENSAGENS (nada é enviado) ==')
        sv.garantir_modelos()
        r = cc.post(f'/acess/colaborador/{c1.id}/whatsapp/', {'modelo': 'liberacao'})
        t('modelo com senha pede a senha', r.status_code == 400 and 'senha' in r.json()['erro'])
        c1.telefone = ''; c1.save()
        r = cc.post(f'/acess/colaborador/{c1.id}/whatsapp/', {'modelo': 'liberacao', 'senha': 'Tmp@123'})
        t('sem telefone na ficha, avisa', r.status_code == 400 and 'Telefone' in r.json()['erro'])
        c1.telefone = '(27) 99999-0000'; c1.save()
        r = cc.post(f'/acess/colaborador/{c1.id}/whatsapp/', {'modelo': 'liberacao', 'senha': 'Tmp@123'})
        j = r.json()
        t('monta o texto com nome, login e senha', j.get('ok') and '81000123' in j['texto'] and 'Tmp@123' in j['texto'], j)
        t('e o link do WhatsApp com 55 + DDD', j.get('link', '').startswith('https://wa.me/5527999990000?text='))
        t('a senha não fica gravada', not LogAcess.objects.filter(descricao__icontains='Tmp@123').exists()
          and 'Tmp@123' not in str(list(LogAcess.objects.values_list('depois', flat=True))))
        try:
            sv.validar_texto('Olá {{apelido}}')
            t('validação de variável', False)
        except sv.RegraAcess:
            t('validação de variável', True)

        print('\n== DIRETÓRIO, DUPLICADOS E EXCEL ==')
        Colaborador.objects.create(nome='ZZ COPIA', cpf=CPF1, local=es, sistemas={})
        ids = sv.duplicados(Colaborador.objects.filter(status='ATIVO'))
        t('mesmo CPF aparece nos duplicados', c1.id in ids)
        r = cc.get('/acess/colaboradores/excel/?col=nome&col=cpf&col=adabas&col=login&q=zz')
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        linhas = [[c.value for c in row] for row in ws.iter_rows()]
        t('Excel com as colunas escolhidas', linhas[0] == ['Nome completo', 'CPF', 'ADABAS', 'Login'], linhas[0])
        t('CPF e login como texto', any(l[1] == CPF1 and isinstance(l[1], str) for l in linhas[1:]))
        t('ADABAS do vínculo/local', any(l[2] == 'SPZZZ01-001' for l in linhas[1:]))

        print('\n== HISTÓRICO E TELAS ==')
        reg = LogAcess.objects.filter(acao='GEFAC_CONCLUIDO', colaborador=c1).first()
        t('o histórico guarda quem e o IP', reg and reg.usuario_id == op.id)
        reg = LogAcess.objects.filter(acao='TRANSFERENCIA', colaborador=c1).first()
        t('a transferência guarda antes e depois', reg and reg.antes.get('local') == 'Loja ZZ Capixaba' and reg.depois.get('local') == 'Loja ZZ Paulista'
          and reg.ip == '10.2.3.4', (reg.antes, reg.depois, reg.ip) if reg else None)
        for u in ('/acess/', '/acess/lojas/', '/acess/colaboradores/', '/acess/colaboradores/?dup=1', f'/acess/colaborador/{c1.id}/',
                  '/acess/pendentes/?etapa=1', '/acess/pendentes/?etapa=4', '/acess/liberados/', '/acess/chamados/?status=todos',
                  '/acess/historico/', '/acess/desligamentos/', '/acess/mensagens/', '/acess/acessos/'):
            r = cc.get(u)
            t(f'abre {u}', r.status_code == 200, r.status_code)
        c4 = sv.criar_colaborador({'nome': 'zz quarto', 'cpf': '86288366757'}, sp, usuario=chefe)
        sv.concluir_gefac(c4, c4.versao, chefe); c4.refresh_from_db()
        sv.salvar_login(c4, '81000444', c4.versao, chefe); c4.refresh_from_db()
        sv.concluir_academia(c4, c4.versao, chefe)
        html = cc.get('/acess/pendentes/?etapa=4').content.decode()
        t('a etapa 4 mostra o checklist com o perfil de SP', 'Perfil regional: Assurant' in html and 'ZZ QUARTO' in html)
        cc.post('/acess/acessos/', {'user': op.id, 'remover': '1'})
        co = Client(); co.force_login(op)
        t('tirado o acesso, não entra mais', co.get('/acess/').status_code == 302)
finally:
    transaction.set_rollback(True)
    marcador.__exit__(None, None, None)
    print('\nrollback: nada gravado; nada enviado.')

print(f'\n{ok} OK / {fail} falhas')
sys.exit(1 if fail else 0)
