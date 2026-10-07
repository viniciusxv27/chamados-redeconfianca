"""Cartões: conciliação geral (vários cartões na mesma fatura) e por cartão.

Pedido (07/10/2026): em /cartoes/, conciliar a fatura inteira — vários cartões
vêm na mesma fatura —; em /cartoes/<id>/fatura/, melhorar a tela, os dados e a
exportação, e conciliar um cartão só pegando os gastos certos pelo final.

A leitura de verdade é conferida na fatura do Itaú de 09/2026 quando o arquivo
existe em ~/Downloads (ele tem dados pessoais e não vai para o repositório):
antes, dois cartões não fechavam (as colunas das páginas cheias não eram
achadas) e os lançamentos internacionais sumiam. As telas e a exportação usam
uma leitura sintética, numa transação desfeita.
"""
import os
import sys
from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest import mock

import django

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'redeconfianca.settings')
os.environ.setdefault('RC_VARREDURA_ROTINA', '0')

from django.conf import settings

settings.CACHES = {
    'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-cf'},
    'local': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-cf2'},
}
django.setup()

from django.test.utils import setup_test_environment

setup_test_environment()
if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS.append('testserver')

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import transaction
from django.test import Client

from cartoes import views as cv
from cartoes.fatura import conciliar, janela_do_portal, ler_fatura
from cartoes.models import Cartao, Gasto

User = get_user_model()
ok = fail = 0
D = Decimal


def t(nome, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'  OK   {nome}')
    else:
        fail += 1
        print(f'  FALHA {nome} {extra}')


FATURA_REAL = os.path.expanduser('~/Downloads/Fatura Itau 092026.pdf')

print('== A FATURA DE VERDADE (Itaú 09/2026) ==')
if os.path.exists(FATURA_REAL):
    r = ler_fatura(FATURA_REAL)
    c = r['cartoes']
    t('o mês sai do vencimento impresso', r['referencia'] == date(2026, 10, 1) and r['vencimento'] == date(2026, 10, 10))
    t('11 cartões na mesma fatura', len(c) == 11, sorted(c))
    t('todos fecham com o total que a fatura declara', all(v['confere'] for v in c.values()),
      {k: v['diferenca'] for k, v in c.items() if not v['confere']})
    t('a fatura inteira fecha: R$ 79.600,66', r['total_lido'] == D('79600.66') and r['confere_total'])
    t('o final 3925 (páginas cheias) fecha: R$ 12.755,25', c['3925']['lido'] == D('12755.25'), c['3925']['lido'])
    t('o final 9503 fecha com o internacional: R$ 37.738,67 + R$ 121,22',
      c['9503']['declarado_nacional'] == D('37738.67') and c['9503']['declarado_internacional'] == D('121.22'))
    inter = [l for l in r['lancamentos'] if l['last4'] == '8865' and l['internacional']]
    t('internacionais do 8865 entram, com o IOF', sum(l['valor'] for l in inter) == D('1236.84')
      and any(l['iof'] and l['valor'] == D('41.95') for l in inter))
    t('o nome impresso de cada cartão vem junto', c['3925']['nome'].startswith('BRUNELLA'))
    # A ARTENEW cobra a parcela 02/10 agora; a 03/10 está nas "próximas faturas" e não pode entrar.
    parcelas = sorted(l['parcela'] for l in r['lancamentos'] if l['estabelecimento'] == 'ARTENEWCOMERCIO')
    t('as parcelas das próximas faturas ficam de fora', parcelas == ['02/10', '02/10'], parcelas)
    com_categoria = sum(1 for l in r['lancamentos'] if l['categoria'] or l['internacional'])
    t('quase todo lançamento nacional traz categoria', com_categoria >= len(r['lancamentos']) - 5,
      (com_categoria, len(r['lancamentos'])))
    so_um = [l for l in r['lancamentos'] if l['last4'] == '6720']
    t('conciliar um cartão só: os lançamentos do final certo', len(so_um) == 40
      and sum(l['valor'] for l in so_um) == D('8871.14'))
else:
    print('  (fatura real não está em ~/Downloads: parte pulada)')

print('\n== A CONTA DA CONCILIAÇÃO ==')
L = lambda d, v, e, **kw: dict({'last4': '4242', 'data': d, 'valor': D(v), 'estabelecimento': e, 'parcela': '',
                                'internacional': False, 'iof': False, 'categoria': 'Alimentação', 'cidade': '',
                                'detalhe': ''}, **kw)
fatura = [L(date(2026, 9, 10), '50.00', 'PADARIA BOA'), L(date(2026, 9, 12), '80.00', 'POSTO X'),
          L(date(2026, 9, 15), '30.00', 'LOJA NOVA'),
          L(date(2025, 11, 4), '19.90', 'ProdutosUOL', parcela='11/12')]
inicio, fim = janela_do_portal(fatura)
t('a janela do portal ignora a parcela de compra antiga', inicio == date(2026, 9, 7) and fim == date(2026, 9, 18),
  (inicio, fim))


class G:
    def __init__(self, d, v, e):
        self.data_gasto, self.valor, self.estabelecimento = d, D(v), e
        self.criado_por, self.ticket_id, self.descricao = None, None, ''


rel = conciliar(fatura, [G(date(2026, 9, 10), '50.00', 'Padaria Boa'), G(date(2026, 9, 12), '85.00', 'Posto X'),
                         G(date(2026, 9, 16), '12.00', 'Café')])
sit = {l['fatura']['estabelecimento'] if l['fatura'] else l['gasto'].estabelecimento: l['situacao']
       for l in rel['linhas']}
t('uma linha por item, com a situação', sit == {'PADARIA BOA': 'conferido', 'POSTO X': 'divergente',
                                                 'LOJA NOVA': 'nao_lancado', 'ProdutosUOL': 'nao_lancado',
                                                 'Café': 'sem_cobranca'}, sit)
t('percentual conciliado pelo valor', rel['percentual_conciliado'] == round(50 * 100 / 179.90, 1),
  rel['percentual_conciliado'])
t('a fatura por categoria', rel['por_categoria'][0]['categoria'] == 'Alimentação')

print('\n== PARCELAS ==')
parc = [L(date(2026, 9, 10), '50.00', 'PADARIA BOA'),
        L(date(2026, 7, 5), '100.00', 'LOJA PARCELADA', parcela='03/03'),     # compra de julho
        L(date(2026, 9, 14), '352.11', 'OBRAMAX', parcela='01/03')]
janela = janela_do_portal(parc)
gastos = [G(date(2026, 9, 10), '50.00', 'Padaria Boa'),
          G(date(2026, 7, 5), '300.00', 'Loja Parcelada'),     # lançada inteira, em julho
          G(date(2026, 7, 6), '45.00', 'Outra de julho'),       # só veio na busca da parcela
          G(date(2026, 9, 14), '1025.53', 'Obramax')]
rel = conciliar(parc, gastos, janela)
por = {l['fatura']['estabelecimento']: l for l in rel['linhas'] if l['fatura']}
t('parcela de compra antiga casa com a compra inteira lançada no mês dela',
  por['LOJA PARCELADA']['situacao'] == 'conferido' and por['LOJA PARCELADA']['parcelado'])
t('parcela contra compra de outro valor: divergência pelo total parcelado',
  por['OBRAMAX']['situacao'] == 'divergente' and por['OBRAMAX']['total_parcelado'] == D('1056.33')
  and por['OBRAMAX']['diferenca'] == D('30.80'), por['OBRAMAX'])
t('gasto antigo trazido só para casar não vira "sem cobrança"',
  not any(l['gasto'] and l['gasto'].estabelecimento == 'Outra de julho' for l in rel['linhas']))
t('"lançado no portal" soma só o período da fatura', rel['total_extrato'] == D('1075.53'), rel['total_extrato'])

print('\n== AS TELAS ==')
leitura = {
    'lancamentos': [
        L(date(2026, 9, 10), '50.00', 'PADARIA BOA'),
        L(date(2026, 9, 12), '80.00', 'POSTO X'),
        L(date(2026, 9, 20), '120.00', 'FACEBK*ADS', internacional=True, categoria=''),
        L(date(2026, 9, 20), '4.20', 'Repasse de IOF (compras internacionais)', internacional=True, iof=True,
          categoria='IOF'),
        dict(L(date(2026, 9, 11), '999.00', 'COISA DO OUTRO'), last4='5151'),
    ],
    'cartoes': {
        '4242': {'nome': 'ZZ TITULAR', 'declarado': D('254.20'), 'declarado_nacional': D('130.00'),
                 'declarado_internacional': D('124.20'), 'lido': D('254.20'), 'lido_nacional': D('130.00'),
                 'lido_internacional': D('124.20'), 'lancamentos': 4, 'internacionais': 2, 'confere': True,
                 'diferenca': D('0')},
        '5151': {'nome': 'ZZ ADICIONAL', 'declarado': D('999.00'), 'declarado_nacional': D('999.00'),
                 'declarado_internacional': None, 'lido': D('999.00'), 'lido_nacional': D('999.00'),
                 'lido_internacional': D('0'), 'lancamentos': 1, 'internacionais': 0, 'confere': True,
                 'diferenca': D('0')},
    },
    'referencia': date(2026, 10, 1), 'vencimento': date(2026, 10, 10), 'emissao': None,
    'total_declarado': D('1253.20'), 'total_lido': D('1253.20'), 'confere_total': True, 'confere': True,
}
pdf = lambda: SimpleUploadedFile('fatura.pdf', b'%PDF-1.4 zz', content_type='application/pdf')

marcador = transaction.atomic()
marcador.__enter__()
try:
    chefe = User.objects.create_user(username='zzcf.chefe', email='zzcf.chefe@exemplo-teste.local',
                                     password='x', hierarchy='SUPERADMIN', is_superuser=True, first_name='Chefe')
    dono = User.objects.create_user(username='zzcf.dono', email='zzcf.dono@exemplo-teste.local',
                                    password='x', first_name='Dono')
    cartao = Cartao.objects.create(apelido='ZZ Cartão Teste', first4='5234', last4='4242', responsavel=dono,
                                   validade_mes=1, validade_ano=2030, bandeira='MASTERCARD')
    Gasto.objects.create(cartao=cartao, criado_por=dono, valor=D('50.00'), estabelecimento='Padaria Boa',
                         data_gasto=date(2026, 9, 10))
    Gasto.objects.create(cartao=cartao, criado_por=dono, valor=D('85.00'), estabelecimento='Posto X',
                         data_gasto=date(2026, 9, 12))
    c = Client()
    c.force_login(chefe)

    with mock.patch.object(cv, 'ler_fatura', return_value=leitura):
        r = c.post('/cartoes/fatura/', {'fatura': pdf()})
    t('a conciliação geral lê e volta para a tela', r.status_code == 302 and r['Location'] == '/cartoes/fatura/')
    r = c.get('/cartoes/fatura/')
    html = r.content.decode()
    t('a tela geral lista os dois finais', 'final 4242' in html and 'final 5151' in html)
    t('o 4242 casa com o cartão cadastrado', 'ZZ Cartão Teste' in html)
    t('o 5151 aparece como não cadastrado', 'Não cadastrado no portal' in html)
    linha = next(l for l in r.context['linhas_geral'] if l['final'] == '4242')
    t('e é conciliado com os gastos dele', linha['relatorio'] and len(linha['relatorio']['conferidos']) == 1
      and len(linha['relatorio']['divergentes']) == 1)
    t('a fatura inteira fecha', 'Fatura inteira conferida' in html)

    r = c.get('/cartoes/fatura/exportar/')
    from openpyxl import load_workbook
    livro = load_workbook(BytesIO(r.content))
    t('exporta a geral em Excel', r.status_code == 200 and livro.sheetnames == ['Resumo', 'Conciliação',
                                                                                'Cartões fora do portal'],
      livro.sheetnames)
    aba = livro['Conciliação']
    cartoes_na_aba = {aba.cell(row=i, column=1).value for i in range(2, aba.max_row + 1)} - {None}
    t('a aba de conciliação tem só os cadastrados, com o cartão', cartoes_na_aba == {'4242 · ZZ Cartão Teste'},
      cartoes_na_aba)

    r = c.get(f'/cartoes/{cartao.pk}/fatura/?geral=1')
    html = r.content.decode()
    t('detalhar abre o cartão sem subir o PDF de novo', r.status_code == 200 and 'relatorio' in r.context)
    estab = [l['fatura']['estabelecimento'] for l in r.context['relatorio']['linhas'] if l['fatura']]
    t('só os lançamentos do final 4242', 'COISA DO OUTRO' not in estab and 'PADARIA BOA' in estab, estab)
    t('o internacional e o IOF entram no cartão', 'FACEBK*ADS' in estab and any('IOF' in e for e in estab))
    t('mostra os outros cartões da mesma fatura', 'Outros cartões nesta fatura' in html and 'final 5151' in html)
    t('filtros por situação na tabela única', 'data-filtro="divergente"' in html and 'concBusca' in html)

    with mock.patch.object(cv, 'ler_fatura', return_value=leitura):
        r = c.post(f'/cartoes/{cartao.pk}/fatura/', {'fatura': pdf()})
    t('subir a fatura de vários cartões na tela do cartão pega só o final dele',
      r.status_code == 200 and r.context['relatorio']['total_fatura'] == D('254.20'),
      r.context['relatorio']['total_fatura'] if r.status_code == 200 else r.status_code)
    r = c.get(f'/cartoes/{cartao.pk}/fatura/exportar/')
    livro = load_workbook(BytesIO(r.content))
    t('exporta o cartão em Excel (resumo + conciliação única)', livro.sheetnames == ['Resumo', 'Conciliação'],
      livro.sheetnames)
    resumo = '\n'.join(str(livro['Resumo'].cell(row=i, column=1).value) for i in range(1, livro['Resumo'].max_row + 1))
    t('o resumo traz a conferência da leitura e as categorias',
      'Leitura conferida' in resumo and 'Categoria (fatura)' in resumo)

    cd = Client()
    cd.force_login(dono)
    r = cd.get('/cartoes/fatura/')
    t('a conciliação geral é de quem gere os cartões', r.status_code == 302)
    with mock.patch.object(cv, 'ler_fatura', return_value=leitura):
        r = cd.post(f'/cartoes/{cartao.pk}/fatura/', {'fatura': pdf()})
    t('o responsável concilia o cartão dele', r.status_code == 200 and r.context['relatorio'] is not None)
    t('sem ver o link para os outros cartões que não são dele',
      all(o['cartao'] is None for o in r.context['outros_cartoes']))

    print('\n== LIMPAR A CONCILIAÇÃO (SUPERADMIN) ==')
    r = cd.post('/cartoes/fatura/limpar/')
    t('quem não é SUPERADMIN não limpa', r.status_code == 302 and 'cartoes_conciliacao' in cd.session)
    with mock.patch.object(cv, 'ler_fatura', return_value=leitura):
        c.post('/cartoes/fatura/', {'fatura': pdf()})
        c.post(f'/cartoes/{cartao.pk}/fatura/', {'fatura': pdf()})
    html = c.get(f'/cartoes/{cartao.pk}/fatura/?geral=1').content.decode()
    t('o SUPERADMIN vê o botão de limpar', 'Limpar conciliação' in html)
    r = c.post('/cartoes/fatura/limpar/', {'cartao': cartao.pk})
    t('limpar um cartão tira só a conciliação dele',
      r['Location'] == f'/cartoes/{cartao.pk}/fatura/'
      and str(cartao.pk) not in (c.session.get('cartoes_conciliacao') or {})
      and 'cartoes_fatura_geral' in c.session)
    r = c.get(f'/cartoes/{cartao.pk}/fatura/exportar/')
    t('e a exportação dele deixa de existir', r.status_code == 302)
    r = c.post('/cartoes/fatura/limpar/')
    t('limpar a fatura tira a conciliação geral', r['Location'] == '/cartoes/fatura/'
      and 'cartoes_fatura_geral' not in c.session)
    html = c.get('/cartoes/fatura/').content.decode()
    t('a tela volta vazia, pronta para outra fatura', 'Cartões da fatura' not in html and 'Limpar conciliação' not in html)
    t('limpar não apaga nada do extrato', Gasto.objects.filter(cartao=cartao).count() == 2)
    t('limpar só por POST', c.get('/cartoes/fatura/limpar/').status_code == 405)
finally:
    transaction.set_rollback(True)
    marcador.__exit__(None, None, None)

print(f'\n{ok} OK / {fail} falhas — rollback: nada deste teste ficou no banco.')
sys.exit(1 if fail else 0)
