"""Assistente de Apresentações — planilha como material base (xlsx, csv, ods).

Gera as planilhas em memória, confere o texto tabular que vai para a IA (abas, cabeçalho,
linhas, valor de fórmula, datas, cortes avisados) e sobe pela tela /apresentacoes/nova/.
Roda dentro de uma transação desfeita no fim: não grava nada no banco. Nenhuma chamada
à OpenAI sai (dublê que falha se for chamado), as tarefas não disparam e os arquivos vão
para um armazenamento em memória.
"""
import io
import os
import re
import sys
import zipfile
from datetime import date, datetime
from unittest import mock

import django

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'redeconfianca.settings')
django.setup()

from django.conf import settings

if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS.append('testserver')
# O bloqueio de cursos acorda a cobrança automática no WhatsApp a cada requisição logada, e o
# menu de cursos lê a configuração deles (coluna nova ainda sem migrate aborta a transação do
# teste): os dois ficam fora deste teste, que é só das apresentações.
settings.MIDDLEWARE = [m for m in settings.MIDDLEWARE if m != 'cursos.middleware.BloqueioCursoMiddleware']
for _motor in settings.TEMPLATES:
    _processadores = _motor.get('OPTIONS', {}).get('context_processors', [])
    _processadores[:] = [p for p in _processadores if not p.startswith('cursos.')]

from django.test.utils import setup_test_environment

setup_test_environment()

from django.contrib.auth import get_user_model
from django.core.cache import caches
from django.core.files.storage import InMemoryStorage
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import transaction
from django.test import Client
from openpyxl import Workbook

from apresentacoes import importacao, roteiro
from apresentacoes.models import Apresentacao, ConfiguracaoApresentacoes, Midia, TarefaIA
from apresentacoes.padrao import garantir_template_padrao
from users.models import Sector

User = get_user_model()
ok = fail = 0


def t(nome, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'  OK   {nome}')
    else:
        fail += 1
        print(f'  FALHA {nome} {extra}')


memoria = InMemoryStorage()
Midia._meta.get_field('arquivo').storage = memoria


# ---------------------------------------------------------------------------
# Planilhas de exemplo (em memória)
# ---------------------------------------------------------------------------
def xlsx(livro, trocas=()):
    """Salva o livro e aplica trocas no XML das abas (o openpyxl não calcula fórmula: o valor
    em cache que o Excel gravaria entra à mão, e dá para estragar a dimensão de propósito)."""
    bruto = io.BytesIO()
    livro.save(bruto)
    entrada, saida = zipfile.ZipFile(io.BytesIO(bruto.getvalue())), io.BytesIO()
    with zipfile.ZipFile(saida, 'w', zipfile.ZIP_DEFLATED) as novo:
        for nome in entrada.namelist():
            dados = entrada.read(nome)
            if nome.startswith('xl/worksheets/'):
                texto = dados.decode()
                for antes, depois in trocas:
                    texto = re.sub(antes, depois, texto)
                dados = texto.encode()
            novo.writestr(nome, dados)
    return saida.getvalue()


def livro_de_precos():
    livro = Workbook()
    aba = livro.active
    aba.title = 'Preços'
    aba.append(['Plano', 'Preço', 'Linhas', 'Total', 'Início', 'Ativo'])
    aba.append(['Vivo Total', 199.9, 2, '=B2*C2', date(2026, 9, 1), True])
    aba.append([None, None, None, None, None, None])          # linha vazia no meio: some
    aba.append(['Controle | Família', 0.15, 4, None, datetime(2026, 9, 15, 14, 30), False])
    aba.append(['Pós\ncom quebra', 59, 1])
    metas = livro.create_sheet('Metas')
    metas.append(['Loja', 'Meta'])
    metas.append(['Centro', 120])
    livro.create_sheet('Vazia')
    # Valor da fórmula como o Excel deixaria salvo: 199,9 × 2.
    return xlsx(livro, [(r'<f>B2\*C2</f><v\s*/?>(</v>)?', '<f>B2*C2</f><v>399.8</v>')])


def ods(linhas_por_aba):
    """ODS mínimo escrito à mão (sem odfpy), com célula repetida e linhas vazias repetidas."""
    corpo = []
    for nome, linhas in linhas_por_aba:
        xml = [f'<table:table table:name="{nome}">']
        for linha in linhas:
            xml.append('<table:table-row>')
            for valor in linha:
                if isinstance(valor, (int, float)):
                    xml.append(f'<table:table-cell office:value-type="float" office:value="{valor}">'
                               f'<text:p>{valor}</text:p></table:table-cell>')
                elif isinstance(valor, tuple):                   # (texto, repetições)
                    xml.append(f'<table:table-cell table:number-columns-repeated="{valor[1]}" '
                               f'office:value-type="string"><text:p>{valor[0]}</text:p></table:table-cell>')
                else:
                    xml.append(f'<table:table-cell office:value-type="string"><text:p>{valor}</text:p>'
                               '</table:table-cell>')
            xml.append('<table:table-cell table:number-columns-repeated="16000"/></table:table-row>')
        xml.append('<table:table-row table:number-rows-repeated="1048000"><table:table-cell/></table:table-row>')
        xml.append('</table:table>')
        corpo.append(''.join(xml))
    content = ('<?xml version="1.0" encoding="UTF-8"?><office:document-content '
               'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
               'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
               'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:spreadsheet>'
               + ''.join(corpo) + '</office:spreadsheet></office:body></office:document-content>')
    saida = io.BytesIO()
    with zipfile.ZipFile(saida, 'w') as pacote:
        pacote.writestr('mimetype', 'application/vnd.oasis.opendocument.spreadsheet')
        pacote.writestr('content.xml', content)
    return saida.getvalue()


CSV_PONTO_E_VIRGULA = 'Loja;Vendas;Valor\nCentro;12;R$ 1.234,56\nNorte;7;R$ 980,00\n'.encode('utf-8-sig')
CSV_LATIN1 = 'Produto,Promoção\nChip,"Grátis, no combo"\n'.encode('latin-1')


print('== XLSX ==')
texto = importacao.texto_de_planilha(livro_de_precos(), '.xlsx')
print('    ' + texto.replace('\n', '\n    '))
t('explica o formato no começo', texto.startswith('Planilha (colunas separadas por " | "'))
t('cada aba com nome e contagem', 'Aba "Preços" (4 linhas com conteúdo):' in texto
  and 'Aba "Metas" (2 linhas com conteúdo):' in texto and 'Aba "Vazia": (vazia)' in texto)
t('cabeçalho marcado', 'Cabeçalho: Plano | Preço | Linhas | Total | Início | Ativo' in texto)
t('fórmula entra pelo valor calculado (data_only), data no formato daqui, bool em português',
  'Linha 2: Vivo Total | 199.9 | 2 | 399.8 | 01/09/2026 | sim' in texto and 'B2*C2' not in texto)
t('linha vazia não conta; "|" da célula não quebra a coluna; data com hora',
  'Linha 3: Controle / Família | 0.15 | 4 |  | 15/09/2026 14:30 | não' in texto)
t('quebra de linha na célula vira espaço e colunas vazias do fim somem', 'Linha 4: Pós com quebra | 59 | 1\n' in texto)
t('número inteiro sem ".0"', 'Linha 2: Centro | 120' in texto)

livro = Workbook()
for i in range(1, 801):
    livro.active.append([f'item {i}', i])
grande = xlsx(livro, [(r'<dimension ref="[^"]+"', '<dimension ref="A1"')])
texto = importacao.texto_de_planilha(grande, '.xlsx')
t('dimensão gravada errada ("A1") não esconde linhas', 'Aba "Sheet" (800 linhas com conteúdo)' in texto, texto[:200])
t('aba longa corta em 500 linhas e avisa', 'Linha 500: item 500 | 500' in texto and 'Linha 501' not in texto
  and '[aba cortada: só as primeiras 500 de 800 linhas foram enviadas]' in texto)

livro = Workbook()
livro.active.append([f'C{i}' for i in range(1, 61)])
texto = importacao.texto_de_planilha(xlsx(livro), '.xlsx')
t('colunas além de 40 cortadas com aviso', '| C40\n' in texto and 'C41' not in texto
  and '[colunas cortadas: só as primeiras 40 foram enviadas]' in texto)

livro = Workbook()
for i in range(25):
    (livro.active if i == 0 else livro.create_sheet()).append([f'aba {i}'])
texto = importacao.texto_de_planilha(xlsx(livro), '.xlsx')
t('mais de 20 abas: corta e avisa', texto.count('Aba "') == 20
  and '[planilha cortada: só as primeiras 20 abas foram enviadas]' in texto)

texto = importacao.texto_de_planilha(grande, '.xlsx', limite=3000)
t('limite de caracteres respeitado, cortando em linha inteira e avisando', len(texto) <= 3000
  and texto.endswith('[planilha cortada: passou de 3.000 caracteres; o restante não foi enviado]')
  and re.search(r'\nLinha \d+: item \d+ \| \d+\n\[planilha cortada', texto), texto[-200:])

print('\n== CSV ==')
texto = importacao.texto_de_planilha(CSV_PONTO_E_VIRGULA, '.csv')
t('";" detectado, BOM some, vírgula decimal fica no valor',
  'Aba "CSV" (3 linhas com conteúdo)' in texto and 'Cabeçalho: Loja | Vendas | Valor' in texto
  and 'Linha 2: Centro | 12 | R$ 1.234,56' in texto, texto)
texto = importacao.texto_de_planilha(CSV_LATIN1, '.csv')
t('latin-1 e "," com aspas', 'Cabeçalho: Produto | Promoção' in texto and 'Linha 2: Chip | Grátis, no combo' in texto, texto)
texto = importacao.texto_de_planilha('Nome\nAna\nBia\n'.encode(), '.csv')
t('uma coluna só (sem delimitador para farejar)', 'Cabeçalho: Nome\nLinha 2: Ana\nLinha 3: Bia' in texto, texto)

print('\n== ODS ==')
texto = importacao.texto_de_planilha(ods([('Ofertas', [['Plano', 'Preço'], ['Total', 199.9], [('x', 3), 2]]),
                                          ('Outra', [['a']])]), '.ods')
t('ODS: abas, números e célula repetida; repetições vazias gigantes não pesam',
  'Aba "Ofertas" (3 linhas com conteúdo)' in texto and 'Linha 2: Total | 199.9' in texto
  and 'Linha 3: x | x | x | 2' in texto and 'Aba "Outra" (1 linhas com conteúdo)' in texto, texto)

print('\n== TELA /apresentacoes/nova/ ==')
marcador = transaction.atomic()
marcador.__enter__()
patches = []
try:
    disparadas = []
    for alvo, novo in (
        ('apresentacoes.ia._cliente', mock.MagicMock(side_effect=AssertionError('OpenAI não pode ser chamada'))),
        ('apresentacoes.ia._http', mock.MagicMock(side_effect=AssertionError('HTTP real não pode sair no teste'))),
        ('apresentacoes.tarefas.disparar', lambda tarefa: disparadas.append(tarefa.pk)),
        ('notifications.services.notification_service._send_in_app', mock.MagicMock()),
    ):
        p = mock.patch(alvo, novo)
        p.start()
        patches.append(p)

    area = Sector.objects.create(name='ZZ Setor Apresentacoes Planilha')
    ana = User.objects.create_user(username='zzapp.ana', email='zzapp.ana@exemplo-teste.local', password='S3nha!teste',
                                   sector=area, first_name='ZZApAna', last_name='Teste')
    ConfiguracaoApresentacoes.get().liberados.add(ana)
    caches['local'].delete(f'apresentacoes:menu:{ana.pk}')
    c_ana = Client()
    c_ana.force_login(ana)
    padrao = garantir_template_padrao()

    html = c_ana.get('/apresentacoes/nova/').content.decode()
    t('o campo de anexos aceita planilha', '.pptx,.xlsx,.xlsm,.csv,.ods"' in html)
    t('e a ajuda lista os formatos', 'planilhas (Excel .xlsx, .csv, .ods' in html and '.xls antigo: salve como .xlsx' in html)

    r = c_ana.post('/apresentacoes/nova/', {
        'pedido': 'Monte a apresentação da tabela de preços para as lojas.', 'template': padrao.pk,
        'anexos': [SimpleUploadedFile('precos.xlsx', livro_de_precos(),
                                      content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                   SimpleUploadedFile('vendas.csv', CSV_PONTO_E_VIRGULA, content_type='text/csv'),
                   SimpleUploadedFile('antiga.xls', b'\xd0\xcf\x11\xe0velho', content_type='application/vnd.ms-excel'),
                   SimpleUploadedFile('quebrada.xlsx', b'nao sou zip', content_type='application/octet-stream')]})
    ap = Apresentacao.objects.filter(dono=ana).order_by('-pk').first()
    t('cria a apresentação e abre o editor', r.status_code == 302 and ap and r['Location'].endswith(f'/apresentacoes/{ap.pk}/'))
    material = (ap.opcoes or {}).get('texto_material', '')
    t('o texto da planilha xlsx vai para o material',
      '[precos.xlsx]\nPlanilha (colunas' in material and 'Linha 2: Vivo Total | 199.9 | 2 | 399.8' in material, material[:300])
    t('e o do csv também', '[vendas.csv]\n' in material and 'Linha 2: Centro | 12 | R$ 1.234,56' in material)
    t('planilha não vira mídia (é só texto)', not Midia.objects.filter(apresentacao=ap).exists())
    avisos = [str(m) for m in r.wsgi_request._messages]
    t('.xls antigo recusado com orientação', any('antiga.xls' in a and 'salve como .xlsx' in a for a in avisos), avisos)
    t('xlsx que não abre fica de fora com aviso', any('"quebrada.xlsx" não pôde ser lido' in a for a in avisos), avisos)
    t('a tarefa de gerar fica na fila (sem disparar a IA aqui)', TarefaIA.objects.filter(apresentacao=ap).exists()
      and disparadas)
    pedido_ia = roteiro.pedido_de_geracao(ap, [], '', ap.opcoes.get('texto_material', ''))
    t('o pedido montado para a IA leva a tabela', 'Texto extraído do material enviado:' in str(pedido_ia)
      and 'Cabeçalho: Plano | Preço' in str(pedido_ia))

    # Vários anexos somados passam de 30.000: o corte fica avisado no próprio texto.
    livro = Workbook()
    for i in range(1, 501):
        livro.active.append([f'linha comprida número {i} ' + 'x' * 60, i])
    pesada = xlsx(livro)
    r = c_ana.post('/apresentacoes/nova/', {
        'pedido': 'Duas planilhas grandes de uma vez.', 'template': padrao.pk,
        'anexos': [SimpleUploadedFile('a.xlsx', pesada), SimpleUploadedFile('b.xlsx', pesada)]})
    ap2 = Apresentacao.objects.filter(dono=ana).order_by('-pk').first()
    material = ap2.opcoes.get('texto_material', '')
    t('soma dos anexos acima de 30.000: corta e avisa', ap2.pk != ap.pk and len(material) <= 30000
      and material.endswith('[material cortado: os anexos somados passaram de 30.000 caracteres]'), len(material))

finally:
    for p in patches:
        p.stop()
    transaction.set_rollback(True)
    marcador.__exit__(None, None, None)
    print('\nrollback: nada deste teste foi gravado no banco.')

print(f'\n{ok} OK / {fail} falhas')
sys.exit(1 if fail else 0)
