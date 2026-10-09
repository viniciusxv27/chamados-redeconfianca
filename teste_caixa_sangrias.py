"""Contagem de caixa: aba Gestão de Sangrias.

Pedido: registrar cada sangria (data, loja, valor, categoria, descrição…), com o
valor entrando sozinho na composição e na conciliação do caixa; relatórios por
loja, período e categoria; achar gastos recorrentes e fora do padrão; conferir
de forma rápida e rastreável.

O que este teste cobre:

- registrar pela tela: o dia de caixa ganha a soma em ``sangria_registrada``,
  a Entrada cai e o saldo dos dias seguintes é refeito;
- sangria em dia ainda sem importação cria o dia, e a importação depois não
  apaga a sangria;
- editar trocando de dia refaz os dois dias; apagar devolve o caixa;
- validações (valor, data futura, descrição, loja de fora);
- permissões: o vendedor da loja registra e altera a própria enquanto não
  conferida; não confere, não vê outra loja nem mexe em categorias;
- conferência em lote com quem e quando, e desfazer;
- relatórios: resumo, por categoria, por loja, recorrentes e fora do padrão;
- a planilha exportada e a coluna na tela da loja.

Caches em memória, transação desfeita no fim: não grava nada no banco.
"""
import os
import sys
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import django

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'redeconfianca.settings')
os.environ.setdefault('RC_VARREDURA_ROTINA', '0')

from django.conf import settings

settings.CACHES = {
    'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-cx-sg'},
    'local': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-cx-sg-2'},
}
django.setup()

from django.test.utils import setup_test_environment

setup_test_environment()
if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS.append('testserver')

from django.contrib.auth import get_user_model
from django.db import transaction
from django.test import Client
from django.utils import timezone

from contagem_caixa import sangrias as svc
from contagem_caixa.models import CategoriaSangria, ContagemCaixaDia, Sangria
from contagem_caixa.servicos import recalcular_saldos
from users.models import Sector

User = get_user_model()
D = Decimal
ok = fail = 0


def t(nome, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'  OK   {nome}')
    else:
        fail += 1
        print(f'  FALHA {nome} {extra}')


def dia(loja, data):
    return ContagemCaixaDia.objects.filter(loja=loja, data=data).first()


marcador = transaction.atomic()
marcador.__enter__()
try:
    loja = Sector.objects.create(name='Loja ZZ Sangria A', adabas='ZZ981')
    outra = Sector.objects.create(name='Loja ZZ Sangria B', adabas='ZZ982')
    chefe = User.objects.create_user(
        username='zz.sg.chefe', email='zz.sg.chefe@exemplo-teste.local', password='S3nha!teste',
        first_name='Zz', last_name='Chefe', sector=loja, hierarchy='SUPERADMIN',
        is_staff=True, is_superuser=True)
    vendedor = User.objects.create_user(
        username='zz.sg.vend', email='zz.sg.vend@exemplo-teste.local', password='S3nha!teste',
        first_name='Zz', last_name='Vendedor', sector=loja, hierarchy='SUPERVISOR')
    colega = User.objects.create_user(
        username='zz.sg.colega', email='zz.sg.colega@exemplo-teste.local', password='S3nha!teste',
        first_name='Zz', last_name='Colega', sector=loja, hierarchy='SUPERVISOR')

    hoje = timezone.localdate()
    d1, d2, d3 = hoje - timedelta(days=3), hoje - timedelta(days=2), hoje - timedelta(days=1)
    for d, sap in ((d1, '1000.00'), (d2, '500.00'), (d3, '800.00')):
        ContagemCaixaDia.objects.create(loja=loja, data=d, valor_sap=D(sap))
    recalcular_saldos(loja.id)
    saldo_d3_antes = dia(loja, d3).saldo

    cat = {c.nome: c for c in CategoriaSangria.objects.all()}
    print('== CATEGORIAS ==')
    t('as 10 categorias da migração existem', len(cat) >= 10, list(cat))
    alim, limpeza, outros = cat['Alimentação'], cat['Limpeza e higiene'], cat['Outros']

    c_chefe, c_vend, c_colega = Client(), Client(), Client()
    c_chefe.force_login(chefe)
    c_vend.force_login(vendedor)
    c_colega.force_login(colega)

    print('== REGISTRAR ==')
    r = c_vend.get('/contagem-caixa/sangrias/')
    t('vendedor abre a aba', r.status_code == 200, r.status_code)
    html = r.content.decode()
    t('formulário oferece só a loja dele', 'Loja ZZ Sangria A' in html and 'Loja ZZ Sangria B' not in html)
    t('aba Sangrias na navegação', 'fas fa-hand-holding-dollar mr-1.5"></i>Sangrias' in html)

    r = c_vend.post('/contagem-caixa/sangrias/registrar/', {
        'loja': loja.id, 'data': d2.isoformat(), 'valor': '45,90', 'categoria': alim.id,
        'descricao': 'Lanche da equipe', 'favorecido': 'Padaria Zz'})
    t('registrar redireciona', r.status_code == 302, r.status_code)
    s1 = Sangria.objects.filter(loja=loja, descricao='Lanche da equipe').first()
    t('sangria gravada com autor', s1 and s1.valor == D('45.90') and s1.registrada_por == vendedor)
    t('dia ganhou sangria_registrada', dia(loja, d2).sangria_registrada == D('45.90'), dia(loja, d2).sangria_registrada)
    t('a Entrada do dia caiu', dia(loja, d2).entrada == D('454.10'), dia(loja, d2).entrada)
    t('o saldo do dia seguinte caiu junto', dia(loja, d3).saldo == saldo_d3_antes - D('45.90'),
      (dia(loja, d3).saldo, saldo_d3_antes))

    c_vend.post('/contagem-caixa/sangrias/registrar/', {
        'loja': loja.id, 'data': d2.isoformat(), 'valor': '10', 'categoria': limpeza.id,
        'descricao': 'Detergente'})
    t('duas no mesmo dia somam', dia(loja, d2).sangria_registrada == D('55.90'))
    t('sangria/erro manual não foi tocada', dia(loja, d2).sangria_erro in (None, D('0')))

    print('== VALIDAÇÕES ==')
    antes = Sangria.objects.count()
    for nome, dados in (
        ('valor zero', {'valor': '0'}),
        ('valor inválido', {'valor': 'abc'}),
        ('data futura', {'data': (hoje + timedelta(days=1)).isoformat()}),
        ('sem descrição', {'descricao': ''}),
        ('loja de fora', {'loja': outra.id}),
        ('sem categoria', {'categoria': ''}),
    ):
        base = {'loja': loja.id, 'data': d1.isoformat(), 'valor': '5', 'categoria': outros.id, 'descricao': 'x'}
        base.update(dados)
        c_vend.post('/contagem-caixa/sangrias/registrar/', base)
        t(f'recusa {nome}', Sangria.objects.count() == antes)

    print('== DIA SEM IMPORTAÇÃO ==')
    d0 = hoje - timedelta(days=10)
    c_vend.post('/contagem-caixa/sangrias/registrar/', {
        'loja': loja.id, 'data': d0.isoformat(), 'valor': '20', 'categoria': outros.id, 'descricao': 'Chaveiro'})
    t('cria o dia de caixa com a sangria', dia(loja, d0) and dia(loja, d0).sangria_registrada == D('20.00'))
    t('e a entrada fica negativa (dinheiro saiu)', dia(loja, d0).entrada == D('-20.00'), dia(loja, d0).entrada)
    # A importação só mexe no Valor SAP.
    x = dia(loja, d0)
    x.valor_sap = D('300.00')
    x.save(update_fields=['valor_sap', 'atualizado_em'])
    t('importação depois não apaga a sangria', dia(loja, d0).sangria_registrada == D('20.00'))

    print('== EDITAR / APAGAR ==')
    r = c_vend.post(f'/contagem-caixa/sangrias/{s1.id}/editar/', {
        'loja': loja.id, 'data': d1.isoformat(), 'valor': '50,00', 'categoria': alim.id,
        'descricao': 'Lanche da equipe', 'favorecido': 'Padaria Zz'})
    s1.refresh_from_db()
    t('autor edita a dele', s1.valor == D('50.00') and s1.data == d1)
    t('dia antigo ficou só com o detergente', dia(loja, d2).sangria_registrada == D('10.00'))
    t('dia novo recebeu', dia(loja, d1).sangria_registrada == D('50.00'))
    t('saldo final confere (60 nos dias + 20 do dia sem importação, que vem antes)', dia(loja, d3).saldo == saldo_d3_antes - D('80.00'),
      (dia(loja, d3).saldo, saldo_d3_antes))

    c_colega.post(f'/contagem-caixa/sangrias/{s1.id}/editar/', {
        'loja': loja.id, 'data': d1.isoformat(), 'valor': '1', 'categoria': alim.id, 'descricao': 'x'})
    s1.refresh_from_db()
    t('colega não edita a sangria de outro', s1.valor == D('50.00'))
    c_colega.post(f'/contagem-caixa/sangrias/{s1.id}/apagar/')
    t('colega não apaga', Sangria.objects.filter(id=s1.id).exists())

    print('== CONFERÊNCIA ==')
    c_vend.post('/contagem-caixa/sangrias/conferir/', {'sangria': [s1.id], 'acao': 'conferir'})
    s1.refresh_from_db()
    t('vendedor não confere', not s1.conferida)
    r = c_chefe.post('/contagem-caixa/sangrias/conferir/', {'sangria': [s1.id], 'acao': 'conferir'})
    s1.refresh_from_db()
    t('gestor confere com quem e quando', s1.conferida and s1.conferida_por == chefe and s1.conferida_em)
    c_vend.post(f'/contagem-caixa/sangrias/{s1.id}/apagar/')
    t('conferida, o autor não apaga mais', Sangria.objects.filter(id=s1.id).exists())
    c_chefe.post('/contagem-caixa/sangrias/conferir/', {'sangria': [s1.id], 'acao': 'desfazer'})
    s1.refresh_from_db()
    t('desfazer conferência', not s1.conferida and s1.conferida_por is None)

    detergente = Sangria.objects.get(descricao='Detergente')
    c_vend.post(f'/contagem-caixa/sangrias/{detergente.id}/apagar/')
    t('autor apaga a dele', not Sangria.objects.filter(id=detergente.id).exists())
    t('e o dia volta a zero', dia(loja, d2).sangria_registrada == D('0.00'))

    print('== CATEGORIAS (GESTOR) ==')
    c_vend.post('/contagem-caixa/sangrias/categorias/', {'nova': 'Zz Categoria Vendedor'})
    t('vendedor não cria categoria', not CategoriaSangria.objects.filter(nome='Zz Categoria Vendedor').exists())
    c_chefe.post('/contagem-caixa/sangrias/categorias/', {'nova': 'Zz Estacionamento'})
    nova = CategoriaSangria.objects.filter(nome='Zz Estacionamento').first()
    t('gestor cria categoria', nova is not None)
    c_chefe.post('/contagem-caixa/sangrias/categorias/', {f'nome_{nova.id}': 'Zz Estacionamento'})
    nova.refresh_from_db()
    t('desmarcar "ativa" desativa', not nova.ativa)

    print('== RELATÓRIOS ==')
    agora = timezone.now()
    for i in range(4):
        Sangria.objects.create(loja=loja, data=d3, valor=D('12.00'), categoria=limpeza,
                               descricao=f'Água sanitária {i}', favorecido='Mercado Zz', registrada_por=vendedor)
    Sangria.objects.create(loja=loja, data=d3, valor=D('300.00'), categoria=limpeza,
                           descricao='Lavagem de carpete', registrada_por=vendedor)
    Sangria.objects.create(loja=outra, data=d3, valor=D('33.00'), categoria=outros,
                           descricao='Da outra loja', registrada_por=chefe)
    qs = Sangria.objects.filter(loja__in=[loja, outra])
    n = svc.resumo(qs)
    t('resumo soma', n['total'] == D('50') + D('20') + D('48') + D('300') + D('33') and n['quantidade'] == 8,
      (n['total'], n['quantidade']))
    t('sem comprovante conta todas', n['sem_comprovante'] == 8, n['sem_comprovante'])
    cats = svc.por_categoria(qs, n['total'])
    t('por categoria ordena pelo maior', cats[0]['nome'] == 'Limpeza e higiene' and cats[0]['barra'] == 100, cats[0])
    lojas_r = svc.por_loja(qs, n['total'])
    t('por loja tem as duas', {l['nome'] for l in lojas_r} == {'Loja ZZ Sangria A', 'Loja ZZ Sangria B'})
    rec = svc.recorrentes(qs)
    t('recorrente: Mercado Zz 4x', any(g['rotulo'] == 'Mercado Zz' and g['n'] == 4 for g in rec), rec)
    fora = svc.fora_do_padrao(qs)
    t('fora do padrão: carpete', [s.descricao for s in fora] == ['Lavagem de carpete'], [s.descricao for s in fora])
    t('com a mediana da categoria', fora and fora[0].mediana_categoria == D('12.00') and fora[0].vezes_mediana == 25.0)

    print('== TELA ==')
    de = (hoje - timedelta(days=30)).isoformat()
    r = c_vend.get(f'/contagem-caixa/sangrias/?de={de}&ate={hoje.isoformat()}')
    html = r.content.decode()
    t('vendedor não vê a outra loja', 'Da outra loja' not in html and 'Lavagem de carpete' in html)
    t('vendedor não vê conferir em lote', 'Conferir marcadas' not in html)
    t('painéis de achados aparecem', 'Gastos recorrentes' in html and 'Fora do padrão' in html)
    r = c_chefe.get(f'/contagem-caixa/sangrias/?de={de}&ate={hoje.isoformat()}')
    html = r.content.decode()
    t('gestor vê as duas lojas e o painel por loja', 'Da outra loja' in html and 'Por loja' in html)
    t('gestor vê conferir e categorias', 'Conferir marcadas' in html and 'Categorias de gasto' in html)
    t('valor com milhar no padrão', 'R$ 300,00' in html or 'R$&nbsp;300,00' in html)
    r = c_chefe.get(f'/contagem-caixa/sangrias/?de={de}&ate={hoje.isoformat()}&situacao=conferidas')
    t('filtro conferidas vazio', 'Nenhuma sangria no período' in r.content.decode())
    r = c_chefe.get(f'/contagem-caixa/sangrias/?de={de}&ate={hoje.isoformat()}&q=carpete')
    html = r.content.decode()
    t('busca pela descrição', 'Lavagem de carpete' in html and 'Água sanitária 0' not in html)

    r = c_chefe.get(f'/contagem-caixa/loja/{loja.id}/?ano={d3.year}&mes={d3.month}')
    html = r.content.decode()
    t('tela da loja abre', r.status_code == 200, r.status_code)
    t('coluna de sangrias registradas na loja', '>sangrias</th>' in html and 'Ver as sangrias do dia' in html)

    print('== EXPORTAR ==')
    r = c_chefe.get(f'/contagem-caixa/sangrias/exportar/?de={de}&ate={hoje.isoformat()}')
    t('exporta xlsx', r.status_code == 200 and 'spreadsheetml' in r['Content-Type'], r.status_code)
    from openpyxl import load_workbook
    livro = load_workbook(BytesIO(r.content))
    t('abas do relatório', livro.sheetnames == ['Resumo', 'Sangrias', 'Recorrentes', 'Fora do padrão'], livro.sheetnames)
    aba = livro['Sangrias']
    t('8 linhas de sangria', aba.max_row >= 9 and aba.cell(row=9, column=4).value is not None)
    r = c_vend.get(f'/contagem-caixa/sangrias/exportar/?de={de}&ate={hoje.isoformat()}')
    livro = load_workbook(BytesIO(r.content))
    nomes = [livro['Sangrias'].cell(row=i, column=2).value for i in range(2, livro['Sangrias'].max_row + 1)]
    t('exportação do vendedor só tem a loja dele', 'Loja ZZ Sangria B' not in nomes)

    print('== FORA DO CAIXA ==')
    sem = User.objects.create_user(
        username='zz.sg.sem', email='zz.sg.sem@exemplo-teste.local', password='S3nha!teste',
        first_name='Zz', last_name='Sem', hierarchy='SUPERVISOR')
    c = Client()
    c.force_login(sem)
    r = c.get('/contagem-caixa/sangrias/')
    t('sem loja não abre a aba', r.status_code in (302, 403), r.status_code)

    print("== ADM DA LOJA (PADRÃO no grupo \"ADM's LOJAS\") ==")
    from communications.models import CommunicationGroup
    grupo = CommunicationGroup.objects.filter(name__icontains="ADM's LOJAS").first()
    t('o grupo existe no banco', grupo is not None)
    adm = User.objects.create_user(
        username='zz.sg.adm', email='zz.sg.adm@exemplo-teste.local', password='S3nha!teste',
        first_name='Zz', last_name='Adm', sector=loja, hierarchy='PADRAO')
    padrao = User.objects.create_user(
        username='zz.sg.padrao', email='zz.sg.padrao@exemplo-teste.local', password='S3nha!teste',
        first_name='Zz', last_name='Padrao', sector=loja, hierarchy='PADRAO')
    grupo.members.add(adm) if hasattr(grupo, 'members') else adm.communication_groups.add(grupo)
    c_adm, c_pad = Client(), Client()
    c_adm.force_login(adm)
    c_pad.force_login(padrao)

    r = c_pad.get('/contagem-caixa/sangrias/')
    t('PADRÃO fora do grupo continua barrado', r.status_code == 302 and r.url == '/', (r.status_code, getattr(r, 'url', '')))
    r = c_adm.get(f'/contagem-caixa/sangrias/?de={de}&ate={hoje.isoformat()}')
    html = r.content.decode()
    t('ADM abre a aba', r.status_code == 200, r.status_code)
    t('vê a loja dele e não a outra', 'Lavagem de carpete' in html and 'Da outra loja' not in html)
    t('a nav mostra só Sangrias', 'Visão geral' not in html and 'Abrir uma loja' not in html)
    t('menu lateral tem o item Sangrias e não o caixa',
      'title="Sangrias"' in html and 'title="Contagem de Caixa"' not in html)
    t('não confere em lote', 'Conferir marcadas' not in html)
    r = c_adm.get('/contagem-caixa/')
    t('o resto do caixa segue fechado', r.status_code == 302 and r.url == '/', (r.status_code, getattr(r, 'url', '')))
    r = c_adm.get(f'/contagem-caixa/loja/{loja.id}/')
    t('a planilha da loja também', r.status_code == 302 and r.url == '/', r.status_code)

    c_adm.post('/contagem-caixa/sangrias/registrar/', {
        'loja': loja.id, 'data': d1.isoformat(), 'valor': '7,50', 'categoria': outros.id, 'descricao': 'Fita adesiva ADM'})
    sa = Sangria.objects.filter(descricao='Fita adesiva ADM').first()
    t('ADM registra na loja dele', sa is not None and sa.registrada_por == adm)
    t('e o caixa do dia recebe', dia(loja, d1).sangria_registrada == D('57.50'), dia(loja, d1).sangria_registrada)
    c_adm.post('/contagem-caixa/sangrias/registrar/', {
        'loja': outra.id, 'data': d1.isoformat(), 'valor': '1', 'categoria': outros.id, 'descricao': 'Fora ADM'})
    t('não registra em outra loja', not Sangria.objects.filter(descricao='Fora ADM').exists())
    c_adm.post(f'/contagem-caixa/sangrias/{sa.id}/apagar/')
    t('apaga a própria', not Sangria.objects.filter(id=sa.id).exists())
    c_pad.post('/contagem-caixa/sangrias/registrar/', {
        'loja': loja.id, 'data': d1.isoformat(), 'valor': '1', 'categoria': outros.id, 'descricao': 'PADRAO barrado'})
    t('POST do PADRÃO fora do grupo não grava', not Sangria.objects.filter(descricao='PADRAO barrado').exists())
    r = c_adm.get(f'/contagem-caixa/sangrias/exportar/?de={de}&ate={hoje.isoformat()}')
    t('ADM exporta', r.status_code == 200, r.status_code)
finally:
    marcador.__exit__(Exception, Exception('rollback'), None)

print(f'\n{ok} OK, {fail} falha(s)')
sys.exit(1 if fail else 0)
