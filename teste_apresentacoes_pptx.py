"""Assistente de Apresentações — exportação para PowerPoint (.pptx nativo e editável).

Monta um documento com todos os tipos de elemento, formas, cores de tema, paginação, notas,
slide oculto, links, todas as transições e todas as animações/gatilhos; gera o .pptx, reabre
com python-pptx e confere posição (EMU), texto, runs, mídia, tabela e o XML de transição e
animação (ordem, ids, alvos, presets, nodeType). Mídia de verdade só em memória
(InMemoryStorage) e dentro de uma transação desfeita no fim.
"""
import io
import logging
import os
import re
import struct
import sys
import zipfile
from unittest import mock

import django

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'redeconfianca.settings')
django.setup()

from django.core.files.storage import InMemoryStorage
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import transaction
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from apresentacoes import exportar_pptx as ex
from apresentacoes.models import Midia

ok = fail = 0


def t(nome, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'  OK   {nome}')
    else:
        fail += 1
        print(f'  FALHA {nome} {extra}')


NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'mc': 'http://schemas.openxmlformats.org/markup-compatibility/2006',
    'p14': 'http://schemas.microsoft.com/office/powerpoint/2010/main',
    'ct': 'http://schemas.openxmlformats.org/package/2006/content-types',
    'rel': 'http://schemas.openxmlformats.org/package/2006/relationships',
}


def E(px):
    return int(round(px * 6350))


def png(largura, altura, cor=(226, 63, 207)):
    saida = io.BytesIO()
    Image.new('RGB', (largura, altura), cor).save(saida, 'PNG')
    return saida.getvalue()


def jpeg_girado():
    saida = io.BytesIO()
    exif = Image.Exif()
    exif[0x0112] = 6                                    # celular "em pé": o navegador gira, o PowerPoint não
    Image.new('RGB', (400, 200), (10, 120, 200)).save(saida, 'JPEG', exif=exif.tobytes())
    return saida.getvalue()


def webp():
    saida = io.BytesIO()
    Image.new('RGB', (300, 300), (0, 200, 100)).save(saida, 'WEBP')
    return saida.getvalue()


def caixa_mp4(tipo, conteudo):
    return struct.pack('>I4s', 8 + len(conteudo), tipo) + conteudo


def mp4(duracao_ms=12500, largura=1280, altura=720):
    ftyp = caixa_mp4(b'ftyp', b'isom' + struct.pack('>I', 512) + b'isomiso2avc1mp41')
    mvhd = caixa_mp4(b'mvhd', struct.pack('>B3xIIII', 0, 0, 0, 1000, duracao_ms) + b'\x00' * 80)
    matriz = struct.pack('>9i', 0x10000, 0, 0, 0, 0x10000, 0, 0, 0, 0x40000000)
    tkhd = caixa_mp4(b'tkhd', struct.pack('>B3xIIIIIQhhhh', 0, 0, 0, 1, 0, duracao_ms, 0, 0, 0, 0, 0) + matriz
                     + struct.pack('>II', largura << 16, altura << 16))
    moov = caixa_mp4(b'moov', mvhd + caixa_mp4(b'trak', tkhd))
    return ftyp + caixa_mp4(b'mdat', b'\x00' * 256) + moov      # moov no fim, como sai de muita câmera


MIDIAS = {
    '/apresentacoes/midia/1/': (png(1600, 900), 'image/png'),
    '/apresentacoes/midia/2/': (png(1920, 1080, (30, 20, 60)), 'image/png'),
    '/apresentacoes/midia/3/': (jpeg_girado(), 'image/jpeg'),
    '/apresentacoes/midia/4/': (webp(), 'image/webp'),
    '/apresentacoes/midia/5/': (mp4(), 'video/mp4'),
    '/apresentacoes/midia/6/': (b'\x1a\x45\xdf\xa3' + b'webm-de-teste' * 40, 'video/webm'),
    '/apresentacoes/midia/7/': (png(640, 360, (200, 200, 0)), 'image/png'),
}
lidas = []


def ler_teste(src):
    lidas.append(src)
    return MIDIAS.get(src)


TEMA = {
    'fonte_titulo': 'Poppins', 'fonte_texto': 'Inter',
    'cores': {'fundo': '#0B0612', 'superficie': '#1A0F24', 'primaria': '#E23FCF', 'secundaria': '#8B3DFF',
              'destaque': '#FFD15C', 'texto': '#FFFFFF', 'texto_suave': '#D9C9E8'},
}


def anim(tipo, gatilho='auto', duracao=600, atraso=0):
    return {'tipo': tipo, 'gatilho': gatilho, 'duracao': duracao, 'atraso': atraso}


def texto(id_, x, y, w, h, html, animacao=None, slot='', **estilo):
    return {'id': id_, 'tipo': 'texto', 'x': x, 'y': y, 'w': w, 'h': h, 'html': html, 'slot': slot,
            'estilo': estilo, 'animacao': animacao or anim('nenhuma')}


slide1 = {
    'id': 's-capa', 'layout': 'capa',
    'fundo': {'cor': '#0B0612', 'gradiente': 'linear-gradient(135deg, tema:fundo 0%, tema:primaria 100%)'},
    'transicao': {'tipo': 'fade', 'duracao': 700},
    'notas': 'Abrir com os números do trimestre.\nDepois passar para o plano.',
    'elementos': [
        texto('e-titulo', 120, 100, 1680, 300,
              'Resultados <b>do trimestre</b><br><i>Rede</i> <u>Confiança</u> <s>velho</s>',
              anim('fade', 'auto', 800, 200), slot='titulo', papel='titulo', tamanho=96, peso=400, cor='#FFFFFF',
              maiusculas=True, espacamento=2, entrelinha=1.1, alinhamento='center', vertical='middle'),
        texto('e-sub', 200, 460, 900, 260,
              '<span style="color:#FFD15C;font-size:48px;font-weight:700;font-style:italic;'
              'text-decoration:underline;font-family:\'Montserrat\', sans-serif;background-color:rgba(0,0,0,.5)">'
              'Destaque</span> e <a href="https://www.vivo.com.br/planos">link</a> e <a href="#slide-4">pular</a>',
              anim('subir', 'clique', 600), papel='texto', tamanho=40, cor='tema:texto_suave', fundo='tema:superficie',
              borda={'cor': 'tema:primaria', 'largura': 4, 'estilo': 'dashed'}, raio=24, preenchimento=20,
              sombra='suave'),
        {'id': 'e-img', 'tipo': 'imagem', 'src': '/apresentacoes/midia/1/', 'x': 1200, 'y': 420, 'w': 600, 'h': 600,
         'ajuste': 'cover', 'foco': {'x': 0.2, 'y': 0.5}, 'raio': 30, 'borda': {'cor': '#FFFFFF', 'largura': 6},
         'sombra': 'forte', 'filtro': 'cinza', 'opacidade': 0.7, 'alt': 'Foto da loja', 'link': '#slide-2',
         'animacao': anim('zoom', 'junto', 500)},
        texto('e-pag1', 1700, 1000, 180, 50, '{n} / {total}', slot='paginacao', tamanho=24),
        {'id': 'e-apagar', 'tipo': 'apagar', 'x': 0, 'y': 0, 'w': 100, 'h': 100},
        {'id': 'e-area', 'tipo': 'forma', 'forma': 'retangulo', 'slot': 'area_conteudo', 'x': 0, 'y': 0, 'w': 50, 'h': 50},
    ],
}
slide1['elementos'][1]['rotacao'] = 5
slide1['elementos'][1]['opacidade'] = 0.8
slide1['elementos'][1]['link'] = 'https://portal.exemplo-teste.local/impulso/'

FORMAS = [
    ('e-ret', 'retangulo', {'preenchimento': 'tema:primaria', 'animacao': anim('entrar-esquerda', 'auto', 500)}),
    ('e-arred', 'retangulo-arredondado', {'preenchimento': '#8B3DFF', 'raio': 40}),
    ('e-pilula', 'pilula', {'preenchimento': 'rgba(255, 209, 92, 0.5)'}),
    ('e-circ', 'circulo', {'preenchimento': '#FFF', 'borda': {'cor': '#000000', 'largura': 3, 'estilo': 'dotted'},
                           'animacao': anim('entrar-direita', 'auto', 400, 300)}),
    ('e-seta', 'seta', {'gradiente': 'linear-gradient(to right, #E23FCF, #8B3DFF 60%, #FFD15C)'}),
    ('e-tri', 'triangulo', {'preenchimento': '#E23FCFB3', 'rotacao': 180}),
    ('e-estrela', 'estrela', {'preenchimento': 'tema:destaque', 'sombra': 'brilho',
                              'animacao': anim('descer', 'junto', 1000)}),
    ('e-hex', 'hexagono', {'preenchimento': '#1A0F24', 'opacidade': 0.5}),
]
elementos2 = []
for indice, (id_, forma, extra) in enumerate(FORMAS):
    elemento = {'id': id_, 'tipo': 'forma', 'forma': forma, 'x': 100 + indice * 200, 'y': 100, 'w': 160, 'h': 120}
    elemento.update(extra)
    elementos2.append(elemento)
elementos2 += [
    {'id': 'e-linha', 'tipo': 'forma', 'forma': 'linha', 'x': 100, 'y': 300, 'w': 800, 'h': 40,
     'borda': {'cor': 'tema:destaque', 'largura': 6}, 'rotacao': 10, 'animacao': anim('revelar', 'clique', 700)},
    {'id': 'e-icone', 'tipo': 'icone', 'icone': 'fa-solid fa-circle-check', 'x': 100, 'y': 400, 'w': 120, 'h': 120,
     'cor': 'tema:destaque', 'animacao': anim('aparecer', 'auto', 999)},
    {'id': 'e-icone-reg', 'tipo': 'icone', 'icone': 'fa-regular fa-star', 'x': 260, 'y': 400, 'w': 120, 'h': 120,
     'cor': '#FFFFFF', 'fundo': 'tema:primaria', 'raio': 24},
    {'id': 'e-icone-marca', 'tipo': 'icone', 'icone': 'fab fa-whatsapp', 'x': 420, 'y': 400, 'w': 120, 'h': 80,
     'cor': '#25D366'},
    {'id': 'e-icone-sumido', 'tipo': 'icone', 'icone': 'fa-solid fa-nao-existe-mesmo', 'x': 580, 'y': 400,
     'w': 120, 'h': 120},
    {'id': 'e-contain', 'tipo': 'imagem', 'src': '/apresentacoes/midia/1/', 'x': 800, 'y': 400, 'w': 400, 'h': 400,
     'ajuste': 'contain'},
    {'id': 'e-exif', 'tipo': 'imagem', 'src': '/apresentacoes/midia/3/', 'x': 1250, 'y': 400, 'w': 200, 'h': 400,
     'ajuste': 'fill', 'filtro': 'escurecer'},
    {'id': 'e-webp', 'tipo': 'imagem', 'src': '/apresentacoes/midia/4/', 'x': 1500, 'y': 400, 'w': 300, 'h': 300,
     'filtro': 'desfocar'},
    {'id': 'e-sumida', 'tipo': 'imagem', 'src': '/apresentacoes/midia/999/', 'x': 100, 'y': 800, 'w': 400, 'h': 200},
    texto('e-lista', 900, 800, 900, 260,
          '<ul><li>Primeiro</li><li><p>Segundo</p></li></ul><ol><li>Um</li><li>Dois</li></ol><div><br></div><div>Fim  do\n texto</div>',
          tamanho=30),
]
slide2 = {
    'id': 's-conteudo', 'layout': 'conteudo',
    'fundo': {'cor': '#101010', 'gradiente': 'radial-gradient(circle at 30% 40%, rgba(226,63,207,.6) 0%, transparent 70%)',
              'imagem': '/apresentacoes/midia/2/', 'ajuste': 'cover'},
    'transicao': {'tipo': 'empurrar-esquerda', 'duracao': 500},
    'notas': '', 'elementos': elementos2,
}
slide3 = {
    'id': 's-quadro', 'layout': 'quadro', 'oculto': True,
    'fundo': {'cor': 'tema:superficie'},
    'transicao': {'tipo': 'revelar', 'duracao': 900},
    'notas': 'Slide reserva.',
    'elementos': [
        {'id': 'e-tabela', 'tipo': 'tabela', 'x': 100, 'y': 100, 'w': 900, 'h': 400, 'cabecalho': True,
         'linhas': [['Plano', 'Preço', 'Dados'], ['Vivo Total', 'R$ 199', '50 GB'], ['Controle', 'R$ 59'],
                    ['Pós', 'R$ 129', '30 GB']],
         'estilo': {'fonte': '', 'tamanho': 28, 'cor': '#FFFFFF', 'fundo_cabecalho': 'tema:primaria',
                    'cor_cabecalho': '#0B0612', 'fundo_linhas': 'rgba(255,255,255,.06)',
                    'fundo_alternado': 'rgba(255,255,255,.12)', 'borda': 'rgba(255,255,255,.2)',
                    'alinhamento': 'center', 'raio': 12},
         'animacao': anim('fade', 'clique', 400)},
        {'id': 'e-video', 'tipo': 'video', 'src': '/apresentacoes/midia/5/', 'poster': '/apresentacoes/midia/7/',
         'x': 1050, 'y': 100, 'w': 800, 'h': 450, 'autoplay': True, 'loop': True, 'mudo': True, 'raio': 16,
         'animacao': anim('fade', 'auto', 300)},
        {'id': 'e-webm', 'tipo': 'video', 'src': '/apresentacoes/midia/6/', 'x': 1050, 'y': 600, 'w': 400, 'h': 400,
         'autoplay': True, 'loop': False, 'mudo': False},
        {'id': 'e-video-sumido', 'tipo': 'video', 'src': '/apresentacoes/midia/998/', 'x': 100, 'y': 600, 'w': 400, 'h': 300},
    ],
}
slide4 = {
    'id': 's-fim', 'layout': 'encerramento', 'fundo': {'cor': '#0B0612'},
    'transicao': {'tipo': 'zoom', 'duracao': 1200},
    'elementos': [
        texto('e-final', 100, 400, 1720, 200, 'Obrigado!', anim('fade', 'junto', 500), papel='titulo', tamanho=120,
              peso=900, alinhamento='center', sombra='forte'),
        texto('e-pag4', 1700, 1000, 180, 50, '{n} / {total}', slot='paginacao', tamanho=24),
    ],
}
extras = [('empurrar-direita', 300), ('empurrar-cima', 800), ('empurrar-baixo', 650), ('cobrir', 1500),
          ('nenhuma', 0)]
slides_extras = [{'id': f's-extra-{i}', 'fundo': {'cor': '#222'}, 'transicao': {'tipo': tipo, 'duracao': dur},
                  'elementos': [texto(f'e-extra-{i}', 100, 100, 800, 200, f'Transição {tipo}')]}
                 for i, (tipo, dur) in enumerate(extras)]
DOCUMENTO = {'formato': 1, 'largura': 1920, 'altura': 1080, 'tema': TEMA,
             'slides': [slide1, slide2, slide3, slide4] + slides_extras}


def xp(elemento, expressao):
    # lxml puro: o xpath dos elementos do python-pptx não aceita namespaces=.
    return etree._Element.xpath(elemento, expressao, namespaces=NS)


def por_nome(slide, nome):
    return next((s for s in slide.shapes if s.name == nome), None)


def xml_do_slide(zip_pptx, numero):
    return etree.fromstring(zip_pptx.read(f'ppt/slides/slide{numero}.xml'))


def efeitos(raiz):
    return xp(raiz, '//p:cTn[@presetClass]')


def spid_do_efeito(ctn):
    return xp(ctn, 'string(.//p:spTgt/@spid)')


class Coletor(logging.Handler):
    def __init__(self):
        super().__init__(level=logging.DEBUG)
        self.registros = []

    def emit(self, record):
        self.registros.append(record)


coletor = Coletor()
logging.getLogger('apresentacoes.exportar_pptx').addHandler(coletor)

memoria = InMemoryStorage()
Midia._meta.get_field('arquivo').storage = memoria

marcador = transaction.atomic()
marcador.__enter__()
try:
    print('== CORES, GRADIENTES E HTML ==')
    cores = TEMA['cores']
    t('#RGB e #RRGGBB', ex.resolver_cor('#fff') == ('FFFFFF', 1.0) and ex.resolver_cor('#e23fcf') == ('E23FCF', 1.0))
    cor = ex.resolver_cor('#E23FCF80')
    t('#RRGGBBAA traz alfa', cor[0] == 'E23FCF' and abs(cor[1] - 128 / 255) < 1e-6, cor)
    t('rgba() com alfa ".06"', ex.resolver_cor('rgba(255,255,255,.06)') == ('FFFFFF', 0.06))
    t('transparent, vazio e lixo não pintam', ex.resolver_cor('transparent') is None and ex.resolver_cor('') is None
      and ex.resolver_cor(None) is None and ex.resolver_cor('javascript:alert(1)') is None)
    t('tema:<chave> resolve pelo tema', ex.resolver_cor('tema:primaria', cores) == ('E23FCF', 1.0)
      and ex.resolver_cor('tema:inexistente', cores) is None)
    gradiente = ex.resolver_gradiente('linear-gradient(135deg, tema:fundo, rgba(0,0,0,.5) 40%, #FFF)', cores)
    t('gradiente linear: ângulo e paradas (posições distribuídas)', gradiente['tipo'] == 'linear'
      and gradiente['angulo'] == 135 and [round(p, 2) for p, _ in gradiente['paradas']] == [0, 0.4, 1]
      and gradiente['paradas'][0][1] == ('0B0612', 1.0) and gradiente['paradas'][1][1] == ('000000', 0.5), gradiente)
    t('"to right" = 90deg', ex.resolver_gradiente('linear-gradient(to right, red, blue)')['angulo'] == 90)
    radial = ex.resolver_gradiente('radial-gradient(circle at 30% 40%, #fff 0%, transparent 70%)')
    t('radial: centro e parada transparente', radial['tipo'] == 'radial' and radial['centro'] == (0.3, 0.4)
      and radial['paradas'][1] == (0.7, ('000000', 0.0)), radial)
    t('gradiente inválido → None', ex.resolver_gradiente('url(x)') is None and ex.resolver_gradiente('') is None)
    base = {'fonte': 'Inter', 'tamanho': 40, 'peso': 400, 'italico': False, 'sublinhado': False, 'tachado': False,
            'cor': '#FFF', 'realce': '', 'link': ''}
    paragrafos = ex.ler_html('<div>a <b>b</b><br></div><div><br></div><ul><li>x</li><li><p>y</p></li></ul>', base)
    t('HTML: blocos, <br> final sem linha fantasma e linha em branco', [
        [t_.get('texto', '⏎') for t_ in p['trechos']] for p in paragrafos] == [['a ', 'b'], [], ['x'], ['y']],
      [[t_.get('texto', '⏎') for t_ in p['trechos']] for p in paragrafos])
    t('HTML: itens de lista com marcador (inclusive <li><p>)', [p['lista'] for p in paragrafos] == [None, None, 'ul', 'ul'])
    t('nome_do_arquivo seguro', ex.nome_do_arquivo('Resultados do mês / Área: Vendas?') == 'Resultados do mes Area Vendas.pptx'
      and ex.nome_do_arquivo('') == 'apresentacao.pptx' and ex.nome_do_arquivo('../../etc/passwd') == 'etc passwd.pptx'
      and ex.nome_do_arquivo('CON') == 'apresentacao-CON.pptx', ex.nome_do_arquivo('../../etc/passwd'))
    t('info_video lê duração e tamanho do MP4 (moov no fim)', ex.info_video(mp4()) == {
        'duracao_ms': 12500, 'largura': 1280.0, 'altura': 720.0}, ex.info_video(mp4()))
    t('ícones: estilo pedido, fallback de estilo e desconhecido', ex.glifo_do_icone('fa-regular fa-star') == ('r', '')
      and ex.glifo_do_icone('fa-regular fa-house')[0] == 's' and ex.glifo_do_icone('fab fa-whatsapp') == ('b', '')
      and ex.glifo_do_icone('fa-solid fa-nao-existe') is None)

    print('\n== GERAÇÃO ==')
    dados = ex.gerar_pptx(DOCUMENTO, 'Resultados do Trimestre — Rede', ler_midia=ler_teste)
    t('gera bytes de um zip', isinstance(dados, bytes) and dados[:2] == b'PK', dados[:4])
    arquivo = zipfile.ZipFile(io.BytesIO(dados))
    t('zip íntegro', arquivo.testzip() is None)
    prs = Presentation(io.BytesIO(dados))
    t('slide 16:9 (12192000 × 6858000 EMU)', prs.slide_width == 12192000 and prs.slide_height == 6858000)
    apresentacao_xml = etree.fromstring(arquivo.read('ppt/presentation.xml'))
    t('sldSz sem "screen4x3"', apresentacao_xml.find('p:sldSz', NS).get('type') is None)
    t('9 slides', len(prs.slides) == 9, len(prs.slides))
    t('título e autor nas propriedades', prs.core_properties.title == 'Resultados do Trimestre — Rede'
      and prs.core_properties.author == 'Portal Rede Confiança' and prs.core_properties.last_modified_by == 'Portal Rede Confiança')
    tema_xml = etree.fromstring(arquivo.read('ppt/theme/theme1.xml'))
    t('tema do PowerPoint com as fontes e a cor primária do documento',
      xp(tema_xml, 'string(//a:majorFont/a:latin/@typeface)') == 'Poppins'
      and xp(tema_xml, 'string(//a:minorFont/a:latin/@typeface)') == 'Inter'
      and xp(tema_xml, 'string(//a:clrScheme/a:accent1/a:srgbClr/@val)') == 'E23FCF')
    tipos = etree.fromstring(arquivo.read('[Content_Types].xml'))
    t('[Content_Types].xml: mp4 como video/mp4', xp(tipos, 'boolean(//ct:Default[@Extension="mp4"][@ContentType="video/mp4"])'))
    t('[Content_Types].xml: webm como video/webm (parte .webm)', xp(
        tipos, 'boolean(//ct:Override[@ContentType="video/webm"][substring(@PartName, string-length(@PartName) - 4) = ".webm"])'))
    midias = [n for n in arquivo.namelist() if n.startswith('ppt/media/')]
    t('vídeos embutidos no pacote', any(n.endswith('.mp4') for n in midias) and any(n.endswith('.webm') for n in midias), midias)

    print('\n== SLIDE 1: TEXTO, IMAGEM, PAGINAÇÃO, NOTAS, LINKS ==')
    s1 = prs.slides[0]
    nomes1 = [s.name for s in s1.shapes]
    t('apagar e area_conteudo não exportam; ordem de pintura mantida', nomes1 == [
        'Texto e-titulo', 'Texto e-sub', 'Imagem e-img', 'Texto e-pag1'], nomes1)
    titulo = por_nome(s1, 'Texto e-titulo')
    t('título é caixa de texto na posição exata (EMU)', titulo.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
      and (titulo.left, titulo.top, titulo.width, titulo.height) == (E(120), E(100), E(1680), E(300)),
      (titulo.left, titulo.top, titulo.width, titulo.height))
    paragrafo = titulo.text_frame.paragraphs[0]
    runs = paragrafo.runs
    t('runs do HTML (b, br, i, u, s)', [r.text for r in runs] == ['Resultados ', 'do trimestre', 'Rede', ' ', 'Confiança', ' ', 'velho']
      and paragrafo.text == 'Resultados do trimestre\vRede Confiança velho', [r.text for r in runs])
    t('negrito só no <b>', [bool(r.font.bold) for r in runs] == [False, True, False, False, False, False, False])
    t('itálico e sublinhado', runs[2].font.italic is True and runs[4].font.underline is True and runs[0].font.italic is False)
    rpr_s = runs[6]._r.find('a:rPr', NS)
    t('tachado e maiúsculas (cap="all") e espaçamento (spc)', rpr_s.get('strike') == 'sngStrike'
      and rpr_s.get('cap') == 'all' and rpr_s.get('spc') == '100', dict(rpr_s.attrib))
    t('fonte do tema pelo papel (título → Poppins), 96 px = 48 pt, cor branca',
      all(r.font.name == 'Poppins' and r.font.size.pt == 48 and str(r.font.color.rgb) == 'FFFFFF' for r in runs))
    corpo = titulo.text_frame._txBody.find('a:bodyPr', NS)
    t('sem autoajuste, com quebra de linha, âncora no meio e sem margens', corpo.get('wrap') == 'square'
      and corpo.find('a:noAutofit', NS) is not None and corpo.get('anchor') == 'ctr' and corpo.get('lIns') == '0')
    ppr = paragrafo._p.find('a:pPr', NS)
    t('alinhamento centralizado e entrelinha exata (96 × 1,1 × 0,5 pt)', ppr.get('algn') == 'ctr'
      and xp(ppr, 'string(a:lnSpc/a:spcPts/@val)') == '5280', etree.tostring(ppr)[:200])

    sub = por_nome(s1, 'Texto e-sub')
    sp = sub._element
    t('texto com fundo/borda/raio vira forma arredondada', sub.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
      and xp(sp, 'string(p:spPr/a:prstGeom/@prst)') == 'roundRect'
      and xp(sp, 'string(p:spPr/a:prstGeom/a:avLst/a:gd/@fmla)') == f'val {round(24 / 260 * 100000)}')
    t('sem p:style (não herda cor do tema do PowerPoint)', sp.find('p:style', NS) is None)
    t('fundo tema:superficie com opacidade 0,8', xp(sp, 'string(p:spPr/a:solidFill/a:srgbClr/@val)') == '1A0F24'
      and xp(sp, 'string(p:spPr/a:solidFill/a:srgbClr/a:alpha/@val)') == '80000')
    t('borda tracejada 4 px na cor primária', xp(sp, 'string(p:spPr/a:ln/@w)') == str(E(4))
      and xp(sp, 'string(p:spPr/a:ln/a:solidFill/a:srgbClr/@val)') == 'E23FCF'
      and xp(sp, 'string(p:spPr/a:ln/a:prstDash/@val)') == 'dash')
    t('sombra suave (outerShdw) e rotação 5°', xp(sp, 'boolean(p:spPr/a:effectLst/a:outerShdw)')
      and sub.rotation == 5.0)
    t('preenchimento = margens internas (20 px + borda 4 px)', xp(sp, 'string(p:txBody/a:bodyPr/@lIns)') == str(E(24)))
    runs_sub = sub.text_frame.paragraphs[0].runs
    destaque = runs_sub[0]
    t('span: cor, tamanho, peso, itálico, sublinhado e família', destaque.text == 'Destaque'
      and str(destaque.font.color.rgb) == 'FFD15C' and destaque.font.size.pt == 24 and destaque.font.bold
      and destaque.font.italic and destaque.font.underline and destaque.font.name == 'Montserrat')
    t('span background-color → realce', xp(destaque._r, 'string(a:rPr/a:highlight/a:srgbClr/@val)') == '000000')
    t('texto herda tema:texto_suave com opacidade', str(runs_sub[1].font.color.rgb) == 'D9C9E8'
      and xp(runs_sub[1]._r, 'string(a:rPr/a:solidFill/a:srgbClr/a:alpha/@val)') == '80000'
      and runs_sub[1].font.name == 'Inter')
    link_run = next(r for r in runs_sub if r.text == 'link')
    t('<a href> vira hyperlink do trecho', link_run.hyperlink.address == 'https://www.vivo.com.br/planos')
    pular = next(r for r in runs_sub if r.text == 'pular')
    hl = pular._r.find('a:rPr/a:hlinkClick', NS)
    alvo = s1.part.rels[hl.get(f'{{{NS["r"]}}}id')].target_part if hl is not None else None
    t('<a href="#slide-4"> salta para o slide 4', hl is not None and hl.get('action') == 'ppaction://hlinksldjump'
      and alvo is prs.slides[3].part)
    t('link do elemento (http) no clique da forma', sub.click_action.hyperlink.address == 'https://portal.exemplo-teste.local/impulso/')

    imagem = por_nome(s1, 'Imagem e-img')
    t('imagem é figura na caixa', imagem.shape_type == MSO_SHAPE_TYPE.PICTURE
      and (imagem.left, imagem.top, imagem.width, imagem.height) == (E(1200), E(420), E(600), E(600)))
    t('cover com foco 0,2: corte 8,75% à esquerda e 35% à direita', abs(imagem.crop_left - 0.0875) < 1e-4
      and abs(imagem.crop_right - 0.35) < 1e-4 and imagem.crop_top == 0 and imagem.crop_bottom == 0,
      (imagem.crop_left, imagem.crop_right))
    pic = imagem._element
    t('raio → roundRect na figura; borda 6 px; sombra forte', xp(pic, 'string(p:spPr/a:prstGeom/@prst)') == 'roundRect'
      and xp(pic, 'string(p:spPr/a:ln/@w)') == str(E(6))
      and xp(pic, 'boolean(p:spPr/a:effectLst/a:outerShdw)'))
    t('filtro cinza (grayscl) e opacidade 0,7 (alphaModFix)', xp(pic, 'boolean(p:blipFill/a:blip/a:grayscl)')
      and xp(pic, 'string(p:blipFill/a:blip/a:alphaModFix/@amt)') == '70000')
    t('alt vira descrição e link "#slide-2" salta de slide', xp(pic, 'string(p:nvPicPr/p:cNvPr/@descr)') == 'Foto da loja'
      and imagem.click_action.target_slide is not None and imagem.click_action.target_slide.slide_id == prs.slides[1].slide_id)
    t('imagem embutida com o tamanho original', imagem.image.size == (1600, 900) and imagem.image.content_type == 'image/png')
    t('paginação só conta slides visíveis (2 dígitos)', por_nome(s1, 'Texto e-pag1').text_frame.text == '01 / 08'
      and por_nome(prs.slides[3], 'Texto e-pag4').text_frame.text == '03 / 08',
      (por_nome(s1, 'Texto e-pag1').text_frame.text, por_nome(prs.slides[3], 'Texto e-pag4').text_frame.text))
    t('notas do apresentador no notes slide', s1.has_notes_slide
      and s1.notes_slide.notes_text_frame.text == 'Abrir com os números do trimestre.\nDepois passar para o plano.')
    t('slide sem notas não ganha notes slide', not prs.slides[1].has_notes_slide)
    raiz1 = xml_do_slide(arquivo, 1)
    t('fundo com gradiente opaco nativo (lin 45° a partir de 135deg CSS)', xp(raiz1, 
        'string(p:cSld/p:bg/p:bgPr/a:gradFill/a:lin/@ang)') == str(45 * 60000)
      and xp(raiz1, 'string(p:cSld/p:bg/p:bgPr/a:gradFill/a:gsLst/a:gs[2]/a:srgbClr/@val)') == 'E23FCF')

    print('\n== SLIDE 2: FUNDO EM CAMADAS, FORMAS, ÍCONES, IMAGENS ==')
    s2 = prs.slides[1]
    raiz2 = xml_do_slide(arquivo, 2)
    t('p:bg com a cor sólida', xp(raiz2, 'string(p:cSld/p:bg/p:bgPr/a:solidFill/a:srgbClr/@val)') == '101010')
    formas2 = list(s2.shapes)
    t('gradiente com transparência é a 1ª camada; imagem de fundo a 2ª, cobrindo o slide',
      formas2[0].name == 'Fundo (gradiente)' and formas2[1].name == 'Fundo (imagem)'
      and (formas2[1].left, formas2[1].top, formas2[1].width, formas2[1].height) == (0, 0, 12192000, 6858000))
    t('radial nativo com centro em 30%/40%', xp(formas2[0]._element, 
        'string(p:spPr/a:gradFill/a:path/a:fillToRect/@l)') == '30000'
      and xp(formas2[0]._element, 'string(p:spPr/a:gradFill/a:path/a:fillToRect/@t)') == '40000')
    esperado = {'e-ret': 'rect', 'e-arred': 'roundRect', 'e-pilula': 'roundRect', 'e-circ': 'ellipse',
                'e-seta': 'rightArrow', 'e-tri': 'triangle', 'e-estrela': 'star5', 'e-hex': 'hexagon'}
    geometrias = {id_: xp(por_nome(s2, f'Forma {id_}')._element, 'string(p:spPr/a:prstGeom/@prst)')
                  for id_ in esperado}
    t('formas nativas com a geometria certa', geometrias == esperado, geometrias)
    t('posição das formas em EMU', all((por_nome(s2, f'Forma {id_}').left, por_nome(s2, f'Forma {id_}').top)
                                        == (E(100 + i * 200), E(100)) for i, (id_, _, _) in enumerate(FORMAS)))
    ajuste = lambda id_: xp(por_nome(s2, f'Forma {id_}')._element, 'string(p:spPr/a:prstGeom/a:avLst/a:gd/@fmla)')
    t('arredondado pelo raio (40/120) e pílula com 50%', ajuste('e-arred') == 'val 33333' and ajuste('e-pilula') == 'val 50000',
      (ajuste('e-arred'), ajuste('e-pilula')))
    val = lambda id_, caminho: xp(por_nome(s2, f'Forma {id_}')._element, f'string({caminho})')
    t('preenchimento de tema e com alfa (rgba e #RRGGBBAA)', val('e-ret', 'p:spPr/a:solidFill/a:srgbClr/@val') == 'E23FCF'
      and val('e-pilula', 'p:spPr/a:solidFill/a:srgbClr/a:alpha/@val') == '50000'
      and val('e-tri', 'p:spPr/a:solidFill/a:srgbClr/a:alpha/@val') == '70196')
    t('gradiente linear na forma (to right → ang 0) com 3 paradas', val('e-seta', 'p:spPr/a:gradFill/a:lin/@ang') == '0'
      and val('e-seta', 'count(p:spPr/a:gradFill/a:gsLst/a:gs)') == '3'
      and val('e-seta', 'p:spPr/a:gradFill/a:gsLst/a:gs[2]/@pos') == '60000')
    t('borda pontilhada', val('e-circ', 'p:spPr/a:ln/a:prstDash/@val') == 'sysDot' and val('e-circ', 'p:spPr/a:ln/@w') == str(E(3)))
    t('forma sem borda tem linha noFill', val('e-ret', 'boolean(p:spPr/a:ln/a:noFill)') == 'true')
    t('brilho → glow na cor primária', val('e-estrela', 'p:spPr/a:effectLst/a:glow/a:srgbClr/@val') == 'E23FCF')
    t('opacidade 0,5 na forma e rotação 180°', val('e-hex', 'p:spPr/a:solidFill/a:srgbClr/a:alpha/@val') == '50000'
      and por_nome(s2, 'Forma e-tri').rotation == 180.0)
    linha = por_nome(s2, 'Linha e-linha')
    t('linha = conector no meio da caixa, espessura e cor da borda', linha is not None
      and linha._element.tag == f'{{{NS["p"]}}}cxnSp'
      and (linha.left, linha.top, linha.width, linha.height) == (E(100), E(320), E(800), 0)
      and xp(linha._element, 'string(p:spPr/a:ln/@w)') == str(E(6))
      and xp(linha._element, 'string(p:spPr/a:ln/a:solidFill/a:srgbClr/@val)') == 'FFD15C'
      and linha.rotation == 10.0)
    icones = [por_nome(s2, f'Ícone {n}') for n in ('e-icone', 'e-icone-reg', 'e-icone-marca')]
    t('ícones viram figuras PNG 2× na caixa', all(i is not None and i.shape_type == MSO_SHAPE_TYPE.PICTURE
                                                   and i.image.content_type == 'image/png' for i in icones)
      and icones[0].image.size == (240, 240) and icones[2].image.size == (240, 160),
      [i.image.size for i in icones if i])
    figura = Image.open(io.BytesIO(icones[0].image.blob)).convert('RGBA')
    pixels = [p for p in figura.getdata() if p[3] > 200]
    t('glifo desenhado na cor pedida sobre fundo transparente', figura.getpixel((0, 0))[3] == 0
      and pixels and all(abs(p[0] - 0xFF) < 8 and abs(p[1] - 0xD1) < 8 and abs(p[2] - 0x5C) < 8 for p in pixels[:50]))
    figura_fundo = Image.open(io.BytesIO(icones[1].image.blob)).convert('RGBA')
    t('ícone com fundo: canto arredondado transparente, meio da borda pintado', figura_fundo.getpixel((0, 0))[3] == 0
      and figura_fundo.getpixel((120, 2))[:3] == (0xE2, 0x3F, 0xCF))
    sumido = por_nome(s2, 'Mídia indisponível e-icone-sumido')
    t('ícone desconhecido vira quadro com legenda', sumido is not None and sumido.text_frame.text == 'Ícone indisponível')
    contain = por_nome(s2, 'Imagem e-contain')
    t('contain encolhe a figura e centraliza (400 × 225 no meio da caixa)',
      (contain.left, contain.top, contain.width, contain.height) == (E(800), E(400 + 87.5), E(400), E(225))
      and contain.crop_left == 0, (contain.left, contain.top, contain.width, contain.height))
    exif = por_nome(s2, 'Imagem e-exif')
    t('JPEG girado pelo EXIF antes de embutir', exif.image.size == (200, 400), exif.image.size)
    t('filtro escurecer → a:lum', xp(exif._element, 'string(p:blipFill/a:blip/a:lum/@bright)') == '-35000')
    t('WEBP convertido para PNG (desfocado)', por_nome(s2, 'Imagem e-webp').image.content_type == 'image/png')
    sumida = por_nome(s2, 'Mídia indisponível e-sumida')
    t('imagem ausente vira quadro cinza com legenda, no lugar dela', sumida is not None
      and sumida.text_frame.text == 'Imagem indisponível'
      and (sumida.left, sumida.top, sumida.width, sumida.height) == (E(100), E(800), E(400), E(200))
      and xp(sumida._element, 'string(p:spPr/a:solidFill/a:srgbClr/@val)') == 'D9D9D9')
    lista = por_nome(s2, 'Texto e-lista').text_frame.paragraphs
    t('listas: marcadores, numeração, linha em branco e espaços colapsados', [p.text for p in lista] == [
        'Primeiro', 'Segundo', 'Um', 'Dois', '', 'Fim do texto']
      and xp(lista[0]._p, 'string(a:pPr/a:buChar/@char)') == '•'
      and xp(lista[2]._p, 'string(a:pPr/a:buAutoNum/@type)') == 'arabicPeriod'
      and xp(lista[5]._p, 'boolean(a:pPr/a:buNone)'), [p.text for p in lista])

    print('\n== SLIDE 3 (OCULTO): TABELA E VÍDEOS ==')
    s3 = prs.slides[2]
    raiz3 = xml_do_slide(arquivo, 3)
    t('slide oculto exportado com show="0"', raiz3.get('show') == '0' and xml_do_slide(arquivo, 1).get('show') is None)
    tabela_forma = por_nome(s3, 'Tabela e-tabela')
    t('tabela nativa na posição', tabela_forma is not None and tabela_forma.has_table
      and (tabela_forma.left, tabela_forma.top, tabela_forma.width, tabela_forma.height) == (E(100), E(100), E(900), E(400)))
    tabela = tabela_forma.table
    t('4 linhas × 3 colunas, linha curta completada', len(tabela.rows) == 4 and len(tabela.columns) == 3
      and tabela.cell(2, 2).text == '' and tabela.cell(1, 1).text == 'R$ 199')
    tbl_pr = xp(tabela_forma._element, './/a:tblPr')[0]
    t('cabeçalho ligado e estilo neutro', tbl_pr.get('firstRow') == '1' and tbl_pr.get('bandRow') == '0'
      and tbl_pr.findtext('a:tableStyleId', namespaces=NS) == ex.ESTILO_TABELA_NENHUM)
    celula = lambda i, j, caminho: xp(tabela.cell(i, j)._tc, f'string({caminho})')
    t('cabeçalho: fundo tema:primaria, texto #0B0612 em negrito', celula(0, 0, 'a:tcPr/a:solidFill/a:srgbClr/@val') == 'E23FCF'
      and str(tabela.cell(0, 0).text_frame.paragraphs[0].runs[0].font.color.rgb) == '0B0612'
      and tabela.cell(0, 0).text_frame.paragraphs[0].runs[0].font.bold)
    t('linhas alternadas com alfa (.06 e .12)', celula(1, 0, 'a:tcPr/a:solidFill/a:srgbClr/a:alpha/@val') == '6000'
      and celula(2, 0, 'a:tcPr/a:solidFill/a:srgbClr/a:alpha/@val') == '12000'
      and celula(3, 0, 'a:tcPr/a:solidFill/a:srgbClr/a:alpha/@val') == '6000')
    corpo_celula = tabela.cell(1, 0).text_frame.paragraphs[0]
    t('fonte do tema, 28 px = 14 pt, cor e alinhamento', corpo_celula.runs[0].font.name == 'Inter'
      and corpo_celula.runs[0].font.size.pt == 14 and str(corpo_celula.runs[0].font.color.rgb) == 'FFFFFF'
      and xp(corpo_celula._p, 'string(a:pPr/@algn)') == 'ctr' and not corpo_celula.runs[0].font.bold)
    t('bordas nas quatro faces com a cor e alfa', all(celula(1, 1, f'a:tcPr/a:{lado}/a:solidFill/a:srgbClr/a:alpha/@val') == '20000'
                                                      for lado in ('lnL', 'lnR', 'lnT', 'lnB')))
    video = por_nome(s3, 'Vídeo e-video')
    t('vídeo mp4 embutido como filme', video is not None and video.shape_type == MSO_SHAPE_TYPE.MEDIA
      and (video.left, video.top, video.width, video.height) == (E(1050), E(100), E(800), E(450)))
    rel_video = xp(video._element, 'string(p:nvPicPr/p:nvPr/a:videoFile/@r:link)')
    parte_video = s3.part.related_part(rel_video)
    t('parte do vídeo .mp4 com video/mp4 e bytes originais', parte_video.partname.endswith('.mp4')
      and parte_video.content_type == 'video/mp4' and parte_video.blob == MIDIAS['/apresentacoes/midia/5/'][0])
    poster = Image.open(io.BytesIO(s3.part.related_part(
        xp(video._element, 'string(p:blipFill/a:blip/@r:embed)')).blob))
    t('pôster = mídia "poster"', poster.size == (640, 360), poster.size)
    t('raio no vídeo → roundRect', xp(video._element, 'string(p:spPr/a:prstGeom/@prst)') == 'roundRect')
    webm = por_nome(s3, 'Vídeo e-webm')
    parte_webm = s3.part.related_part(xp(webm._element, 'string(p:nvPicPr/p:nvPr/a:videoFile/@r:link)'))
    t('webm com extensão e mime certos', parte_webm.partname.endswith('.webm') and parte_webm.content_type == 'video/webm')
    poster_padrao = Image.open(io.BytesIO(s3.part.related_part(
        xp(webm._element, 'string(p:blipFill/a:blip/@r:embed)')).blob)).convert('RGB')
    largura_p, altura_p = poster_padrao.size
    meio = poster_padrao.crop((largura_p // 2 - altura_p // 6, altura_p // 3, largura_p // 2 + altura_p // 6, altura_p * 2 // 3))
    brancos = sum(1 for p in meio.getdata() if min(p) > 200)
    t('sem pôster: quadro gerado (escuro com play branco no meio)', poster_padrao.getpixel((5, 5)) == (17, 17, 24)
      and brancos > meio.size[0] * meio.size[1] * 0.05, (poster_padrao.getpixel((5, 5)), brancos))
    t('vídeo ausente vira quadro com legenda', por_nome(s3, 'Mídia indisponível e-video-sumido').text_frame.text == 'Vídeo indisponível')

    print('\n== TRANSIÇÕES ==')
    esperadas = {1: ('fade', {}, 'med', '700'), 2: ('push', {'dir': 'l'}, 'fast', '500'),
                 3: ('wipe', {'dir': 'r'}, 'slow', '900'), 4: ('zoom', {'dir': 'in'}, 'slow', '1200'),
                 5: ('push', {'dir': 'r'}, 'fast', '300'), 6: ('push', {'dir': 'u'}, 'med', '800'),
                 7: ('push', {'dir': 'd'}, 'med', '650'), 8: ('cover', {'dir': 'l'}, 'slow', '1500')}
    for numero in range(1, 10):
        raiz = xml_do_slide(arquivo, numero)
        filhos = [etree.QName(f).localname for f in raiz]
        alternativas = raiz.findall('mc:AlternateContent', NS)
        if numero == 9:
            t('slide 9: "nenhuma" não grava transição', not alternativas and raiz.find('p:transition', NS) is None, filhos)
            continue
        esperado_tag, atributos, velocidade, duracao = esperadas[numero]
        escolha = alternativas[0].find('mc:Choice', NS) if alternativas else None
        reserva = alternativas[0].find('mc:Fallback', NS) if alternativas else None
        trans_escolha = escolha.find('p:transition', NS) if escolha is not None else None
        trans_reserva = reserva.find('p:transition', NS) if reserva is not None else None
        efeito = trans_escolha[0] if trans_escolha is not None and len(trans_escolha) else None
        t(f'slide {numero}: {esperado_tag} {atributos} spd={velocidade} p14:dur={duracao}',
          len(alternativas) == 1 and escolha.get('Requires') == 'p14'
          and trans_escolha.get(f'{{{NS["p14"]}}}dur') == duracao and trans_escolha.get('spd') == velocidade
          and efeito is not None and etree.QName(efeito).localname == esperado_tag and dict(efeito.attrib) == atributos
          and trans_reserva.get('spd') == velocidade and trans_reserva.get(f'{{{NS["p14"]}}}dur') is None
          and etree.QName(trans_reserva[0]).localname == esperado_tag,
          etree.tostring(alternativas[0]) if alternativas else filhos)
        ordem = [f for f in filhos if f in ('cSld', 'clrMapOvr', 'AlternateContent', 'timing')]
        t(f'slide {numero}: ordem cSld → clrMapOvr → transição → timing', ordem[:3] == ['cSld', 'clrMapOvr', 'AlternateContent']
          and (len(ordem) == 3 or ordem[3] == 'timing') and filhos[-1] in ('AlternateContent', 'timing'), filhos)

    print('\n== ANIMAÇÕES ==')
    def validar_ids_e_alvos(numero):
        raiz = xml_do_slide(arquivo, numero)
        ids = [int(c.get('id')) for c in raiz.iter(f'{{{NS["p"]}}}cTn')]
        formas_ids = {c.get('id') for c in xp(raiz, '//p:cNvPr')}
        alvos = {s.get('spid') for s in xp(raiz, '//p:spTgt')}
        construcoes = {b.get('spid') for b in xp(raiz, '//p:bldLst/*')}
        t(f'slide {numero}: ids de cTn únicos e crescentes', ids == sorted(set(ids)) and ids[:2] == [1, 2], ids[:12])
        t(f'slide {numero}: todo spTgt e bldLst aponta para forma existente', alvos and alvos <= formas_ids
          and construcoes <= formas_ids, (alvos - formas_ids, construcoes - formas_ids))
        return raiz

    raiz = validar_ids_e_alvos(1)
    id_titulo, id_sub, id_img = (str(por_nome(s1, n).shape_id) for n in ('Texto e-titulo', 'Texto e-sub', 'Imagem e-img'))
    grupos = xp(raiz, '//p:cTn[@nodeType="mainSeq"]/p:childTnLst/p:par/p:cTn')
    t('slide 1: 2 grupos (auto inicial + clique)', len(grupos) == 2, len(grupos))
    t('1º auto inicia sozinho após a transição (onBegin → mainSeq)', xp(
        grupos[0], 'boolean(p:stCondLst/p:cond[@delay="indefinite"]) and boolean(p:stCondLst/p:cond[@evt="onBegin"]/p:tn[@val="2"])')
      and xp(grupos[1], 'count(p:stCondLst/p:cond)') == 1)
    efs = efeitos(raiz)
    resumo = [(e.get('presetID'), e.get('presetSubtype'), e.get('nodeType'), spid_do_efeito(e), e.get('grpId')) for e in efs]
    t('fade auto (10/0 afterEffect), subir clique (42/0 clickEffect), zoom junto (53/16 withEffect)', resumo == [
        ('10', '0', 'afterEffect', id_titulo, '0'), ('42', '0', 'clickEffect', id_sub, '0'), ('53', '16', 'withEffect', id_img, None)],
      resumo)
    t('atraso do fade (200 ms) e duração (800 ms)', xp(efs[0], 'string(p:stCondLst/p:cond/@delay)') == '200'
      and xp(efs[0], 'string(.//p:animEffect[@filter="fade"]/p:cBhvr/p:cTn/@dur)') == '800')
    t('subir = fade + ppt_y de "#ppt_y+.1"', xp(efs[1], 'boolean(.//p:animEffect[@filter="fade"])')
      and xp(efs[1], 'string(.//p:anim[p:cBhvr/p:attrNameLst/p:attrName="ppt_y"]/p:tavLst/p:tav[1]//p:strVal/@val)') == '#ppt_y+.1')
    t('zoom = ppt_w/ppt_h de 0 + fade', xp(efs[2], 'string(.//p:anim[p:cBhvr/p:attrNameLst/p:attrName="ppt_w"]/p:tavLst/p:tav[1]//p:fltVal/@val)') == '0'
      and xp(efs[2], 'boolean(.//p:animEffect[@filter="fade"])'))
    t('todo efeito torna a forma visível (set style.visibility)', all(xp(e, 'boolean(p:childTnLst/p:set[p:cBhvr/p:attrNameLst/p:attrName="style.visibility"])') for e in efs))
    t('bldLst: bldP das formas de texto (a figura não entra)', sorted(b.get('spid') for b in xp(raiz, '//p:bldLst/p:bldP')) == sorted([id_titulo, id_sub]))
    t('sequência com prevCondLst/nextCondLst', xp(raiz, 'boolean(//p:seq/p:prevCondLst/p:cond[@evt="onPrev"]) and boolean(//p:seq/p:nextCondLst/p:cond[@evt="onNext"])'))

    raiz = validar_ids_e_alvos(2)
    grupos = xp(raiz, '//p:cTn[@nodeType="mainSeq"]/p:childTnLst/p:par/p:cTn')
    etapas = [[(xp(e, 'string(p:stCondLst/p:cond/@delay)'),
                [(ef.get('presetID'), ef.get('presetSubtype'), ef.get('nodeType'))
                 for ef in xp(e, 'p:childTnLst/p:par/p:cTn')])
               for e in xp(g, 'p:childTnLst/p:par/p:cTn')] for g in grupos]
    t('slide 2: etapas encadeadas (after começa quando a anterior termina)', etapas == [
        [('0', [('2', '8', 'afterEffect')]), ('500', [('2', '2', 'afterEffect'), ('47', '0', 'withEffect')])],
        [('0', [('22', '8', 'clickEffect')]), ('700', [('1', '0', 'afterEffect')])],
    ], etapas)
    efs = efeitos(raiz)
    por_preset = {(e.get('presetID'), e.get('presetSubtype')): e for e in efs}
    ppt_x = lambda e: xp(e, 'string(.//p:anim[p:cBhvr/p:attrNameLst/p:attrName="ppt_x"]/p:tavLst/p:tav[1]//p:strVal/@val)')
    t('entrar-esquerda: Fly In 2/8 de "0-#ppt_w/2"; entrar-direita: 2/2 de "1+#ppt_w/2"',
      ppt_x(por_preset[('2', '8')]) == '0-#ppt_w/2' and ppt_x(por_preset[('2', '2')]) == '1+#ppt_w/2'
      and xp(por_preset[('2', '8')], 'string(.//p:anim/p:cBhvr/@additive)') == 'base')
    t('descer = "#ppt_y-.1"', xp(por_preset[('47', '0')], 'string(.//p:anim[p:cBhvr/p:attrNameLst/p:attrName="ppt_y"]/p:tavLst/p:tav[1]//p:strVal/@val)') == '#ppt_y-.1')
    t('revelar = Wipe 22/8 com filter="wipe(right)" no conector', xp(por_preset[('22', '8')], 'string(.//p:animEffect/@filter)') == 'wipe(right)'
      and spid_do_efeito(por_preset[('22', '8')]) == str(por_nome(s2, 'Linha e-linha').shape_id))
    t('aparecer = só o set (dur 1)', xp(por_preset[('1', '0')], 'count(p:childTnLst/*)') == 1
      and xp(por_preset[('1', '0')], 'string(.//p:set/p:cBhvr/p:cTn/@dur)') == '1')
    t('atraso do entrar-direita (300 ms) no próprio efeito', xp(por_preset[('2', '2')], 'string(p:stCondLst/p:cond/@delay)') == '300')

    raiz = validar_ids_e_alvos(3)
    grupos = xp(raiz, '//p:cTn[@nodeType="mainSeq"]/p:childTnLst/p:par/p:cTn')
    id_tabela, id_video, id_webm = (str(por_nome(s3, n).shape_id) for n in ('Tabela e-tabela', 'Vídeo e-video', 'Vídeo e-webm'))
    t('slide 3: vídeo sem entrada toca ao abrir (grupo onBegin antes do 1º clique)', len(grupos) == 2
      and xp(grupos[0], 'boolean(p:stCondLst/p:cond[@evt="onBegin"])')
      and not xp(grupos[1], 'boolean(p:stCondLst/p:cond[@evt="onBegin"])'))
    midia = xp(raiz, '//p:cTn[@presetClass="mediacall"]')
    t('mediacall playFrom(0.0) para os dois vídeos com autoplay', sorted(spid_do_efeito(m) for m in midia) == sorted([id_video, id_webm])
      and all(xp(m, 'string(.//p:cmd/@cmd)') == 'playFrom(0.0)' and m.get('presetID') == '1' for m in midia))
    mp4_call = next(m for m in midia if spid_do_efeito(m) == id_video)
    t('mp4 toca junto com a entrada, pela duração lida do arquivo (12,5 s)', mp4_call.get('nodeType') == 'withEffect'
      and xp(mp4_call, 'string(.//p:cmd/p:cBhvr/p:cTn/@dur)') == '12500')
    etapas = [(xp(e, 'string(p:stCondLst/p:cond/@delay)'),
               [(ef.get('presetID'), ef.get('presetClass'), ef.get('nodeType'), spid_do_efeito(ef))
                for ef in xp(e, 'p:childTnLst/p:par/p:cTn')])
              for e in xp(grupos[1], 'p:childTnLst/p:par/p:cTn')]
    t('clique (tabela) → auto (vídeo) 400 ms depois', etapas == [
        ('0', [('10', 'entr', 'clickEffect', id_tabela)]),
        ('400', [('10', 'entr', 'afterEffect', id_video), ('1', 'mediacall', 'withEffect', id_video)])], etapas)
    t('tabela no bldLst como bldGraphic (bldAsOne)', xp(raiz, f'boolean(//p:bldLst/p:bldGraphic[@spid="{id_tabela}"]/p:bldAsOne)'))
    nos_video = {xp(n, 'string(p:cMediaNode/p:tgtEl/p:spTgt/@spid)'): n for n in xp(raiz, '//p:video')}
    t('p:video para cada vídeo: mudo e loop no mp4, som no webm', set(nos_video) == {id_video, id_webm}
      and xp(nos_video[id_video], 'string(p:cMediaNode/@mute)') == '1'
      and xp(nos_video[id_video], 'string(p:cMediaNode/p:cTn/@repeatCount)') == 'indefinite'
      and xp(nos_video[id_webm], 'string(p:cMediaNode/@mute)') == ''
      and xp(nos_video[id_webm], 'string(p:cMediaNode/p:cTn/@repeatCount)') == '')

    raiz = validar_ids_e_alvos(4)
    grupos = xp(raiz, '//p:cTn[@nodeType="mainSeq"]/p:childTnLst/p:par/p:cTn')
    efs = efeitos(raiz)
    t('slide 4: "junto" como primeiro também inicia sozinho (withEffect no grupo onBegin)', len(grupos) == 1
      and xp(grupos[0], 'boolean(p:stCondLst/p:cond[@evt="onBegin"])')
      and [e.get('nodeType') for e in efs] == ['withEffect'])
    final = por_nome(prs.slides[3], 'Texto e-final')
    t('texto sem caixa: sombra vai nas letras (rPr/effectLst)', final.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
      and xp(final._element, 'boolean(.//a:r/a:rPr/a:effectLst/a:outerShdw)')
      and not xp(final._element, 'boolean(p:spPr/a:effectLst)'))
    t('slides sem animação e sem vídeo não têm p:timing', all(xml_do_slide(arquivo, n).find('p:timing', NS) is None for n in range(5, 10)))

    print('\n== ROBUSTEZ ==')
    for documento in (None, {}, {'slides': 'isso não é lista'}, {'slides': [None, 3, {'elementos': 'x'}]}):
        try:
            gerado = ex.gerar_pptx(documento, '', ler_midia=ler_teste)
            valido = len(Presentation(io.BytesIO(gerado)).slides) >= 0
        except Exception as exc:                        # noqa: BLE001
            valido = repr(exc)
        t(f'documento malformado {str(documento)[:40]!r} gera arquivo válido', valido is True, valido)
    quebrado = {'tema': TEMA, 'slides': [{
        'fundo': {'cor': 'cor-que-nao-existe', 'gradiente': 'linear-gradient(', 'imagem': 'https://externo/x.png'},
        'transicao': {'tipo': 'efeito-inventado', 'duracao': 'abc'},
        'elementos': [
            texto('ok-antes', 10, 10, 100, 100, 'antes', anim('fade', 'auto')),
            {'id': 'forma-bomba', 'tipo': 'forma', 'forma': 'estrela', 'x': 'x', 'y': None, 'w': -5, 'h': 'NaN',
             'animacao': anim('zoom', 'clique')},
            {'id': 'tabela-ruim', 'tipo': 'tabela', 'linhas': 'não é lista', 'animacao': anim('fade', 'clique')},
            {'id': 'tipo-estranho', 'tipo': 'grafico3d'},
            'não é dict',
            texto('ok-depois', 10, 200, 100, 100, '<b>depois<', anim('revelar', 'clique'), estilo_invalido=True),
        ]}]}
    coletor.registros.clear()
    original = ex._Exportador._forma

    def forma_que_quebra(self, slide, elemento):
        if elemento.get('id') == 'forma-bomba':
            slide.shapes.add_shape(1, 0, 0, 10, 10)     # deixa lixo na árvore antes de quebrar
            raise RuntimeError('quebra proposital')
        return original(self, slide, elemento)

    with mock.patch.object(ex._Exportador, '_forma', forma_que_quebra):
        gerado = ex.gerar_pptx(quebrado, 'Robustez', ler_midia=ler_teste)
    s = Presentation(io.BytesIO(gerado)).slides[0]
    t('elementos válidos saem; o que quebrou some sem deixar forma pela metade', [f.name for f in s.shapes] == [
        'Texto ok-antes', 'Texto ok-depois'], [f.name for f in s.shapes])
    avisos = [r for r in coletor.registros if r.levelno >= logging.WARNING]
    t('falhas registradas no log', any('forma-bomba' in r.getMessage() for r in avisos)
      and any('tabela-ruim' in r.getMessage() for r in avisos), [r.getMessage() for r in avisos])
    raiz = etree.fromstring(zipfile.ZipFile(io.BytesIO(gerado)).read('ppt/slides/slide1.xml'))
    alvos = {x.get('spid') for x in xp(raiz, '//p:spTgt')}
    t('animação de elemento que falhou não entra na sequência', alvos == {str(f.shape_id) for f in s.shapes}, alvos)
    t('transição desconhecida é ignorada e cor inválida cai no tema', raiz.find('mc:AlternateContent', NS) is None
      and xp(raiz, 'string(p:cSld/p:bg/p:bgPr/a:solidFill/a:srgbClr/@val)') == '0B0612')
    t('URL externa de mídia nunca é lida pelo leitor padrão', ex.ler_midia_padrao('https://externo/x.png') is None)
    t('texto com HTML quebrado ainda sai (como o navegador mostra)', s.shapes[1].text_frame.text == 'depois<',
      s.shapes[1].text_frame.text)

    print('\n== LEITOR PADRÃO DE MÍDIA (Midia + estático) ==')
    bytes_foto = png(320, 180, (1, 2, 3))
    foto = Midia.objects.create(tipo=Midia.Tipo.IMAGEM, nome='foto.png', mime='image/png',
                                arquivo=SimpleUploadedFile('foto.png', bytes_foto, content_type='image/png'))
    t('Midia salva no armazenamento em memória', memoria.exists(foto.arquivo.name))
    t('/apresentacoes/midia/<id>/ lê o arquivo da Midia', ex.ler_midia_padrao(f'/apresentacoes/midia/{foto.pk}/') == (bytes_foto, 'image/png'))
    t('permitir() pode barrar a mídia', ex.ler_midia_padrao(f'/apresentacoes/midia/{foto.pk}/', permitir=lambda m: False) is None)
    t('id inexistente → None', ex.ler_midia_padrao('/apresentacoes/midia/0/') is None)
    estatico = ex.ler_midia_padrao('static:apresentacoes/fontes.json')
    t('static:apresentacoes/... pelo staticfiles', estatico is not None and estatico[0][:1] == b'{' and estatico[1] == 'application/json')
    t('static fora de apresentacoes/ ou com ".." é recusado', ex.ler_midia_padrao('static:apresentacoes/../../redeconfianca/settings.py') is None
      and ex.ler_midia_padrao('static:css/responsive.css') is None)
    documento_real = {'tema': TEMA, 'slides': [{'fundo': {'cor': '#000'}, 'elementos': [
        {'id': 'foto', 'tipo': 'imagem', 'src': f'/apresentacoes/midia/{foto.pk}/', 'x': 0, 'y': 0, 'w': 320, 'h': 180}]}]}
    gerado = ex.gerar_pptx(documento_real, 'Com mídia do banco')
    figura_real = Presentation(io.BytesIO(gerado)).slides[0].shapes[0]
    t('gerar_pptx(documento, titulo) usa o leitor padrão', figura_real.shape_type == MSO_SHAPE_TYPE.PICTURE
      and figura_real.image.blob == bytes_foto)

finally:
    transaction.set_rollback(True)
    marcador.__exit__(None, None, None)
    logging.getLogger('apresentacoes.exportar_pptx').removeHandler(coletor)
    print('\nrollback: nada deste teste foi gravado no banco.')

print(f'\n{ok} OK / {fail} falhas')
sys.exit(1 if fail else 0)
