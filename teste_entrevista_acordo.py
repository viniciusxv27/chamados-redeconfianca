"""Entrevista de desligamento: campo "Acordo" para colaborador inativo.

Pedido: "Adicionar 'Acordo' na entrevista de desligamento e esse campo quando o
usuário estiver inativo".

Também cobre a correção do "Desligar": a data de demissão não era gravada
(o código punha a data em admission_date e salvava só demission_date).

Caches em memória, transação desfeita no fim: não grava nada no banco. A
migração nova do feedback é aplicada dentro da transação (e desfeita junto).
"""
import io
import os
import sys
from datetime import date

import django

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'redeconfianca.settings')
os.environ.setdefault('RC_VARREDURA_ROTINA', '0')

from django.conf import settings

settings.CACHES = {
    'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-fb-acordo'},
    'local': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-fb-acordo-2'},
}
django.setup()

from django.test.utils import setup_test_environment

setup_test_environment()
if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS.append('testserver')

from django.contrib.auth import get_user_model
from django.db import connection, transaction
from django.db.migrations.executor import MigrationExecutor
from django.test import Client

User = get_user_model()
ok = fail = 0


def t(nome, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'  OK   {nome}')
    else:
        fail += 1
        print(f'  FALHA {nome} {extra}')


marcador = transaction.atomic()
marcador.__enter__()
try:
    executor = MigrationExecutor(connection)
    alvos = executor.loader.graph.leaf_nodes('feedback')
    if executor.migration_plan(alvos):
        executor.migrate(alvos)

    import openpyxl

    from feedback.models import ExitInterviewParticipation
    from feedback.views import EXIT_INTERVIEW_KEY

    def pessoa(username, **kw):
        return User.objects.create_user(
            username=username, email=f'{username}@exemplo-teste.local', password='S3nha!teste',
            first_name='Zz', last_name=kw.pop('last_name', username), **kw)

    rh = pessoa('zz.ac.rh', hierarchy='SUPERADMIN', is_staff=True, is_superuser=True)
    comum = pessoa('zz.ac.comum')
    ativo = pessoa('zz.ac.ativo', last_name='AtivoAcordo')
    saindo = pessoa('zz.ac.saindo', last_name='SaindoAcordo')
    inativo = pessoa('zz.ac.inativo', last_name='InativoAcordo')
    inativo.is_active = False
    inativo.status = 'INATIVO'
    inativo.save(update_fields=['is_active', 'status'])
    # Situação Inativo no cadastro, mesmo com login ainda ativo, também conta.
    so_status = pessoa('zz.ac.status', last_name='StatusInativoAcordo')
    so_status.status = 'INATIVO'
    so_status.save(update_fields=['status'])

    for u in (ativo, saindo, inativo, so_status):
        ExitInterviewParticipation.objects.create(
            survey_key=EXIT_INTERVIEW_KEY, user=u, status='COMPLETED', last_step='Concluída',
            dismissal_date=date(2026, 10, 1))

    c = Client()
    c.force_login(rh)
    url = '/feedback/entrevista-desligamento/'
    url_acordo = lambda u: f'{url}{u.id}/acordo/'  # noqa: E731

    print('== TELA DA ENTREVISTA ==')
    html = c.get(url).content.decode()
    t('inativo aparece na lista de acordo', url_acordo(inativo) in html)
    t('situação Inativo também aparece', url_acordo(so_status) in html)
    t('ativo não aparece na lista de acordo', url_acordo(ativo) not in html and url_acordo(saindo) not in html)
    t('o bloco tem as opções', 'Sim, foi acordo' in html and 'Não foi acordo' in html)

    print('\n== SALVAR ACORDO ==')
    r = c.post(url_acordo(inativo), {'acordo': 'SIM', 'acordo_observacao': 'Acordo assinado em 01/10'})
    p = ExitInterviewParticipation.objects.get(user=inativo, survey_key=EXIT_INTERVIEW_KEY)
    t('grava SIM + observação', p.acordo == 'SIM' and p.acordo_observacao == 'Acordo assinado em 01/10',
      (p.acordo, p.acordo_observacao))
    t('volta para a tela', r.status_code == 302)
    c.post(url_acordo(so_status), {'acordo': 'NAO', 'acordo_observacao': ''})
    t('grava NAO', ExitInterviewParticipation.objects.get(user=so_status).acordo == 'NAO')
    c.post(url_acordo(inativo), {'acordo': 'TALVEZ', 'acordo_observacao': 'x'})
    t('opção inválida não grava', ExitInterviewParticipation.objects.get(user=inativo).acordo == 'SIM')
    c.post(url_acordo(inativo), {'acordo': 'SIM', 'acordo_observacao': 'y' * 900})
    t('observação cortada no limite',
      len(ExitInterviewParticipation.objects.get(user=inativo).acordo_observacao) == 500)
    c.post(url_acordo(ativo), {'acordo': 'SIM'})
    t('ativo não recebe acordo', ExitInterviewParticipation.objects.get(user=ativo).acordo == '')
    html = c.get(url).content.decode()
    t('a tela mostra o valor salvo', 'value="SIM" selected' in html or "value=\"SIM\" selected" in html)

    print('\n== DESLIGAR (correção da data de demissão) ==')
    admissao_antes = User.objects.get(pk=saindo.pk).admission_date
    t('o formulário de desligar tem o acordo',
      f'{url}{saindo.id}/desligar/' in html and 'placeholder="Obs. do acordo"' in html)
    r = c.post(f'{url}{saindo.id}/desligar/', {'dismissal_date': '2026-10-20', 'acordo': 'SIM',
                                              'acordo_observacao': 'Acordo no desligamento'})
    saindo.refresh_from_db()
    t('login inativado', not saindo.is_active and saindo.status == 'INATIVO')
    t('data de demissão gravada', saindo.demission_date == date(2026, 10, 20), saindo.demission_date)
    t('data de admissão intacta', saindo.admission_date == admissao_antes)
    p = ExitInterviewParticipation.objects.get(user=saindo)
    t('acordo gravado no desligamento', p.acordo == 'SIM' and p.acordo_observacao == 'Acordo no desligamento',
      (p.acordo, p.acordo_observacao))
    t('desligamento registrado', p.dismissal_executed_at is not None and p.dismissal_date == date(2026, 10, 20))
    html = c.get(url).content.decode()
    t('depois de desligado entra na lista de acordo', url_acordo(saindo) in html)

    # Desligar sem mandar o campo não apaga o acordo que já existia.
    outro = pessoa('zz.ac.outro')
    ExitInterviewParticipation.objects.create(survey_key=EXIT_INTERVIEW_KEY, user=outro, status='REFUSED',
                                              acordo='NAO', acordo_observacao='antes')
    c.post(f'{url}{outro.id}/desligar/', {'dismissal_date': '2026-10-21'})
    p = ExitInterviewParticipation.objects.get(user=outro)
    outro.refresh_from_db()
    t('sem o campo, mantém o acordo', p.acordo == 'NAO' and p.acordo_observacao == 'antes')
    t('e grava a demissão', outro.demission_date == date(2026, 10, 21))

    print('\n== RELATÓRIO ==')
    r = c.get(url + 'relatorio/')
    linhas = {row['user'].id: row for row in r.context['detailed']}
    html = r.content.decode()
    t('coluna Acordo no relatório', '>Acordo</th>' in html)
    t('inativo mostra o acordo', linhas[inativo.id]['acordo'] == 'SIM'
      and linhas[inativo.id]['acordo_label'] == 'Sim, foi acordo')
    t('ativo não mostra acordo', linhas[ativo.id]['acordo'] is None and linhas[ativo.id]['acordo_label'] == '')
    r = c.get(url + 'relatorio/?acordo=sim')
    ids = {row['user'].id for row in r.context['detailed']}
    t('filtro sim', inativo.id in ids and saindo.id in ids and so_status.id not in ids and ativo.id not in ids)
    t('filtro sim só traz inativos com SIM',
      all(row['inativo'] and row['acordo'] == 'SIM' for row in r.context['detailed']))
    r = c.get(url + 'relatorio/?acordo=nao')
    ids = {row['user'].id for row in r.context['detailed']}
    t('filtro não', so_status.id in ids and outro.id in ids and inativo.id not in ids)
    r = c.get(url + 'relatorio/?acordo=nao_informado')
    t('filtro não informado ignora ativos',
      all(row['inativo'] and row['acordo'] == '' for row in r.context['detailed'])
      and ativo.id not in {row['user'].id for row in r.context['detailed']})
    r = c.get(url + 'relatorio/?acordo=bobagem')
    t('filtro inválido é ignorado', r.status_code == 200 and r.context['acordo_filter'] == '')

    print('\n== EXCEL ==')
    r = c.get(url + 'relatorio/exportar/?acordo=sim')
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb['Participação']
    cab = [c_.value for c_ in ws[1]]
    t('coluna Acordo na planilha', 'Acordo' in cab and 'Observação do acordo' in cab, cab)
    i_ac, i_obs = cab.index('Acordo'), cab.index('Observação do acordo')
    dados = {row[0]: row for row in ws.iter_rows(min_row=2, values_only=True)}
    nome_inativo = inativo.get_full_name()
    t('linha do inativo com o acordo', nome_inativo in dados and dados[nome_inativo][i_ac] == 'Sim, foi acordo'
      and dados[nome_inativo][i_obs] == 'y' * 500)
    t('filtro aplicado no Excel', ativo.get_full_name() not in dados)
    t('filtro no resumo', any('Acordo: Sim' in str(v) for v in [wb['Resumo']['A2'].value]))

    print('\n== PERMISSÃO ==')
    c2 = Client()
    c2.force_login(comum)
    r = c2.post(url_acordo(inativo), {'acordo': 'NAO', 'acordo_observacao': 'invasor'})
    t('sem permissão: 403', r.status_code == 403, r.status_code)
    t('sem permissão: não altera', ExitInterviewParticipation.objects.get(user=inativo).acordo == 'SIM')
    r = c2.get(url)
    t('sem permissão: não vê a tela', r.status_code == 403, r.status_code)
finally:
    transaction.set_rollback(True)
    marcador.__exit__(None, None, None)

print(f'\n{ok} OK / {fail} falhas')
sys.exit(1 if fail else 0)
