"""Contagem de caixa: Contas a pagar · PIS/Cofins.

Pedido: categoria de PIS/Cofins em Contas a Pagar para lançar no dia a dia
notas, boletos e recibos de aluguel com mês de competência, fornecedor/CNPJ e
valor; no fechamento, a planilha de PIS/Cofins da competência gerada sozinha.
Retorno da aprovação: buscar na API da Receita e linkar.

O que este teste cobre:

- quem entra (gestor do caixa e a liberação individual caixa.piscofins);
- consulta de CNPJ (dublê da Receita): achou, não existe, Receita fora do ar
  com e sem razão social digitada, CNPJ inválido, cadastro guardado;
- lançar, validar (arquivo, competência futura, valor, número repetido),
  editar trocando o arquivo e apagar;
- a tela: competência, totais por tipo, alertas (CNPJ inativo, duplicado);
- a planilha (abas, fórmulas que apontam para as alíquotas) e o pacote .zip.

Caches em memória e transação desfeita no fim; os arquivos que subirem para o
armazenamento são apagados no fim.
"""
import os
import sys
import zipfile
from datetime import date
from decimal import Decimal
from io import BytesIO

import django

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'redeconfianca.settings')
os.environ.setdefault('RC_VARREDURA_ROTINA', '0')

from django.conf import settings

settings.CACHES = {
    'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-pc'},
    'local': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-pc-2'},
}
django.setup()

from django.test.utils import setup_test_environment

setup_test_environment()
if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS.append('testserver')

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import transaction
from django.test import Client
from django.utils import timezone

from contagem_caixa import cnpj as receita
from contagem_caixa.models import DocumentoPisCofins, FornecedorPisCofins
from users.models import Sector

User = get_user_model()
D = Decimal
ok = fail = 0


def t(nome, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'  OK   {nome}')
    else:
        fail += 1
        print(f'  FALHA {nome} {extra}')


# CNPJs com dígito verificador certo, gerados para o teste.
def cnpj_de(base12):
    d = [int(c) for c in base12]
    for tamanho in (12, 13):
        pesos = list(range(tamanho - 7, 1, -1)) + list(range(9, 1, -1))
        resto = sum(d[i] * pesos[i] for i in range(tamanho)) % 11
        d.append(0 if resto < 2 else 11 - resto)
    return ''.join(map(str, d))


ATIVO, BAIXADO, INEXISTENTE, FORA = cnpj_de('112223330001'), cnpj_de('445556660001'), cnpj_de('778889990001'), cnpj_de('121212120001')
RECEITA = {
    ATIVO: {'cnpj': ATIVO, 'razao_social': 'ZZ IMOBILIARIA TESTE LTDA', 'nome_fantasia': 'ZZ Imóveis',
            'situacao': 'ATIVA', 'atividade': 'Aluguel de imóveis próprios', 'simples': False,
            'municipio': 'Vitória', 'uf': 'ES'},
    BAIXADO: {'cnpj': BAIXADO, 'razao_social': 'ZZ FORNECEDOR BAIXADO ME', 'nome_fantasia': '',
              'situacao': 'BAIXADA', 'atividade': 'Comércio', 'simples': True, 'municipio': 'Serra', 'uf': 'ES'},
}
consultas = []


def buscar_falso(valor):
    d = receita.so_digitos(valor)
    consultas.append(d)
    if d == FORA:
        raise receita.CnpjIndisponivel('fora do ar')
    return RECEITA.get(d)


receita.buscar = buscar_falso
pdf = lambda nome='nota.pdf': SimpleUploadedFile(nome, b'%PDF-1.4 zz teste', content_type='application/pdf')
subidos = []

hoje = timezone.localdate()
# Os lançamentos do teste vão numa competência antiga: o mês atual já tem
# documentos de verdade (a tela está em uso).
mes = date(2019, 5, 1)
comp = f'{mes:%Y-%m}'
comp_hoje = f'{hoje:%Y-%m}'

marcador = transaction.atomic()
marcador.__enter__()
try:
    loja = Sector.objects.create(name='Loja ZZ PisCofins', adabas='ZZ971')
    chefe = User.objects.create_user(
        username='zz.pc.chefe', email='zz.pc.chefe@exemplo-teste.local', password='S3nha!teste',
        first_name='Zz', last_name='Chefe', sector=loja, hierarchy='SUPERADMIN', is_superuser=True)
    contas = User.objects.create_user(
        username='zz.pc.contas', email='zz.pc.contas@exemplo-teste.local', password='S3nha!teste',
        first_name='Zz', last_name='Contas', hierarchy='PADRAO')
    vendedor = User.objects.create_user(
        username='zz.pc.vend', email='zz.pc.vend@exemplo-teste.local', password='S3nha!teste',
        first_name='Zz', last_name='Vend', sector=loja, hierarchy='SUPERVISOR')
    c_chefe, c_contas, c_vend = Client(), Client(), Client()
    c_chefe.force_login(chefe)
    c_contas.force_login(contas)
    c_vend.force_login(vendedor)

    print('== QUEM ENTRA ==')
    r = c_vend.get('/contagem-caixa/pis-cofins/')
    t('vendedor de loja não entra', r.status_code == 302, r.status_code)
    r = c_vend.get(f'/contagem-caixa/pis-cofins/cnpj/?cnpj={ATIVO}')
    t('nem na consulta de CNPJ (403 em JSON)', r.status_code == 403, r.status_code)
    r = c_contas.get('/contagem-caixa/pis-cofins/')
    t('PADRÃO sem liberação não entra', r.status_code == 302, r.status_code)
    from users.module_access import set_user_modules
    set_user_modules(contas, ['caixa.piscofins'], granted_by=chefe)
    contas = User.objects.get(id=contas.id)
    c_contas.force_login(contas)
    r = c_contas.get('/contagem-caixa/pis-cofins/')
    t('com a liberação caixa.piscofins entra', r.status_code == 200, r.status_code)
    html = r.content.decode()
    t('sem as abas do caixa (Visão geral, Sangrias)', 'Visão geral' not in html and 'fa-hand-holding-dollar mr-1.5' not in html)
    t('item no menu Financeiro', 'Contas a pagar · PIS/Cofins</div>' in html)
    r = c_chefe.get('/contagem-caixa/pis-cofins/')
    t('gestor do caixa entra', r.status_code == 200)
    r = c_chefe.get(f'/contagem-caixa/pis-cofins/?competencia={comp}')
    t('mês vazio abre o formulário', 'Nenhum documento lançado' in r.content.decode())

    print('== CONSULTA DE CNPJ ==')
    j = c_contas.get(f'/contagem-caixa/pis-cofins/cnpj/?cnpj={receita.formatar(ATIVO)}').json()
    t('achou na Receita', j['ok'] and j['fornecedor']['razao_social'] == 'ZZ IMOBILIARIA TESTE LTDA', j)
    t('com link para o comprovante', ATIVO in j['fornecedor']['link_receita'])
    t('e guardou o fornecedor', FornecedorPisCofins.objects.filter(cnpj=ATIVO, situacao='ATIVA').exists())
    j = c_contas.get(f'/contagem-caixa/pis-cofins/cnpj/?cnpj={INEXISTENTE}').json()
    t('CNPJ que não existe', not j['ok'] and 'não consta' in j['erro'], j)
    j = c_contas.get('/contagem-caixa/pis-cofins/cnpj/?cnpj=11.111.111/1111-11').json()
    t('CNPJ inválido', not j['ok'] and 'inválido' in j['erro'], j)
    j = c_contas.get(f'/contagem-caixa/pis-cofins/cnpj/?cnpj={FORA}').json()
    t('Receita fora do ar pede a razão social', not j['ok'] and j.get('indisponivel'), j)

    print('== LANÇAR ==')

    def lancar(cliente=c_contas, **campos):
        dados = {'tipo': 'NF', 'competencia': comp, 'cnpj': ATIVO, 'valor': '1.250,00',
                 'numero': '123', 'data_documento': hoje.isoformat(), 'arquivo': pdf()}
        dados.update(campos)
        dados = {k: v for k, v in dados.items() if v is not None}
        return cliente.post('/contagem-caixa/pis-cofins/lancar/', dados)

    r = lancar(tipo='ALUGUEL', numero='', valor='3.500,00', descricao='Aluguel da loja', loja=loja.id)
    d1 = DocumentoPisCofins.objects.filter(descricao='Aluguel da loja').first()
    if d1:
        subidos.append(d1.arquivo.name)
    t('recibo de aluguel lançado', d1 is not None and d1.valor == D('3500.00') and d1.tipo == 'ALUGUEL')
    t('com fornecedor, competência, loja e autor',
      d1 and d1.fornecedor.cnpj == ATIVO and d1.competencia == mes and d1.loja == loja and d1.registrado_por == contas)
    t('nome original do arquivo guardado', d1 and d1.nome_arquivo == 'nota.pdf')
    lancar(numero='123')
    d2 = DocumentoPisCofins.objects.filter(numero='123').first()
    if d2:
        subidos.append(d2.arquivo.name)
    t('nota fiscal lançada', d2 is not None)

    antes = DocumentoPisCofins.objects.count()
    for nome, campos in (
        ('número repetido do mesmo fornecedor', {'numero': '0123'.lstrip('0')}),
        ('sem arquivo', {'arquivo': None, 'numero': '9001'}),
        ('arquivo de tipo estranho', {'arquivo': SimpleUploadedFile('x.exe', b'MZ'), 'numero': '9002'}),
        ('competência futura', {'competencia': f'{hoje.year + 1}-01', 'numero': '9003'}),
        ('valor zero', {'valor': '0', 'numero': '9004'}),
        ('CNPJ inexistente', {'cnpj': INEXISTENTE, 'numero': '9005'}),
        ('Receita fora sem razão social', {'cnpj': FORA, 'numero': '9006'}),
    ):
        r = lancar(**campos)
        t(f'recusa {nome}', DocumentoPisCofins.objects.count() == antes)
    r = lancar(cnpj=FORA, numero='777', razao_social='ZZ DIGITADA LTDA', valor='80,00')
    d3 = DocumentoPisCofins.objects.filter(numero='777').first()
    if d3:
        subidos.append(d3.arquivo.name)
    t('Receita fora + razão digitada: lança', d3 is not None and d3.fornecedor.razao_social == 'ZZ DIGITADA LTDA')
    t('e marca o fornecedor como não conferido', d3 and d3.fornecedor.consultado_em is None)
    lancar(cnpj=BAIXADO, numero='55', valor='200,00', tipo='BOLETO')
    d4 = DocumentoPisCofins.objects.filter(numero='55').first()
    if d4:
        subidos.append(d4.arquivo.name)
    t('fornecedor baixado lança, com aviso', d4 is not None and not d4.fornecedor.situacao_ok)
    lancar(cnpj=BAIXADO, numero='', valor='200,00', tipo='OUTRO')
    d5 = DocumentoPisCofins.objects.filter(tipo='OUTRO').first()
    if d5:
        subidos.append(d5.arquivo.name)

    print('== TELA ==')
    r = c_contas.get(f'/contagem-caixa/pis-cofins/?competencia={comp}')
    html = r.content.decode()
    t('lista os documentos', 'ZZ IMOBILIARIA TESTE LTDA' in html and 'ZZ DIGITADA LTDA' in html)
    t('total do mês com milhar', 'R$ 5.230,00' in html, 'total esperado 3500+1250+80+200+200')
    t('selo de CNPJ baixado', 'CNPJ baixada' in html)
    t('selo de não conferido', 'não conferido na Receita' in html)
    t('alerta de CNPJ fora de ATIVA', 'fora da situação ATIVA' in html)
    t('alerta de possível duplicado (mesmo fornecedor e valor sem número)', 'possível duplicado' in html)
    t('link do comprovante na Receita', receita.LINK_RECEITA.format(cnpj=ATIVO).replace('&', '&amp;') in html
      or receita.LINK_RECEITA.format(cnpj=ATIVO) in html)
    r = c_contas.get(f'/contagem-caixa/pis-cofins/?competencia={comp}&tipo=ALUGUEL')
    html = r.content.decode()
    t('filtro por tipo', 'Aluguel da loja' in html and 'nº 777' not in html and 'nº 123' not in html)
    r = c_contas.get(f'/contagem-caixa/pis-cofins/?competencia={comp}&q={BAIXADO[:8]}')
    html = r.content.decode()
    t('busca pelo CNPJ', 'ZZ FORNECEDOR BAIXADO ME' in html and 'Aluguel da loja' not in html)
    r = c_contas.get('/contagem-caixa/pis-cofins/?competencia=2020-03')
    t('outra competência vazia', 'Nenhum documento lançado em 03/2020' in r.content.decode())

    print('== EDITAR / APAGAR ==')
    nome_antigo = d2.arquivo.name
    r = c_contas.post(f'/contagem-caixa/pis-cofins/{d2.id}/editar/', {
        'tipo': 'NF', 'competencia': comp, 'cnpj': ATIVO, 'valor': '1.300,00', 'numero': '123',
        'descricao': 'Energia', 'arquivo': pdf('nova.pdf')})
    d2.refresh_from_db()
    subidos.append(d2.arquivo.name)
    t('edita valor e descrição', d2.valor == D('1300.00') and d2.descricao == 'Energia')
    t('troca o arquivo e apaga o antigo', d2.nome_arquivo == 'nova.pdf' and d2.arquivo.name != nome_antigo
      and not d2.arquivo.storage.exists(nome_antigo))
    r = c_contas.post(f'/contagem-caixa/pis-cofins/{d2.id}/editar/', {
        'tipo': 'NF', 'competencia': comp, 'cnpj': ATIVO, 'valor': '1.300,00', 'numero': '123', 'descricao': 'Sem trocar'})
    d2.refresh_from_db()
    t('editar sem arquivo mantém o atual', d2.descricao == 'Sem trocar' and d2.arquivo.name)
    nome5 = d5.arquivo.name
    c_contas.post(f'/contagem-caixa/pis-cofins/{d5.id}/apagar/')
    t('apaga o documento e o arquivo', not DocumentoPisCofins.objects.filter(id=d5.id).exists()
      and not d1.arquivo.storage.exists(nome5))

    print('== PLANILHA E PACOTE ==')
    r = c_contas.get(f'/contagem-caixa/pis-cofins/planilha/?competencia={comp}')
    t('planilha baixa', r.status_code == 200 and 'spreadsheetml' in r['Content-Type'], r.status_code)
    t('nome com a competência', f'pis_cofins_{comp}.xlsx' in r['Content-Disposition'])
    from openpyxl import load_workbook
    livro = load_workbook(BytesIO(r.content))
    t('abas', livro.sheetnames == ['Resumo', 'Documentos', 'Por fornecedor'], livro.sheetnames)
    docs = livro['Documentos']
    linhas = [[c.value for c in row] for row in docs.iter_rows(min_row=2, max_row=5)]
    t('4 documentos na planilha', all(l[5] for l in linhas) and docs.cell(row=6, column=6).value is None,
      [l[5] for l in linhas])
    t('PIS e Cofins são fórmulas sobre as alíquotas', linhas[0][11] == '=ROUND(K2*Resumo!$B$5,2)'
      and linhas[0][12] == '=ROUND(K2*Resumo!$B$6,2)', linhas[0][11:13])
    resumo = livro['Resumo']
    t('alíquotas 1,65% e 7,6% editáveis', resumo['B5'].value == 0.0165 and resumo['B6'].value == 0.076)
    t('resumo por tipo com SUMIF', str(resumo['C9'].value).startswith('=SUMIF('), resumo['C9'].value)
    atencao = [l[13] for l in linhas]
    t('coluna Atenção avisa CNPJ baixado e não conferido',
      any('baixada' in (a or '') for a in atencao) and any('não conferido' in (a or '') for a in atencao), atencao)
    soma = sum(l[10] for l in linhas)
    t('soma dos valores', abs(soma - 5080.0) < 0.01, soma)

    r = c_contas.get(f'/contagem-caixa/pis-cofins/pacote/?competencia={comp}')
    t('pacote baixa', r.status_code == 200 and r['Content-Type'] == 'application/zip', r.status_code)
    zf = zipfile.ZipFile(BytesIO(r.content))
    nomes = zf.namelist()
    t('planilha na raiz', f'pis_cofins_{comp}.xlsx' in nomes, nomes)
    t('os 4 arquivos em documentos/', len([n for n in nomes if n.startswith('documentos/')]) == 4, nomes)
    t('sem arquivo faltando', 'LEIA-ME_arquivos_faltando.txt' not in nomes,
      zf.read('LEIA-ME_arquivos_faltando.txt').decode() if 'LEIA-ME_arquivos_faltando.txt' in nomes else '')
    t('arquivo íntegro no pacote', zf.read([n for n in nomes if n.startswith('documentos/')][0]) == b'%PDF-1.4 zz teste')

    print('== LEITURA DO ARQUIVO ==')
    from contagem_caixa import leitura_documento as leitura
    nfe = (f'<?xml version="1.0"?><nfeProc xmlns="http://www.portalfiscal.inf.br/nfe"><NFe><infNFe>'
           f'<ide><nNF>4521</nNF><dhEmi>{hoje.isoformat()}T10:00:00-03:00</dhEmi></ide>'
           f'<emit><CNPJ>{ATIVO}</CNPJ><xNome>ZZ IMOBILIARIA TESTE LTDA</xNome></emit>'
           f'<det nItem="1"><prod><xProd>Cadeira giratória</xProd></prod></det>'
           f'<total><ICMSTot><vNF>1234.50</vNF></ICMSTot></total></infNFe></NFe></nfeProc>').encode()
    r = c_contas.post('/contagem-caixa/pis-cofins/ler/', {'arquivo': SimpleUploadedFile('nota.xml', nfe, 'text/xml')},
                      HTTP_ACCEPT='application/json')
    j = r.json()
    c = j.get('campos', {})
    t('XML da NF-e: lê tudo sem IA', j['ok'] and j['fonte'] == 'xml' and c.get('tipo') == 'NF' and c.get('cnpj') == ATIVO
      and c.get('numero') == '4521' and c.get('valor') == '1234,50' and c.get('data_documento') == hoje.isoformat(), j)
    t('competência = mês da emissão', c.get('competencia') == comp_hoje, c.get('competencia'))
    t('descrição do produto', c.get('descricao') == 'Cadeira giratória', c.get('descricao'))
    t('não grava nada', not DocumentoPisCofins.objects.filter(numero='4521').exists())

    nfse = (f'<CompNfse xmlns="http://www.abrasf.org.br/nfse.xsd"><Nfse><InfNfse><Numero>778</Numero>'
            f'<DataEmissao>{hoje.isoformat()}T09:00:00</DataEmissao><Competencia>{comp_hoje}-01</Competencia>'
            f'<ValoresNfse><ValorLiquidoNfse>980.00</ValorLiquidoNfse></ValoresNfse>'
            f'<PrestadorServico><IdentificacaoPrestador><CpfCnpj><Cnpj>{BAIXADO}</Cnpj></CpfCnpj></IdentificacaoPrestador>'
            f'<RazaoSocial>ZZ FORNECEDOR BAIXADO ME</RazaoSocial></PrestadorServico>'
            f'<DeclaracaoPrestacaoServico><InfDeclaracaoPrestacaoServico><Servico><Valores><ValorServicos>1000.00</ValorServicos></Valores>'
            f'<Discriminacao>Manutenção do ar-condicionado</Discriminacao></Servico></InfDeclaracaoPrestacaoServico>'
            f'</DeclaracaoPrestacaoServico></InfNfse></Nfse></CompNfse>').encode()
    lido = leitura.ler_documento(nfse, 'nfse.xml')
    c = lido['campos']
    t('XML de NFS-e (ABRASF)', c.get('cnpj') == BAIXADO and c.get('numero') == '778' and c.get('valor') == '980.00'
      and c.get('competencia') == comp_hoje and 'ar-condicionado' in c.get('descricao', ''), c)
    lido = leitura.ler_documento(b'<?xml version="1.0"?><qualquer><coisa/></qualquer>', 'x.xml')
    t('XML que não é nota avisa', not lido['campos'] and lido['avisos'], lido)

    linha = '23790.12301 60000.000053 25000.456704 1 99990000150050'
    t('linha digitável do boleto dá o valor', leitura.valor_da_linha_digitavel(f'Pague até\n{linha}\n') == '1500.50')
    t('sem separador também', leitura.valor_da_linha_digitavel(linha.replace('.', '').replace(' ', '')) == '1500.50')
    conta = '836200000015 245600480001 123456789012 345678901234'
    t('conta de consumo (48 dígitos)', leitura.valor_da_linha_digitavel(conta) == '124.56',
      leitura.valor_da_linha_digitavel(conta))

    # PDF/foto pela IA (dublê): a IA leu o NOSSO CNPJ como fornecedor.
    NOSSO = cnpj_de('393939390001')
    leitura_cnpjs, leitura_ia, leitura_anexo = leitura.cnpjs_da_rede, leitura.ler_com_ia, None
    import cartoes.ai as cartoes_ai
    preparar_original = cartoes_ai.preparar_anexo
    try:
        leitura.cnpjs_da_rede = lambda: {NOSSO}
        cartoes_ai.preparar_anexo = lambda conteudo, mime='': {
            'tipo': 'pdf', 'imagens': [], 'texto': f'BENEFICIARIO ZZ IMOBILIARIA CNPJ {receita.formatar(ATIVO)}\n'
                                                   f'PAGADOR REDE CNPJ {receita.formatar(NOSSO)}\n{linha}'}
        leitura.ler_com_ia = lambda anexo: ({'tipo': 'BOLETO', 'fornecedor_cnpj': NOSSO, 'pagador_cnpj': ATIVO,
                                             'fornecedor_nome': 'ZZ IMOBILIARIA', 'numero': '000123',
                                             'data_emissao': hoje.strftime('%d/%m/%Y'), 'valor': '1500,00',
                                             'descricao': 'Aluguel de outubro', 'competencia': comp_hoje}, '')
        r = c_contas.post('/contagem-caixa/pis-cofins/ler/',
                          {'arquivo': SimpleUploadedFile('boleto.pdf', b'%PDF-1.4 boleto', 'application/pdf')},
                          HTTP_ACCEPT='application/json')
        j = r.json()
        c = j['campos']
        t('PDF pela IA: lados trocados são corrigidos (fornecedor não é o nosso CNPJ)', c.get('cnpj') == ATIVO, c)
        t('valor da linha digitável prevalece, com aviso', c.get('valor') == '1500,50'
          and any('linha digitável' in a for a in j['avisos']), (c.get('valor'), j['avisos']))
        t('data dd/mm/aaaa vira ISO', c.get('data_documento') == hoje.isoformat(), c.get('data_documento'))
        t('tipo, número, descrição e competência', c.get('tipo') == 'BOLETO' and c.get('numero') == '000123'
          and c.get('descricao') == 'Aluguel de outubro' and c.get('competencia') == comp_hoje, c)

        leitura.ler_com_ia = lambda anexo: ({'fornecedor_cnpj': NOSSO, 'pagador_cnpj': '', 'valor': ''}, '')
        cartoes_ai.preparar_anexo = lambda conteudo, mime='': {'tipo': 'imagem', 'imagens': [(b'x', 'image/png')], 'texto': ''}
        lido = leitura.ler_documento(b'\x89PNG foto', 'foto.png', 'image/png')
        t('só o nosso CNPJ: não preenche e avisa', 'cnpj' not in lido['campos'] and lido['avisos'], lido)

        leitura.ler_com_ia = lambda anexo: (None, 'sem rede')
        lido = leitura.ler_documento(b'\x89PNG foto', 'foto.png', 'image/png')
        t('IA fora: pede para preencher à mão', not lido['campos'] and 'sem rede' in lido['avisos'][0], lido)

        leitura.ler_com_ia = lambda anexo: ({'tipo': 'NF', 'fornecedor_cnpj': ATIVO, 'numero': '123', 'valor': '1300'}, '')
        r = c_contas.post('/contagem-caixa/pis-cofins/ler/',
                          {'arquivo': SimpleUploadedFile('nota.pdf', b'%PDF-1.4 nota', 'application/pdf')},
                          HTTP_ACCEPT='application/json')
        t('arquivo já lançado é avisado na leitura', any('já foi lançado' in a for a in r.json()['avisos']), r.json())
    finally:
        leitura.cnpjs_da_rede, leitura.ler_com_ia = leitura_cnpjs, leitura_ia
        cartoes_ai.preparar_anexo = preparar_original
    r = c_contas.post('/contagem-caixa/pis-cofins/ler/', {'arquivo': SimpleUploadedFile('x.exe', b'MZ')},
                      HTTP_ACCEPT='application/json')
    t('leitura recusa tipo estranho', not r.json()['ok'])
    r = c_vend.post('/contagem-caixa/pis-cofins/ler/', {'arquivo': SimpleUploadedFile('n.xml', nfe)},
                    HTTP_ACCEPT='application/json')
    t('leitura fechada para quem não tem acesso', r.status_code == 403, r.status_code)

    print('== FORMATO DO CNPJ ==')
    t('formatar', receita.formatar(ATIVO) == f'{ATIVO[:2]}.{ATIVO[2:5]}.{ATIVO[5:8]}/{ATIVO[8:12]}-{ATIVO[12:]}')
    t('valido recusa repetido', not receita.valido('11111111111111') and receita.valido(ATIVO))
finally:
    marcador.__exit__(Exception, Exception('rollback'), None)
    from contagem_caixa.models import DocumentoPisCofins as _D
    storage = _D._meta.get_field('arquivo').storage
    for nome in set(subidos):
        try:
            if nome and storage.exists(nome):
                storage.delete(nome)
        except Exception as exc:  # noqa: BLE001
            print('  (não apagou do armazenamento)', nome, exc)

print(f'\n{ok} OK, {fail} falha(s)')
sys.exit(1 if fail else 0)
