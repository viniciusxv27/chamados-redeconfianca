"""Conciliação de cartões: o que a adquirente diz que vendemos × as vendas no portal.

Lê o relatório de vendas da adquirente (Cielo, Rede, Stone, Getnet, PagSeguro…)
em CSV ou Excel. Cada uma escreve o cabeçalho de um jeito; as colunas são
reconhecidas pelo nome (NSU/DOC, data da venda, valor bruto, líquido, taxa,
parcelas, bandeira, previsão de pagamento, estabelecimento).

O par é pelo NSU (sem zeros à esquerda), com a data a até 1 dia: a venda do
portal guarda o NSU desde 15/10/2026. Valor igual (até 5 centavos) concilia;
valor diferente fica "divergente" para alguém olhar.
"""
import csv
import hashlib
import io
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone

from . import auditoria
from .extrato import _sem_acento, data_br, valor_br
from .models import EstabelecimentoCartao, RecebivelCartao

TOLERANCIA = Decimal('0.05')

COLUNAS = {
    'nsu': ('nsu', 'nsu/doc', 'nsu doc', 'doc', 'numero do documento', 'nsu/cv', 'cv', 'numero cv', 'nsu host',
            'codigo da transacao', 'nsu da transacao', 'id da transacao', 'stone id', 'nsu sitef'),
    'autorizacao': ('autorizacao', 'codigo de autorizacao', 'cod. autorizacao', 'cod autorizacao', 'numero autorizacao'),
    'data': ('data da venda', 'data venda', 'data', 'data da transacao', 'data transacao', 'data de venda',
             'data/hora da venda', 'data e hora'),
    'bruto': ('valor bruto', 'valor da venda', 'valor venda', 'valor', 'valor total', 'valor da transacao',
              'valor bruto da venda'),
    'liquido': ('valor liquido', 'valor liquido da venda', 'liquido', 'valor a receber'),
    'taxa': ('taxa', 'valor da taxa', 'taxa (r$)', 'mdr', 'valor descontado', 'desconto'),
    'parcelas': ('parcelas', 'qtd parcelas', 'quantidade de parcelas', 'numero de parcelas', 'plano', 'nº parcelas'),
    'bandeira': ('bandeira', 'brand'),
    'modalidade': ('modalidade', 'produto', 'tipo', 'forma de pagamento', 'tipo de transacao'),
    'previsao': ('previsao de pagamento', 'data prevista', 'data de pagamento', 'previsao', 'data do pagamento',
                 'data prevista de pagamento'),
    'estabelecimento': ('estabelecimento', 'ec', 'numero do estabelecimento', 'codigo do estabelecimento',
                        'stonecode', 'stone code', 'pv', 'numero pv', 'ponto de venda'),
}


class RelatorioInvalido(Exception):
    pass


def so_digitos(texto):
    return ''.join(c for c in str(texto or '') if c.isdigit())


def nsu_normal(texto):
    digitos = so_digitos(texto)
    return digitos.lstrip('0') or digitos if digitos else str(texto or '').strip().upper()


def _registros(nome, conteudo):
    if nome.lower().endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook
        livro = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        aba = livro.worksheets[0]
        return [[('' if c is None else c) for c in linha] for linha in aba.iter_rows(values_only=True)]
    texto = conteudo.decode('utf-8-sig', errors='replace')
    if texto.count('�') > 3:
        texto = conteudo.decode('latin-1')
    separador = ';' if texto[:4000].count(';') >= texto[:4000].count(',') else ','
    return list(csv.reader(io.StringIO(texto), delimiter=separador))


def _data(valor):
    if hasattr(valor, 'date') and not isinstance(valor, str):
        return valor.date() if isinstance(valor, datetime) else valor
    return data_br(valor)


def _decimal(valor):
    if isinstance(valor, (int, float, Decimal)):
        return Decimal(str(valor)).quantize(Decimal('0.01'))
    return valor_br(valor)


def ler_relatorio(nome, conteudo):
    registros = _registros(nome, conteudo)
    for inicio, cab in enumerate(registros[:20]):
        nomes = [_sem_acento(c) for c in cab]
        indices = {chave: next((i for i, n in enumerate(nomes) if n in opcoes), None)
                   for chave, opcoes in COLUNAS.items()}
        if indices['nsu'] is not None and indices['data'] is not None and indices['bruto'] is not None:
            break
    else:
        raise RelatorioInvalido('Não reconheci o relatório: ele precisa das colunas de NSU (ou DOC), '
                                'data da venda e valor bruto.')
    vendas = []
    for registro in registros[inicio + 1:]:
        def pega(chave):
            i = indices[chave]
            return registro[i] if i is not None and i < len(registro) else ''
        if not str(pega('nsu')).strip():
            continue
        try:
            data = _data(pega('data'))
            bruto = _decimal(pega('bruto'))
        except (ValueError, InvalidOperation):
            continue
        try:
            liquido = _decimal(pega('liquido')) if str(pega('liquido')).strip() else Decimal('0')
        except InvalidOperation:
            liquido = Decimal('0')
        try:
            taxa = abs(_decimal(pega('taxa'))) if str(pega('taxa')).strip() else Decimal('0')
        except InvalidOperation:
            taxa = Decimal('0')
        if not taxa and liquido:
            taxa = bruto - liquido
        try:
            previsao = _data(pega('previsao')) if str(pega('previsao')).strip() else None
        except ValueError:
            previsao = None
        parcelas = so_digitos(str(pega('parcelas')).split('/')[-1]) or '1'
        vendas.append({
            'nsu': str(pega('nsu')).strip()[:40], 'autorizacao': str(pega('autorizacao')).strip()[:40],
            'data_venda': data, 'valor_bruto': abs(bruto), 'valor_liquido': abs(liquido), 'taxa': taxa,
            'parcelas': max(1, min(int(parcelas), 99)), 'bandeira': str(pega('bandeira')).strip()[:40],
            'modalidade': str(pega('modalidade')).strip()[:40], 'previsao_pagamento': previsao,
            'estabelecimento': so_digitos(pega('estabelecimento')) or str(pega('estabelecimento')).strip()[:40],
        })
    if not vendas:
        raise RelatorioInvalido('O relatório não tem nenhuma venda legível.')
    return vendas


@transaction.atomic
def importar(nome, conteudo, adquirente='', loja=None, usuario=None, request=None):
    """Grava as vendas da adquirente (sem repetir) e já tenta conciliar."""
    vendas = ler_relatorio(nome, conteudo)
    mapa = dict(EstabelecimentoCartao.objects.values_list('codigo', 'loja_id'))
    novas = repetidas = 0
    criados = []
    for v in vendas:
        ident = hashlib.sha1(f"{adquirente}|{v['estabelecimento']}|{nsu_normal(v['nsu'])}|{v['data_venda']}|"
                             f"{v['valor_bruto']}".encode()).hexdigest()
        loja_id = mapa.get(v['estabelecimento']) or (loja.id if loja else None)
        obj, criado = RecebivelCartao.objects.get_or_create(
            identificador=ident, defaults=dict(v, adquirente=adquirente[:40], loja_id=loja_id, importado_por=usuario))
        if criado:
            novas += 1
            criados.append(obj)
        else:
            repetidas += 1
    auditoria.registrar(request, 'IMPORTAR', f'Importou relatório de cartões {adquirente or ""} ({nome}): '
                        f'{novas} venda(s) nova(s)', area='cartoes', loja=loja, usuario=usuario,
                        depois={'arquivo': nome, 'novas': novas, 'repetidas': repetidas})
    conciliados = conciliar_pendentes(RecebivelCartao.objects.filter(pk__in=[c.pk for c in criados]),
                                      usuario=usuario, request=request)
    return novas, repetidas, conciliados


def vendas_no_cartao():
    from vendas.models import Venda
    return Venda.objects.filter(forma_pagamento='CARTAO').exclude(nsu='')


def _par(recebivel, usadas):
    alvo = nsu_normal(recebivel.nsu)
    inicio, fim = recebivel.data_venda - timedelta(days=1), recebivel.data_venda + timedelta(days=1)
    for venda in vendas_no_cartao().filter(data_venda__date__gte=inicio, data_venda__date__lte=fim):
        if venda.pk in usadas or nsu_normal(venda.nsu) != alvo:
            continue
        return venda
    return None


def _status_pelo_valor(recebivel, venda):
    return 'CONCILIADO' if abs(venda.total - recebivel.valor_bruto) <= TOLERANCIA else 'DIVERGENTE'


def conciliar_pendentes(qs=None, usuario=None, request=None):
    qs = qs if qs is not None else RecebivelCartao.objects.filter(status='PENDENTE')
    usadas = set(RecebivelCartao.objects.exclude(venda__isnull=True).values_list('venda_id', flat=True))
    feitos = 0
    for rec in qs.filter(status='PENDENTE', venda__isnull=True):
        venda = _par(rec, usadas)
        if not venda:
            continue
        rec.venda = venda
        rec.status = _status_pelo_valor(rec, venda)
        rec.conciliado_em, rec.conciliado_por = timezone.now(), usuario
        if not rec.loja_id and venda.loja_id:
            rec.loja_id = venda.loja_id
        rec.save()
        usadas.add(venda.pk)
        feitos += 1
    if feitos:
        auditoria.registrar(request, 'CONCILIAR', f'Conciliou {feitos} venda(s) no cartão pelo NSU', area='cartoes',
                            usuario=usuario)
    return feitos


@transaction.atomic
def vincular(recebivel, venda, usuario=None, request=None):
    if RecebivelCartao.objects.filter(venda=venda).exclude(pk=recebivel.pk).exists():
        raise ValueError('Esta venda já está ligada a outro registro da adquirente.')
    recebivel.venda = venda
    recebivel.status = _status_pelo_valor(recebivel, venda)
    recebivel.conciliado_em, recebivel.conciliado_por = timezone.now(), usuario
    recebivel.save()
    auditoria.registrar(request, 'CONCILIAR', f'Ligou o NSU {recebivel.nsu} à venda #{venda.pk}', area='cartoes',
                        obj=recebivel, valor=recebivel.valor_bruto, loja=recebivel.loja, usuario=usuario)


@transaction.atomic
def marcar(recebivel, status, observacao='', usuario=None, request=None):
    """Aceitar a divergência (vira conciliado), ignorar ou soltar o par."""
    antes = {'status': recebivel.status, 'venda': recebivel.venda_id}
    if status == 'PENDENTE':
        recebivel.venda = None
    recebivel.status = status
    recebivel.observacao = (observacao or recebivel.observacao)[:255]
    recebivel.conciliado_em, recebivel.conciliado_por = timezone.now(), usuario
    recebivel.save()
    acao = 'DESCONCILIAR' if status == 'PENDENTE' else 'CONCILIAR'
    auditoria.registrar(request, acao, f'Marcou o NSU {recebivel.nsu} como {recebivel.get_status_display().lower()}'
                        + (f' — {observacao}' if observacao else ''), area='cartoes', obj=recebivel,
                        valor=recebivel.valor_bruto, loja=recebivel.loja, antes=antes,
                        depois={'status': status, 'venda': recebivel.venda_id}, usuario=usuario)
