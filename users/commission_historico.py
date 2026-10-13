"""Histórico de comissão por cargo e loja, mês a mês (/users/commission/projecao/historico/).

Responde perguntas como "quanto o GG do Montserrat recebeu desde que temos
histórico no portal". A fonte é o retrato de cada linha das planilhas que o
portal guarda a cada mês (``CommissionUserReferenceHistory``) — o
``CommissionMonthlyTotal`` só existe para os meses em que alguém clicou em
"Sincronizar tabela", então não serve de série.

O valor recebido é a mesma coluna que a tela de comissionamento mostra
("REMUNERAÇÃO FINAL TOTAL"; no coordenador, "REMUNERAÇÃO FINAL COO/SNP"). O
detalhamento sai das colunas da própria planilha:

    comissão + premiação alto desempenho + bônus da loja      (TOTAL COMISSÃO + PREMIAÇÃO)
    + aceleradores                                            (REMUNERAÇÃO TOTAL + ACELERADORES)
    + ajuste do IQ (deflator/acelerador de qualidade)         (REMUNERAÇÃO COM IQ)
    − descontos                                               → REMUNERAÇÃO FINAL

Escopo de visão igual ao do comissionamento: superadmin vê a rede, o
coordenador as lojas que coordena, o gerente a própria loja e os demais só a
própria linha.
"""
import unicodedata
from collections import OrderedDict, defaultdict
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import redirect, render

from .commission_views import MONTH_NAMES_PT, get_user_role
from .models import CommissionSpreadsheetVersion, CommissionUserReferenceHistory

CARGOS = OrderedDict([
    ('GG', 'Gerente Geral (GG)'),
    ('GO', 'Gerente Operacional (GO)'),
    ('CN', 'Consultor de Negócios (CN)'),
    ('COO', 'Coordenador / SNP'),
])
PARTES = [
    ('comissao', 'Comissão'),
    ('premiacao', 'Premiação'),
    ('bonus_loja', 'Bônus loja'),
    ('aceleradores', 'Aceleradores'),
    ('ajuste_iq', 'Ajuste IQ'),
    ('descontos', 'Descontos'),
]


def normalizar(texto):
    """'Loja Glória' → 'GLORIA'; 'GLÓRIA' → 'GLORIA'. Casa PDV da planilha com o setor."""
    texto = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore').decode()
    texto = ' '.join(texto.upper().split())
    return texto[5:] if texto.startswith('LOJA ') else texto


def _num(valor):
    try:
        return float(valor or 0)
    except (TypeError, ValueError):
        return 0.0


def _coluna(dados, *trechos):
    """Valor da primeira coluna cujo nome (sem acento/espaço) contém todos os trechos."""
    for chave, valor in dados.items():
        nome = normalizar(chave).replace('_', ' ')
        if all(t in nome for t in trechos):
            return _num(valor)
    return 0.0


def _detalhe(papel, dados):
    """As partes do valor do mês e o valor recebido, a partir da linha da planilha."""
    if papel == 'coordenador':
        comissao = _num(dados.get('COMISSÃO_COO/SNP'))
        premiacao = _num(dados.get('PREMIAÇÃO_COO/SNP'))
        bruta = _num(dados.get('REMUNERÇÃO BRUTA')) or comissao + premiacao
        com_iq = _num(dados.get('REMUNERÇÃO COM IQ')) or bruta
        final = _coluna(dados, 'REMUNERA', 'FINAL', 'COO')
        bonus = 0.0
        acel = bruta - (comissao + premiacao)
    else:
        comissao = _num(dados.get('Total Comissão'))
        premiacao = _num(dados.get('TOTAL PREMIAÇÃO ALTO DESEMPENHO'))
        bonus = _num(dados.get('TOTAL BONUS LOJA'))
        base = _num(dados.get('TOTAL COMISSÃO + PREMIAÇÃO')) or comissao + premiacao + bonus
        com_acel = _num(dados.get('REMUNERAÇÃO TOTAL + ACELERADORES')) or base
        com_iq = _num(dados.get('REMUNERÇÃO COM IQ')) or com_acel
        final = _coluna(dados, 'REMUNERA', 'FINAL', 'TOTAL')
        acel = com_acel - base
        bruta = com_acel
    descontos = _num(dados.get('DESCONTO_TOTAL'))
    if not final:
        final = com_iq - descontos
    return {
        'comissao': comissao, 'premiacao': premiacao, 'bonus_loja': bonus,
        'aceleradores': acel, 'ajuste_iq': com_iq - bruta, 'descontos': descontos,
        'recebido': final,
    }


def _cargo(entrada):
    if entrada.role == 'coordenador':
        return 'COO'
    cargo = str((entrada.row_data or {}).get('CARGO') or '').strip().upper()
    if cargo in CARGOS:
        return cargo
    return 'CN' if entrada.role == 'cn' else ('GG' if entrada.role == 'gerente' else cargo or '?')


def linhas_do_historico(ver_rascunho=False):
    """Uma linha por pessoa e mês, já com cargo, loja e o detalhamento."""
    qs = CommissionUserReferenceHistory.objects.select_related('source_version')
    if not ver_rascunho:
        qs = qs.filter(source_version__status=CommissionSpreadsheetVersion.STATUS_RELEASED)
    linhas = []
    for e in qs.order_by('year', 'month', 'user_name'):
        dados = e.row_data or {}
        pdv = str(dados.get('PDV') or '').strip()
        versao = e.source_version
        linhas.append({
            'ano': e.year, 'mes': e.month, 'chave_mes': f'{e.year}-{e.month:02d}',
            'rotulo_mes': f'{MONTH_NAMES_PT.get(e.month, e.month)[:3]}/{e.year}',
            'pessoa': (e.user_name or '').strip(), 'user_id': e.user_id,
            'cargo': _cargo(e), 'loja': normalizar(pdv) if pdv else '',
            'loja_rotulo': pdv.title() if pdv else ('Coordenação' if e.role == 'coordenador' else '—'),
            'coordenador': str(dados.get('COORDENADOR (A)') or '').strip(),
            'fase': versao.get_contestacao_phase_display() if versao else '',
            'previa': bool(versao and versao.contestacao_phase == 'antes'),
            'rascunho': bool(versao and not versao.is_released),
            **_detalhe(e.role, dados),
        })
    return linhas


def _escopo(user, papel, linhas):
    """O que cada um pode ver — o mesmo recorte do comissionamento."""
    if papel == 'superadmin':
        return linhas
    meu_nome = normalizar(user.get_full_name())

    def e_minha(l):
        return l['user_id'] == user.id or (meu_nome and normalizar(l['pessoa']) == meu_nome)
    if papel == 'coordenador':
        from simulator.services import get_coordinator_sectors
        lojas = {normalizar(s.name) for s in get_coordinator_sectors(user)}
        return [l for l in linhas if e_minha(l) or l['loja'] in lojas
                or (l['coordenador'] and normalizar(l['coordenador']) == meu_nome)]
    if papel == 'gerente':
        minha_loja = normalizar(user.sector.name) if user.sector_id else ''
        return [l for l in linhas if e_minha(l) or (minha_loja and l['loja'] == minha_loja)]
    return [l for l in linhas if e_minha(l)]


def _somar(linhas):
    total = {k: 0.0 for k, _ in PARTES}
    total['recebido'] = 0.0
    for l in linhas:
        for k in total:
            total[k] += l[k]
    return total


def _montar(linhas):
    """Meses (com subtotal), pessoas × meses e os números do topo."""
    por_mes = OrderedDict()
    for l in sorted(linhas, key=lambda x: (x['chave_mes'], -x['recebido'])):
        por_mes.setdefault(l['chave_mes'], {'rotulo': l['rotulo_mes'], 'linhas': []})['linhas'].append(l)
    meses = []
    for chave, m in por_mes.items():
        meses.append({'chave': chave, 'rotulo': m['rotulo'], 'linhas': m['linhas'], 'total': _somar(m['linhas']),
                      'previa': any(l['previa'] for l in m['linhas']),
                      'rascunho': any(l['rascunho'] for l in m['linhas'])})

    maior = max((m['total']['recebido'] for m in meses), default=0) or 1
    for m in meses:
        m['pct'] = round(max(m['total']['recebido'], 0) / maior * 100, 1)

    pessoas = defaultdict(lambda: {'valores': {}, 'total': 0.0})
    for l in linhas:
        p = pessoas[(l['pessoa'], l['cargo'], l['loja_rotulo'])]
        p['valores'][l['chave_mes']] = p['valores'].get(l['chave_mes'], 0.0) + l['recebido']
        p['total'] += l['recebido']
    chaves = [m['chave'] for m in meses]
    grade = [{'pessoa': k[0], 'cargo': k[1], 'loja': k[2], 'total': v['total'],
              'meses_com_valor': sum(1 for c in chaves if c in v['valores']),
              'celulas': [v['valores'].get(c) for c in chaves]}
             for k, v in sorted(pessoas.items(), key=lambda kv: -kv[1]['total'])]

    total = _somar(linhas)
    return {
        'meses': meses, 'grade': grade, 'total': total,
        'media_mes': total['recebido'] / len(meses) if meses else 0,
        'melhor_mes': max(meses, key=lambda m: m['total']['recebido']) if meses else None,
        'n_pessoas': len(pessoas),
    }


def _exportar(dados, filtros_texto):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = 'Por mês'
    ws.append([f'Histórico de comissão — {filtros_texto}'])
    ws['A1'].font = Font(bold=True, size=13)
    cab = ['Mês', 'Pessoa', 'Cargo', 'Loja'] + [r for _, r in PARTES] + ['Recebido', 'Fase']
    ws.append([])
    ws.append(cab)
    for c in ws[3]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='FF6B35')
    for m in dados['meses']:
        for l in m['linhas']:
            ws.append([m['rotulo'], l['pessoa'], CARGOS.get(l['cargo'], l['cargo']), l['loja_rotulo']]
                      + [round(l[k], 2) for k, _ in PARTES] + [round(l['recebido'], 2), l['fase']])
        ws.append([f'Total {m["rotulo"]}', '', '', ''] + [round(m['total'][k], 2) for k, _ in PARTES]
                  + [round(m['total']['recebido'], 2), ''])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    for linha in ws.iter_rows(min_row=4, min_col=5, max_col=11):
        for c in linha:
            c.number_format = 'R$ #,##0.00'
    for col, largura in zip('ABCDEFGHIJKL', (12, 34, 26, 22, 13, 13, 13, 13, 13, 13, 14, 22)):
        ws.column_dimensions[col].width = largura

    ws2 = wb.create_sheet('Pessoa x mês')
    rotulos = [m['rotulo'] for m in dados['meses']]
    ws2.append(['Pessoa', 'Cargo', 'Loja'] + rotulos + ['Total'])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for g in dados['grade']:
        ws2.append([g['pessoa'], CARGOS.get(g['cargo'], g['cargo']), g['loja']]
                   + [round(v, 2) if v is not None else None for v in g['celulas']] + [round(g['total'], 2)])
    for linha in ws2.iter_rows(min_row=2, min_col=4):
        for c in linha:
            c.number_format = 'R$ #,##0.00'
            c.alignment = Alignment(horizontal='right')
    ws2.column_dimensions['A'].width = 34
    ws2.column_dimensions['B'].width = 26
    ws2.column_dimensions['C'].width = 22

    buf = BytesIO()
    wb.save(buf)
    resposta = HttpResponse(buf.getvalue(),
                            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resposta['Content-Disposition'] = 'attachment; filename="historico_comissao.xlsx"'
    return resposta


@login_required
def commission_history_view(request):
    from .commission_views import pode_ver_comissionamento
    user = request.user
    if not pode_ver_comissionamento(user):
        messages.error(request, 'O comissionamento fica disponível para gerentes e coordenadores.')
        return redirect('home')
    papel = get_user_role(user)
    linhas = _escopo(user, papel, linhas_do_historico(ver_rascunho=papel == 'superadmin'))

    # Opções dos filtros: só o que existe no escopo de quem está vendo.
    cargos = [(k, v) for k, v in CARGOS.items() if any(l['cargo'] == k for l in linhas)]
    lojas = sorted({(l['loja'], l['loja_rotulo']) for l in linhas if l['loja']}, key=lambda x: x[1])
    chaves_meses = sorted({l['chave_mes'] for l in linhas})

    f_cargo = [c for c in request.GET.getlist('cargo') if c in CARGOS]
    f_loja = [normalizar(x) for x in request.GET.getlist('setor') if x]
    f_pessoa = (request.GET.get('pessoa') or '').strip()
    f_de = request.GET.get('de') or ''
    f_ate = request.GET.get('ate') or ''

    filtradas = [l for l in linhas
                 if (not f_cargo or l['cargo'] in f_cargo)
                 and (not f_loja or l['loja'] in f_loja)
                 and (not f_pessoa or normalizar(l['pessoa']) == normalizar(f_pessoa))
                 and (not f_de or l['chave_mes'] >= f_de)
                 and (not f_ate or l['chave_mes'] <= f_ate)]
    pessoas = sorted({l['pessoa'] for l in linhas
                      if (not f_cargo or l['cargo'] in f_cargo) and (not f_loja or l['loja'] in f_loja)})
    dados = _montar(filtradas)

    partes_texto = []
    if f_cargo:
        partes_texto.append(', '.join(CARGOS[c].split(' (')[0] if '(' in CARGOS[c] else CARGOS[c] for c in f_cargo))
    if f_loja:
        nomes = dict(lojas)
        partes_texto.append(', '.join(nomes.get(x, x.title()) for x in f_loja))
    if f_pessoa:
        partes_texto.append(f_pessoa.title())
    filtros_texto = ' · '.join(partes_texto) or 'toda a rede no seu escopo'

    if request.GET.get('exportar') == 'xlsx':
        return _exportar(dados, filtros_texto)

    return render(request, 'users/commission_historico.html', {
        **dados,
        'partes': PARTES,
        'cargos': cargos, 'lojas': lojas, 'pessoas': pessoas,
        'meses_opcoes': [(c, f'{MONTH_NAMES_PT.get(int(c[5:]), c[5:])[:3]}/{c[:4]}') for c in chaves_meses],
        'f_cargo': f_cargo, 'f_loja': f_loja, 'f_pessoa': f_pessoa, 'f_de': f_de, 'f_ate': f_ate,
        'filtros_texto': filtros_texto,
        'tem_filtro': bool(f_cargo or f_loja or f_pessoa or f_de or f_ate),
        'primeiro_mes': chaves_meses[0] if chaves_meses else '',
        'papel': papel,
        'cargos_rotulo': dict(CARGOS),
        'query_export': request.GET.urlencode(),
    })
