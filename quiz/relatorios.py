"""Resultados e relatórios do Quiz — na tela e em Excel.

- ``resultado_da_sala``: o que o responsável vê de uma sala — quem participou e
  quem não, acertos, erros, sem resposta, pontos, posição e as perguntas com mais
  erros.
- ``resultado_do_participante``: o que cada um vê de si — acertos, erros,
  pontos, posição e pergunta por pergunta (a certa, a dele, a explicação).
- ``relatorio_geral``: as salas encerradas no período — participação, média de
  acertos, desempenho por colaborador e as perguntas com maior índice de erro.
- ``planilha_*``: o mesmo em .xlsx.

"Participou" = entrou na sala. "Erro" = respondeu errado; quem não respondeu a
tempo conta como "sem resposta" (e também pesa no índice de erro da pergunta).
"""
from io import BytesIO

from django.db.models import Count
from django.utils import timezone

from . import jogo
from .models import Participante, Pergunta, Resposta, Sala


def _nome(user):
    return user.get_full_name() or user.username


def _setor(user):
    return user.sector.name if getattr(user, 'sector_id', None) else ''


def degraus(podio):
    """O pódio na ordem da tela — 2º, 1º, 3º —, com None na vaga que ninguém ocupou."""
    por_lugar = dict(enumerate(podio, start=1))
    return [(lugar, por_lugar.get(lugar)) for lugar in (2, 1, 3)]


def resultado_da_sala(sala):
    total = len(jogo.perguntas_da_sala(sala))
    placar = jogo.ranking(sala)
    for p in placar:
        p.sem_resposta = max(total - p.acertos - p.erros, 0)
        p.aproveitamento = (p.acertos / total) if total else 0.0
    ausentes = list(sala.participantes.filter(entrou_em__isnull=True).select_related('user', 'user__sector')
                    .order_by('user__first_name', 'user__last_name'))
    perguntas = jogo.acertos_por_pergunta(sala)
    convidados = len(placar) + len(ausentes)
    return {
        'total_perguntas': total,
        'placar': placar,
        'ausentes': ausentes,
        'convidados': convidados,
        'participacao': (len(placar) / convidados) if convidados else 0.0,
        'media_acertos': (sum(p.acertos for p in placar) / len(placar)) if placar else 0.0,
        'aproveitamento_medio': (sum(p.aproveitamento for p in placar) / len(placar)) if placar else 0.0,
        'perguntas': perguntas,
        'mais_erradas': sorted([l for l in perguntas if l['erros'] or l['sem_resposta']],
                               key=lambda l: (-l['taxa_erro'], l['n']))[:5],
        'degraus': degraus(placar[:3]),
    }


def resultado_do_participante(sala, participante):
    perguntas = jogo.perguntas_da_sala(sala)
    respostas = {r.pergunta_id: r for r in Resposta.objects.filter(participante=participante)
                 .select_related('alternativa')}
    linhas = []
    for n, pergunta in enumerate(perguntas, start=1):
        r = respostas.get(pergunta.pk)
        respondida = bool(r and r.respondida_em)
        linhas.append({'n': n, 'pergunta': pergunta, 'certa': pergunta.alternativa_correta(),
                       'escolhida': r.alternativa if respondida else None, 'respondida': respondida,
                       'correta': bool(respondida and r.correta), 'pontos': r.pontos if r else 0,
                       'tempo_s': (r.tempo_ms / 1000) if respondida and r.tempo_ms is not None else None})
    placar = jogo.ranking(sala)
    posicao = next((p.posicao for p in placar if p.pk == participante.pk), None)
    return {
        'linhas': linhas, 'posicao': posicao, 'presentes': len(placar),
        'acertos': sum(1 for l in linhas if l['correta']),
        'erros': sum(1 for l in linhas if l['respondida'] and not l['correta']),
        'sem_resposta': sum(1 for l in linhas if not l['respondida']),
        'pontos': sum(l['pontos'] for l in linhas),
        'podio': placar[:3],
        'degraus': degraus(placar[:3]),
    }


def salas_do_periodo(inicio=None, fim=None, quiz_id=None):
    salas = Sala.objects.filter(fase=Sala.Fase.ENCERRADA).select_related('quiz', 'responsavel')
    if inicio:
        salas = salas.filter(agendada_para__date__gte=inicio)
    if fim:
        salas = salas.filter(agendada_para__date__lte=fim)
    if quiz_id:
        salas = salas.filter(quiz_id=quiz_id)
    return salas.order_by('-agendada_para')


def relatorio_geral(inicio=None, fim=None, quiz_id=None):
    salas = list(salas_do_periodo(inicio, fim, quiz_id).annotate(
        convidados=Count('participantes', distinct=True)))
    ids = [s.pk for s in salas]
    perguntas_do_quiz = dict(Pergunta.objects.filter(quiz__salas__in=ids).values('quiz_id')
                             .annotate(n=Count('id', distinct=True)).values_list('quiz_id', 'n'))
    participantes = list(Participante.objects.filter(sala_id__in=ids).select_related('user', 'user__sector', 'sala'))

    por_sala = {}
    for p in participantes:
        por_sala.setdefault(p.sala_id, []).append(p)
    linhas_salas = []
    for sala in salas:
        gente = por_sala.get(sala.pk, [])
        presentes = [p for p in gente if p.entrou_em]
        total = perguntas_do_quiz.get(sala.quiz_id, 0)
        acertos = sum(p.acertos for p in presentes)
        linhas_salas.append({
            'sala': sala, 'convidados': len(gente), 'presentes': len(presentes),
            'participacao': (len(presentes) / len(gente)) if gente else 0.0,
            'media_acertos': (acertos / len(presentes)) if presentes else 0.0,
            'aproveitamento': (acertos / (total * len(presentes))) if total and presentes else 0.0,
            'total_perguntas': total,
        })

    colaboradores = {}
    for p in participantes:
        c = colaboradores.setdefault(p.user_id, {'user': p.user, 'convites': 0, 'salas': 0, 'acertos': 0,
                                                 'erros': 0, 'pontos': 0, 'perguntas': 0})
        c['convites'] += 1
        if p.entrou_em:
            c['salas'] += 1
            c['acertos'] += p.acertos
            c['erros'] += p.erros
            c['pontos'] += p.pontos
            c['perguntas'] += perguntas_do_quiz.get(p.sala.quiz_id, 0)
    for c in colaboradores.values():
        c['sem_resposta'] = max(c['perguntas'] - c['acertos'] - c['erros'], 0)
        c['aproveitamento'] = (c['acertos'] / c['perguntas']) if c['perguntas'] else 0.0
        c['participacao'] = c['salas'] / c['convites'] if c['convites'] else 0.0
    ordem_colab = sorted(colaboradores.values(), key=lambda c: (-c['aproveitamento'], -c['pontos'],
                                                                _nome(c['user'])))

    erros_por_pergunta = {}
    for sala in salas:
        for linha in jogo.acertos_por_pergunta(sala):
            pergunta = linha['pergunta']
            item = erros_por_pergunta.setdefault(pergunta.pk, {'pergunta': pergunta, 'quiz': sala.quiz,
                                                               'acertos': 0, 'erros': 0, 'sem_resposta': 0})
            item['acertos'] += linha['acertos']
            item['erros'] += linha['erros']
            item['sem_resposta'] += linha['sem_resposta']
    for item in erros_por_pergunta.values():
        total = item['acertos'] + item['erros'] + item['sem_resposta']
        item['taxa_erro'] = ((item['erros'] + item['sem_resposta']) / total) if total else 0.0
    mais_erradas = sorted((i for i in erros_por_pergunta.values() if i['erros'] or i['sem_resposta']),
                          key=lambda i: (-i['taxa_erro'], -(i['erros'] + i['sem_resposta'])))[:15]

    presentes_total = sum(l['presentes'] for l in linhas_salas)
    convidados_total = sum(l['convidados'] for l in linhas_salas)
    acertos_total = sum(l['media_acertos'] * l['presentes'] for l in linhas_salas)
    return {
        'salas': linhas_salas,
        'colaboradores': ordem_colab,
        'mais_erradas': mais_erradas,
        'resumo': {
            'salas': len(linhas_salas),
            'convidados': convidados_total,
            'presentes': presentes_total,
            'participacao': (presentes_total / convidados_total) if convidados_total else 0.0,
            'media_acertos': (acertos_total / presentes_total) if presentes_total else 0.0,
            'aproveitamento': (sum(l['aproveitamento'] * l['presentes'] for l in linhas_salas) / presentes_total)
            if presentes_total else 0.0,
        },
    }


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def _aba(livro, titulo, cabecalho, linhas, larguras):
    from openpyxl.styles import Alignment, Font, PatternFill

    aba = livro.create_sheet(titulo)
    aba.append(cabecalho)
    for celula in aba[1]:
        celula.font = Font(bold=True, color='FFFFFF')
        celula.fill = PatternFill('solid', fgColor='F26522')
        celula.alignment = Alignment(vertical='center')
    for linha in linhas:
        aba.append(list(linha))
    for coluna, largura in zip('ABCDEFGHIJKLMNOP', larguras):
        aba.column_dimensions[coluna].width = largura
    aba.freeze_panes = 'A2'
    return aba


def _xlsx(livro):
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


def _pct(valor):
    return round(valor * 100, 1)


def planilha_da_sala(sala):
    from openpyxl import Workbook

    dados = resultado_da_sala(sala)
    livro = Workbook()
    livro.remove(livro.active)
    ranking = [(p.posicao, _nome(p.user), p.user.job_title, _setor(p.user), p.acertos, p.erros, p.sem_resposta,
                _pct(p.aproveitamento), p.pontos) for p in dados['placar']]
    ranking += [('—', _nome(p.user), p.user.job_title, _setor(p.user), 'não participou', '', '', '', '')
                for p in dados['ausentes']]
    _aba(livro, 'Ranking', ['Posição', 'Nome', 'Cargo', 'Setor', 'Acertos', 'Erros', 'Sem resposta',
                            '% de acerto', 'Pontos'], ranking, [9, 34, 26, 26, 10, 8, 13, 12, 10])
    _aba(livro, 'Perguntas', ['Nº', 'Pergunta', 'Acertos', 'Erros', 'Sem resposta', 'Índice de erro (%)'],
         [(l['n'], l['pergunta'].enunciado, l['acertos'], l['erros'], l['sem_resposta'], _pct(l['taxa_erro']))
          for l in dados['perguntas']], [6, 70, 10, 8, 13, 18])
    respostas = (Resposta.objects.filter(participante__sala=sala)
                 .select_related('participante__user', 'pergunta', 'alternativa')
                 .order_by('participante__user__first_name', 'pergunta__ordem', 'pergunta_id'))
    numero = {p.pk: n for n, p in enumerate(jogo.perguntas_da_sala(sala), start=1)}
    _aba(livro, 'Respostas', ['Nome', 'Nº', 'Pergunta', 'Resposta', 'Correta?', 'Tempo (s)', 'Pontos'],
         [(_nome(r.participante.user), numero.get(r.pergunta_id), r.pergunta.enunciado,
           r.alternativa.texto if r.alternativa else '(sem resposta)',
           ('sim' if r.correta else 'não') if r.respondida_em else '—',
           round(r.tempo_ms / 1000, 1) if r.tempo_ms is not None else '', r.pontos) for r in respostas],
         [30, 6, 60, 40, 10, 10, 9])
    return _xlsx(livro)


def planilha_geral(dados):
    from openpyxl import Workbook

    livro = Workbook()
    livro.remove(livro.active)
    _aba(livro, 'Salas', ['Data', 'Sala', 'Quiz', 'Código', 'Responsável', 'Convidados', 'Participaram',
                          'Participação (%)', 'Média de acertos', 'Aproveitamento (%)'],
         [(timezone.localtime(l['sala'].agendada_para).strftime('%d/%m/%Y %H:%M'), l['sala'].titulo,
           l['sala'].quiz.titulo, l['sala'].codigo,
           _nome(l['sala'].responsavel) if l['sala'].responsavel else '', l['convidados'], l['presentes'],
           _pct(l['participacao']), round(l['media_acertos'], 1), _pct(l['aproveitamento']))
          for l in dados['salas']], [17, 34, 34, 9, 26, 11, 13, 16, 16, 18])
    _aba(livro, 'Colaboradores', ['Nome', 'Cargo', 'Setor', 'Chamado para', 'Participou de', 'Acertos', 'Erros',
                                  'Sem resposta', 'Aproveitamento (%)', 'Pontos'],
         [(_nome(c['user']), c['user'].job_title, _setor(c['user']), c['convites'], c['salas'], c['acertos'],
           c['erros'], c['sem_resposta'], _pct(c['aproveitamento']), c['pontos'])
          for c in dados['colaboradores']], [32, 26, 26, 13, 14, 9, 8, 13, 18, 10])
    _aba(livro, 'Perguntas com mais erro', ['Quiz', 'Pergunta', 'Acertos', 'Erros', 'Sem resposta',
                                             'Índice de erro (%)'],
         [(i['quiz'].titulo, i['pergunta'].enunciado, i['acertos'], i['erros'], i['sem_resposta'],
           _pct(i['taxa_erro'])) for i in dados['mais_erradas']], [30, 70, 10, 8, 13, 18])
    return _xlsx(livro)
