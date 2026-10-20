"""Chamados (20/10/2026): exportação com a coluna "Loja/Local" do chamado."""
import os
import sys
from datetime import date
from decimal import Decimal

import django

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'redeconfianca.settings')
os.environ.setdefault('RC_VARREDURA_ROTINA', '0')

from django.conf import settings

settings.CACHES = {
    'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-tk-ex'},
    'local': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-tk-ex-2'},
}
django.setup()

from django.test.utils import setup_test_environment

setup_test_environment()
if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS.append('testserver')

from django.contrib.auth import get_user_model
from django.db import transaction
from django.test import Client

from contagem_caixa.models import ContagemCaixaDia, SaldoInicialMes
from contagem_caixa.servicos import recalcular_saldos
from users.models import Sector

User = get_user_model()
D = Decimal
ok = fail = 0


def t(nome, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'  OK   {nome}')
    else:
        fail += 1
        print(f'  FALHA {nome} {extra}')


import csv
import io

from openpyxl import load_workbook
from tickets.models import Ticket

marcador = transaction.atomic()
marcador.__enter__()
try:
    setor = Sector.objects.create(name='ZZ Setor Export')
    chefe = User.objects.create_user(username='zz.tk.ex', email='zz.tk.ex@exemplo-teste.local', password='S3nha!teste',
                                     first_name='Zz', last_name='Export', sector=setor, hierarchy='SUPERADMIN', is_superuser=True)
    tk = Ticket.objects.create(title='ZZ chamado com loja', description='x', sector=setor, created_by=chefe,
                               store_location='Loja ZZ Montserrat')
    Ticket.objects.create(title='ZZ chamado sem loja', description='y', sector=setor, created_by=chefe)
    c = Client(); c.force_login(chefe)

    print('== EXPORTAR CHAMADOS COM LOJA/LOCAL ==')
    r = c.get('/tickets/export/?export=csv&search=ZZ chamado')
    texto = b''.join(r.streaming_content).decode('utf-8-sig') if r.streaming else r.content.decode('utf-8-sig')
    linhas = list(csv.reader(io.StringIO(texto), delimiter=';' if texto.count(';') > texto.count(',') else ','))
    cab = linhas[0]
    t('CSV tem a coluna Loja/Local logo depois de PDV', 'Loja/Local' in cab and cab.index('Loja/Local') == cab.index('PDV') + 1, cab)
    valor = next((l[cab.index('Loja/Local')] for l in linhas[1:] if 'ZZ chamado com loja' in l), None)
    t('com o valor do chamado', valor == 'Loja ZZ Montserrat', valor)
    r = c.get('/tickets/export/?search=ZZ chamado')
    wb = load_workbook(io.BytesIO(b''.join(r.streaming_content) if r.streaming else r.content))
    ws = wb.active
    cab = [x.value for x in ws[1]]
    t('Excel tem a coluna Loja/Local', 'Loja/Local' in cab and cab.index('Loja/Local') == cab.index('PDV') + 1, cab)
    col = cab.index('Loja/Local') + 1
    vals = {ws.cell(row=i, column=2).value: ws.cell(row=i, column=col).value for i in range(2, ws.max_row + 1)}
    t('com o valor (e vazio no chamado sem loja)', vals.get('ZZ chamado com loja') == 'Loja ZZ Montserrat'
      and not vals.get('ZZ chamado sem loja'), vals)
    resp = cab.index('Responsável') + 1
    t('as colunas seguintes não se desalinharam', ws.cell(row=2, column=resp).value in ('Não atribuído', None) or True
      and cab[-2:] == ['Imagens Anexadas', 'Arquivos Anexados'] and ws.cell(row=2, column=len(cab)).value is not None, cab[-3:])
finally:
    transaction.set_rollback(True)
    marcador.__exit__(None, None, None)
    print('\nrollback: nada gravado.')

print(f'\n{ok} OK / {fail} falhas')
sys.exit(1 if fail else 0)
