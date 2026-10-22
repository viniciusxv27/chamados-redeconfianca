"""Telas do Acess."""
import io
from datetime import timedelta

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.dateparse import parse_date, parse_datetime
from django.views.decorators.http import require_POST

from users.models import Sector

from . import servicos as sv
from .models import (AvisoInativacao, MOTIVOS_CHAMADO, SISTEMAS, SISTEMAS_CHAMADO, STATUS_SISTEMA, AcessoAcess, Chamado, Colaborador,
                     LogAcess, ModeloMensagem, Solicitacao)
from .permissoes import FUNCOES, completar, e_superadmin, exige, pode, trata_chamados

User = get_user_model()
ABAS = [('dashboard', 'Dashboard', 'fa-gauge-high', None), ('lojas', 'Lojas & Sede', 'fa-store', 'colaboradores.ver'),
        ('colaboradores', 'Colaboradores', 'fa-users', 'colaboradores.ver'),
        ('pendentes', 'Acessos pendentes', 'fa-list-check', 'colaboradores.ver'),
        ('liberados', 'Logins liberados', 'fa-circle-check', 'colaboradores.ver'),
        ('chamados', 'Chamados', 'fa-headset', 'chamados.ver'),
        ('solicitacoes', 'Solicitações', 'fa-inbox', 'chamados.ver'),
        ('historico', 'Histórico', 'fa-clock-rotate-left', 'historico.ver'),
        ('desligamentos', 'Desligamentos', 'fa-user-xmark', 'desligamentos.gerenciar'),
        ('mensagens', 'Mensagens', 'fa-comment-dots', None)]
ETAPAS = [(1, 'Criação do GEFAC'), (2, 'Login inicial'), (3, 'Academia V'), (4, 'Liberação e envio')]
CHAVES_TELA = [f[0] for f in FUNCOES]


# ── apoio ───────────────────────────────────────────────────────────────────
def _ctx(request, aba, **extra):
    u = request.user
    perms = {c.replace('.', '_'): pode(u, c) for c in CHAVES_TELA}
    abas = [a for a in ABAS if a[3] is None or perms[a[3].replace('.', '_')]]
    contadores = {}
    if perms['colaboradores_ver']:
        ativos = Colaborador.objects.filter(status=Colaborador.ATIVO)
        contadores['pendentes'] = ativos.filter(etapa__in=(1, 2, 3, 4)).count()
    if perms['chamados_ver']:
        contadores['chamados'] = Chamado.objects.filter(status=Chamado.ANDAMENTO).count()
        contadores['solicitacoes'] = Solicitacao.objects.filter(status=Solicitacao.ABERTA).count()
    if perms['desligamentos_gerenciar']:
        contadores['inativados'] = AvisoInativacao.objects.filter(status=AvisoInativacao.PENDENTE).count()
        contadores['desligamentos'] = Colaborador.objects.filter(
            status=Colaborador.ATIVO, desligamento_em__lte=timezone.localdate()).count() + contadores['inativados']
    sv.garantir_modelos()
    extra.setdefault('modelos', ModeloMensagem.objects.all())
    return dict(aba=aba, abas=abas, p=perms, e_superadmin=e_superadmin(u), contadores=contadores, **extra)


def lojas_do_acess():
    """Unidades: as lojas (com ADABAS ou "Loja" no nome) e a Sede Administrativa."""
    return Sector.objects.filter(Q(name__istartswith='loja') | ~Q(adabas='') | Q(name__icontains='sede administrativa')) \
        .order_by('name')


def _voltar(request, padrao):
    destino = request.POST.get('voltar') or ''
    return redirect(destino if destino.startswith('/acess/') else padrao)


def _int(valor):
    return int(valor) if str(valor or '').isdigit() else None


def _setor(valor):
    i = _int(valor)
    return Sector.objects.filter(pk=i).first() if i else None


def filtrar(request, qs, prefixo=''):
    """Pesquisa, cargo, estado, local, vínculo e situação — os filtros comuns das listas."""
    g = lambda k: (request.GET.get(prefixo + k) or '').strip()  # noqa: E731
    f = {k: g(k) for k in ('q', 'cargo', 'uf', 'local', 'vinculo', 'acessos', 'dup')}
    if f['q']:
        d = sv.so_digitos(f['q'])
        filtro = Q(nome__icontains=f['q']) | Q(email__icontains=f['q']) | Q(login__icontains=f['q'])
        if len(d) >= 3:
            filtro |= Q(cpf__contains=d)
        filtro |= Q(local__adabas__icontains=f['q']) | Q(loja_vinculo__adabas__icontains=f['q'])
        qs = qs.filter(filtro)
    if f['cargo'] == '-':
        qs = qs.filter(cargo='')
    elif f['cargo']:
        qs = qs.filter(cargo=f['cargo'])
    if f['local'].isdigit():
        qs = qs.filter(local_id=int(f['local']))
    if f['vinculo'] == '-':
        qs = qs.filter(loja_vinculo__isnull=True)
    elif f['vinculo'].isdigit():
        qs = qs.filter(loja_vinculo_id=int(f['vinculo']))
    lista = list(qs.select_related('local', 'loja_vinculo', 'user'))
    if f['uf']:
        lista = [c for c in lista if sv.uf_da_loja(c.local) == f['uf']
                 or (sv.e_sede(c.local) and sv.uf_da_loja(c.loja_vinculo) == f['uf'])]
    if f['acessos'] == 'com':
        lista = [c for c in lista if c.pendencias]
    elif f['acessos'] == 'sem':
        lista = [c for c in lista if not c.pendencias]
    if f['dup'] == '1':
        ids = sv.duplicados(Colaborador.objects.filter(status=Colaborador.ATIVO))
        lista = [c for c in lista if c.id in ids]
    return lista, f


def _opcoes_filtro():
    return {'cargos': sorted(set(Colaborador.objects.exclude(cargo='').values_list('cargo', flat=True))),
            'unidades': list(lojas_do_acess()), 'ufs': ['ES', 'RJ', 'SP']}


def _erro(request, exc):
    messages.error(request, str(exc))


# ── Dashboard ───────────────────────────────────────────────────────────────
@login_required
@exige()
def dashboard(request):
    ctx = _ctx(request, 'dashboard')
    if ctx['p']['colaboradores_ver']:
        ativos = Colaborador.objects.filter(status=Colaborador.ATIVO)
        lista, filtro = filtrar(request, ativos)
        com_login = [c for c in lista if c.login_valido and c.pendencias and c.entrega_login != 'PENDENTE']
        ctx.update(filtro=filtro, **_opcoes_filtro(), pendentes_login=sorted(com_login, key=lambda c: c.nome),
                   k_pendencias=sum(1 for c in lista if c.pendencias),
                   k_logins=sum(1 for c in lista if c.login_valido),
                   k_unidades=len({c.local_id for c in lista if c.local_id}))
    if ctx['p']['chamados_ver']:
        ctx['chamados_abertos'] = list(Chamado.objects.filter(status=Chamado.ANDAMENTO)
                                       .select_related('colaborador', 'responsavel').order_by('proxima_verificacao')[:30])
    return render(request, 'acess/dashboard.html', ctx)


@login_required
@exige('sistemas.concluir')
@require_POST
def concluir_lote(request):
    try:
        feitos, pulados = sv.concluir_pendencias_em_lote(request.POST.getlist('ids'), usuario=request.user,
                                                         request=request)
        messages.success(request, f'Pendências concluídas de {feitos} colaborador(es).')
        if pulados:
            messages.warning(request, 'Ficaram de fora (etapa 4 ou sem login): ' + ', '.join(pulados[:10]))
    except (sv.RegraAcess, ValueError) as exc:
        _erro(request, exc)
    return redirect('acess:dashboard')


# ── Lojas & Sede ────────────────────────────────────────────────────────────
@login_required
@exige('colaboradores.ver')
def lojas(request):
    lista, filtro = filtrar(request, Colaborador.objects.filter(status=Colaborador.ATIVO))
    chamados = set(Chamado.objects.filter(status=Chamado.ANDAMENTO).values_list('colaborador_id', flat=True))
    por_unidade = {}
    for c in lista:
        por_unidade.setdefault(c.local, []).append(c)
        c.tem_chamado = c.id in chamados
    situacao = request.GET.get('situacao') or ''
    unidades = []
    for setor, gente in sorted(por_unidade.items(), key=lambda kv: sv.norm(kv[0].name if kv[0] else 'ZZZ')):
        gente.sort(key=lambda c: (not c.pendencias, c.nome))
        if situacao == 'pendentes' and not any(c.pendencias for c in gente):
            continue
        unidades.append({'setor': setor, 'gente': gente, 'pendentes': sum(1 for c in gente if c.pendencias),
                         'chamados': sum(1 for c in gente if c.tem_chamado), 'uf': sv.uf_da_loja(setor),
                         'sede': sv.e_sede(setor)})
    return render(request, 'acess/lojas.html', _ctx(request, 'lojas', blocos=unidades, filtro=filtro,
                                                    situacao=situacao, **_opcoes_filtro()))


# ── Colaboradores ───────────────────────────────────────────────────────────
@login_required
@exige('colaboradores.ver')
def colaboradores(request):
    ativos = Colaborador.objects.filter(status=Colaborador.ATIVO)
    lista, filtro = filtrar(request, ativos)
    dups = sv.duplicados(ativos)
    usuarios_sem_ficha = User.objects.filter(is_active=True, ficha_acess__isnull=True).order_by('first_name') \
        if pode(request.user, 'colaboradores.cadastrar') else []
    query = request.GET.copy()
    return render(request, 'acess/colaboradores.html', _ctx(
        request, 'colaboradores', colaboradores=lista, filtro=filtro, n_duplicados=len(dups),
        k_total=len(lista), k_ok=sum(1 for c in lista if not c.pendencias), k_pend=sum(1 for c in lista if c.pendencias),
        k_unidades=len({c.local_id for c in lista if c.local_id}), usuarios_sem_ficha=usuarios_sem_ficha,
        colunas_excel=COLUNAS_EXCEL, query=query.urlencode(), **_opcoes_filtro()))


@login_required
@exige('colaboradores.cadastrar')
@require_POST
def novo(request):
    p = request.POST
    try:
        user = User.objects.filter(pk=_int(p.get('user'))).first() if p.get('user') else None
        dados = sv.dados_do_usuario(user) if user else p
        if user and p.get('cpf'):
            dados = dict(dados, cpf=p.get('cpf'))
        local = _setor(p.get('local')) or (user.sector if user else None)
        c = sv.criar_colaborador(dados, local, _setor(p.get('loja_vinculo')), usuario=request.user, request=request,
                                 origem='PORTAL' if user else 'MANUAL', user=user)
    except sv.RegraAcess as exc:
        _erro(request, exc)
        return redirect('acess:colaboradores')
    messages.success(request, f'{c.nome} cadastrado e na Etapa {c.etapa} da esteira.')
    return redirect('acess:ficha', c.pk)


COLUNAS_EXCEL = [('nome_proprio', 'Nome'), ('sobrenome', 'Sobrenome'), ('nome', 'Nome completo'), ('cpf', 'CPF'),
                 ('nascimento', 'Data de nascimento'), ('email', 'E-mail'), ('telefone', 'Telefone'), ('cargo', 'Cargo'),
                 ('local', 'Local de trabalho'), ('vinculo', 'Loja de vínculo'), ('adabas', 'ADABAS'), ('uf', 'Estado'),
                 ('login', 'Login')]


@login_required
@exige('colaboradores.ver')
def exportar(request):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    escolhidas = [c for c in COLUNAS_EXCEL if c[0] in request.GET.getlist('col')]
    lista, _f = filtrar(request, Colaborador.objects.filter(status=Colaborador.ATIVO))
    if not escolhidas or not lista:
        messages.error(request, 'Escolha ao menos uma coluna, com ao menos um colaborador na lista.')
        return redirect(reverse('acess:colaboradores') + '?' + request.GET.urlencode())

    def valor(c, chave):
        partes = c.nome.split(' ', 1)
        return {
            'nome_proprio': partes[0], 'sobrenome': partes[1] if len(partes) > 1 else '', 'nome': c.nome,
            'cpf': c.cpf, 'nascimento': c.nascimento.strftime('%d/%m/%Y') if c.nascimento else '', 'email': c.email,
            'telefone': c.telefone, 'cargo': c.cargo, 'local': c.local.name if c.local else '',
            'vinculo': c.loja_vinculo.name if c.loja_vinculo else '', 'adabas': c.unidade_adabas,
            'uf': sv.uf_da_loja(c.loja_vinculo if sv.e_sede(c.local) else c.local), 'login': c.login,
        }[chave]

    wb = Workbook()
    ws = wb.active
    ws.title = 'Colaboradores'
    ws.append([r for _k, r in escolhidas])
    for c in lista:
        ws.append([str(valor(c, k) or '') for k, _r in escolhidas])
    for linha in ws.iter_rows(min_row=1):
        for cel in linha:
            cel.data_type = 's'  # texto sempre: CPF com zero à esquerda e nada vira fórmula
    for i, (_k, rotulo) in enumerate(escolhidas, start=1):
        largura = max([len(rotulo)] + [len(str(valor(c, escolhidas[i - 1][0]) or '')) for c in lista[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(largura + 2, 8), 50)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    sv.registrar(request, 'EXPORTAR_EXCEL', 'ADMIN', f'Exportou {len(lista)} colaborador(es) em Excel')
    resp = HttpResponse(buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="colaboradores-{timezone.localdate():%Y-%m-%d}.xlsx"'
    return resp


# ── Ficha ───────────────────────────────────────────────────────────────────
@login_required
@exige('colaboradores.ver')
def ficha(request, pk):
    c = get_object_or_404(Colaborador.objects.select_related('local', 'loja_vinculo', 'user'), pk=pk)
    return render(request, 'acess/ficha.html', _ctx(
        request, 'colaboradores', c=c, sistemas=[(s, c.status_de(s)) for s in SISTEMAS], status_opcoes=STATUS_SISTEMA,
        unidades=lojas_do_acess(), setores=Sector.objects.order_by('name'),
        chamados=list(c.chamados.select_related('responsavel').order_by('-aberto_em')),
        historico=list(c.historico.all()[:40]) if pode(request.user, 'historico.ver') else None,
        perfil_requerido=sv.rotulo_perfil(sv.perfil_regional_requerido(c)), etapas=ETAPAS,
        modelos=ModeloMensagem.objects.all(), motivos=MOTIVOS_CHAMADO, sistemas_chamado=SISTEMAS_CHAMADO))


@login_required
@exige('colaboradores.ver')
@require_POST
def ficha_acao(request, pk):
    """Ações da ficha. Pela tela (fetch) responde JSON: salva sem recarregar e devolve a versão nova,
    para o próximo bloco salvo na mesma tela não esbarrar no "alterada por outra pessoa"."""
    c = get_object_or_404(Colaborador, pk=pk)
    p, u = request.POST, request.user
    acao, versao = p.get('acao'), p.get('versao')
    por_fetch = request.headers.get('X-Requested-With') == 'XMLHttpRequest'
    precisa = {'dados': 'colaboradores.editar', 'transferir': 'colaboradores.transferir',
               'sistemas': 'sistemas.concluir', 'excluir': 'colaboradores.excluir'}.get(acao)
    if not precisa or not pode(u, precisa):
        aviso = 'Esta ação não está liberada para o seu usuário.'
        if por_fetch:
            return JsonResponse({'ok': False, 'erro': aviso}, status=403)
        messages.error(request, aviso)
        return redirect('acess:ficha', pk)
    try:
        if acao == 'dados':
            c = sv.editar_dados(c, {k: p.get(k) for k in ('nome', 'cpf', 'nascimento', 'email', 'telefone', 'cargo')},
                                versao, u, request)
            msg = 'Dados salvos.'
        elif acao == 'transferir':
            c = sv.transferir(c, _setor(p.get('local')), _setor(p.get('loja_vinculo')), versao, u, request)
            msg = 'Local e vínculo salvos.'
        elif acao == 'sistemas':
            c = sv.alterar_sistemas(c, {s: p.get(f'sis_{i}', '') for i, s in enumerate(SISTEMAS)}, versao, u, request)
            msg = 'Sistemas salvos.'
        elif acao == 'excluir':
            if not u.check_password(p.get('senha') or ''):
                raise sv.RegraAcess('Senha incorreta: para excluir, confirme a sua senha do portal.')
            sv.arquivar(c, versao, u, request, motivo=(p.get('motivo') or '').strip())
            messages.success(request, f'{c.nome} arquivado (o histórico fica; os acessos externos não são revogados).')
            return redirect('acess:colaboradores')
    except sv.RegraAcess as exc:
        if por_fetch:
            return JsonResponse({'ok': False, 'erro': str(exc), 'conflito': isinstance(exc, sv.Conflito)}, status=400)
        _erro(request, exc)
        return redirect('acess:ficha', pk)
    if por_fetch:
        c = Colaborador.objects.select_related('local', 'loja_vinculo').get(pk=c.pk)
        return JsonResponse({'ok': True, 'msg': msg, 'versao': c.versao, 'quando': timezone.localtime().strftime('%H:%M'),
                             'dados': {'nome': c.nome, 'cpf': c.cpf_formatado, 'email': c.email, 'telefone': c.telefone,
                                       'cargo': c.cargo, 'nascimento': c.nascimento.isoformat() if c.nascimento else ''},
                             'sistemas': {s: c.status_de(s) for s in SISTEMAS},
                             'copia': {k: v for _r, k, v in c.dados_para_copiar()}, 'copia_tudo': c.texto_para_copiar,
                             'recarregar': acao == 'transferir'})
    messages.success(request, msg)
    return redirect('acess:ficha', pk)


# ── Esteira ─────────────────────────────────────────────────────────────────
@login_required
@exige('colaboradores.ver')
def pendentes(request):
    etapa = _int(request.GET.get('etapa')) or 1
    ativos = Colaborador.objects.filter(status=Colaborador.ATIVO)
    lista, filtro = filtrar(request, ativos)
    por_etapa = {n: [] for n, _r in ETAPAS}
    for c in lista:
        e = c.etapa_calculada
        if e in por_etapa:
            por_etapa[e].append(c)
    for grupo in por_etapa.values():
        grupo.sort(key=lambda c: c.pendente_desde)
    linhas = por_etapa.get(etapa, [])
    # Admissão no futuro (cadastro do portal): a data em que é para criar no GEFAC aparece embaixo do nome.
    hoje = timezone.localdate()
    admissoes = dict(User.objects.filter(id__in=[c.user_id for c in linhas if c.user_id])
                     .values_list('id', 'admission_date'))
    for c in linhas:
        admissao = admissoes.get(c.user_id)
        c.admissao_futura = admissao if admissao and admissao > hoje else None
        c.dias_ate_admissao = (admissao - hoje).days if c.admissao_futura else None
        c.chamado_lib = sv.chamado_de_liberacao(c) if etapa == 4 else None
        c.perfil_requerido = sv.perfil_regional_requerido(c)
        c.checklist_ok = sv.checklist_completo(c) if etapa == 4 else False
        if etapa == 4:
            perfil = sv.rotulo_perfil(c.perfil_requerido)
            c.itens_checklist = [
                ('pool', 'Pool de cobrança', c.check_pool, 'Criação do Pool de cobrança realizada.'),
                ('perfil', f'Perfil regional: {perfil}', c.check_perfil == c.perfil_requerido,
                 'Assurant para SP; Renova Allied para os demais estados.'
                 + (' Marcado para outro perfil: confira de novo.' if c.check_perfil and c.check_perfil != c.perfil_requerido else '')),
                ('login', 'Teste do login', c.check_login, 'Acesso testado fora do portal.'),
                ('sap', 'Senha SAP', c.check_sap, 'Senha SAP criada fora do portal.'),
            ]
    return render(request, 'acess/pendentes.html', _ctx(
        request, 'pendentes', etapa=etapa, etapas=[(n, r, len(por_etapa[n])) for n, r in ETAPAS], linhas=linhas,
        filtro=filtro, modelos=ModeloMensagem.objects.all(), **_opcoes_filtro()))


@login_required
@exige('colaboradores.ver')
@require_POST
def esteira_acao(request, pk):
    c = get_object_or_404(Colaborador, pk=pk)
    p, u, acao, versao = request.POST, request.user, request.POST.get('acao'), request.POST.get('versao')
    precisa = {'gefac': ('gefac.concluir', 'sistemas.concluir'), 'login': ('login.preencher',),
               'academia': ('sistemas.concluir',), 'checklist': ('sistemas.concluir',),
               'voltar': ('sistemas.concluir',)}.get(acao, ())
    if acao == 'finalizar':
        ok = pode(u, 'sistemas.concluir') and pode(u, 'chamados.tratar')
    else:
        ok = bool(precisa) and any(pode(u, x) for x in precisa)
    if not ok:
        messages.error(request, 'Esta ação não está liberada para o seu usuário.')
        return _voltar(request, reverse('acess:pendentes'))
    try:
        if acao == 'gefac':
            sv.concluir_gefac(c, versao, u, request)
            messages.success(request, f'GEFAC de {c.nome} concluído: agora aguarda o login.')
        elif acao == 'login':
            sv.salvar_login(c, p.get('login'), versao, u, request)
            messages.success(request, f'Login de {c.nome} salvo: agora é a Academia V.')
        elif acao == 'academia':
            if p.get('confirmo') != '1':
                raise sv.RegraAcess('Confirme que a capacitação foi realizada.')
            sv.concluir_academia(c, versao, u, request)
            messages.success(request, f'Academia V de {c.nome} concluída: falta a liberação.')
        elif acao == 'checklist':
            sv.salvar_checklist(c, p.get('item'), p.get('marcado') == '1', versao, u, request)
        elif acao == 'finalizar':
            if p.get('confirmo') != '1':
                raise sv.RegraAcess('Confirme que enviou a orientação ao colaborador.')
            sv.finalizar_liberacao(c, versao, u, request, chamado_versao=p.get('chamado_versao'))
            messages.success(request, f'Liberação de {c.nome} confirmada: está em Logins liberados.')
        elif acao == 'voltar':
            sv.voltar_etapa(c, versao, u, request, pode_chamados=pode(u, 'chamados.tratar'))
            messages.success(request, f'{c.nome} voltou uma etapa.')
    except sv.RegraAcess as exc:
        _erro(request, exc)
    return _voltar(request, reverse('acess:pendentes'))


@login_required
@exige('colaboradores.ver')
@require_POST
def whatsapp(request, pk):
    """Prepara a mensagem (texto + link do WhatsApp). Não envia nada: quem envia é o operador."""
    c = get_object_or_404(Colaborador.objects.select_related('local'), pk=pk)
    sv.garantir_modelos()
    modelo = ModeloMensagem.objects.filter(chave=request.POST.get('modelo') or '').first()
    if not modelo:
        return JsonResponse({'ok': False, 'erro': 'Modelo de mensagem não encontrado.'}, status=400)
    try:
        texto = sv.montar_mensagem(modelo.texto, c, senha=(request.POST.get('senha') or '').strip())
        return JsonResponse({'ok': True, 'texto': texto, 'link': sv.link_whatsapp(c.telefone, texto)})
    except sv.RegraAcess as exc:
        return JsonResponse({'ok': False, 'erro': str(exc)}, status=400)


# ── Logins liberados ────────────────────────────────────────────────────────
@login_required
@exige('colaboradores.ver')
def liberados(request):
    lista, filtro = filtrar(request, Colaborador.objects.filter(entrega_login='ENVIADO'))
    lista.sort(key=lambda c: (c.liberado_em is None, -(c.liberado_em.timestamp() if c.liberado_em else 0)))
    return render(request, 'acess/liberados.html', _ctx(request, 'liberados', colaboradores=lista, filtro=filtro,
                                                        **_opcoes_filtro()))


# ── Chamados ────────────────────────────────────────────────────────────────
@login_required
@exige('chamados.ver')
def chamados(request):
    st = request.GET.get('status') or 'andamento'
    qs = Chamado.objects.select_related('colaborador', 'colaborador__local', 'colaborador__loja_vinculo',
                                        'responsavel', 'concluido_por')
    agora = timezone.now()
    if st == 'andamento':
        qs = qs.filter(status=Chamado.ANDAMENTO)
    elif st == 'vencido':
        qs = qs.filter(status=Chamado.ANDAMENTO, sla_original__lt=agora)
    elif st == 'concluidos':
        qs = qs.filter(status=Chamado.CONCLUIDO)
    q = (request.GET.get('q') or '').strip()
    if q:
        qs = qs.filter(Q(codigo__icontains=sv.codigo_chamado(q) or q) | Q(colaborador__nome__icontains=q)
                       | Q(observacao__icontains=q)
                       | Q(colaborador__local__name__icontains=q) | Q(colaborador__login__icontains=q))
    pagina = Paginator(qs.prefetch_related('eventos__usuario'), 50).get_page(request.GET.get('p'))
    responsaveis = [x for x in User.objects.filter(is_active=True).filter(
        Q(acesso_acess__trata_chamados=True) | Q(is_superuser=True) | Q(hierarchy='SUPERADMIN')).distinct()
        if trata_chamados(x)]
    colabs = Colaborador.objects.filter(status=Colaborador.ATIVO).order_by('nome') \
        if pode(request.user, 'chamados.criar') else []
    return render(request, 'acess/chamados.html', _ctx(
        request, 'chamados', pagina=pagina, chamados=pagina.object_list, st=st, q=q, motivos=MOTIVOS_CHAMADO,
        sistemas_chamado=SISTEMAS_CHAMADO, responsaveis=responsaveis, colabs=colabs, agora=agora,
        pode_tratar=pode(request.user, 'chamados.tratar') and (trata_chamados(request.user) or e_superadmin(request.user))))


def _prazo_do_post(p):
    quando = parse_datetime(p.get('quando') or '') if p.get('quando') else None
    return p.get('minutos'), quando


@login_required
@exige('chamados.criar')
@require_POST
def chamado_novo(request):
    p = request.POST
    c = get_object_or_404(Colaborador, pk=_int(p.get('colaborador')) or 0)
    minutos, quando = _prazo_do_post(p)
    resp = User.objects.filter(pk=_int(p.get('responsavel'))).first() if p.get('responsavel') else None
    try:
        if resp is not None and not trata_chamados(resp):
            raise sv.RegraAcess('O responsável precisa tratar chamados.')
        sol = Solicitacao.objects.filter(pk=_int(p.get('solicitacao')) or 0).first() if p.get('solicitacao') else None
        with transaction.atomic():
            ch = sv.criar_chamado(c, p.get('codigo'), p.get('motivo'), p.getlist('sistemas'), p.get('observacao'),
                                  minutos=minutos, quando=quando, responsavel=resp, usuario=request.user,
                                  request=request)
            if sol is not None:
                sv.vincular_solicitacao(sol, ch, request.user)
        messages.success(request, f'Chamado #{ch.codigo} aberto.' + (f' O pedido {sol.numero} foi avisado.' if sol else ''))
    except sv.RegraAcess as exc:
        _erro(request, exc)
    return _voltar(request, reverse('acess:chamados'))


@login_required
@exige('chamados.tratar')
@require_POST
def chamado_acao(request, pk):
    ch = get_object_or_404(Chamado, pk=pk)
    p, u = request.POST, request.user
    try:
        if p.get('acao') == 'concluir':
            sv.concluir_chamado(ch, p.get('versao'), u, request)
            messages.success(request, f'Chamado #{ch.codigo} concluído.')
        elif p.get('acao') == 'reagendar':
            minutos, quando = _prazo_do_post(p)
            sv.reagendar_chamado(ch, p.get('versao'), minutos, quando, u, request)
            messages.success(request, f'Chamado #{ch.codigo}: nova verificação registrada.')
        elif p.get('acao') == 'atribuir':
            alvo = u if p.get('responsavel') == 'eu' else (
                User.objects.filter(pk=_int(p.get('responsavel'))).first() if p.get('responsavel') else None)
            sv.atribuir_chamado(ch, p.get('versao'), alvo, u, request)
            messages.success(request, f'Responsável do chamado #{ch.codigo} atualizado.')
    except sv.RegraAcess as exc:
        _erro(request, exc)
    return _voltar(request, reverse('acess:chamados'))


# ── Solicitações (janela de Suporte do portal) ─────────────────────────────
@login_required
def solicitar(request):
    """A aba "Acessos" da janela de Suporte: qualquer pessoa logada pede um chamado de acesso.
    GET devolve o formulário (motivos, sistemas, a ficha e os pedidos dela); POST cria o pedido."""
    import json

    u = request.user
    if request.method == 'POST':
        try:
            corpo = json.loads(request.body or b'{}') if request.content_type == 'application/json' else request.POST
            sistemas = corpo.get('sistemas') if isinstance(corpo, dict) else corpo.getlist('sistemas')
            sol = sv.criar_solicitacao(u, corpo.get('motivo'), sistemas or [], corpo.get('descricao'),
                                       corpo.get('contato'), request)
        except (ValueError, sv.RegraAcess) as exc:
            return JsonResponse({'ok': False, 'erro': str(exc) if isinstance(exc, sv.RegraAcess) else 'Envio inválido.'},
                                status=400)
        return JsonResponse({'ok': True, 'numero': sol.numero,
                             'msg': f'Pedido {sol.numero} enviado ao time de Acessos. Você recebe um aviso quando o '
                                    f'chamado for aberto.'})
    ficha = sv.ficha_do_usuario(u)
    minhas = Solicitacao.objects.filter(solicitante=u).select_related('chamado')[:8]
    return JsonResponse({
        'motivos': [{'valor': k, 'rotulo': r} for k, r in MOTIVOS_CHAMADO],
        'sistemas': SISTEMAS_CHAMADO,
        'ficha': {'nome': ficha.nome, 'login': ficha.login, 'loja': (ficha.local.name if ficha.local else '')}
        if ficha else None,
        'contato': getattr(u, 'phone', '') or '',
        'minhas': [{'numero': s.numero, 'motivo': s.get_motivo_display(), 'status': s.status,
                    'status_rotulo': s.get_status_display(), 'quando': timezone.localtime(s.criado_em).strftime('%d/%m %H:%M'),
                    'chamado': s.chamado.codigo if s.chamado else '', 'resposta': s.resposta} for s in minhas],
    })


@login_required
@exige('chamados.ver')
def solicitacoes(request):
    st = request.GET.get('status') or 'abertas'
    qs = Solicitacao.objects.select_related('solicitante', 'solicitante__sector', 'colaborador', 'colaborador__local',
                                            'colaborador__loja_vinculo', 'chamado', 'tratado_por')
    if st == 'abertas':
        qs = qs.filter(status=Solicitacao.ABERTA).order_by('criado_em')
    pagina = Paginator(qs, 40).get_page(request.GET.get('p'))
    return render(request, 'acess/solicitacoes.html', _ctx(
        request, 'solicitacoes', pagina=pagina, solicitacoes=pagina.object_list, st=st, motivos=MOTIVOS_CHAMADO,
        sistemas_chamado=SISTEMAS_CHAMADO,
        pode_tratar=pode(request.user, 'chamados.tratar'), pode_criar=pode(request.user, 'chamados.criar')))


@login_required
@exige('chamados.tratar')
@require_POST
def solicitacao_responder(request, pk):
    sol = get_object_or_404(Solicitacao, pk=pk)
    try:
        sv.responder_solicitacao(sol, request.POST.get('status'), request.POST.get('resposta'), request.user, request)
        messages.success(request, f'Pedido {sol.numero} respondido.')
    except sv.RegraAcess as exc:
        _erro(request, exc)
    return _voltar(request, reverse('acess:solicitacoes'))


# ── Histórico ───────────────────────────────────────────────────────────────
@login_required
@exige('historico.ver')
def historico(request):
    qs = LogAcess.objects.select_related('colaborador', 'colaborador__local')
    f = {k: (request.GET.get(k) or '').strip() for k in ('categoria', 'q', 'de', 'ate', 'local')}
    if f['categoria'] in dict(LogAcess.CATEGORIAS):
        qs = qs.filter(categoria=f['categoria'])
    if f['categoria'] == 'ADMIN' and not e_superadmin(request.user):
        qs = qs.none()
    elif not e_superadmin(request.user):
        qs = qs.exclude(categoria='ADMIN')
    if f['q']:
        qs = qs.filter(Q(descricao__icontains=f['q']) | Q(usuario_nome__icontains=f['q'])
                       | Q(colaborador__nome__icontains=f['q']))
    if f['local'].isdigit():
        qs = qs.filter(colaborador__local_id=int(f['local']))
    if parse_date(f['de'] or ''):
        qs = qs.filter(quando__date__gte=parse_date(f['de']))
    if parse_date(f['ate'] or ''):
        qs = qs.filter(quando__date__lte=parse_date(f['ate']))
    pagina = Paginator(qs, 100).get_page(request.GET.get('p'))
    query = request.GET.copy()
    query.pop('p', None)
    return render(request, 'acess/historico.html', _ctx(
        request, 'historico', pagina=pagina, registros=pagina.object_list, filtro=f, categorias=LogAcess.CATEGORIAS,
        query=query.urlencode(), unidades=lojas_do_acess()))


# ── Desligamentos ───────────────────────────────────────────────────────────
@login_required
@exige('desligamentos.gerenciar')
def desligamentos(request):
    hoje = timezone.localdate()
    base = Colaborador.objects.select_related('local')
    return render(request, 'acess/desligamentos.html', _ctx(
        request, 'desligamentos', hoje=hoje,
        pendentes_execucao=list(base.filter(status=Colaborador.ATIVO, desligamento_em__lte=hoje).order_by('desligamento_em')),
        agendados=list(base.filter(status=Colaborador.ATIVO, desligamento_em__gt=hoje).order_by('desligamento_em')),
        arquivados=list(base.filter(status=Colaborador.DESLIGADO).order_by('-desligado_em')[:200]),
        inativados=list(AvisoInativacao.objects.filter(status=AvisoInativacao.PENDENTE)
                        .select_related('colaborador', 'colaborador__local', 'user').order_by('inativado_em')),
        ativos=base.filter(status=Colaborador.ATIVO, desligamento_em__isnull=True).order_by('nome')))


@login_required
@exige('desligamentos.gerenciar')
@require_POST
def desligamento_acao(request, pk):
    c = get_object_or_404(Colaborador, pk=pk)
    p, u = request.POST, request.user
    try:
        if p.get('acao') == 'agendar':
            sv.agendar_desligamento(c, parse_date(p.get('saida') or ''), parse_date(p.get('execucao') or ''),
                                    p.get('motivo'), p.get('versao'), u, request)
            messages.success(request, f'Desligamento de {c.nome} agendado.')
        elif p.get('acao') == 'cancelar':
            sv.cancelar_desligamento(c, p.get('versao'), u, request)
            messages.success(request, f'Desligamento de {c.nome} cancelado.')
        elif p.get('acao') == 'confirmar':
            if p.get('confirmo') != '1':
                raise sv.RegraAcess('Confirme que a execução foi feita nos sistemas externos.')
            sv.arquivar(c, p.get('versao'), u, request, exige_agendamento=True)
            messages.success(request, f'Desligamento de {c.nome} confirmado e ficha arquivada.')
    except sv.RegraAcess as exc:
        _erro(request, exc)
    return _voltar(request, reverse('acess:desligamentos'))


@login_required
@exige('desligamentos.gerenciar')
@require_POST
def inativacao_decidir(request, pk):
    """Usuário inativado no portal: desligar já, agendar (PPL) ou não desligar."""
    aviso = get_object_or_404(AvisoInativacao.objects.select_related('colaborador'), pk=pk)
    p = request.POST
    try:
        sv.decidir_inativacao(aviso, p.get('decisao'), parse_date(p.get('execucao') or ''), request.user, request,
                              observacao=p.get('observacao'))
        nome = aviso.colaborador.nome
        messages.success(request, {'imediato': f'Desligamento de {nome} para hoje: está em "Ação manual pendente".',
                                   'agendar': f'Desligamento de {nome} agendado; no dia chega a notificação.',
                                   'descartar': f'Aviso de {nome} encerrado sem desligamento.'}.get(p.get('decisao'), 'Feito.'))
    except sv.RegraAcess as exc:
        _erro(request, exc)
    return _voltar(request, reverse('acess:desligamentos'))


# ── Mensagens ───────────────────────────────────────────────────────────────
@login_required
@exige()
def mensagens(request):
    sv.garantir_modelos()
    if request.method == 'POST':
        if not pode(request.user, 'mensagens.gerenciar'):
            messages.error(request, 'Editar as mensagens não está liberado para o seu usuário.')
            return redirect('acess:mensagens')
        m = get_object_or_404(ModeloMensagem, pk=_int(request.POST.get('id')) or 0)
        try:
            texto = sv.validar_texto(request.POST.get('texto'))
            antes = m.texto
            m.texto, m.atualizado_por = texto, request.user
            m.save()
            sv.registrar(request, 'MENSAGEM_EDITADA', 'ADMIN', f'Editou o modelo "{m.titulo}"',
                         antes={'texto': antes}, depois={'texto': texto})
            messages.success(request, f'Modelo "{m.titulo}" salvo.')
        except sv.RegraAcess as exc:
            _erro(request, exc)
        return redirect('acess:mensagens')
    return render(request, 'acess/mensagens.html', _ctx(request, 'mensagens', modelos=ModeloMensagem.objects.all(),
                                                        variaveis=sv.VARIAVEIS))


# ── Acessos (só SUPERADMIN) ─────────────────────────────────────────────────
@login_required
def acessos(request):
    if not e_superadmin(request.user):
        messages.error(request, 'Só o SUPERADMIN define quem entra no Acess.')
        return redirect('acess:dashboard')
    if request.method == 'POST':
        p = request.POST
        pessoa = User.objects.filter(pk=_int(p.get('user')), is_active=True).first()
        if pessoa is None:
            messages.error(request, 'Escolha a pessoa.')
        elif e_superadmin(pessoa):
            messages.info(request, 'SUPERADMIN já tem o Acess inteiro.')
        elif p.get('remover'):
            a = AcessoAcess.objects.filter(user=pessoa).first()
            if a:
                sv.registrar(request, 'ACESSO_REMOVIDO', 'ADMIN', f'Tirou o acesso de {pessoa.get_full_name()} ao Acess',
                             antes={'permissoes': a.permissoes})
                a.delete()
            messages.success(request, f'{pessoa.get_full_name()} não entra mais no Acess.')
        else:
            chaves = completar(p.getlist('permissoes'))
            if not chaves:
                messages.error(request, 'Marque ao menos uma função (ou use "Tirar acesso").')
            else:
                a, criado = AcessoAcess.objects.get_or_create(user=pessoa, defaults={'concedido_por': request.user})
                antes = None if criado else {'permissoes': a.permissoes, 'trata_chamados': a.trata_chamados}
                a.permissoes, a.ativo = chaves, p.get('ativo', 'on') == 'on'
                a.trata_chamados = p.get('trata_chamados') == 'on'
                a.observacao, a.concedido_por = (p.get('observacao') or '').strip()[:200], request.user
                a.save()
                sv.registrar(request, 'ACESSO_ALTERADO', 'ADMIN',
                             f'{"Liberou" if criado else "Mudou"} o acesso de {pessoa.get_full_name()} ao Acess',
                             antes=antes, depois={'permissoes': chaves, 'trata_chamados': a.trata_chamados})
                messages.success(request, f'Acesso de {pessoa.get_full_name()} salvo: {len(chaves)} função(ões).')
        return redirect('acess:acessos')
    liberados_ = list(AcessoAcess.objects.select_related('user', 'user__sector').order_by('-ativo', 'user__first_name'))
    for a in liberados_:
        a.chaves = set(a.permissoes or [])
    candidatos = (User.objects.filter(is_active=True).exclude(id__in=[a.user_id for a in liberados_])
                  .exclude(Q(is_superuser=True) | Q(hierarchy='SUPERADMIN')).select_related('sector').order_by('first_name'))
    return render(request, 'acess/acessos.html', _ctx(request, 'acessos', liberados=liberados_, candidatos=candidatos,
                                                      funcoes=FUNCOES))
