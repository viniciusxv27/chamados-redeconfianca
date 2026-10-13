"""Documentos criados pelo assistente: Markdown → Word, Excel e Markdown.

O Claude escreve o documento em Markdown (títulos, listas, negrito, tabelas com
``|``) e, quando há números para planilha, em tabelas estruturadas. Daqui saem:

- **Word (.docx)**: títulos, parágrafos, listas, negrito/itálico e tabelas de verdade;
- **Excel (.xlsx)**: uma aba por tabela (as estruturadas e as do texto), números como número;
- **Markdown (.md)**: o texto como foi escrito;
- a **página** do documento (``/assistente/documentos/<id>/``) imprime como PDF.
"""
import io
import re
from decimal import Decimal, InvalidOperation

from django.utils import timezone

MAX_TABELAS = 20
MAX_LINHAS = 2000
MAX_COLUNAS = 40

_NEGRITO_ITALICO = re.compile(r'(\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*|_[^_\s][^_]*_|`[^`]+`|\[[^\]]+\]\([^)]+\))')


# ---------------------------------------------------------------------------
# Entrada
# ---------------------------------------------------------------------------
def limpar_tabelas(tabelas):
    """Tabelas como o Claude mandou → lista segura de {titulo, colunas, linhas}."""
    saida = []
    for tabela in (tabelas or [])[:MAX_TABELAS]:
        if not isinstance(tabela, dict):
            continue
        colunas = [str(c)[:120] for c in (tabela.get('colunas') or [])][:MAX_COLUNAS]
        linhas = []
        for linha in (tabela.get('linhas') or [])[:MAX_LINHAS]:
            if isinstance(linha, (list, tuple)):
                linhas.append([_celula(v) for v in linha][:MAX_COLUNAS])
        if colunas or linhas:
            saida.append({'titulo': str(tabela.get('titulo') or '')[:120], 'colunas': colunas, 'linhas': linhas})
    return saida


def _celula(valor):
    if valor is None:
        return ''
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return valor
    return str(valor)[:2000]


def tabelas_do_markdown(texto):
    """As tabelas escritas com | no texto, para virarem abas no Excel."""
    tabelas, atual, titulo = [], None, ''
    for linha in (texto or '').splitlines():
        crua = linha.strip()
        if crua.startswith('|') and crua.endswith('|') and crua.count('|') >= 2:
            celulas = [c.strip() for c in crua.strip('|').split('|')]
            if all(re.fullmatch(r':?-{2,}:?', c) for c in celulas if c):
                continue                                    # linha separadora do cabeçalho
            if atual is None:
                atual = {'titulo': titulo, 'colunas': [_sem_marcas(c) for c in celulas], 'linhas': []}
                tabelas.append(atual)
            else:
                atual['linhas'].append([_sem_marcas(c) for c in celulas])
        else:
            atual = None
            if crua.startswith('#'):
                titulo = crua.lstrip('#').strip()
    return tabelas


def _sem_marcas(texto):
    texto = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', texto)
    return re.sub(r'(\*\*|__|`)', '', texto).strip()


def _numero(texto):
    """'R$ 1.234,56' / '12,5%' / '1234' → número (para o Excel somar); senão None."""
    if isinstance(texto, (int, float)):
        return texto
    bruto = str(texto or '').strip()
    if not bruto or len(bruto) > 30:
        return None
    pct = bruto.endswith('%')
    limpo = re.sub(r'^(R\$|\$)\s*', '', bruto.rstrip('%').strip())
    if not re.fullmatch(r'-?[\d.,]+', limpo):
        return None
    if ',' in limpo:
        limpo = limpo.replace('.', '').replace(',', '.')
    elif limpo.count('.') > 1:
        limpo = limpo.replace('.', '')
    try:
        valor = Decimal(limpo)
    except InvalidOperation:
        return None
    valor = float(valor)
    return valor / 100 if pct else valor


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------
def _inline(paragrafo, texto):
    """Negrito, itálico, código e links dentro de um parágrafo do Word."""
    for pedaco in _NEGRITO_ITALICO.split(texto):
        if not pedaco:
            continue
        if (pedaco.startswith('**') and pedaco.endswith('**')) or (pedaco.startswith('__') and pedaco.endswith('__')):
            paragrafo.add_run(pedaco[2:-2]).bold = True
        elif pedaco.startswith('`') and pedaco.endswith('`'):
            run = paragrafo.add_run(pedaco[1:-1])
            run.font.name = 'Consolas'
        elif pedaco.startswith('[') and '](' in pedaco:
            rotulo, url = re.match(r'\[([^\]]+)\]\(([^)]+)\)', pedaco).groups()
            paragrafo.add_run(rotulo)
            if url.startswith('http') or url.startswith('/'):
                paragrafo.add_run(f' ({url})').italic = True
        elif len(pedaco) > 2 and pedaco[0] in '*_' and pedaco[-1] == pedaco[0]:
            paragrafo.add_run(pedaco[1:-1]).italic = True
        else:
            paragrafo.add_run(pedaco)


def _tabela_word(documento, colunas, linhas):
    largura = max([len(colunas)] + [len(l) for l in linhas]) if (colunas or linhas) else 0
    if not largura:
        return
    tabela = documento.add_table(rows=0, cols=largura)
    tabela.style = 'Table Grid'
    if colunas:
        celulas = tabela.add_row().cells
        for i, valor in enumerate(colunas):
            celulas[i].text = ''
            celulas[i].paragraphs[0].add_run(_sem_marcas(str(valor))).bold = True
    for linha in linhas:
        celulas = tabela.add_row().cells
        for i, valor in enumerate(linha[:largura]):
            celulas[i].text = ''
            _inline(celulas[i].paragraphs[0], str(valor))
    documento.add_paragraph()


def gerar_docx(doc):
    from docx import Document
    from docx.shared import Pt

    documento = Document()
    estilo = documento.styles['Normal']
    estilo.font.name = 'Calibri'
    estilo.font.size = Pt(11)
    documento.add_heading(doc.titulo, level=0)
    rodape = documento.add_paragraph()
    rodape.add_run(f'Gerado pelo Assistente do portal Rede Confiança em '
                   f'{timezone.localtime(doc.atualizado_em):%d/%m/%Y %H:%M} · '
                   f'{doc.user.get_full_name() or doc.user.get_username()}').italic = True

    linhas = (doc.conteudo or '').splitlines()
    i = 0
    paragrafo = []

    def fechar_paragrafo():
        if paragrafo:
            _inline(documento.add_paragraph(), ' '.join(paragrafo))
            paragrafo.clear()

    while i < len(linhas):
        crua = linhas[i].rstrip()
        texto = crua.strip()
        if not texto:
            fechar_paragrafo()
            i += 1
            continue
        if texto.startswith('|') and texto.endswith('|'):
            fechar_paragrafo()
            bloco = []
            while i < len(linhas) and linhas[i].strip().startswith('|'):
                bloco.append(linhas[i].strip())
                i += 1
            celulas = [[c.strip() for c in l.strip('|').split('|')] for l in bloco]
            celulas = [c for c in celulas if not all(re.fullmatch(r':?-{2,}:?', x) for x in c if x)]
            if celulas:
                _tabela_word(documento, celulas[0], celulas[1:])
            continue
        cabecalho = re.match(r'^(#{1,6})\s+(.*)$', texto)
        if cabecalho:
            fechar_paragrafo()
            documento.add_heading(_sem_marcas(cabecalho.group(2)), level=min(len(cabecalho.group(1)), 4))
        elif re.match(r'^[-*+]\s+', texto):
            fechar_paragrafo()
            _inline(documento.add_paragraph(style='List Bullet'), re.sub(r'^[-*+]\s+', '', texto))
        elif re.match(r'^\d+[.)]\s+', texto):
            fechar_paragrafo()
            _inline(documento.add_paragraph(style='List Number'), re.sub(r'^\d+[.)]\s+', '', texto))
        elif texto.startswith('>'):
            fechar_paragrafo()
            p = documento.add_paragraph()
            p.add_run(texto.lstrip('> ').strip()).italic = True
        elif re.fullmatch(r'(-{3,}|\*{3,}|_{3,})', texto):
            fechar_paragrafo()
        else:
            paragrafo.append(texto)
        i += 1
    fechar_paragrafo()

    for tabela in doc.tabelas or []:
        if tabela.get('titulo'):
            documento.add_heading(tabela['titulo'], level=2)
        _tabela_word(documento, tabela.get('colunas') or [], tabela.get('linhas') or [])

    saida = io.BytesIO()
    documento.save(saida)
    return saida.getvalue()


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def gerar_xlsx(doc):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    tabelas = list(doc.tabelas or []) + tabelas_do_markdown(doc.conteudo)
    livro = Workbook()
    livro.remove(livro.active)
    usados = set()
    for n, tabela in enumerate(tabelas, start=1):
        nome = re.sub(r'[\[\]\*\?/\\:]', ' ', tabela.get('titulo') or f'Tabela {n}').strip()[:28] or f'Tabela {n}'
        base, k = nome, 2
        while nome.lower() in usados:
            nome = f'{base[:25]} {k}'
            k += 1
        usados.add(nome.lower())
        aba = livro.create_sheet(nome)
        colunas = tabela.get('colunas') or []
        if colunas:
            aba.append([_sem_marcas(str(c)) for c in colunas])
            for celula in aba[1]:
                celula.font = Font(bold=True, color='FFFFFF')
                celula.fill = PatternFill('solid', fgColor='4B2A7B')
                celula.alignment = Alignment(vertical='center')
            aba.freeze_panes = 'A2'
        for linha in tabela.get('linhas') or []:
            valores = []
            for valor in linha:
                numero = _numero(valor)
                valores.append(numero if numero is not None else _sem_marcas(str(valor)))
            aba.append(valores)
        for i, coluna in enumerate(aba.columns, start=1):
            largura = max((len(str(c.value)) for c in coluna if c.value is not None), default=8)
            aba.column_dimensions[get_column_letter(i)].width = min(max(largura + 2, 10), 60)
    if not tabelas:
        aba = livro.create_sheet('Documento')
        aba.append([doc.titulo])
        aba['A1'].font = Font(bold=True, size=14)
        for linha in (doc.conteudo or '').splitlines():
            aba.append([_sem_marcas(linha.lstrip('#-* ').strip())])
        aba.column_dimensions['A'].width = 120
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def gerar_md(doc):
    partes = [f'# {doc.titulo}', '', doc.conteudo or '']
    for tabela in doc.tabelas or []:
        partes += ['', f"## {tabela.get('titulo') or 'Tabela'}", '']
        colunas = tabela.get('colunas') or []
        if colunas:
            partes.append('| ' + ' | '.join(str(c) for c in colunas) + ' |')
            partes.append('|' + '---|' * len(colunas))
        for linha in tabela.get('linhas') or []:
            partes.append('| ' + ' | '.join(str(v) for v in linha) + ' |')
    return '\n'.join(partes).encode('utf-8')


def nome_de_arquivo(doc, extensao):
    import unicodedata
    base = unicodedata.normalize('NFKD', doc.titulo).encode('ascii', 'ignore').decode()
    base = re.sub(r'[^A-Za-z0-9]+', '-', base).strip('-').lower()[:60] or 'documento'
    return f'{base}.{extensao}'
