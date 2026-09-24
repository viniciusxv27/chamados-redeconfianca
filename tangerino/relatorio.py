"""Relatório de ponto em Excel: uma linha por pessoa e dia, com a pendência em texto.

As colunas são as que a gestão pediu para conferir o mês: nome, loja, data, as
quatro batidas, previsto, trabalhado, intervalo, horas extras e a pendência.
Quem monta as linhas é ``tangerino/pendencias.py`` — aqui só entra a planilha e a
leitura dos filtros da tela (período, setor e pessoa).
"""
from datetime import date, timedelta
from io import BytesIO

from django.utils import timezone

from . import pendencias as svc

COLUNAS = [
    ('Nome', 32), ('Loja', 26), ('Data', 12),
    ('1ª batida', 11), ('2ª batida', 11), ('3ª batida', 11), ('4ª batida', 11),
    ('Horas previstas', 15), ('Horas trabalhadas', 17), ('Intervalo', 11),
    ('Horas extras', 13), ('Pendência', 70),
]
MAXIMO_DE_DIAS = 186            # seis meses: o período pedido tem limite para a tela não travar


def periodo_padrao(hoje=None):
    """Do primeiro dia do mês até hoje — o recorte que a gestão abre primeiro."""
    hoje = hoje or timezone.localdate()
    return hoje.replace(day=1), hoje


def ler_data(texto, padrao):
    try:
        return date.fromisoformat((texto or '').strip())
    except (TypeError, ValueError):
        return padrao


def ler_filtros(pedido, hoje=None):
    """Período, setor e pessoa do GET, já arrumados (e com o período em ordem)."""
    inicio_padrao, fim_padrao = periodo_padrao(hoje)
    inicio = ler_data(pedido.get('de'), inicio_padrao)
    fim = ler_data(pedido.get('ate'), fim_padrao)
    if fim < inicio:
        inicio, fim = fim, inicio
    if (fim - inicio).days > MAXIMO_DE_DIAS:
        inicio = fim - timedelta(days=MAXIMO_DE_DIAS)
    setor = (pedido.get('setor') or '').strip()
    usuario = (pedido.get('usuario') or '').strip()
    return {
        'de': inicio,
        'ate': fim,
        'setor': int(setor) if setor.isdigit() else None,
        'usuario': int(usuario) if usuario.isdigit() else None,
        'pendencias': pedido.get('pendencias') == '1',
    }


def linhas(filtros, aviso=None):
    """As linhas do período pedido.

    ``aviso`` é preenchido com o que aconteceu ao completar o período no
    Tangerino (quantas pessoas foram buscadas, ou o erro) — a tela conta isso
    para ninguém ler um relatório curto achando que é a realidade.
    """
    return svc.linhas_do_periodo(
        filtros['de'], filtros['ate'], setor_id=filtros['setor'],
        usuarios=[filtros['usuario']] if filtros['usuario'] else None,
        apenas_com_pendencia=filtros['pendencias'], aviso=aviso)


def nome_do_arquivo(filtros):
    return f"ponto-{filtros['de']:%Y-%m-%d}-a-{filtros['ate']:%Y-%m-%d}.xlsx"


def planilha(linhas_do_relatorio, filtros, rotulos=None):
    """A planilha pronta para download, com o cabeçalho laranja do portal."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    livro = Workbook()
    aba = livro.active
    aba.title = 'Ponto'

    rotulos = rotulos or {}
    recorte = [f"Período: {filtros['de']:%d/%m/%Y} a {filtros['ate']:%d/%m/%Y}"]
    if rotulos.get('setor'):
        recorte.append(f"Loja/setor: {rotulos['setor']}")
    if rotulos.get('usuario'):
        recorte.append(f"Pessoa: {rotulos['usuario']}")
    if filtros['pendencias']:
        recorte.append('Somente dias com pendência')
    aba.append([' · '.join(recorte)])
    aba['A1'].font = Font(bold=True, size=12)
    aba.append([])

    aba.append([titulo for titulo, _ in COLUNAS])
    for celula in aba[3]:
        celula.font = Font(bold=True, color='FFFFFF')
        celula.fill = PatternFill('solid', fgColor='F26522')
        celula.alignment = Alignment(vertical='center')

    for linha in linhas_do_relatorio:
        aba.append([
            linha['nome'], linha['loja'], linha['data'].strftime('%d/%m/%Y'),
            *linha['batidas'],
            linha['previsto'], linha['trabalhado'], linha['intervalo'], linha['horas_extras'],
            linha['pendencia'],
        ])

    for coluna, (_, largura) in zip('ABCDEFGHIJKL', COLUNAS):
        aba.column_dimensions[coluna].width = largura
    aba.freeze_panes = 'A4'
    aba.auto_filter.ref = f'A3:L{max(3, len(linhas_do_relatorio) + 3)}'

    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()
