"""/vendas/precos/importar/ (29/10/2026): aviso no WhatsApp do que subiu e do que baixou.

Pedido: "Ao atualizar tabela de preços, SUPERADMIN deve poder selecionar quem vai
receber um aviso no WhatsApp pelo WhatsApp configurado com produtos que subiram e
produtos que diminuíram valor."

Cobre:
- a comparação (subiu / diminuiu / novo / saiu, % de variação, PIX decide quando o
  Preço não muda, ajuste só nas parcelas vira "outros");
- a mensagem: contagens, maiores altas e quedas, "e mais N", e o limite de tamanho;
- uma mensagem por pessoa marcada; quem não tem telefone fica de fora e aparece
  marcado na tela e no resultado; falha no envio não desfaz a importação;
- só o SUPERADMIN escolhe quem recebe (de outro perfil o campo é ignorado);
- a última seleção volta marcada na próxima vez.

Transação desfeita no fim (a migration 0010 só existe dentro dela). O envio é um
dublê de `core.evolution.enviar_texto` — e, de todo jeito, processo de teste
nunca manda WhatsApp (core.utils.processo_de_teste).
"""
import os
import re
import sys
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from unittest import mock

import django

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'redeconfianca.settings')
os.environ.setdefault('RC_VARREDURA_ROTINA', '0')

from django.conf import settings

settings.CACHES = {
    'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-aviso-precos'},
    'local': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-aviso-precos-2'},
}
django.setup()

from django.test.utils import setup_test_environment

setup_test_environment()
if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS.append('testserver')

import openpyxl
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection, transaction
from django.db.migrations.executor import MigrationExecutor
from django.test import Client

User = get_user_model()
D = Decimal
ok = fail = 0


def t(nome, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'  OK   {nome}')
    else:
        fail += 1
        print(f'  FALHA {nome} {extra}')


# ── Planilha no formato oficial (o mesmo de teste_vendas_slv.py) ───────────────

def bloco(base):
    if base == '-':
        return ['-'] * 22
    return [base, round(base * 0.9, 2)] + [base] * 11 + [round(base * (1 + 0.01 * n), 2) for n in range(1, 10)]


def cabecalho_regular(grupos):
    cab = []
    for g in grupos:
        cab += [g, 'PIX\ne\nVivo Pay'] + [f'{n}x' for n in range(2, 22)]
    return cab


def aba_regular(wb, nome, grupos, linhas):
    ws = wb.create_sheet(nome)
    ws.append([None] * 4 + ['TITULO'])
    largura = len(cabecalho_regular(grupos))
    ws.append([None] * 4 + ['Vigência', 'TABELA REGULAR'] + [None] * (largura - 1) + ['TABELA RENOVA'])
    ws.append(['PORTFÓLIO', 'CATEGORIA', 'MARCA', 'OFERTA DE COMUNICAÇÃO', 'Nome Comercial']
              + cabecalho_regular(grupos) + ['DESCONTO', 'PRÉ-RENOVA'])
    for nome_item, bases in linhas:
        ws.append(['In', 'Smartphone', 'Zz', 'Oferta', nome_item] + sum((bloco(b) for b in bases), []) + [0, 999999])


def planilha(smartphones, pvp):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    aba_regular(wb, 'SMARTPHONES', ['PRÉ', 'PÓS INDIVIDUAL'], smartphones)
    aba_regular(wb, 'ELETRÔNICOS_LP_Conectados', ['PRÉ', 'MULTIVIVO'], [('ZZAV Modem', [400, 300])])
    ws = wb.create_sheet('ELETRÔNICOS_LP_Não Conectados')
    ws.append([None, None, None, 'TITULO'])
    ws.append([None, None, None, 'Vigência'])
    ws.append(['PORTFÓLIO', 'CATEGORIA', 'MARCA', 'Nome Comercial', 'PVP Base', 'PIX\ne\nVivo Pay', '2x'])
    for nome_item, valor in pvp:
        ws.append(['In', 'Essenciais', 'Zz', nome_item, valor, None, None])
    arquivo = BytesIO()
    wb.save(arquivo)
    return arquivo.getvalue()


marcador = transaction.atomic()
marcador.__enter__()
try:
    executor = MigrationExecutor(connection)
    alvos = executor.loader.graph.leaf_nodes('vendas')
    if executor.migration_plan(alvos):
        executor.migrate(alvos)

    from vendas import aviso_precos
    from vendas.aviso_precos import calcular_variacao, montar_mensagem, precos_do_item
    from vendas.models import AcessoVendas, ConfiguracaoVendas, ImportacaoPrecos, ItemPreco

    print('== COMPARAÇÃO ==')
    tab = lambda pre, pix, pos=None: {'tabela': {'PRÉ': {'Preço': pre, 'PIX e Vivo Pay': pix, '12x': pre},
                                                 **({'PÓS': {'Preço': pos, 'PIX e Vivo Pay': pos}} if pos else {})}}
    p = precos_do_item(D('3000'), tab('3000.00', '2700.00', '2500.00'))
    t('Preço = menor preço entre os grupos; PIX = menor PIX', p['Preço'] == D('2500.00') and p['PIX'] == D('2500.00'), p)
    p = precos_do_item(D('39.90'), {'CATEGORIA': 'Essenciais'})
    t('Não Conectados: Preço = PVP, sem PIX', p['Preço'] == D('39.90') and p['PIX'] is None, p)

    S = 'SMARTPHONES'
    antes = {(S, 'Sobe'): precos_do_item(None, tab('1000.00', '900.00')),
             (S, 'Desce'): precos_do_item(None, tab('2000.00', '1800.00')),
             (S, 'So PIX'): precos_do_item(None, tab('500.00', '450.00')),
             (S, 'Misto'): precos_do_item(None, tab('800.00', '720.00')),
             (S, 'Parcelas'): precos_do_item(None, tab('300.00', '270.00')),
             (S, 'Igual'): precos_do_item(None, tab('100.00', '90.00')),
             (S, 'Saiu'): precos_do_item(None, tab('700.00', '630.00'))}
    parcela_nova = tab('300.00', '270.00')
    parcela_nova['tabela']['PRÉ']['12x'] = '330.00'
    depois = {(S, 'Sobe'): precos_do_item(None, tab('1100.00', '990.00')),
              (S, 'Desce'): precos_do_item(None, tab('1500.00', '1350.00')),
              (S, 'So PIX'): precos_do_item(None, tab('500.00', '400.00')),
              (S, 'Misto'): precos_do_item(None, tab('840.00', '700.00')),
              (S, 'Parcelas'): precos_do_item(None, parcela_nova),
              (S, 'Igual'): precos_do_item(None, tab('100.00', '90.00')),
              (S, 'Novo'): precos_do_item(None, tab('999.00', '899.00'))}
    v = calcular_variacao(antes, depois)
    nomes = lambda lista: [i['nome'] for i in lista]
    t('subiram: maior alta primeiro', nomes(v['subiram']) == ['Sobe', 'Misto'], nomes(v['subiram']))
    t('diminuíram: maior queda primeiro', nomes(v['diminuiram']) == ['Desce', 'So PIX'], nomes(v['diminuiram']))
    t('novos e saíram', nomes(v['novos']) == ['Novo'] and nomes(v['removidos']) == ['Saiu'], (v['novos'], v['removidos']))
    t('ajuste só nas parcelas: "outros", fora das listas', v['outros'] == 1, v['outros'])
    sobe = v['subiram'][0]['mudancas']
    t('% de variação (Preço e PIX)', [(m['coluna'], m['de'], m['para'], m['pct']) for m in sobe]
      == [('Preço', D('1000.00'), D('1100.00'), D('10.0')), ('PIX', D('900.00'), D('990.00'), D('10.0'))], sobe)
    t('queda de 25%', v['diminuiram'][0]['mudancas'][0]['pct'] == D('-25.0'))
    so_pix = v['diminuiram'][1]['mudancas']
    t('Preço igual: a direção vem do PIX', len(so_pix) == 1 and so_pix[0]['coluna'] == 'PIX' and so_pix[0]['pct'] == D('-11.1'), so_pix)
    misto = v['subiram'][1]['mudancas']
    t('Preço subiu e PIX baixou: conta como subiu, com as duas colunas', [m['coluna'] for m in misto] == ['Preço', 'PIX']
      and misto[1]['pct'] < 0, misto)

    print('\n== MENSAGEM ==')
    msg = montar_mensagem(v, arquivo='tabela.xlsx', quando=datetime(2026, 10, 29, 14, 30), nome='Ana')
    t('cabeçalho, nome, data e arquivo', msg.startswith('Olá, Ana!') and '29/10/2026 14:30' in msg and 'tabela.xlsx' in msg, msg)
    t('contagens', '⬆️ Subiram: 2' in msg and '⬇️ Diminuíram: 2' in msg and '🆕 Novos: 1' in msg
      and 'Saíram da tabela: 1' in msg and 'Outros ajustes (parcelas/grupos): 1' in msg, msg)
    t('linha com de → para e %', '• Sobe: Preço R$ 1.000,00 → R$ 1.100,00 (+10,0%); PIX R$ 900,00 → R$ 990,00 (+10,0%)' in msg, msg)
    t('queda com sinal', '• Desce: Preço R$ 2.000,00 → R$ 1.500,00 (-25,0%)' in msg)
    t('novos e saíram pelo nome', '*🆕 Novos (1):* Novo' in msg and '*❌ Saíram (1):* Saiu' in msg)
    print('     ---- exemplo ----\n' + '\n'.join('     ' + l for l in msg.splitlines()) + '\n     -----------------')

    grande_antes = {(S, f'Aparelho {n:03d} ' + 'X' * 80): {'Preço': D('1000.00'), 'PIX': D('900.00'), 'tabela': None}
                    for n in range(300)}
    grande_depois = {k: {'Preço': D('1000.00') + n + 1, 'PIX': D('900.00'), 'tabela': None}
                     for n, k in enumerate(sorted(grande_antes))}
    for n in range(150):
        grande_depois[(S, f'Novo {n:03d} ' + 'Y' * 80)] = {'Preço': D('10'), 'PIX': None, 'tabela': None}
    vg = calcular_variacao(grande_antes, grande_depois)
    msg = montar_mensagem(vg, arquivo='x.xlsx')
    t('300 altas: só as 10 maiores e "e mais 290"', msg.count('\n• ') == 10 and '… e mais 290.' in msg
      and 'Aparelho 299' in msg, msg[:400])
    t('mensagem cabe no limite', len(msg) <= aviso_precos.LIMITE_MENSAGEM, len(msg))
    pequena = montar_mensagem(vg, limite=900)
    t('limite apertado: corta produtos até caber', len(pequena) <= 900 and '⬆️ Subiram: 300' in pequena, len(pequena))
    t('limite absurdo nunca passa do teto', len(montar_mensagem(vg, limite=120)) <= 120)
    t('nada mudou: diz isso', 'Nenhum preço mudou.' in montar_mensagem(calcular_variacao(antes, antes)))

    print('\n== IMPORTAÇÃO + AVISO ==')

    def pessoa(username, **kw):
        return User.objects.create_user(username=username, email=f'{username}@exemplo-teste.local', password='S3nha!teste',
                                        first_name=kw.pop('first_name', 'ZZ'), last_name=kw.pop('last_name', username), **kw)

    chefe = pessoa('zzav.super', first_name='Chefe', hierarchy='SUPERADMIN')
    comum = pessoa('zzav.comum', first_name='Comum')
    AcessoVendas.objects.create(user=comum, permissoes=['modulo.ver', 'produtos.importar'])
    r1 = pessoa('zzav.r1', first_name='Rita', phone='(27) 99999-0001')
    r2 = pessoa('zzav.r2', first_name='Rui', phone='27 99999-0002')
    r3 = pessoa('zzav.r3', first_name='Sem', last_name='Fone', phone='')
    cs, cc = Client(), Client()
    cs.force_login(chefe)
    cc.force_login(comum)

    html = cs.get('/vendas/precos/importar/').content.decode()
    t('SUPERADMIN vê o seletor de quem recebe', 'id="vdAviso"' in html and 'name="aviso_destinatarios"' in html)
    t('quem não tem telefone aparece sinalizado e não marcável',
      re.search(rf'value="{r3.id}"[^>]*disabled', html) is not None and 'vd-sem-telefone' in html)
    t('quem tem telefone aparece com telefone', 'com telefone' in html)
    t('outro perfil não vê o seletor', 'id="vdAviso"' not in cc.get('/vendas/precos/importar/').content.decode())

    base_sm = [('ZZAV Phone Sobe', [3000, 2500]), ('ZZAV Phone Desce', [2000, 1800]), ('ZZAV Phone Igual', [1000, 900]),
               ('ZZAV Phone Sai', [500, 450])]
    base_pvp = [('ZZAV Capa', 39.9), ('ZZAV Fone', 299)]

    enviar = mock.Mock(return_value=(True, 'ok'))
    with mock.patch('core.evolution.enviar_texto', enviar):
        r = cs.post('/vendas/precos/importar/', {'arquivo': SimpleUploadedFile('base.xlsx', planilha(base_sm, base_pvp)),
                                                 'aviso_form': '1'})
    t('importação base sem ninguém marcado: nada enviado', r.status_code == 302 and enviar.call_count == 0, enviar.call_count)
    t('seleção vazia gravada', ConfiguracaoVendas.atual().aviso_precos_destinatarios.count() == 0)

    novo_sm = [('ZZAV Phone Sobe', [3300, 2750]), ('ZZAV Phone Desce', [1800, 1500]), ('ZZAV Phone Igual', [1000, 900]),
               ('ZZAV Phone Novo', [4000, 3500])]
    novo_pvp = [('ZZAV Capa', 49.9), ('ZZAV Fone', 299)]
    chamadas = []

    def dublê(numero, texto, timeout=10):
        chamadas.append((numero, texto))
        if numero.endswith('0002'):
            raise RuntimeError('Evolution fora do ar')
        return True, '{"key": "zz"}'

    with mock.patch('core.evolution.enviar_texto', side_effect=dublê) as m:
        r = cs.post('/vendas/precos/importar/', {'arquivo': SimpleUploadedFile('nova.xlsx', planilha(novo_sm, novo_pvp)),
                                                 'aviso_form': '1',
                                                 'aviso_destinatarios': [str(r1.id), str(r2.id), str(r3.id)]})
    imp = ImportacaoPrecos.objects.order_by('-id').first()
    t('falha no envio não desfaz a importação', r.status_code == 302 and imp.arquivo == 'nova.xlsx' and not imp.erro
      and ItemPreco.objects.get(categoria='SMARTPHONES', nome='ZZAV Phone Sobe').valor == D('3300.00'), imp.erro)
    t('uma mensagem por pessoa com telefone (2), nenhuma para quem não tem', m.call_count == 2
      and sorted(n for n, _ in chamadas) == ['5527999990001', '5527999990002'], chamadas)
    texto = dict(chamadas).get('5527999990001', '')
    t('mensagem com o que subiu e o que baixou', 'Olá, Rita!' in texto and '• ZZAV Phone Sobe: Preço R$ 2.500,00 → R$ 2.750,00 (+10,0%)' in texto
      and 'ZZAV Phone Desce: Preço R$ 1.800,00 → R$ 1.500,00 (-16,7%)' in texto
      and 'ZZAV Capa: Preço R$ 39,90 → R$ 49,90 (+25,1%)' in texto, texto[:1500])
    t('mensagem dentro do limite', 0 < len(texto) <= aviso_precos.LIMITE_MENSAGEM, len(texto))
    t('a seleção ficou salva', set(ConfiguracaoVendas.atual().aviso_precos_destinatarios.values_list('id', flat=True))
      == {r1.id, r2.id, r3.id})

    r = cs.get('/vendas/precos/importar/')
    html = r.content.decode()
    res = r.context['resultado']
    zz = lambda chave: [i['nome'] for i in res['variacao'][chave] if i['nome'].startswith('ZZAV')]
    t('resultado: subiram / diminuíram / novos / saíram', zz('subiram') == ['ZZAV Capa', 'ZZAV Phone Sobe']
      and zz('diminuiram') == ['ZZAV Phone Desce'] and zz('novos') == ['ZZAV Phone Novo'] and 'ZZAV Phone Sai' in zz('removidos'),
      (zz('subiram'), zz('diminuiram'), zz('novos'), zz('removidos')))
    envios = {e['id']: e for e in res['envios']}
    t('resultado do envio por pessoa', envios[r1.id]['status'] == 'enviado' and envios[r2.id]['status'] == 'falhou'
      and 'Evolution fora do ar' in envios[r2.id]['detalhe'] and envios[r3.id]['status'] == 'sem_telefone'
      and envios[r3.id]['detalhe'] == 'sem telefone no cadastro', res['envios'])
    t('a tela mostra o resumo e quem recebeu', 'O que mudou nesta importação' in html and 'ZZAV Phone Sobe' in html
      and 'Aviso no WhatsApp' in html and '>enviado<' in html and '>falhou<' in html and '>sem telefone<' in html)
    t('o resultado aparece uma vez só', cs.get('/vendas/precos/importar/').context['resultado'] is None)
    marcados = [p.id for p in r.context['aviso_pessoas'] if p.escolhido]
    t('a última seleção volta marcada (e primeiro na lista)', set(marcados) == {r1.id, r2.id, r3.id}
      and [p.id for p in r.context['aviso_pessoas'][:3]] and set(p.id for p in r.context['aviso_pessoas'][:3]) == {r1.id, r2.id, r3.id}
      and r.context['aviso_escolhidos'] == 3, marcados)
    t('quem sem telefone foi marcado antes continua marcado (para poder desmarcar)',
      re.search(rf'value="{r3.id}"[^>]*checked', html) is not None)

    print('\n== SÓ O SUPERADMIN ==')
    with mock.patch('core.evolution.enviar_texto', return_value=(True, 'ok')) as m:
        r = cc.post('/vendas/precos/importar/', {'arquivo': SimpleUploadedFile('comum.xlsx', planilha(base_sm, base_pvp)),
                                                 'aviso_form': '1', 'aviso_destinatarios': [str(r1.id)]})
    imp = ImportacaoPrecos.objects.order_by('-id').first()
    t('outro perfil importa normalmente', r.status_code == 302 and imp.arquivo == 'comum.xlsx' and not imp.erro
      and ItemPreco.objects.get(categoria='SMARTPHONES', nome='ZZAV Phone Sobe').valor == D('3000.00'), imp.erro)
    t('...mas o campo de quem recebe é ignorado: nada enviado', m.call_count == 0, m.call_count)
    t('...e a seleção salva não muda', set(ConfiguracaoVendas.atual().aviso_precos_destinatarios.values_list('id', flat=True))
      == {r1.id, r2.id, r3.id})
    r = cc.get('/vendas/precos/importar/')
    res = r.context['resultado']
    t('outro perfil também vê o que mudou, sem envio', res and res['envios'] is None
      and 'ZZAV Phone Sobe' in [i['nome'] for i in res['variacao']['diminuiram']], res and res.get('envios'))

    print('\n== SEM MARCAR NINGUÉM ==')
    with mock.patch('core.evolution.enviar_texto', return_value=(True, 'ok')) as m:
        cs.post('/vendas/precos/importar/', {'arquivo': SimpleUploadedFile('vazia.xlsx', planilha(novo_sm, novo_pvp)),
                                             'aviso_form': '1'})
    t('SUPERADMIN desmarca todos: nada sai e a seleção fica vazia', m.call_count == 0
      and ConfiguracaoVendas.atual().aviso_precos_destinatarios.count() == 0)
    with mock.patch('core.evolution.enviar_texto', return_value=(True, 'ok')) as m:
        cs.post('/vendas/precos/importar/', {'arquivo': SimpleUploadedFile('ruim.xlsx', b'nao e planilha'),
                                             'aviso_form': '1', 'aviso_destinatarios': [str(r1.id)]})
    t('importação que falha não manda aviso', m.call_count == 0 and ImportacaoPrecos.objects.order_by('-id').first().erro)
finally:
    transaction.set_rollback(True)
    marcador.__exit__(None, None, None)

print(f'\n{ok} OK, {fail} falha(s)')
sys.exit(1 if fail else 0)
