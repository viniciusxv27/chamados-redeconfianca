"""Ativos legado (/assets/legado/): Setor e PDV novos, e ações em vários ativos de uma vez.

Pedidos:
- na lista, selecionar vários com caixa de seleção (e fazer algo com eles: editar,
  exportar ou excluir; ou todos do filtro, em todas as páginas);
- na edição, o Setor é Loja ou Escritório; na loja o PDV é um setor com "Loja" no
  nome, no escritório qualquer setor.

Nada é gravado: roda numa transação desfeita no fim. Essas telas não mandam aviso
nenhum (nem sino, nem push, nem e-mail).
"""
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile

import django

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'redeconfianca.settings')
os.environ.setdefault('RC_VARREDURA_ROTINA', '0')
django.setup()

from django.conf import settings

if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS.append('testserver')

import openpyxl
from django.contrib.auth import get_user_model
from django.db import transaction
from django.test import Client

from assets.forms import AssetForm
from assets.models import Asset
from assets.setor_pdv import ESCRITORIO, LOJA, conferir, pdvs_por_setor, sugerir
from core.models import SystemLog
from users.models import Sector

User = get_user_model()
ok = fail = 0


def t(nome, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'  OK   {nome}')
    else:
        fail += 1
        print(f'  FALHA {nome} {extra}')


def js_valido(html, trecho, rotulo):
    node = shutil.which('node')
    script = next((s for s in re.findall(r'<script>(.*?)</script>', html, flags=re.S) if trecho in s), '')
    if not (node and script):
        print(f'  (node ou script ausente: {rotulo} não conferido)')
        return
    with tempfile.NamedTemporaryFile('w', suffix='.js', delete=False, encoding='utf-8') as arquivo:
        arquivo.write(script)
    checagem = subprocess.run([node, '--check', arquivo.name], capture_output=True, text=True, timeout=30)
    os.unlink(arquivo.name)
    t(f'o JS {rotulo} é válido (node --check)', checagem.returncode == 0, checagem.stderr[-300:])


marcador = transaction.atomic()
marcador.__enter__()
try:
    loja = Sector.objects.create(name='Loja ZZ Teste Ativos')
    outra_loja = Sector.objects.create(name='Loja ZZ Outra Teste')
    almox = Sector.objects.create(name='ZZ Almoxarifado Teste')
    vix = Sector.objects.create(name='Loja ZZ Centro VIX')
    pessoa = User.objects.create_user(username='zzativos.pessoa', email='zzativos.pessoa@exemplo-teste.local',
                                      password='S3nha!teste', first_name='ZZAtivos', last_name='Pessoa')
    c = Client()
    c.force_login(pessoa)

    def ativo(numero, **extra):
        dados = dict(patrimonio_numero=numero, nome=f'Ativo {numero}', localizado='SIM', setor='Salão',
                     pdv='ZZ Teste Ativos', estado_fisico='bom', created_by=pessoa)
        dados.update(extra)
        return Asset.objects.create(**dados)

    a1, a2, a3 = ativo('ZZTESTE-001'), ativo('ZZTESTE-002'), ativo('ZZTESTE-003', setor='Loja', pdv='ZZ Almoxarifado Teste')
    fora = Asset.objects.exclude(patrimonio_numero__startswith='ZZTESTE').order_by('pk').first()

    print('== SETOR E PDV ==')
    pdvs = pdvs_por_setor()
    t('na loja, o PDV é um setor com "Loja" no nome; no escritório, qualquer setor',
      loja.name in pdvs[LOJA] and almox.name not in pdvs[LOJA] and {loja.name, almox.name} <= set(pdvs[ESCRITORIO])
      and all('loja' in nome.lower() for nome in pdvs[LOJA]))
    t('combinações válidas passam', conferir(LOJA, loja.name, pdvs) == '' and conferir(ESCRITORIO, almox.name, pdvs) == ''
      and conferir(ESCRITORIO, loja.name, pdvs) == '')
    t('e as inválidas não', all(conferir(*par, pdvs) for par in (
        (LOJA, almox.name), (ESCRITORIO, 'Setor que não existe'), ('Salão', loja.name), (LOJA, ''))))

    print('\n== ATIVOS DO CADASTRO ANTIGO ==')
    t('PDV antigo sem o "Loja" acha a loja ("ZZ Teste Ativos" → "Loja ZZ Teste Ativos")',
      sugerir('Salão', 'ZZ Teste Ativos', pdvs) == (LOJA, loja.name), sugerir('Salão', 'ZZ Teste Ativos', pdvs))
    t('sem acento e caixa diferente também, e um setor que não é loja vira Escritório',
      sugerir('DP/ Financeiro', 'zz ALMOXARIFADO teste', pdvs) == (ESCRITORIO, almox.name))
    t('um nome dentro do outro vale só quando é um setor só ("Outra Teste" → "Loja ZZ Outra Teste")',
      sugerir('', 'Outra Teste', pdvs) == (LOJA, outra_loja.name) and sugerir('', 'Teste', pdvs) == ('', ''))
    t('"de", "Sh." e o apelido VIX não atrapalham ("ZZ Centro de Vitória" → "Loja ZZ Centro VIX")',
      sugerir('Retaguarda', 'ZZ Centro de Vitória', pdvs) == (LOJA, vix.name)
      and sugerir('Salão', 'Sh. ZZ Outra Teste', pdvs) == (LOJA, outra_loja.name))
    t('PDV que não é setor nenhum fica em branco (a pessoa escolhe)', sugerir('Loja', 'Pessoa Qualquer', pdvs) == (LOJA, ''))
    t('marcado como Loja com PDV de escritório: o PDV fica para escolher de novo',
      sugerir('Loja', 'ZZ Almoxarifado Teste', pdvs) == (LOJA, ''))

    print('\n== TELA DE EDIÇÃO ==')
    r = c.get(f'/assets/legado/{a1.pk}/edit/')
    html = r.content.decode()
    t('abre com Setor e PDV em lista (não mais texto livre)', r.status_code == 200
      and re.search(r'<select name="setor"[^>]*>', html) is not None and re.search(r'<select name="pdv"[^>]*>', html) is not None
      and 'name="setor" value=' not in html)
    t('o Setor tem só Loja e Escritório', re.findall(r'<select name="setor".*?</select>', html, re.S)
      and re.findall(r'<option value="([^"]*)"', re.findall(r'<select name="setor".*?</select>', html, re.S)[0])
      == ['', 'Loja', 'Escritório'])
    t('ativo antigo abre já no formato novo, com o aviso do que estava antes',
      'veio do cadastro antigo' in html and '“Salão”' in html and '“ZZ Teste Ativos”' in html
      and '<option value="Loja" selected>' in html and f'<option value="{loja.name}" selected>' in html)
    t('a lista de PDVs por setor vai para a tela (o PDV acompanha o Setor)', 'id="asset-pdvs"' in html
      and 'ligarSetorPdv(' in html)
    js_valido(html, 'ligarSetorPdv(document', 'da edição')

    base = {'patrimonio_numero': 'ZZTESTE-001', 'nome': 'Ativo editado', 'imei_serial': '', 'localizado': 'SIM',
            'estado_fisico': 'regular', 'observacoes': ''}
    r = c.post(f'/assets/legado/{a1.pk}/edit/', {**base, 'setor': LOJA, 'pdv': almox.name})
    a1.refresh_from_db()
    t('loja com PDV de escritório é recusada', r.status_code == 200 and 'Na loja, o PDV é uma das lojas.' in r.content.decode()
      and a1.setor == 'Salão')
    r = c.post(f'/assets/legado/{a1.pk}/edit/', {**base, 'setor': 'Salão', 'pdv': loja.name})
    t('Setor fora de Loja/Escritório é recusado', r.status_code == 200 and Asset.objects.get(pk=a1.pk).setor == 'Salão')
    r = c.post(f'/assets/legado/{a1.pk}/edit/', {**base, 'setor': ESCRITORIO, 'pdv': almox.name})
    a1.refresh_from_db()
    t('Escritório com qualquer setor salva', r.status_code == 302 and (a1.setor, a1.pdv, a1.nome) == (ESCRITORIO, almox.name, 'Ativo editado'),
      (r.status_code, a1.setor, a1.pdv))
    html = c.get(f'/assets/legado/{a1.pk}/edit/').content.decode()
    t('já no formato novo, a edição não mostra mais o aviso', 'veio do cadastro antigo' not in html
      and f'<option value="{almox.name}" selected>' in html)
    r = c.post('/assets/legado/create/', {**base, 'patrimonio_numero': 'ZZTESTE-NOVO', 'setor': LOJA, 'pdv': loja.name})
    novo = Asset.objects.filter(patrimonio_numero='ZZTESTE-NOVO').first()
    t('o cadastro novo usa o mesmo Setor/PDV', r.status_code == 302 and novo and (novo.setor, novo.pdv) == (LOJA, loja.name))
    comprido = Sector.objects.create(name='Loja ' + 'Z' * 60)
    r = c.post(f'/assets/legado/{a2.pk}/edit/', {**base, 'patrimonio_numero': 'ZZTESTE-002', 'setor': LOJA, 'pdv': comprido.name})
    t('o PDV guarda nome de setor comprido (até 100)', r.status_code == 302 and Asset.objects.get(pk=a2.pk).pdv == comprido.name)
    comprido.delete()
    Asset.objects.filter(pk=a2.pk).update(setor='Salão', pdv='ZZ Teste Ativos')

    print('\n== LISTA: SELECIONAR VÁRIOS ==')
    html = c.get('/assets/legado/', {'q': 'ZZTESTE'}).content.decode()
    t('cada ativo tem a caixa de seleção', all(f'name="ids" value="{a.pk}"' in html for a in (a1, a2, a3, novo)))
    t('e a barra: todos da página, todos do filtro e as ações', all(p in html for p in (
        'id="bulk-pagina"', 'id="bulk-todos-filtro"', 'id="bulk-editar"', 'id="bulk-exportar"', 'id="bulk-excluir"',
        'action="/assets/legado/em-massa/"', 'name="filtro_q" value="ZZTESTE"')))
    t('o formulário de edição em conjunto tem Setor (Loja/Escritório) e o PDV que depende dele',
      'id="bulk-setor"' in html and 'id="bulk-pdv"' in html and 'id="asset-pdvs"' in html and 'name="estado_fisico"' in html)
    js_valido(html, 'bulk-form', 'da lista')

    def acao(**dados):
        return c.post('/assets/legado/em-massa/', {'voltar': '/assets/legado/?q=ZZTESTE', **dados}, follow=True)

    print('\n== EDITAR OS SELECIONADOS ==')
    antes_fora = (fora.setor, fora.pdv, fora.estado_fisico) if fora else None
    logs = SystemLog.objects.count()
    r = acao(acao='editar', ids=[a2.pk, a3.pk], setor=LOJA, pdv=loja.name, estado_fisico='ruim', localizado='  NÃO  ')
    a2.refresh_from_db(); a3.refresh_from_db(); a1.refresh_from_db()
    t('muda Setor, PDV, estado e localizado só dos marcados',
      all((x.setor, x.pdv, x.estado_fisico, x.localizado) == (LOJA, loja.name, 'ruim', 'NÃO') for x in (a2, a3))
      and a1.setor == ESCRITORIO and '2 ativos atualizados' in r.content.decode(), r.content.decode()[:0])
    t('volta para a lista de onde saiu, com o filtro', r.redirect_chain and r.redirect_chain[-1][0] == '/assets/legado/?q=ZZTESTE')
    t('marca a hora da mudança e fica no log', a2.updated_at > a2.created_at
      and SystemLog.objects.filter(action_type='ADMIN_ACTION', description__contains='Editou 2 ativo').exists()
      and SystemLog.objects.count() == logs + 1)
    r = acao(acao='editar', ids=[a2.pk], estado_fisico='excelente')
    a2.refresh_from_db()
    t('o que não foi preenchido fica como estava', a2.estado_fisico == 'excelente' and a2.pdv == loja.name and a2.localizado == 'NÃO')
    r = acao(acao='editar', ids=[a2.pk], setor=LOJA, pdv=almox.name)
    a2.refresh_from_db()
    t('Setor/PDV que não combinam: nada muda', a2.pdv == loja.name and 'Nada foi alterado' in r.content.decode())
    r = acao(acao='editar', ids=[a2.pk], setor=LOJA)
    t('Setor sem PDV também não', Asset.objects.get(pk=a2.pk).pdv == loja.name and 'Escolha o PDV' in r.content.decode())
    r = acao(acao='editar', ids=[a2.pk])
    t('sem nenhum campo preenchido, avisa e não mexe', 'preencha pelo menos um campo' in r.content.decode())
    r = acao(acao='editar', ids=[], estado_fisico='bom')
    t('sem ativo marcado, avisa', 'Nenhum ativo selecionado' in r.content.decode())

    print('\n== TODOS DO FILTRO (TODAS AS PÁGINAS) ==')
    zz = Asset.objects.filter(patrimonio_numero__startswith='ZZTESTE')
    r = acao(acao='editar', todos='1', filtro_q='ZZTESTE', ids=[a1.pk], estado_fisico='pessimo')
    t('vale para todos os ativos do filtro, não só os marcados', zz.filter(estado_fisico='pessimo').count() == zz.count() == 4,
      (zz.filter(estado_fisico='pessimo').count(), zz.count()))
    t('e só para eles', fora is None or (Asset.objects.get(pk=fora.pk).setor, Asset.objects.get(pk=fora.pk).pdv,
                                          Asset.objects.get(pk=fora.pk).estado_fisico) == antes_fora)

    print('\n== EXPORTAR OS SELECIONADOS ==')
    r = c.post('/assets/legado/em-massa/', {'acao': 'exportar', 'ids': [a1.pk, a3.pk]})
    t('baixa um Excel', r.status_code == 200 and r['Content-Type'].startswith('application/vnd.openxmlformats')
      and 'ativos_selecionados_' in r['Content-Disposition'])
    planilha = openpyxl.load_workbook(io.BytesIO(r.content)).active
    numeros = [linha[0] for linha in planilha.iter_rows(min_row=2, values_only=True)]
    t('só com os marcados, com o Setor e o PDV novos', numeros == ['ZZTESTE-001', 'ZZTESTE-003']
      and [linha[4:6] for linha in planilha.iter_rows(min_row=2, values_only=True)] == [(ESCRITORIO, almox.name), (LOJA, loja.name)],
      numeros)
    r = c.get('/assets/legado/export-excel/')
    t('a exportação de tudo continua igual', r.status_code == 200 and 'ativos_' in r['Content-Disposition']
      and openpyxl.load_workbook(io.BytesIO(r.content)).active.max_row == Asset.objects.count() + 1)

    print('\n== EXCLUIR OS SELECIONADOS ==')
    r = acao(acao='excluir', ids=[a1.pk, a2.pk])
    t('sem a confirmação, nada sai', Asset.objects.filter(pk__in=[a1.pk, a2.pk]).count() == 2
      and 'Confirme a exclusão' in r.content.decode())
    r = acao(acao='excluir', ids=[a1.pk, a2.pk], confirmar='sim')
    t('confirmado, exclui só os marcados', not Asset.objects.filter(pk__in=[a1.pk, a2.pk]).exists()
      and Asset.objects.filter(pk=a3.pk).exists() and '2 ativos excluídos' in r.content.decode())
    t('e deixa no log quais foram', SystemLog.objects.filter(
        action_type='ADMIN_ACTION', description__contains='ZZTESTE-001, ZZTESTE-002').exists())

    print('\n== SEGURANÇA ==')
    r = c.post('/assets/legado/em-massa/', {'acao': 'editar', 'ids': [a3.pk], 'estado_fisico': 'bom',
                                             'voltar': 'https://exemplo.com/fora'})
    t('não redireciona para fora do portal', r.status_code == 302 and r['Location'] == '/assets/legado/', r.get('Location'))
    t('só aceita POST', c.get('/assets/legado/em-massa/').status_code == 405)
    r = Client().post('/assets/legado/em-massa/', {'acao': 'excluir', 'ids': [a3.pk], 'confirmar': 'sim'})
    t('sem login não faz nada', r.status_code == 302 and '/login' in r['Location'] and Asset.objects.filter(pk=a3.pk).exists())
    r = acao(acao='apagar-tudo', ids=[a3.pk])
    t('ação desconhecida é recusada', 'Escolha o que fazer' in r.content.decode() and Asset.objects.filter(pk=a3.pk).exists())
finally:
    transaction.set_rollback(True)
    marcador.__exit__(None, None, None)
    print('\nrollback: nada deste teste foi gravado no banco.')

print(f'\n{ok} OK / {fail} falhas')
sys.exit(1 if fail else 0)
