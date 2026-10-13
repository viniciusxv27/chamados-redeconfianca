"""Assistente: o portal inteiro (leitura como a pessoa), análises com fôlego e documentos.

Pedido: "deixe ele ainda mais inteligente: deve abranger todas as áreas do portal da
pessoa, fora os resultados que ele já tem, gerar análises detalhadas e criar documentos".

- ler_pagina abre qualquer página como o próprio usuário (mesmas permissões), só leitura;
- o laço tem fôlego (30 rodadas), raciocínio adaptativo, execução de código e fallback;
- criar_documento gera Word, Excel, Markdown e página para PDF — só do dono.

O Claude é um dublê: nenhuma chamada sai. Roda dentro de uma transação desfeita no fim
(inclusive a migração nova).
"""
import io
import json
import os
import sys
from types import SimpleNamespace
from unittest import mock

import django

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'redeconfianca.settings')
os.environ.setdefault('RC_VARREDURA_ROTINA', '0')
django.setup()

from django.conf import settings

if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS.append('testserver')

from django.contrib.auth import get_user_model
from django.db import connection, transaction
from django.db.migrations.executor import MigrationExecutor
from django.test import Client
from django.test.utils import setup_test_environment

setup_test_environment()

User = get_user_model()
ok = fail = 0


def t(nome, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'  OK   {nome}')
    else:
        fail += 1
        print(f'  FALHA {nome} {extra}')


marcador = transaction.atomic()
marcador.__enter__()
try:
    executor = MigrationExecutor(connection)
    alvo = [('assistente', '0002_documentos_e_modelos_novos')]
    if executor.migration_plan(alvo):
        executor.migrate(alvo)

    from assistente import claude_client, ferramentas, portal
    from assistente import documentos as gerador
    from assistente.models import AssistenteDocumento

    chefe = User.objects.create_user(username='zzas.chefe', email='zzas.chefe@exemplo-teste.local',
                                     password='S3nha!teste', first_name='Chefe', last_name='Teste',
                                     hierarchy='SUPERADMIN')
    comum = User.objects.create_user(username='zzas.comum', email='zzas.comum@exemplo-teste.local',
                                     password='S3nha!teste', first_name='Comum', last_name='Teste',
                                     hierarchy='PADRAO')

    print('== O PORTAL INTEIRO: MAPA E LEITURA COMO A PESSOA ==')
    from django.urls import Resolver404, resolve
    quebrados = []
    for area in portal.mapa():
        try:
            resolve(area['caminho'])
        except Resolver404:
            quebrados.append(area['caminho'])
    t('todo caminho do mapa existe no portal', not quebrados and len(portal.mapa()) >= 40, quebrados)
    for caminho in ('/files/delete-file/3/', '/sap/atualizar/', '/vendas/nova/', '/logout/',
                    '/commission/export/', '/assistente/limpar/', '/admin/', 'https://exemplo.com/x'):
        _, _, erro = portal.caminho_seguro(caminho)
        t(f'endereço de ação/arquivo recusado: {caminho}', bool(erro), erro)
    for caminho in ('/sap/painel/', '/tickets/', '/vendas/?de=2026-10-01', '/impulso/acompanhamento/5/'):
        _, _, erro = portal.caminho_seguro(caminho)
        t(f'página de leitura aceita: {caminho}', not erro, erro)

    html = ('<html><body><nav>MENU LATERAL</nav><main id="mainContent"><h1>Título</h1>'
            '<script>var x = "segredo";</script><div class="hidden">escondido</div>'
            '<table><tr><th>Loja</th><th>Valor</th></tr><tr><td>Glória</td><td>R$ 10,00</td></tr></table>'
            '<ul><li>um</li><li>dois</li></ul><a href="/tickets/9/">Chamado 9</a></main></body></html>')
    texto = portal.html_para_texto(html)
    t('HTML vira texto: sem menu, script nem escondido', 'MENU LATERAL' not in texto and 'segredo' not in texto
      and 'escondido' not in texto, texto)
    t('tabelas viram linhas com |, listas e links preservados', '| Loja | Valor |' in texto
      and '| Glória | R$ 10,00 |' in texto and '- um' in texto and '[Chamado 9](/tickets/9/)' in texto, texto)

    r = portal.ler_pagina(chefe, '/sap/painel/')
    t('SUPERADMIN lê o painel do SAP como texto', r.get('conteudo') and 'Pendências por loja' in r['conteudo'], r.keys())
    r = portal.ler_pagina(comum, '/sap/painel/')
    t('PADRÃO sem acesso ao SAP: a página manda embora (sem conteúdo)', not r.get('conteudo')
      and (r.get('redirecionou_para') or r.get('erro')), r)
    r = portal.ler_pagina(chefe, '/pagina-que-nao-existe/')
    t('página inexistente avisa', 'não encontrada' in (r.get('erro') or ''), r)
    r = portal.ler_pagina(chefe, '/tickets/')
    t('página comum abre e diz em quantas partes veio', r.get('parte') == 1 and r.get('total_partes') >= 1)
    saida = json.loads(ferramentas.executar('ler_pagina', {'caminho': '/files/delete-file/1/'}, chefe))
    t('pela ferramenta, ação também é recusada (nada roda)', 'só lê' in saida.get('erro', ''), saida)
    t('mapa_do_portal pela ferramenta', 'areas' in json.loads(ferramentas.executar('mapa_do_portal', {}, chefe)))

    print('\n== DOCUMENTOS ==')
    conteudo = ('# Resumo\n\nAs vendas subiram **12%** no mês.\n\n## Por loja\n\n'
                '| Loja | Vendas | Atingimento |\n|---|---|---|\n| Glória | R$ 1.234,56 | 98,5% |\n| Iconha | 900 | 87% |\n\n'
                '- Ponto 1\n- Ponto 2\n\n1. Ligar para a loja\n2. Revisar a meta\n')
    tabelas = [{'titulo': 'Ranking', 'colunas': ['Loja', 'Pontos'], 'linhas': [['Glória', 98.5], ['Iconha', 87]]}]
    saida = json.loads(ferramentas.executar('criar_documento', {'titulo': 'Análise de vendas — outubro',
                                                                'conteudo': conteudo, 'tabelas': tabelas}, chefe))
    doc = AssistenteDocumento.objects.get(pk=saida['id'])
    t('criar_documento grava e devolve os links', saida['ok'] and doc.user == chefe
      and set(saida['links']) == {'pagina', 'word', 'excel', 'markdown', 'pdf'}, saida)
    t('documento vazio é recusado', 'vazio' in ferramentas.executar('criar_documento', {'titulo': 'X', 'conteudo': ' '}, chefe))
    from docx import Document
    word = Document(io.BytesIO(gerador.gerar_docx(doc)))
    textos = [p.text for p in word.paragraphs]
    t('Word: título, seções, listas e o texto', 'Análise de vendas — outubro' in textos and 'Por loja' in textos
      and 'Ponto 1' in textos and 'Ligar para a loja' in textos, textos[:12])
    t('Word: a tabela do texto e a tabela estruturada viram tabelas de verdade', len(word.tables) == 2
      and word.tables[0].cell(1, 0).text == 'Glória' and word.tables[1].cell(2, 1).text == '87')
    negrito = [r.text for p in word.paragraphs for r in p.runs if r.bold]
    t('Word: o **negrito** vira negrito', '12%' in negrito, negrito)
    from openpyxl import load_workbook
    livro = load_workbook(io.BytesIO(gerador.gerar_xlsx(doc)))
    t('Excel: uma aba por tabela', livro.sheetnames == ['Ranking', 'Por loja'], livro.sheetnames)
    aba = livro['Por loja']
    t('Excel: R$ e % viram número (dá para somar)', aba['B2'].value == 1234.56 and abs(aba['C2'].value - 0.985) < 1e-9
      and aba['B3'].value == 900, (aba['B2'].value, aba['C2'].value))
    t('Markdown baixado leva título e tabelas', gerador.gerar_md(doc).decode().startswith('# Análise de vendas'))
    saida = json.loads(ferramentas.executar('revisar_documento', {'id': doc.pk, 'conteudo': '# Novo\n\nTexto novo.'}, chefe))
    doc.refresh_from_db()
    t('revisar_documento troca o conteúdo e sobe a versão', saida['ok'] and doc.versao == 2 and 'Texto novo' in doc.conteudo)
    t('ninguém revisa documento dos outros', 'não encontrado' in ferramentas.executar(
        'revisar_documento', {'id': doc.pk, 'conteudo': 'invasão'}, comum))
    t('meus_documentos_assistente lista só os do usuário',
      len(json.loads(ferramentas.executar('meus_documentos_assistente', {}, chefe))) == 1
      and json.loads(ferramentas.executar('meus_documentos_assistente', {}, comum)) == [])

    c_chefe, c_comum = Client(), Client()
    c_chefe.force_login(chefe)
    c_comum.force_login(comum)
    r = c_chefe.get(f'/assistente/documentos/{doc.pk}/')
    t('a página do documento abre para o dono, com Word, Excel e PDF', r.status_code == 200
      and 'Imprimir / PDF' in r.content.decode() and 'Texto novo' in r.content.decode())
    t('para outra pessoa, não existe', c_comum.get(f'/assistente/documentos/{doc.pk}/').status_code == 404
      and c_comum.get(f'/assistente/documentos/{doc.pk}/baixar/docx/').status_code == 404)
    r = c_chefe.get(f'/assistente/documentos/{doc.pk}/baixar/docx/')
    t('download do Word com nome de arquivo', r.status_code == 200 and 'wordprocessingml' in r['Content-Type']
      and 'analise-de-vendas-outubro.docx' in r['Content-Disposition'], r.get('Content-Disposition'))
    r = c_chefe.get(f'/assistente/documentos/{doc.pk}/baixar/xlsx/')
    t('download do Excel', r.status_code == 200 and 'spreadsheetml' in r['Content-Type'])
    t('formato desconhecido: 404', c_chefe.get(f'/assistente/documentos/{doc.pk}/baixar/exe/').status_code == 404)
    r = c_chefe.get('/assistente/documentos/')
    t('a lista de documentos mostra o documento', 'Novo' in r.content.decode() or doc.titulo in r.content.decode())
    t('o chat tem o atalho para os documentos', 'Documentos' in c_chefe.get('/assistente/').content.decode())
    c_comum.post(f'/assistente/documentos/{doc.pk}/excluir/')
    t('outra pessoa não exclui', AssistenteDocumento.objects.filter(pk=doc.pk).exists())

    print('\n== O LAÇO DO CLAUDE (DUBLÊ) ==')
    pedidos, roteiro = [], []

    class Mensagens:
        def create(self, **kwargs):
            # Cópia da lista: o laço continua completando a mesma lista depois desta chamada.
            pedidos.append({**kwargs, 'messages': list(kwargs['messages'])})
            proximo = roteiro.pop(0)
            if isinstance(proximo, Exception):
                raise proximo
            return proximo

    def resposta(*blocos, parar='end_turn'):
        return SimpleNamespace(stop_reason=parar, content=list(blocos))

    def uso(nome, entrada, ident):
        return SimpleNamespace(type='tool_use', name=nome, input=entrada, id=ident)

    def texto_(valor):
        return SimpleNamespace(type='text', text=valor)

    cliente = SimpleNamespace(messages=Mensagens(), beta=SimpleNamespace(messages=Mensagens()))
    with mock.patch.object(claude_client, '_cliente', return_value=(cliente, 'claude-opus-5-5')):
        roteiro[:] = [
            resposta(uso('ler_pagina', {'caminho': '/sap/painel/'}, 'u1'), parar='tool_use'),
            resposta(SimpleNamespace(type='server_tool_use', name='code_execution', id='s1', input={}), parar='pause_turn'),
            resposta(uso('criar_documento', {'titulo': 'Relatório SAP', 'conteudo': '# SAP\n\nTudo certo.'}, 'u2'),
                     parar='tool_use'),
            resposta(texto_('Pronto: [Relatório SAP](/assistente/documentos/1/)')),
        ]
        final = claude_client.responder(chefe, 'analise o SAP e me dê um relatório', [])
        primeiro = pedidos[0]
        t('o laço lê a página, retoma o pause_turn e cria o documento', final.startswith('Pronto')
          and AssistenteDocumento.objects.filter(user=chefe, titulo='Relatório SAP').exists(), final)
        t('com fôlego: 16 mil tokens de resposta', primeiro['max_tokens'] == 16000)
        t('raciocínio adaptativo e esforço alto no Opus 5.5', primeiro['thinking'] == {'type': 'adaptive'}
          and primeiro['output_config'] == {'effort': 'high'})
        t('execução de código e ferramentas novas na lista', any(x.get('type') == 'code_execution_20260521'
                                                                for x in primeiro['tools'])
          and {'ler_pagina', 'criar_documento', 'mapa_do_portal'} <= {x.get('name') for x in primeiro['tools']})
        t('fallback do servidor em caso de recusa', primeiro['fallbacks'] == 'default'
          and primeiro['betas'] == ['server-side-fallback-2026-07-01'])
        resultado_leitura = [b['content'] for m in pedidos[1]['messages'] if isinstance(m['content'], list)
                             for b in m['content'] if isinstance(b, dict) and b.get('type') == 'tool_result']
        t('o resultado da página voltou ao Claude', resultado_leitura and 'Pendências por loja' in resultado_leitura[0])
        t('o pause_turn não ganha mensagem extra do usuário (a API retoma sozinha)',
          pedidos[2]['messages'][-1]['role'] == 'assistant')
        t('o SYSTEM explica portal, análises e documentos', 'ler_pagina' in primeiro['system'][0]['text']
          and 'criar_documento' in primeiro['system'][0]['text'] and 'ANÁLISES' in primeiro['system'][0]['text'])

        pedidos.clear()
        roteiro[:] = [Exception('tools.1: code_execution is not supported on this model'), resposta(texto_('ok sem código'))]
        t('modelo/conta sem execução de código: segue sem ela', claude_client.responder(chefe, 'oi', []) == 'ok sem código'
          and not any(x.get('type') == 'code_execution_20260521' for x in pedidos[-1]['tools']))
        roteiro[:] = [resposta(texto_('parte'), parar='max_tokens')]
        t('resposta cortada avisa e sugere documento', 'gerar um documento' in claude_client.responder(chefe, 'oi', []))
        roteiro[:] = [resposta(parar='refusal')]
        t('recusa vira mensagem clara', 'reformular' in claude_client.responder(chefe, 'oi', []))

    with mock.patch.object(claude_client, '_cliente', return_value=(cliente, 'claude-haiku-4-5-20251001')):
        pedidos.clear()
        roteiro[:] = [resposta(texto_('rápido'))]
        claude_client.responder(chefe, 'oi', [])
        t('Haiku 4.5: sem raciocínio, sem código e sem fallback (não aceita)', 'thinking' not in pedidos[0]
          and 'fallbacks' not in pedidos[0] and not any(x.get('type') == 'code_execution_20260521'
                                                         for x in pedidos[0]['tools']))
finally:
    transaction.set_rollback(True)
    marcador.__exit__(None, None, None)
    print('\nrollback: nada deste teste foi gravado no banco (nem a migração); nenhuma chamada ao Claude saiu.')

print(f'\n{ok} OK / {fail} falhas')
sys.exit(1 if fail else 0)
