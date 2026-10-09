"""Ativos legado: log de alterações (/assets/legado/{id}).

Pedido: "tudo que for editado fique registrado em um log, com quem fez a
alteração, o que mudou, data e horário".

O que este teste cobre: cadastrar, editar pela tela (só o que mudou, com
antes → depois e rótulo legível), editar sem mudar nada (não grava), edição em
massa, importação de planilha (cria e atualiza), exclusão (some o ativo, fica o
log com o número do patrimônio), o admin, a tela do ativo e o log geral com
filtro e Excel.

Caches em memória e transação desfeita no fim (sem foto: nada sobe ao MinIO).
"""
import os
import sys
from io import BytesIO

import django

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'redeconfianca.settings')
os.environ.setdefault('RC_VARREDURA_ROTINA', '0')

from django.conf import settings

settings.CACHES = {
    'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-ah'},
    'local': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'zz-ah-2'},
}
django.setup()

from django.test.utils import setup_test_environment

setup_test_environment()
if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS.append('testserver')

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import transaction
from django.test import Client

from assets.models import Asset, AssetHistorico
from users.models import Sector

User = get_user_model()
ok = fail = 0


def t(nome, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'  OK   {nome}')
    else:
        fail += 1
        print(f'  FALHA {nome} {extra}')


marcador = transaction.atomic()
marcador.__enter__()
try:
    loja = Sector.objects.create(name='Loja ZZ Historico Ativo')
    outra = Sector.objects.create(name='Loja ZZ Historico Outra')
    ana = User.objects.create_user(username='zz.ah.ana', email='zz.ah.ana@exemplo-teste.local', password='S3nha!teste',
                                   first_name='Ana', last_name='Patrimonio', hierarchy='SUPERADMIN',
                                   is_superuser=True, is_staff=True)
    beto = User.objects.create_user(username='zz.ah.beto', email='zz.ah.beto@exemplo-teste.local', password='S3nha!teste',
                                    first_name='Beto', last_name='Conferente', hierarchy='SUPERADMIN', is_superuser=True)
    c_ana, c_beto = Client(), Client()
    c_ana.force_login(ana)
    c_beto.force_login(beto)

    base = {'patrimonio_numero': 'ZZ-9001', 'nome': 'Notebook ZZ', 'imei_serial': 'SN123', 'categoria': 'eletronico',
            'localizado': 'SIM', 'setor': 'Loja', 'pdv': loja.name, 'estado_fisico': 'bom', 'observacoes': ''}

    print('== CADASTRAR ==')
    r = c_ana.post('/assets/legado/create/', base)
    asset = Asset.objects.filter(patrimonio_numero='ZZ-9001').first()
    t('ativo cadastrado', asset is not None, getattr(r, 'context', None) and r.context.get('form').errors if r.status_code == 200 else '')
    h = AssetHistorico.objects.filter(patrimonio_numero='ZZ-9001').first()
    t('log do cadastro com quem fez', h and h.acao == 'criado' and h.usuario == ana and h.usuario_nome == 'Ana Patrimonio')
    t('com os valores iniciais (rótulo legível)', h and {'rotulo': 'Estado físico', 'antes': '', 'depois': 'Bom', 'campo': 'estado_fisico'} in h.mudancas, h and h.mudancas)

    print('== EDITAR PELA TELA ==')
    r = c_beto.post(f'/assets/legado/{asset.pk}/edit/', dict(base, nome='Notebook ZZ Pro', estado_fisico='regular', pdv=outra.name))
    h = AssetHistorico.objects.filter(asset=asset, acao='editado').first()
    t('edição registrada por quem editou', h and h.usuario == beto and h.origem == 'tela', r.status_code)
    campos = {m['campo']: m for m in (h.mudancas if h else [])}
    t('só os campos que mudaram', set(campos) == {'nome', 'estado_fisico', 'pdv'}, set(campos))
    t('antes → depois', campos.get('nome', {}).get('antes') == 'Notebook ZZ' and campos.get('nome', {}).get('depois') == 'Notebook ZZ Pro')
    t('escolha mostra o rótulo (Bom → Regular)', campos.get('estado_fisico', {}).get('antes') == 'Bom'
      and campos.get('estado_fisico', {}).get('depois') == 'Regular')
    t('data e horário gravados', h and h.quando is not None)
    antes = AssetHistorico.objects.count()
    c_beto.post(f'/assets/legado/{asset.pk}/edit/', dict(base, nome='Notebook ZZ Pro', estado_fisico='regular', pdv=outra.name))
    t('salvar sem mudar nada não grava', AssetHistorico.objects.count() == antes)

    print('== EM MASSA ==')
    outro = Asset.objects.create(patrimonio_numero='ZZ-9002', nome='Cadeira ZZ', categoria='movel', localizado='SIM',
                                 setor='Loja', pdv=loja.name, estado_fisico='bom', created_by=ana)
    c_ana.post('/assets/legado/em-massa/', {'acao': 'editar', 'ids': [asset.pk, outro.pk], 'estado_fisico': 'ruim',
                                           'voltar': '/assets/legado/'})
    massa = AssetHistorico.objects.filter(origem='massa', acao='editado')
    t('uma linha por ativo alterado em massa', massa.count() == 2, massa.count())
    t('com o estado antes de cada um', {(m.patrimonio_numero, m.mudancas[0]['antes']) for m in massa}
      == {('ZZ-9001', 'Regular'), ('ZZ-9002', 'Bom')}, [(m.patrimonio_numero, m.mudancas) for m in massa])

    print('== PLANILHA ==')
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['Nº', 'Nome', 'IMEI', 'Localizado', 'Setor', 'PDV', 'Estado', 'Obs', 'Criado por', 'Criado', 'Atualizado', 'Categoria'])
    ws.append(['ZZ-9001', 'Notebook ZZ Ultra', 'SN123', 'SIM', 'Loja', outra.name, 'Ruim', 'veio da planilha', '', '', '', 'Eletrônico'])
    ws.append(['ZZ-9003', 'Mesa ZZ', '', 'NÃO', 'Loja', loja.name, 'Bom', '', '', '', '', 'Móvel'])
    arquivo = BytesIO()
    wb.save(arquivo)
    c_ana.post('/assets/legado/import-excel/', {'excel_file': SimpleUploadedFile('ativos.xlsx', arquivo.getvalue())})
    pl = AssetHistorico.objects.filter(origem='planilha')
    t('planilha: cadastro e edição no log', set(pl.values_list('patrimonio_numero', 'acao'))
      == {('ZZ-9001', 'editado'), ('ZZ-9003', 'criado')}, list(pl.values_list('patrimonio_numero', 'acao')))
    ed = pl.filter(acao='editado').first()
    t('planilha: o que mudou', ed and {m['campo'] for m in ed.mudancas} == {'nome', 'observacoes'}, ed and ed.mudancas)

    print('== ADMIN ==')
    r = c_ana.post(f'/admin/assets/asset/{outro.pk}/change/', {
        'patrimonio_numero': 'ZZ-9002', 'nome': 'Cadeira ZZ Gamer', 'imei_serial': '', 'categoria': 'movel',
        'localizado': 'SIM', 'setor': 'Loja', 'pdv': loja.name, 'estado_fisico': 'ruim', 'observacoes': '',
        'created_by': ana.pk, '_save': 'Salvar'})
    adm = AssetHistorico.objects.filter(origem='admin', asset=outro).first()
    t('edição pelo admin também entra', adm and adm.mudancas[0]['depois'] == 'Cadeira ZZ Gamer',
      (r.status_code, adm and adm.mudancas))

    print('== TELAS ==')
    r = c_ana.get(f'/assets/legado/{asset.pk}/')
    html = r.content.decode()
    t('tela do ativo mostra o histórico', 'Histórico de alterações' in html and 'Beto Conferente' in html
      and 'Notebook ZZ Pro' in html and 'edição em massa' in html and 'importação de planilha' in html)
    r = c_ana.get('/assets/legado/historico/?q=ZZ-900')
    html = r.content.decode()
    t('log geral lista tudo', r.status_code == 200 and 'ZZ-9001' in html and 'ZZ-9002' in html and 'ZZ-9003' in html)
    r = c_ana.get(f'/assets/legado/historico/?q=ZZ-900&usuario={beto.pk}')
    t('filtro por usuário', 'Beto Conferente' in r.content.decode() and 'ZZ-9003' not in r.content.decode())
    r = c_ana.get('/assets/legado/historico/?q=ZZ-900&exportar=xlsx')
    livro = openpyxl.load_workbook(BytesIO(r.content))
    linhas = list(livro.active.iter_rows(min_row=2, values_only=True))
    t('Excel com uma linha por campo alterado', any(l[2] == 'ZZ-9001' and l[7] == 'Nome' and l[8] == 'Notebook ZZ'
                                                   and l[9] == 'Notebook ZZ Pro' and l[6] == 'Beto Conferente' for l in linhas))
    r = c_ana.get('/assets/legado/')
    t('botão do log na lista', 'Log de alterações' in r.content.decode())

    print('== EXCLUIR ==')
    c_ana.post(f'/assets/legado/{asset.pk}/delete/')
    ex = AssetHistorico.objects.filter(patrimonio_numero='ZZ-9001', acao='excluido').first()
    t('exclusão registrada com o que o ativo tinha', ex and ex.usuario == ana and
      any(m['campo'] == 'nome' and m['antes'] == 'Notebook ZZ Ultra' for m in ex.mudancas))
    t('o log antigo continua depois do ativo apagado', AssetHistorico.objects.filter(patrimonio_numero='ZZ-9001').count() >= 4)
    c_ana.post('/assets/legado/em-massa/', {'acao': 'excluir', 'confirmar': 'sim', 'ids': [outro.pk], 'voltar': '/assets/legado/'})
    t('exclusão em massa também', AssetHistorico.objects.filter(patrimonio_numero='ZZ-9002', acao='excluido', origem='massa').exists())
finally:
    marcador.__exit__(Exception, Exception('rollback'), None)

print(f'\n{ok} OK, {fail} falha(s)')
sys.exit(1 if fail else 0)
